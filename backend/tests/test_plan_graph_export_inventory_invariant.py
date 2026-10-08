"""Инвариант инвентаризации плана-графика (владелец 09.10.2026, прод ФАДМ
2026_2, id 89 — «строки закупок теряются»). Фикстура держит плановые позиции
и закупки на категориях уровня 1, 2, 3 И 4 (дерево местами глубже 3 уровней),
закупку, чья позиция привязана к НЕАКТИВНОЙ плановой позиции (раньше
пропадала — ни purchased_by_item, ни purchased_by_cat), голову рамочного
договора (должна быть исключена) и закупку без категории ФЭО (секция «без
категории»).

Проверяется:
  1. Каждая из 8 закупленных позиций встречается РОВНО ОДИН раз и в
     collect_plan_graph_rows, и в листе «по направлениям», и в листе
     «по порядку» (строки уровня «Закупка»).
  2. Σ по стадиям (Запланировано/Договор/Заказано/Поставлено/Оплачено) по
     этим строкам совпадает с цифрами, посчитанными вручную по статусам/
     суммам фикстуры (независимая проверка, не вызов cascade_by_stage).
  3. На листе «по направлениям» верхняя строка 1 (=SUMIFS по «Уровень»=
     "Закупка") и строка 2 (=SUBTOTAL(109,...) по всей колонке) дают ОДНО И
     ТО ЖЕ число — вычислено маленьким интерпретатором формул (вложенные
     SUBTOTAL игнорируются, как в Excel/Google Sheets)."""
import re
import uuid

import pytest
from openpyxl.utils import column_index_from_string

from app.models.feo_category import FeoCategory
from app.models.feo_planned_item import FeoPlannedItem
from app.models.purchase import Purchase
from app.models.purchase_item import PurchaseItem
from app.services.plan_graph_export_columns import resolve_selected_keys
from app.services.plan_graph_export_data import gather_live_plan_graph_data
from app.services.plan_graph_export_flat_sheet import write_flat_plan_graph_sheet
from app.services.plan_graph_export_rows import collect_plan_graph_rows
from app.services.plan_graph_export_xlsx import build_live_plan_graph_xlsx


async def _make_subsidy(db_session, budget=5_000_000):
    from app.models.subsidy import Subsidy
    s = Subsidy(
        name=f"TestSubsidy-{uuid.uuid4().hex[:8]}", year=2026, budget=budget,
        require_planned_dates=False, status="approved",
    )
    db_session.add(s)
    await db_session.commit()
    await db_session.refresh(s)
    return s


async def _make_cat(db_session, subsidy_id, level, parent_id=None, name="Статья"):
    cat = FeoCategory(subsidy_id=subsidy_id, parent_id=parent_id, level=level, name=name)
    db_session.add(cat)
    await db_session.commit()
    await db_session.refresh(cat)
    return cat


async def _make_item(db_session, feo_category_id, name, amount=0, is_active=True):
    item = FeoPlannedItem(feo_category_id=feo_category_id, name=name, amount=amount, is_active=is_active)
    db_session.add(item)
    await db_session.commit()
    await db_session.refresh(item)
    return item


async def _make_purchase(db_session, **kwargs) -> Purchase:
    defaults = dict(
        status="ordered", purchase_contract_type=None, parent_purchase_id=None,
        contract_id=None, contract_number=None, feo_category_id=None,
    )
    defaults.update(kwargs)
    p = Purchase(**defaults)
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)
    return p


async def _make_purchase_item(db_session, purchase_id, **kwargs) -> PurchaseItem:
    defaults = dict(item_name="Позиция", quantity=1, unit_price=0, total_price=0)
    defaults.update(kwargs)
    pi = PurchaseItem(purchase_id=purchase_id, **defaults)
    db_session.add(pi)
    await db_session.commit()
    await db_session.refresh(pi)
    return pi


async def _build_fixture(db_session):
    """Возвращает (subsidy, positions, head_id). positions —
    {label: (purchase_id, effective_amount, status)} — 8 РЕАЛЬНЫХ позиций,
    которые обязаны пройти сквозь выгрузку; head_id — рамочная голова,
    обязана быть ИСКЛЮЧЕНА."""
    sub = await _make_subsidy(db_session)
    cat1 = await _make_cat(db_session, sub.id, 1, name="Направление А")
    cat2 = await _make_cat(db_session, sub.id, 2, cat1.id, name="Тип Б")
    cat3 = await _make_cat(db_session, sub.id, 3, cat2.id, name="Статья В")
    cat4 = await _make_cat(db_session, sub.id, 4, cat3.id, name="Подстатья Г")

    positions = {}

    # 1) Плановая позиция прямо на level1 (cat1) — 100 ₽, paid.
    item_l1 = await _make_item(db_session, cat1.id, "Позиция на направлении", amount=100)
    p_l1 = await _make_purchase(db_session, subsidy_id=sub.id, status="paid")
    await _make_purchase_item(
        db_session, p_l1.id, feo_planned_item_id=item_l1.id, item_name="Товар L1", unit_price=100, total_price=100,
    )
    positions["l1"] = (p_l1.id, 100.0, "paid")

    # 2) Плановая позиция на level2 (cat2) — 200 ₽, ordered.
    item_l2 = await _make_item(db_session, cat2.id, "Позиция на типе", amount=200)
    p_l2 = await _make_purchase(db_session, subsidy_id=sub.id, status="ordered")
    await _make_purchase_item(
        db_session, p_l2.id, feo_planned_item_id=item_l2.id, item_name="Товар L2", unit_price=200, total_price=200,
    )
    positions["l2"] = (p_l2.id, 200.0, "ordered")

    # 3) Обычная позиция на level3 (cat3) — 300 ₽, delivered.
    item_l3 = await _make_item(db_session, cat3.id, "Позиция на статье", amount=300)
    p_l3 = await _make_purchase(db_session, subsidy_id=sub.id, status="delivered")
    await _make_purchase_item(
        db_session, p_l3.id, feo_planned_item_id=item_l3.id, item_name="Товар L3", unit_price=300, total_price=300,
    )
    positions["l3"] = (p_l3.id, 300.0, "delivered")

    # 3b) Закупка, чья позиция привязана к НЕАКТИВНОЙ плановой позиции (cat3)
    # — владелец: «пропадают закупки 909/975/2044/3218» — 400 ₽, paid.
    item_l3_inactive = await _make_item(db_session, cat3.id, "Удалённая позиция", amount=0, is_active=False)
    p_l3b = await _make_purchase(db_session, subsidy_id=sub.id, feo_category_id=cat3.id, status="paid")
    await _make_purchase_item(
        db_session, p_l3b.id, feo_planned_item_id=item_l3_inactive.id,
        item_name="Товар на неактивной позиции", unit_price=400, total_price=400,
    )
    positions["l3_inactive_link"] = (p_l3b.id, 400.0, "paid")

    # 4) Плановая позиция на level4 (cat4, глубже 3 уровней) — 500 ₽, contracted.
    item_l4 = await _make_item(db_session, cat4.id, "Позиция на подстатье", amount=500)
    p_l4 = await _make_purchase(db_session, subsidy_id=sub.id, status="contracted")
    await _make_purchase_item(
        db_session, p_l4.id, feo_planned_item_id=item_l4.id, item_name="Товар L4", unit_price=500, total_price=500,
    )
    positions["l4"] = (p_l4.id, 500.0, "contracted")

    # 5) Закупка прямо на level1 (cat1), БЕЗ плановой позиции (itemless) —
    # 600 ₽, work_in_progress (стадия "Запланировано").
    p_cat1_direct = await _make_purchase(
        db_session, subsidy_id=sub.id, feo_category_id=cat1.id, status="work_in_progress",
        item_name="Прямая закупка на направлении", planned_total_price=600,
    )
    positions["cat1_direct"] = (p_cat1_direct.id, 600.0, "work_in_progress")

    # 6) Голова рамочного договора (cat3) — ДОЛЖНА быть исключена целиком —
    # и её заказ (itemless, 700 ₽, ordered) — должен остаться.
    head = await _make_purchase(
        db_session, subsidy_id=sub.id, feo_category_id=cat3.id, status="ordered",
        purchase_contract_type="framework_with_amount", item_name="Рамочный договор",
    )
    order = await _make_purchase(
        db_session, subsidy_id=sub.id, feo_category_id=cat3.id, status="ordered",
        parent_purchase_id=head.id, item_name="Заказ по рамочному", planned_total_price=700,
    )
    positions["order_under_head"] = (order.id, 700.0, "ordered")

    # 7) Закупка без категории ФЭО вовсе (только subsidy_id) — 800 ₽, paid.
    p_unlinked = await _make_purchase(
        db_session, subsidy_id=sub.id, status="paid",
        item_name="Закупка без категории", planned_total_price=800,
    )
    positions["unlinked"] = (p_unlinked.id, 800.0, "paid")

    return sub, positions, head.id


def _expected_stage_sums(positions: dict) -> dict:
    """Вручную — НЕ через cascade_by_stage — по статусу/сумме каждой
    позиции (независимая проверка)."""
    order_by_status = {
        "wishes": 0, "plan_schedule": 0, "work_in_progress": 0,
        "contracted": 1, "ordered": 2, "delivered": 3, "paid": 4,
    }
    stage_keys = ("planned", "contract", "ordered", "delivered", "paid")
    totals = {s: 0.0 for s in stage_keys}
    for _pid, amount, status in positions.values():
        idx = order_by_status[status]
        for i, s in enumerate(stage_keys):
            if i <= idx:
                totals[s] += amount
    return totals


class _FormulaEvaluator:
    """Крошечный интерпретатор =SUBTOTAL(109,...)/=SUMIFS(...) по уже
    построенной книге openpyxl — для проверки, что верхние строки «Итого
    всего»/«Итого по фильтру» дают одинаковое число. Ячейка, сама содержащая
    формулу SUBTOTAL, пропускается при суммировании внешним SUBTOTAL — ровно
    правило Excel/Google Sheets («вложенные SUBTOTAL игнорируются») — её
    собственные листовые потомки уже стоят отдельными строками того же
    диапазона и посчитаются сами."""

    def __init__(self, ws):
        self.ws = ws

    def _cell(self, row, col):
        return self.ws.cell(row=row, column=col).value

    @staticmethod
    def _split_ref(ref: str):
        m = re.match(r"([A-Z]+)(\d+)", ref)
        return m.group(1), int(m.group(2))

    def subtotal_109(self, formula: str) -> float:
        inner = formula[len("=SUBTOTAL(109,"):-1]
        start_ref, end_ref = inner.split(":")
        col_letter, start_row = self._split_ref(start_ref)
        _, end_row = self._split_ref(end_ref)
        col = column_index_from_string(col_letter)
        total = 0.0
        for r in range(start_row, end_row + 1):
            v = self._cell(r, col)
            if isinstance(v, str) and v.startswith("=SUBTOTAL("):
                continue
            if isinstance(v, (int, float)):
                total += v
        return total

    def sumifs_level(self, formula: str, level_col: int) -> float:
        inner = formula[len('=SUMIFS('):-1]
        sum_range = inner.split(",", 1)[0]
        col_letter, start_row = self._split_ref(sum_range.split(":")[0])
        _, end_row = self._split_ref(sum_range.split(":")[1])
        col = column_index_from_string(col_letter)
        total = 0.0
        for r in range(start_row, end_row + 1):
            if self._cell(r, level_col) == "Закупка":
                v = self._cell(r, col)
                if isinstance(v, (int, float)):
                    total += v
        return total


@pytest.mark.asyncio
async def test_positions_at_any_level_appear_exactly_once_and_sums_match(db_session):
    sub, positions, head_id = await _build_fixture(db_session)
    expected_ids = {pid for pid, _amount, _status in positions.values()}
    expected_sums = _expected_stage_sums(positions)
    assert expected_sums == {
        "planned": 3600.0, "contract": 3000.0, "ordered": 2500.0, "delivered": 1600.0, "paid": 1300.0,
    }

    data = await gather_live_plan_graph_data(db_session, sub.id)
    rows_info = collect_plan_graph_rows(data)

    collected_ids = []
    stage_keys = ("planned", "contract", "ordered", "delivered", "paid")
    actual_sums = {s: 0.0 for s in stage_keys}

    def _consume(rows):
        for r in rows:
            collected_ids.append(r["_purchase_id"])
            for i, k in enumerate(stage_keys):
                actual_sums[k] += r["_stages"][i]

    for rows in rows_info["fact_rows_by_item"].values():
        _consume(rows)
    for rows in rows_info["fact_rows_by_cat"].values():
        _consume(rows)
    _consume(rows_info["fact_rows_unlinked"])

    assert len(collected_ids) == len(expected_ids) == 8
    assert set(collected_ids) == expected_ids  # каждая позиция — РОВНО один раз, ничего лишнего/дублей
    assert head_id not in collected_ids  # голова рамочного исключена
    assert actual_sums == expected_sums  # Σ по стадиям == посчитанному вручную по БД

    # ── Лист «по направлениям» ───────────────────────────────────────────
    wb = build_live_plan_graph_xlsx(sub, "http://example.test", data)
    ws = wb["План закупок (по направлениям)"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[3])}
    level_col = headers["Уровень"]
    rows_h = list(ws.iter_rows(min_row=4, values_only=False))

    leaf_rows = [r for r in rows_h if r[level_col - 1].value == "Закупка"]
    assert len(leaf_rows) == 8

    paid_col = headers["Оплачено, ₽"]
    evaluator = _FormulaEvaluator(ws)
    row1_formula = ws.cell(row=1, column=paid_col).value
    row2_formula = ws.cell(row=2, column=paid_col).value
    assert str(row1_formula).startswith("=SUMIFS(")
    assert str(row2_formula).startswith("=SUBTOTAL(109,")
    row1_value = evaluator.sumifs_level(row1_formula, level_col)
    row2_value = evaluator.subtotal_109(row2_formula)
    assert row1_value == row2_value == expected_sums["paid"] == 1300.0

    # ── Лист «по порядку» — те же 8 позиций, тоже ровно по одному разу ────
    import openpyxl
    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    selected = resolve_selected_keys(None)
    write_flat_plan_graph_sheet(wb2, data, selected, "http://example.test")
    ws2 = wb2["План закупок (по порядку)"]
    headers2 = {cell.value: idx + 1 for idx, cell in enumerate(ws2[3])}
    level_col2 = headers2["Уровень"]
    rows2 = list(ws2.iter_rows(min_row=4, values_only=False))
    leaf_rows2 = [r for r in rows2 if r[level_col2 - 1].value == "Закупка"]
    assert len(leaf_rows2) == 8


@pytest.mark.asyncio
async def test_plan_amount_matches_tree_with_substitution_and_manual_sum(db_session, test_org):
    """Владелец 09.10.2026, сверка на проде ФАДМ 2026_2 (id 89) — Σ «План» в
    выгрузке (15 883 686,50) расходилась с деревом/subsidy_type_totals
    (15 880 300,00). Причина: дерево (feo_plan_tree._manual_plan_for) считает
    план УЗЛА с plan_source='planned_items' как Σ КОНТРИБЬЮЦИЙ позиций
    (committed_amounts.planned_item_contributions — закрытая количеством
    позиция замещается её законтрактованной суммой, экономия высвобождается),
    а экспорт раньше брал голое FeoPlannedItem.amount. Фикстура — ОДИН К
    ОДНОМУ с test_money_committed.py::test_2_both_items_closed_releases_savings
    (ПРАВИЛО №6 — не второй сценарий): позиция quantity=2/amount=200000,
    закрыта двумя contracted-закупками 101000+80000=181000 → её вклад в план
    181000, не 200000. Плюс контрольная позиция в категории plan_source=
    'manual_sum' (amount=50000), тоже полностью законтрактованная, — там
    замещение НЕ применяется (дерево явно читает голую Σ amount для этого
    режима) — её вклад остаётся 50000."""
    from app.services.type_totals import subsidy_type_totals
    from tests.test_feo_plan_tree_scenarios import _make_category, _make_planned_item, _make_subsidy
    from tests.test_money_committed import _make_linked_purchase

    subsidy = await _make_subsidy(db_session, test_org.id, budget=1_000_000)
    leaf = await _make_category(db_session, subsidy.id, name="Бензопила (лист)")
    fpi = await _make_planned_item(db_session, leaf.id, "Бензопила", 2, 200_000)
    await _make_linked_purchase(db_session, subsidy.id, leaf.id, fpi.id, quantity=1, contract_price=101_000)
    await _make_linked_purchase(db_session, subsidy.id, leaf.id, fpi.id, quantity=1, contract_price=80_000)

    # manual_plan_amount == Σ позиций (50000), иначе дерево видело бы
    # несогласованное превышение и держало бы план на manual_plan_amount
    # (здесь — 0, т.к. он не введён) — см. _manual_plan_for в feo_plan_tree.py.
    from decimal import Decimal as _D
    manual_leaf = await _make_category(
        db_session, subsidy.id, name="Ручная сумма (лист)",
        plan_source="manual_sum", manual_plan_amount=_D("50000"),
    )
    fpi_manual = await _make_planned_item(db_session, manual_leaf.id, "Ручная позиция", 1, 50_000)
    await _make_linked_purchase(db_session, subsidy.id, manual_leaf.id, fpi_manual.id, quantity=1, contract_price=50_000)

    _totals = (await subsidy_type_totals(db_session, [subsidy.id]))[subsidy.id]
    expected_plan_total = sum(
        _totals[k] for k in ("plan_goods", "plan_services", "plan_payroll", "plan_unspecified")
    )
    assert expected_plan_total == 231_000.0  # 181000 (замещено) + 50000 (manual_sum, не замещено)

    data = await gather_live_plan_graph_data(db_session, subsidy.id)
    assert data["item_plan_map"][fpi.id] == 181_000.0
    assert data["item_plan_map"][fpi_manual.id] == 50_000.0

    # ── Лист «по направлениям» — верхняя строка SUMIFS(level="Позиция"). ────
    wb = build_live_plan_graph_xlsx(subsidy, "http://example.test", data)
    ws = wb["План закупок (по направлениям)"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[3])}
    plan_ci = headers["План (плановые позиции), ₽"]
    level_col = headers["Уровень"]
    rows_h = list(ws.iter_rows(min_row=4, values_only=False))
    evaluator = _FormulaEvaluator(ws)
    row1_formula = ws.cell(row=1, column=plan_ci).value
    assert str(row1_formula).startswith("=SUMIFS(")
    row1_value = sum(
        r[plan_ci - 1].value for r in rows_h
        if r[level_col - 1].value == "Позиция" and isinstance(r[plan_ci - 1].value, (int, float))
    )
    assert row1_value == expected_plan_total

    # ── Лист «по порядку» — простой SUM по столбцу (владелец, часть 3). ─────
    import openpyxl
    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    selected = resolve_selected_keys(None)
    write_flat_plan_graph_sheet(wb2, data, selected, "http://example.test")
    ws2 = wb2["План закупок (по порядку)"]
    headers2 = {cell.value: idx + 1 for idx, cell in enumerate(ws2[3])}
    plan_ci2 = headers2["План (плановые позиции), ₽"]
    rows2_h = list(ws2.iter_rows(min_row=4, values_only=False))
    flat_total = sum(
        r[plan_ci2 - 1].value for r in rows2_h if isinstance(r[plan_ci2 - 1].value, (int, float))
    )
    assert flat_total == expected_plan_total
