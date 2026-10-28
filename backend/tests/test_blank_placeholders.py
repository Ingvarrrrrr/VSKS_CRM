# -*- coding: utf-8 -*-
"""Тест модуля app/services/documents/blank_placeholders.py (ПРАВИЛО №6).

Владелец (28.10.2026): если дата/номер в шаблоне есть, а в данных пусто —
печатать линию подчёркивания ("__.__.____" для даты, 21 "_" для номера),
чтобы вписать от руки после распечатки, а не "" (пустое место без ориентира,
где писать) и не падать на отсутствующем значении.

Три уровня проверки:
1. Blank — печатается линией, но ложен в Python/Jinja (bool/len/==).
2. apply_blank_placeholders — какие ключи гасятся, какие нет (денежные/
   служебные/today, item.num, вложенные items/items_list).
3. Сквозной рендер маленького docx-шаблона (собран здесь же через
   python-docx) через реальный docxtpl — пустые значения печатаются линией,
   {% if %}/{% else %} работает как и раньше с "".

Плюс статическая проверка: все найденные ранее места рендера документов
зовут render_docx(), а не tpl.render()/doc.render() напрямую — одна точка
применения гашения (ПРАВИЛО №6), не вторая копия логики по месту рендера.
"""
import os
import re
import sys
from io import BytesIO

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import pytest
from docx import Document
from docxtpl import DocxTemplate

from app.services.documents.blank_placeholders import (
    Blank,
    BLANK_DATE,
    BLANK_NUMBER,
    BLANK_DAY,
    BLANK_MONTH,
    BLANK_YEAR,
    apply_blank_placeholders,
    render_docx,
)


# ── 1. Blank ──────────────────────────────────────────────────────────────

def test_blank_prints_as_line_but_falsy():
    b = Blank(BLANK_DATE)
    assert str(b) == "__.__.____"
    assert bool(b) is False
    assert len(b) == 0
    assert (b == "") is True
    assert (b == None) is True  # noqa: E711 — проверяем именно __eq__, не `is`
    assert not b  # as in `if not value`


def test_blank_in_if_condition_takes_else_branch():
    b = Blank(BLANK_DATE)
    if b:
        taken = "then"
    else:
        taken = "else"
    assert taken == "else"


# ── 2. apply_blank_placeholders — выбор ключей ──────────────────────────────

def test_empty_date_field_blanked():
    out = apply_blank_placeholders({"delivery_date": None, "contract_end_date": ""})
    assert str(out["delivery_date"]) == BLANK_DATE
    assert str(out["contract_end_date"]) == BLANK_DATE
    assert not out["delivery_date"]
    assert not out["contract_end_date"]


def test_filled_date_field_not_touched():
    out = apply_blank_placeholders({"delivery_date": "20.01.2026"})
    assert out["delivery_date"] == "20.01.2026"
    assert not isinstance(out["delivery_date"], Blank)


def test_empty_number_field_blanked():
    out = apply_blank_placeholders({
        "procurement_order_number": "",
        "contract_number": None,
        "registry_number": "   ",
    })
    assert str(out["procurement_order_number"]) == BLANK_NUMBER
    assert str(out["contract_number"]) == BLANK_NUMBER
    assert str(out["registry_number"]) == BLANK_NUMBER


def test_contract_date_parts_use_their_own_blank_text():
    out = apply_blank_placeholders({
        "contract_date_day": "",
        "contract_date_month": None,
        "contract_date_year": "",
    })
    assert str(out["contract_date_day"]) == BLANK_DAY
    assert str(out["contract_date_month"]) == BLANK_MONTH
    assert str(out["contract_date_year"]) == BLANK_YEAR


def test_money_num_fields_never_blanked():
    """contract_price_num/doc_amount_num/vat_amount_num оканчиваются на
    "_num", как и номера, но это суммы цифрами — денежные поля правило не
    трогает (владелец: «денежные/текстовые поля не трогать»)."""
    out = apply_blank_placeholders({
        "contract_price_num": "",
        "doc_amount_num": None,
        "vat_amount_num": "",
    })
    assert out["contract_price_num"] == ""
    assert out["doc_amount_num"] is None
    assert out["vat_amount_num"] == ""


def test_today_never_blanked():
    out = apply_blank_placeholders({"today": "", "today_iso": ""})
    assert out["today"] == ""
    assert out["today_iso"] == ""


def test_submission_deadline_datetime_not_treated_as_plain_date():
    """Составной формат "дата+время" — линия "__.__.____" испортит формат,
    общее правило по имени его не должно ловить."""
    out = apply_blank_placeholders({"submission_deadline_datetime": ""})
    assert out["submission_deadline_datetime"] == ""


def test_item_num_field_excluded_but_item_dates_numbers_blanked():
    out = apply_blank_placeholders({
        "items": [
            {"num": "", "name": "Товар", "item_number": "", "delivery_date": ""},
        ],
    })
    item = out["items"][0]
    assert item["num"] == ""  # «кроме item 'num'»
    assert str(item["item_number"]) == BLANK_NUMBER
    assert str(item["delivery_date"]) == BLANK_DATE


def test_items_list_alias_also_processed():
    out = apply_blank_placeholders({
        "items_list": [{"delivery_date": None}],
    })
    assert str(out["items_list"][0]["delivery_date"]) == BLANK_DATE


def test_context_not_mutated():
    ctx = {"delivery_date": None, "items": [{"delivery_date": None}]}
    apply_blank_placeholders(ctx)
    assert ctx["delivery_date"] is None
    assert ctx["items"][0]["delivery_date"] is None


# ── 3. Сквозной рендер через docxtpl ────────────────────────────────────────

def _build_minimal_template() -> BytesIO:
    doc = Document()
    doc.add_paragraph("Дата поставки: {{ delivery_date }}")
    doc.add_paragraph("Приказ № {{ procurement_order_number }}")
    doc.add_paragraph(
        "Договор действует до {% if contract_end_date %}{{ contract_end_date }}"
        "{% else %}полного исполнения обязательств{% endif %}"
    )
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _all_text(tpl: DocxTemplate) -> str:
    return "\n".join(p.text for p in tpl.docx.paragraphs)


def test_render_docx_blanks_empty_date_and_number_and_keeps_if_else():
    tpl = DocxTemplate(_build_minimal_template())
    render_docx(tpl, {
        "delivery_date": None,
        "procurement_order_number": "",
        "contract_end_date": "",
    })
    text = _all_text(tpl)
    assert "Дата поставки: __.__.____" in text
    assert f"Приказ № {BLANK_NUMBER}" in text
    assert "до полного исполнения обязательств" in text


def test_render_docx_keeps_filled_values():
    tpl = DocxTemplate(_build_minimal_template())
    render_docx(tpl, {
        "delivery_date": "20.01.2026",
        "procurement_order_number": "17",
        "contract_end_date": "31.12.2026",
    })
    text = _all_text(tpl)
    assert "Дата поставки: 20.01.2026" in text
    assert "Приказ № 17" in text
    assert "до 31.12.2026" in text
    assert "полного исполнения" not in text


# ── 4. Все точки рендера зовут render_docx, не tpl.render() напрямую ───────

_BACKEND_APP = os.path.join(os.path.dirname(__file__), "..", "app")

# Файлы, где .render( на DocxTemplate — НЕ печать реального документа:
# - subsidy_templates.py: tpl.render({}) — пробный рендер шаблона, проверка
#   что шаблон не падает (пустой context, гасить в нём нечего);
# - smoke_and_seeds.py: стартовый smoke-тест на фиктивных данных, не выдаётся
#   пользователю.
_ALLOWED_DIRECT_RENDER = {
    os.path.normpath(os.path.join(_BACKEND_APP, "routers", "subsidy_templates.py")),
    os.path.normpath(os.path.join(_BACKEND_APP, "startup", "smoke_and_seeds.py")),
    os.path.normpath(os.path.join(_BACKEND_APP, "services", "documents", "blank_placeholders.py")),
}

_RENDER_CALL_RE = re.compile(r"\b(\w+)\.render\(")

# Положительная проверка: конкретные файлы, где раньше жил прямой
# tpl.render(context)/doc.render(context) на реальном документе — теперь
# должны импортировать render_docx (иначе предыдущая регрессионная проверка
# была бы бессмысленна — ничего не находит, потому что ничего не искала).
_KNOWN_CONVERTED_FILES = (
    os.path.join(_BACKEND_APP, "services", "documents", "stages_render.py"),
    os.path.join(_BACKEND_APP, "services", "documents", "stages_merge.py"),
    os.path.join(_BACKEND_APP, "services", "documents", "fabrikant_package.py"),
    os.path.join(_BACKEND_APP, "routers", "wish_documents.py"),
    os.path.join(_BACKEND_APP, "routers", "trips_waybill.py"),
    os.path.join(_BACKEND_APP, "services", "waybill_docx.py"),
    os.path.join(_BACKEND_APP, "services", "plan_graph_export_docx.py"),
)


@pytest.mark.parametrize("fpath", _KNOWN_CONVERTED_FILES)
def test_known_render_point_imports_render_docx(fpath):
    with open(fpath, encoding="utf-8") as fh:
        src = fh.read()
    assert "render_docx" in src, f"{fpath} должен использовать render_docx()"
    assert ".render(context)" not in src and ".render(ctx)" not in src


def test_all_known_render_points_use_render_docx():
    """Регрессионный guard: любой БУДУЩИЙ `tpl.render(context)`/`doc.render(...)`
    на объекте docxtpl, добавленный мимо render_docx(), должен провалить этот
    тест. Сейчас таких мест не осталось (все переведены на render_docx) —
    цикл просто ничего не найдёт, это ожидаемо и является признаком успеха,
    а не поводом считать проверку нерабочей."""
    for dirpath, _dirs, files in os.walk(_BACKEND_APP):
        for fname in files:
            if not fname.endswith(".py"):
                continue
            fpath = os.path.normpath(os.path.join(dirpath, fname))
            if fpath in _ALLOWED_DIRECT_RENDER:
                continue
            with open(fpath, encoding="utf-8") as fh:
                src = fh.read()
            for lineno, line in enumerate(src.splitlines(), start=1):
                m = _RENDER_CALL_RE.search(line)
                if not m:
                    continue
                var = m.group(1)
                if var in ("render_docx", "self", "cls"):
                    continue
                # Эвристика: интересуют только вызовы .render( на объекте,
                # похожем на DocxTemplate/Document (tpl/doc/sec_tpl/...),
                # не произвольные .render() других классов проекта.
                if not re.search(r"(tpl|doc)\w*\.render\(", line, re.IGNORECASE):
                    continue
                assert "render_docx(" in line or "def render_docx" in src, (
                    f"{fpath}:{lineno} зовёт .render() напрямую, минуя "
                    f"render_docx() — строка: {line.strip()!r}"
                )
