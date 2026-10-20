"""GET /api/feo-categories/export?with_summary=1 (владелец 20.10.2026, часть
3) — app.services.feo_export_summary: листы «Сводная»/«Дашборд».

Без with_summary — листов нет (поведение по умолчанию не меняется). С
with_summary=1 — «Сводная» ссылается формулами на строки категорий верхнего
уровня листа данных, и на ней минимум 3 графика."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.test_feo_export_outline import _make_category, _make_subsidy


@pytest.mark.asyncio
async def test_no_summary_sheet_without_flag(client, db_session, admin_headers):
    subsidy = await _make_subsidy(db_session)
    await _make_category(db_session, subsidy.id, level=1, name="Направление 1", budget=100000)

    r = await client.get("/api/feo-categories/export", params={"subsidy_id": subsidy.id}, headers=admin_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert "Сводная" not in wb.sheetnames
    assert "Дашборд" not in wb.sheetnames


@pytest.mark.asyncio
async def test_summary_sheet_references_category_rows(client, db_session, admin_headers):
    subsidy = await _make_subsidy(db_session)
    root1 = await _make_category(db_session, subsidy.id, level=1, name="Экипировка", budget=100000)
    root2 = await _make_category(db_session, subsidy.id, level=1, name="Транспорт", budget=50000)

    r = await client.get(
        "/api/feo-categories/export",
        params={"subsidy_id": subsidy.id, "with_summary": "1"},
        headers=admin_headers,
    )
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    assert "Сводная" in wb.sheetnames
    sm = wb["Сводная"]
    data_ws = wb.active if wb.active.title != "Сводная" else wb[wb.sheetnames[1]]

    # Шапка «Сводной»: строка 1 — титул, строка 2 — заголовки, данные с 3.
    headers = [c.value for c in sm[2]]
    assert "Направление" in headers
    assert any(h and "ФЭО" in h for h in headers)

    names = [sm.cell(r_, 1).value for r_ in range(3, sm.max_row + 1)]
    assert "Экипировка" in names
    assert "Транспорт" in names

    row_exk = next(r_ for r_ in range(3, sm.max_row + 1) if sm.cell(r_, 1).value == "Экипировка")
    feo_col_idx = headers.index(next(h for h in headers if h and "ФЭО" in h)) + 1
    cell_val = sm.cell(row_exk, feo_col_idx).value
    assert isinstance(cell_val, str) and cell_val.startswith("='")
    assert data_ws.title in cell_val

    # ИТОГО — последняя строка данных, формула SUM.
    total_row = sm.max_row
    assert sm.cell(total_row, 1).value == "ИТОГО"
    total_formula = sm.cell(total_row, feo_col_idx).value
    assert isinstance(total_formula, str) and total_formula.startswith("=SUM(")

    n_charts = len(getattr(sm, "_charts", []))
    assert n_charts >= 3
