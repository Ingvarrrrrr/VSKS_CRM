# -*- coding: utf-8 -*-
"""plan_graph_export_dashboard.py — интерактивный дашборд на листе «Сводная»
(владелец 08.10.2026, план binary-crunching-island.md, раздел 3): «В Сводную
Excel добавить интерактивный дашборд с графиками на средствах Excel».

НИЧЕГО не считает сам — ТОЛЬКО формулы Excel (INDEX/MATCH + DataValidation)
над уже написанными write_summary_sheet числами, и ТОЛЬКО 4 графика
(openpyxl.chart) над этими формулами/над готовым списком направлений,
который передаёт вызывающий код (ПРАВИЛО №6 — вторая формула ничего из
«Сводной» здесь не считается).

Блок начинается с колонки N (1 пустая колонка M после таблицы «Сводная»,
последняя колонка которой — L, 12-я). Раскладка (row — 1-индекс листа):
  N1    «Дашборд»
  N2    «Выберите вид в ячейке N3 — графики обновятся»
  N3    выпадающий список (DataValidation) — источник $A$<first_data_row>:
        $A$<total_row> (названия видов + «Итого»), по умолчанию значение
        ячейки total_row («ИТОГО»/«ИТОГО (без ФОТ)» и т.п. — что бы там ни
        было написано).
  N5..N13  таблица-помощник (label|value) — INDEX/MATCH на таблицу «Сводная»
        по выбранному в N3 виду.
  N15.. (если directions непусто) — таблица «По направлениям» (ФЭО/План/
        Заказано/Оплачено), числами, не формулами (роллап направлений
        считает вызывающий код, те же числа, что лист «по направлениям»).
  Графики — якорем в колонке T (не пересекаются с таблицами N..R)."""
from __future__ import annotations

from typing import Optional

try:
    from openpyxl.chart import BarChart, DoughnutChart, Reference, Series
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover — openpyxl отсутствует, роутер уже отказал раньше
    pass

_COL_N = 14  # колонка N — начало блока дашборда
_COL_O = 15
_COL_T = 20  # колонка T — якорь графиков, без пересечения с таблицами N..R

_HELPER_ROWS = (
    # (row_offset от _HELPER_FIRST_ROW, label, source_col_letter_in_table)
    (0, "План по ФЭО, ₽", "B"),
    (1, "План закупок (плановые позиции), ₽", "C"),
    (2, "Заказано, ₽", "H"),
    (3, "Поставлено, ₽", "I"),
    (4, "Оплачено, ₽", "J"),
    (5, "Поставлено, не оплачено, ₽", "K"),
    (6, "Заказано, не поставлено, ₽", None),   # = Заказано - Поставлено
    (7, "Планируется заказать, ₽", "D"),
    (8, "Свободно, ₽", None),                   # = MAX(0, ФЭО - План закупок)
)
_HELPER_FIRST_ROW = 5

_DIR_HEADERS = ("Направление", "ФЭО, ₽", "План, ₽", "Заказано, ₽", "Оплачено, ₽")

NUM_FMT = "#,##0"


def write_dashboard(
    sm, *, first_data_row: int, last_data_row: int, total_row: int,
    directions: Optional[list] = None,
) -> None:
    """Пишет блок дашборда на лист `sm` («Сводная», уже заполненный
    write_summary_sheet — таблица строк по видам + ИТОГО в строке
    `total_row`, первая строка данных `first_data_row`, последняя — до ИТОГО
    `last_data_row`).

    directions — список {"name","feo","plan","ordered","paid"} по
    направлениям (роллап поддерева, та же величина, что лист «по
    направлениям» — считает вызывающий код, plan_graph_export_xlsx.py; см.
    докстринг модуля). Пусто/None — таблица «По направлениям» и график 4
    пишутся с одной строкой-заглушкой «Нет данных», чтобы дашборд всё равно
    нёс 4 графика."""
    HEADER_FONT = Font(bold=True, size=11, color="1E3A5F")
    SUB_FONT = Font(size=9, italic=True, color="6B7280")
    LABEL_FONT = Font(size=9)
    TABLE_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
    TABLE_HEADER_FONT = Font(color="FFFFFF", bold=True, size=9)
    LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
    RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")

    a_range = f"$A${first_data_row}:$A${total_row}"
    default_label_cell = sm.cell(row=total_row, column=1).value or "ИТОГО"

    # ── N1/N2/N3 — заголовок, подсказка, выпадающий список ────────────────
    sm.cell(row=1, column=_COL_N, value="Дашборд").font = HEADER_FONT
    sm.cell(row=2, column=_COL_N, value="Выберите вид в ячейке N3 — графики обновятся").font = SUB_FONT
    n3 = sm.cell(row=3, column=_COL_N, value=default_label_cell)
    n3.font = Font(bold=True, size=10)
    dv = DataValidation(type="list", formula1=f"={a_range}", allow_blank=False)
    dv.add(n3)
    sm.add_data_validation(dv)

    # ── N5..O13 — таблица-помощник INDEX/MATCH ─────────────────────────────
    n3_ref = f"$N$3"
    for offset, label, src_col in _HELPER_ROWS:
        r = _HELPER_FIRST_ROW + offset
        sm.cell(row=r, column=_COL_N, value=label).font = LABEL_FONT
        cell = sm.cell(row=r, column=_COL_O)
        if src_col is not None:
            cell.value = f"=INDEX(${src_col}${first_data_row}:${src_col}${total_row},MATCH({n3_ref},{a_range},0))"
        elif offset == 6:  # Заказано, не поставлено = Заказано(row+_HELPER_FIRST_ROW+2) - Поставлено(row+3)
            cell.value = f"=O{_HELPER_FIRST_ROW + 2}-O{_HELPER_FIRST_ROW + 3}"
        elif offset == 8:  # Свободно = MAX(0, ФЭО - План закупок)
            cell.value = f"=MAX(0,O{_HELPER_FIRST_ROW}-O{_HELPER_FIRST_ROW + 1})"
        cell.number_format = NUM_FMT
        cell.alignment = RIGHT_ALIGN

    helper_last_row = _HELPER_FIRST_ROW + len(_HELPER_ROWS) - 1  # = 13

    # ── Таблица «По направлениям» (числа, не формулы) ──────────────────────
    dir_header_row = helper_last_row + 2  # строка 15
    rows = directions or [{"name": "Нет данных", "feo": 0.0, "plan": 0.0, "ordered": 0.0, "paid": 0.0}]
    for ci, h in enumerate(_DIR_HEADERS):
        c = sm.cell(row=dir_header_row, column=_COL_N + ci, value=h)
        c.fill = TABLE_HEADER_FILL
        c.font = TABLE_HEADER_FONT
        c.alignment = LEFT_ALIGN if ci == 0 else RIGHT_ALIGN
    dir_first_row = dir_header_row + 1
    for i, d in enumerate(rows):
        r = dir_first_row + i
        sm.cell(row=r, column=_COL_N, value=d.get("name") or "").alignment = LEFT_ALIGN
        for ci, key in enumerate(("feo", "plan", "ordered", "paid"), start=1):
            c = sm.cell(row=r, column=_COL_N + ci, value=round(float(d.get(key, 0.0) or 0.0), 2))
            c.number_format = NUM_FMT
            c.alignment = RIGHT_ALIGN
    dir_last_row = dir_first_row + len(rows) - 1

    # ── Графики — якорь в колонке T, вертикально друг под другом ──────────
    anchor_col_letter = sm.cell(row=1, column=_COL_T).column_letter

    chart1 = BarChart()
    chart1.type = "col"
    chart1.title = "Деньги по стадиям"
    chart1.y_axis.title = "₽"
    chart1.legend = None
    chart1.width, chart1.height = 16, 8
    cats1 = Reference(sm, min_col=_COL_N, min_row=_HELPER_FIRST_ROW, max_row=_HELPER_FIRST_ROW + 4)
    data1 = Reference(sm, min_col=_COL_O, min_row=_HELPER_FIRST_ROW, max_row=_HELPER_FIRST_ROW + 4)
    chart1.add_data(data1, titles_from_data=False)
    chart1.set_categories(cats1)
    sm.add_chart(chart1, f"{anchor_col_letter}2")

    chart2 = DoughnutChart()
    chart2.title = "Куда ушёл бюджет ФЭО"
    chart2.width, chart2.height = 16, 8
    ring_first = _HELPER_FIRST_ROW + 4  # «Оплачено» — общая строка с графиком 1
    ring_last = _HELPER_FIRST_ROW + 8   # «Свободно»
    cats2 = Reference(sm, min_col=_COL_N, min_row=ring_first, max_row=ring_last)
    data2 = Reference(sm, min_col=_COL_O, min_row=ring_first, max_row=ring_last)
    chart2.add_data(data2, titles_from_data=False)
    chart2.set_categories(cats2)
    sm.add_chart(chart2, f"{anchor_col_letter}20")

    chart3 = BarChart()
    chart3.type = "col"
    chart3.grouping = "stacked"
    chart3.overlap = 100
    chart3.title = "Запланированные траты по видам"
    chart3.width, chart3.height = 16, 8
    cats3 = Reference(sm, min_col=1, min_row=first_data_row, max_row=last_data_row)
    for col, label in ((5, "Ежемесячные"), (6, "Реальные"), (7, "Выдуманные")):
        ref = Reference(sm, min_col=col, min_row=first_data_row, max_row=last_data_row)
        series = Series(ref, title=label)
        chart3.series.append(series)
    chart3.set_categories(cats3)
    sm.add_chart(chart3, f"{anchor_col_letter}38")

    chart4 = BarChart()
    chart4.type = "bar"  # горизонтальные столбцы
    chart4.title = "По направлениям"
    chart4.width, chart4.height = 16, 8
    cats4 = Reference(sm, min_col=_COL_N, min_row=dir_first_row, max_row=dir_last_row)
    for col, label in ((_COL_N + 1, "ФЭО"), (_COL_N + 2, "План"), (_COL_N + 3, "Заказано"), (_COL_N + 4, "Оплачено")):
        ref = Reference(sm, min_col=col, min_row=dir_first_row, max_row=dir_last_row)
        series = Series(ref, title=label)
        chart4.series.append(series)
    chart4.set_categories(cats4)
    sm.add_chart(chart4, f"{anchor_col_letter}56")
