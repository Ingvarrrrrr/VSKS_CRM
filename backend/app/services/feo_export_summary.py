"""feo_export_summary.py — листы «Сводная» и «Дашборд» выгрузки «Выгрузить
ФЭО» (владелец 20.10.2026, часть 3: «как у выгрузки плана-графика»).

ПРАВИЛО №6/№5: ничего не считает заново — «Сводная» пишет ТОЛЬКО формулы-
ссылки на уже готовые ячейки строк категорий верхнего уровня листа данных
(app.services.feo_export_xlsx.build_feo_export_workbook — эти ячейки сами
либо литерал, либо =SUM(...) по прямым детям, см. докстрин _money_subtotal
там), а графики строит общими фабриками
app.services.export_dashboard_charts (те же BarChart/DoughnutChart, что
app.services.plan_graph_export_dashboard — второй раз создание графика не
заводим)."""
from __future__ import annotations

try:
    from openpyxl.chart import Reference
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover — openpyxl отсутствует, роутер уже отказал раньше
    pass

from app.services.export_dashboard_charts import build_bar_chart, build_doughnut_chart
from app.services.feo_export_columns import STAGE_DEFS, column_index

# (key, label) — порядок столбцов «Сводной»; «Финансирование по ФЭО» и т.д.
# показываются, только если ключ есть в selected_keys выгрузки (владелец:
# «суммы стадий — только те, что выбраны в выгрузке» — это касается и пяти
# денежных столбцов, не только стадий).
_MONEY_COLS = (
    ("feo", "Финансирование по ФЭО, ₽"),
    ("plan", "Запланировано, ₽"),
    ("ordered", "Заказано, ₽"),
    ("to_order", "Надо заказать, ₽"),
    ("residual", "Остаток, ₽"),
)
NUM_FMT = "#,##0.00"
_LABEL_COL = 1  # столбец A — название направления


def _escape_sheet_name(name: str) -> str:
    return name.replace("'", "''")


def write_feo_summary_and_dashboard(wb, sub, data_ws_title: str, selected_keys: list, top_rows: list[dict]) -> None:
    """Пишет лист «Сводная» (index 0) и блок «Дашборд» на нём.

    `top_rows` — [{"name": имя_категории_верхнего_уровня, "row": номер_её_
    строки_на_листе_данных}], в порядке corneвых категорий (roots) из
    build_feo_export_workbook. `selected_keys` — те же ключи, что выбраны для
    листа данных (app.services.feo_export_columns) — отбирают, какие из пяти
    денежных и шести стадийных столбцов попадут в «Сводную»."""
    cols = [(k, label) for k, label in _MONEY_COLS if k in selected_keys]
    cols += [(f"{key}_amt", f"{label}, ₽") for key, label in STAGE_DEFS if f"{key}_amt" in selected_keys]
    if not cols:
        # Нечего сводить (владелец ни одного денежного столбца не выбрал) —
        # «Сводная» всё равно нужна (галочка включена), пишем лист с одной
        # колонкой-предупреждением, без дашборда (графикам нечего рисовать).
        sm = wb.create_sheet("Сводная", 0)
        sm.cell(row=1, column=1, value="Нет выбранных денежных столбцов для сводки")
        return

    HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=9)
    TOTAL_FILL = PatternFill("solid", fgColor="DBEAFE")
    THIN_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    RIGHT = Alignment(horizontal="right", vertical="center")
    LEFT = Alignment(horizontal="left", vertical="center")

    sm = wb.create_sheet("Сводная", 0)
    n_cols = 1 + len(cols)
    last_col_letter = get_column_letter(n_cols)

    title = sm.cell(row=1, column=1, value=f"СВОДНАЯ ПО ФЭО — {sub.name} ({sub.year})")
    title.font = Font(bold=True, size=12, color="1E3A5F")
    sm.merge_cells(f"A1:{last_col_letter}1")
    title.alignment = CENTER
    sm.row_dimensions[1].height = 24

    HEADERS = ["Направление"] + [label for _k, label in cols]
    WIDTHS = [28] + [18] * len(cols)
    for ci, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = sm.cell(row=2, column=ci, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        sm.column_dimensions[cell.column_letter].width = w

    esc_name = _escape_sheet_name(data_ws_title)
    first_data_row = 3
    row = first_data_row
    for entry in top_rows:
        data_row = entry["row"]
        label_cell = sm.cell(row=row, column=1, value=entry["name"])
        label_cell.alignment = LEFT
        label_cell.border = THIN_BORDER
        for ci, (key, _label) in enumerate(cols, 2):
            data_col = column_index(selected_keys, key)
            data_col_letter = get_column_letter(data_col)
            c = sm.cell(row=row, column=ci, value=f"='{esc_name}'!{data_col_letter}{data_row}")
            c.number_format = NUM_FMT
            c.alignment = RIGHT
            c.border = THIN_BORDER
        row += 1
    last_data_row = row - 1
    total_row = row

    sm.cell(row=total_row, column=1, value="ИТОГО").font = Font(bold=True, size=10)
    for ci in range(1, n_cols + 1):
        sm.cell(row=total_row, column=ci).fill = TOTAL_FILL
        sm.cell(row=total_row, column=ci).border = THIN_BORDER
    for ci in range(2, n_cols + 1):
        col_letter = get_column_letter(ci)
        formula = f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})" if last_data_row >= first_data_row else 0
        c = sm.cell(row=total_row, column=ci, value=formula)
        c.number_format = NUM_FMT
        c.alignment = RIGHT

    sm.freeze_panes = f"A{first_data_row}"

    _write_dashboard(sm, cols=cols, first_data_row=first_data_row, last_data_row=last_data_row, total_row=total_row, n_cols=n_cols)


def _write_dashboard(sm, *, cols: list, first_data_row: int, last_data_row: int, total_row: int, n_cols: int) -> None:
    """4 графика (владелец 20.10.2026, часть 3):
    (1) ФЭО/Заказано/Остаток по направлениям; (2) накопительные столбцы по
    стадиям по направлениям; (3) доля ФЭО по направлениям (круговая);
    (4) общий итог по стадиям."""
    anchor_col = n_cols + 2  # один пустой столбец после таблицы
    anchor_col_letter = get_column_letter(anchor_col)
    key_to_ci = {key: ci for ci, (key, _label) in enumerate(cols, 2)}
    stage_cols = [(key, label, ci) for ci, (key, label) in zip(range(2, 2 + len(cols)), cols) if key.endswith("_amt")]

    sm.cell(row=1, column=anchor_col - 1, value="Дашборд").font = Font(bold=True, size=11, color="1E3A5F")

    cats = Reference(sm, min_col=1, min_row=first_data_row, max_row=last_data_row)

    # (1) ФЭО / Заказано / Остаток по направлениям.
    main_keys = [("feo", "ФЭО"), ("ordered", "Заказано"), ("residual", "Остаток")]
    data_refs1 = [
        (Reference(sm, min_col=key_to_ci[k], min_row=first_data_row, max_row=last_data_row), label)
        for k, label in main_keys if k in key_to_ci
    ]
    if data_refs1:
        chart1 = build_bar_chart(title="ФЭО / Заказано / Остаток по направлениям", cats_ref=cats, data_refs=data_refs1)
        sm.add_chart(chart1, f"{anchor_col_letter}2")

    # (2) накопительные столбцы по стадиям по направлениям.
    if stage_cols:
        data_refs2 = [
            (Reference(sm, min_col=ci, min_row=first_data_row, max_row=last_data_row), label)
            for _key, label, ci in stage_cols
        ]
        chart2 = build_bar_chart(
            title="Стадии по направлениям", cats_ref=cats, data_refs=data_refs2, stacked=True,
        )
        sm.add_chart(chart2, f"{anchor_col_letter}20")

    # (3) доля ФЭО по направлениям (круговая); если «feo» не выбран — берём
    # первый выбранный денежный столбец, чтобы дашборд всё равно нёс 4 графика.
    pie_key = "feo" if "feo" in key_to_ci else (cols[0][0] if cols else None)
    if pie_key:
        pie_ci = key_to_ci[pie_key]
        data3 = Reference(sm, min_col=pie_ci, min_row=first_data_row, max_row=last_data_row)
        chart3 = build_doughnut_chart(title="Доля по направлениям", cats_ref=cats, data_ref=data3)
        sm.add_chart(chart3, f"{anchor_col_letter}38")

    # (4) общий итог по стадиям — строка ИТОГО, столбцы стадий (данные идут
    # по строке, from_rows=True); если стадии не выбраны — по пяти денежным.
    totals_source = stage_cols or [(k, l, key_to_ci[k]) for k, l in cols]
    if totals_source:
        min_ci = min(ci for _k, _l, ci in totals_source)
        max_ci = max(ci for _k, _l, ci in totals_source)
        cats4 = Reference(sm, min_col=min_ci, max_col=max_ci, min_row=2, max_row=2)
        data4 = Reference(sm, min_col=min_ci, max_col=max_ci, min_row=total_row, max_row=total_row)
        chart4 = build_bar_chart(
            title="Общий итог по стадиям" if stage_cols else "Общий итог",
            cats_ref=cats4, data_refs=[(data4, None)], show_legend=False, from_rows=True,
        )
        sm.add_chart(chart4, f"{anchor_col_letter}56")
