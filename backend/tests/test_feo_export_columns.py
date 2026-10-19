"""GET /api/feo-categories/export/columns и /export (столбцы по стадиям,
задача владельца 20.10.2026) — app.services.feo_export_columns/feo_export_xlsx.

Объём сокращён владельцем (20.10.2026, часть 2): переиспользует план-график
не копируя его тесты — проверяет только то, что ново здесь: каталог столбцов,
один сценарий «Частично» (количество выбрано наполовину) и формулу категории.
"""
import uuid
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.test_feo_export_outline import _make_category, _make_subsidy


@pytest.mark.asyncio
async def test_columns_catalog(client, admin_headers):
    r = await client.get("/api/feo-categories/export/columns", headers=admin_headers)
    assert r.status_code == 200, r.text
    cols = r.json()
    keys = {c["key"] for c in cols}
    assert {"name", "feo", "plan", "ordered", "to_order", "residual", "ordered_flag", "ordered_flag_amt"} <= keys
    by_key = {c["key"]: c for c in cols}
    assert by_key["name"]["default"] is True
    assert by_key["code"]["default"] is False


@pytest.mark.asyncio
async def test_item_partial_ordered_and_category_subtotal(client, db_session, admin_headers):
    from app.models.feo_planned_item import FeoPlannedItem
    from app.models.purchase import Purchase
    from app.models.purchase_item import PurchaseItem

    subsidy = await _make_subsidy(db_session)
    root = await _make_category(db_session, subsidy.id, level=1, name="Направление", budget=None)
    leaf = await _make_category(db_session, subsidy.id, parent_id=root.id, level=2, name="Спецодежда", budget=None)

    planned_item = FeoPlannedItem(
        feo_category_id=leaf.id, name="Комбинезон", quantity=Decimal("200"),
        unit_price=Decimal("1000"), amount=Decimal("200000"),
        feo_amount=Decimal("200000"), is_active=True,
    )
    db_session.add(planned_item)
    await db_session.commit()
    await db_session.refresh(planned_item)

    # 100 из 200 заказано (ordered), остаток не закрыт — доля 0.5 → "Частично".
    p = Purchase(
        subsidy_id=subsidy.id, feo_category_id=leaf.id, item_name="Комбинезон",
        status="ordered", contract_price=Decimal("100000"), total_nmck=Decimal("100000"),
        nmck=Decimal("100000"),
    )
    db_session.add(p)
    await db_session.flush()
    pi = PurchaseItem(
        purchase_id=p.id, item_name="Комбинезон", quantity=Decimal("100"), unit="шт",
        unit_price=Decimal("1000"), total_price=Decimal("100000"),
        feo_category_id=leaf.id, feo_planned_item_id=planned_item.id, over_plan=False,
    )
    db_session.add(pi)
    await db_session.commit()

    r = await client.get("/api/feo-categories/export", params={"subsidy_id": subsidy.id}, headers=admin_headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active

    def _find_row(exact_name):
        for row in ws.iter_rows(min_row=4):
            v = row[0].value
            if v and v.strip() == exact_name:
                return row[0].row
        raise AssertionError(f"row not found: {exact_name}")

    headers = [c.value for c in ws[3]]
    ci = {h: i + 1 for i, h in enumerate(headers)}

    r_item = _find_row("Комбинезон")
    r_leaf = _find_row("Спецодежда")

    # «Товар заказан» — доля 0.5 (100 из 200) → "Частично"; сумма — 100 000.
    assert ws.cell(r_item, ci["Товар заказан"]).value == "Частично"
    assert ws.cell(r_item, ci["Товар заказан, ₽"]).value == pytest.approx(100000.0)
    assert ws.cell(r_item, ci["Заказано (₽)"]).value == pytest.approx(100000.0)
    # Надо заказать = План (200000) − Заказано (100000).
    assert ws.cell(r_item, ci["Надо заказать (₽)"]).value == pytest.approx(100000.0)

    # Категория-лист (без детей-категорий, но с плановой позицией) —
    # SUBTOTAL(9, ...) по строке позиции, не литерал.
    feo_leaf = ws.cell(r_leaf, ci["Финансирование по ФЭО (₽)"]).value
    assert isinstance(feo_leaf, str) and feo_leaf.startswith("=SUBTOTAL(9,")
    assert f"{ws.cell(r_leaf, ci['Финансирование по ФЭО (₽)']).column_letter}{r_item}" in feo_leaf
