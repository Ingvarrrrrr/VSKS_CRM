"""Bulk-импорт сотрудников из Excel: шаблон + разбор файла.

ПЕРЕНЕСЕНО (не изменено) из app/routers/users.py при разрезании монолитного
роутера (Правило №5, сессия 2026-09-08). Тот же префикс /api/users. Пути
здесь двухсегментные (/import/template, /import/excel) — конфликта по форме
с GET /{user_id} core-роутера нет, порядок регистрации относительно него
не важен.
"""
from io import BytesIO
from urllib.parse import quote as _url_quote
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.user import User
from app.services.fio import resolve_user_name_input
from app.auth.jwt import hash_password, get_single_org_id
from app.auth.permissions import (
    require_action, ensure_user_org_access,
    assert_new_user_role_allowed, split_new_user_role, ROLE_LABELS_RU,
)
from app.services.password_policy import validate_new_password

router = APIRouter(prefix="/api/users", tags=["users"])

# org_admin добавлен задачей «Администратор организации в списке ролей при
# создании сотрудника» (2026-10-16) — везде, где создаётся сотрудник, включая
# импорт. superadmin сюда НЕ входит намеренно (служебная SaaS-роль, строкой
# файла её назначить нельзя) — см. assert_new_user_role_allowed.
VALID_ROLES = ("employee", "manager", "org_admin", "admin", "account_owner")

# Русское название → код роли (для колонки «Роль» в файле импорта, Правило №6 —
# подписи только из ROLE_LABELS_RU, здесь просто обратный индекс, не вторая
# таблица с текстами).
_ROLE_RU_TO_CODE = {
    label.lower(): code
    for code, label in ROLE_LABELS_RU.items()
    if code != "superadmin"
}


def _normalize_import_role(raw: str) -> Optional[str]:
    """Код роли ('org_admin') ИЛИ русское название ('Администратор организации',
    регистр/пробелы не важны) → код роли. None — если не распозналось."""
    v = (raw or "").strip()
    if not v:
        return None
    low = v.lower()
    if low in VALID_ROLES:
        return low
    return _ROLE_RU_TO_CODE.get(low)


@router.get("/import/template")
async def users_import_template(_=Depends(require_action('user.manage'))):
    """Download xlsx template for bulk user import."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise HTTPException(500, "openpyxl не установлен")

    wb = Workbook()
    ws = wb.active
    ws.title = "Пользователи"

    headers = ["ФИО", "Email", "Логин", "Пароль", "Роль", "Город"]
    required = {"ФИО", "Email", "Логин", "Пароль"}

    header_fill = PatternFill("solid", fgColor="1E40AF")
    req_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(bold=True, color="FFFFFF")

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = req_fill if h in required else header_fill
        cell.alignment = Alignment(horizontal="center")

    # Example row. Роль в примере — русским названием (как ждёт выпадающий
    # список ниже), а не кодом 'employee': владелец попросил писать словами
    # владельца, не внутренними названиями полей.
    example = [
        "Иванов Иван Иванович", "ivanov@example.com", "ivanov",
        "Password123", ROLE_LABELS_RU["employee"], "Москва",
    ]
    for ci, val in enumerate(example, 1):
        ws.cell(row=2, column=ci, value=val)

    # Допустимые роли — русские названия (Правило №6: подписи только из
    # ROLE_LABELS_RU, тот же порядок, что и в getCreateRoleItems на фронте).
    role_labels_ordered = [
        ROLE_LABELS_RU["account_owner"], ROLE_LABELS_RU["admin"],
        ROLE_LABELS_RU["org_admin"], ROLE_LABELS_RU["manager"],
        ROLE_LABELS_RU["employee"],
    ]
    ws.cell(row=3, column=5, value=f"Допустимые роли: {', '.join(role_labels_ordered)}")

    # Выпадающий список в колонке «Роль» (E2:E500) — подсказывает те же 5
    # названий без необходимости помнить их наизусть.
    dv_formula = '"' + ','.join(role_labels_ordered) + '"'
    dv = DataValidation(type="list", formula1=dv_formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E500")

    col_widths = [35, 30, 20, 20, 20, 20]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_url_quote('Шаблон_импорта_сотрудников.xlsx', safe='-_.~')}"},
    )


@router.post("/import/excel")
async def import_users_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_action('user.manage')),
):
    """Bulk import users from Excel file."""
    if not (file.filename or '').lower().endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "Поддерживаются только файлы .xlsx / .xls")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(500, "openpyxl не установлен")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")

    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise HTTPException(400, "Файл пустой")

    def _norm(v) -> str:
        return str(v).strip().lower() if v else ''

    # Fuzzy column mapping
    col: dict[str, int] = {}
    for i, h in enumerate(header_row):
        h_str = _norm(h)
        if any(x in h_str for x in ('фио', 'фамил', 'имя', 'full_name', 'full name', 'ф.и.о')):
            col.setdefault('full_name', i)
        elif any(x in h_str for x in ('email', 'e-mail', 'почта', 'mail')):
            col.setdefault('email', i)
        elif any(x in h_str for x in ('логин', 'login', 'username', 'пользовател')):
            col.setdefault('username', i)
        elif any(x in h_str for x in ('пароль', 'password', 'pass')):
            col.setdefault('password', i)
        elif any(x in h_str for x in ('роль', 'role', 'должност')):
            col.setdefault('role', i)
        elif any(x in h_str for x in ('город', 'city', 'населённ')):
            col.setdefault('city', i)

    if 'full_name' not in col:
        raise HTTPException(400, "Не найдена колонка ФИО. Убедитесь что заголовок содержит «ФИО» или «full_name».")

    def _cell(row, field):
        idx = col.get(field)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        if v is None:
            return None
        s = str(v).strip()
        if not s or s.lower() in ('none', 'null', '-', '—'):
            return None
        return s

    # Collect existing emails and usernames for dedup
    email_result = await db.execute(select(User.email).where(User.email.isnot(None)))
    existing_emails = {r[0].lower() for r in email_result if r[0]}
    username_result = await db.execute(select(User.username))
    existing_usernames = {r[0].lower() for r in username_result if r[0]}

    org_id = get_single_org_id(current_user)

    created = 0
    skipped = 0
    errors_list = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        full_name = _cell(row, 'full_name')
        if not full_name:
            skipped += 1
            continue

        email = _cell(row, 'email')
        username = _cell(row, 'username')
        password = _cell(row, 'password')
        role = _cell(row, 'role') or 'employee'
        city = _cell(row, 'city')

        # Validate required fields
        if not email:
            errors_list.append({"row": row_idx, "error": f"Нет email для «{full_name}»"})
            continue
        if not username:
            # Auto-generate username from email
            username = email.split('@')[0]
        if not password:
            errors_list.append({"row": row_idx, "error": f"Нет пароля для «{full_name}»"})
            continue
        # Пароль берётся из файла (не генерируется системой) — политика
        # (app/services/password_policy.py) та же, что при регистрации/сбросе;
        # ошибка — по строке, не роняет весь импорт.
        try:
            validate_new_password(str(password), identifiers=[email, username])
        except HTTPException as _pwd_exc:
            detail = _pwd_exc.detail if isinstance(_pwd_exc.detail, str) else "Пароль не проходит политику безопасности"
            errors_list.append({"row": row_idx, "error": f"{detail} («{full_name}»)"})
            continue

        # Validate role — код ('org_admin') или русское название ('Администратор
        # организации', регистр/пробелы не важны), см. _normalize_import_role.
        role_code = _normalize_import_role(role)
        if role_code is None:
            allowed_ru = ', '.join(
                ROLE_LABELS_RU[r] for r in ("account_owner", "admin", "org_admin", "manager", "employee")
            )
            errors_list.append({"row": row_idx, "error": f"Недопустимая роль «{role}» для «{full_name}». Допустимые роли: {allowed_ru}"})
            continue

        # Единая проверка «вправе ли current_user назначить эту роль в этой
        # орге» (Правило №6, та же функция, что в create_user) — раньше импорт
        # вообще не проверял assert_role_assignable/ранг для org_admin, строка
        # падала бы в errors только через необработанное исключение.
        try:
            await assert_new_user_role_allowed(current_user, role_code, org_id, db)
        except HTTPException as _role_exc:
            detail = _role_exc.detail if isinstance(_role_exc.detail, str) else "Роль недопустима для вас"
            errors_list.append({"row": row_idx, "error": f"{detail} («{full_name}»)"})
            continue

        # Dedup
        if email.lower() in existing_emails:
            skipped += 1
            continue
        if username.lower() in existing_usernames:
            skipped += 1
            continue

        global_role, org_role = split_new_user_role(role_code)

        row_last, row_first, row_middle, row_full = resolve_user_name_input(None, None, None, full_name)
        user = User(
            username=username,
            password_hash=hash_password(password),
            role=global_role,
            last_name=row_last,
            first_name=row_first,
            middle_name=row_middle,
            full_name=row_full,
            city=city,
            email=email,
            is_email_confirmed=True,
            org_id=org_id,
        )
        db.add(user)
        await db.flush()
        # Phase 17.1-05 (выровнено с create_user, Правило №6): импорт раньше
        # вообще не писал user_org_access — сотрудник, импортированный с ролью
        # org_admin, не получал орг-роль, только пустой глобальный 'employee'.
        if org_id:
            await ensure_user_org_access(user.id, org_id, org_role, db)
        existing_emails.add(email.lower())
        existing_usernames.add(username.lower())
        created += 1

    await db.commit()
    return {"created": created, "skipped": skipped, "errors": errors_list}
