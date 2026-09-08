"""Phase 22 — построчный разбор xlsx-выписки: ParsedRow + parse_workbook.

Разрезание bank_statement_parser.py (Правило №5, сессия 2026-09-08):
ParsedRow (dataclass), поиск строки заголовков/данных с учётом merged cells
и служебных строк (нумерация/sub-headers), сборка одной ParsedRow из ячеек
(_build_row) и главная функция parse_workbook — весь конвейер «свежего»
разбора xlsx целиком, без внутренних границ (как compute_feo_plan_tree).
Отдельный путь от bank_statement_parser_reparse (typed reparse уже
сохранённого raw_json) — общие карты/хелперы вынесены в _maps/_helpers.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Optional

from openpyxl import load_workbook

from app.services.bank_statement_parser_helpers import (
    _norm_header,
    _norm_inn,
    _to_date,
    _to_datetime,
    _to_decimal,
    compute_row_hash,
    extract_all_documents,
    parse_basis_doc,
    parse_purpose,
)
from app.services.bank_statement_parser_maps import (
    EXECUTED_STATUSES,
    HEADER_MAP,
    RX_INN,
    _EXPECTED_HEADERS,
    _normalize_lookup_key,
)

# Логгер сохраняет исходное имя "app.services.bank_statement_parser" (а не
# __name__ этого файла) — поведение/фильтрация логов не должны измениться
# при разрезании модуля (Правило №5).
_parser_log = logging.getLogger("app.services.bank_statement_parser")


# ---------------------------------------------------------------------------
# ParsedRow
# ---------------------------------------------------------------------------

@dataclass
class ParsedRow:
    payment_number: Optional[str] = None
    payment_date: Optional[date] = None
    execution_datetime: Optional[datetime] = None
    status: Optional[str] = None
    amount: Optional[Decimal] = None

    payer_inn: Optional[str] = None
    payer_kpp: Optional[str] = None
    payer_name: Optional[str] = None
    payer_account: Optional[str] = None

    payee_inn: Optional[str] = None
    payee_kpp: Optional[str] = None
    payee_name: Optional[str] = None
    payee_account: Optional[str] = None
    payee_bik: Optional[str] = None
    payee_bank: Optional[str] = None

    purpose_text: Optional[str] = None
    parsed_contract_number: Optional[str] = None
    parsed_contract_date: Optional[date] = None
    parsed_kbk: Optional[str] = None
    parsed_documents: Optional[dict] = None

    basis_doc_text: Optional[str] = None
    basis_doc_number: Optional[str] = None
    basis_doc_date: Optional[date] = None
    subsidy_code: Optional[str] = None
    external_doc_id: Optional[str] = None

    # Этап 3: код расходов (КРЦС). *_col — сырые значения колонок выписки
    # (транзитные, в BankPayment не сохраняются), expense_code — итоговое
    # значение (см. app.services.payment_basis.expense_code).
    expense_code_short_col: Optional[str] = None
    expense_code_detail_col: Optional[str] = None
    expense_code: Optional[str] = None

    raw_json: dict = field(default_factory=dict)
    source_row_hash: Optional[str] = None
    is_executed: bool = False
    skip_reason: Optional[str] = None


# ---------------------------------------------------------------------------
# _find_header_row
# ---------------------------------------------------------------------------

def _find_header_row(ws) -> int:
    """Возвращает 1-based номер строки с заголовками таблицы.

    В реальных выписках первые 1-4 строки — metadata («Тип документа»,
    «Дата формирования», пустая строка и т.д.).  Ищем строку, где хотя бы
    одно из ожидаемых ключевых слов совпадает с нормализованным значением
    ячейки.  Fallback — строка 1 (совместимость с тестовыми файлами).
    """
    for row_idx in range(1, 15):  # ищем в первых 14 строках
        try:
            cells = [str(c.value or '').strip().upper() for c in ws[row_idx]]
        except Exception:
            break
        normalized = {re.sub(r'\s+', ' ', c) for c in cells if c}
        if normalized & _EXPECTED_HEADERS:
            _parser_log.info(f"bank_statement_parser: headers found at row {row_idx}")
            return row_idx
    _parser_log.warning("bank_statement_parser: header row not detected, using row 1 as fallback")
    return 1


# ---------------------------------------------------------------------------
# _extract_headers
# ---------------------------------------------------------------------------

def _expand_merged_row(ws, row_idx: int, max_col: int) -> list:
    """Возвращает значения row_idx с разворачиванием merged cells.

    Для merged range — все колонки получают value верхне-левой ячейки.
    """
    values = [None] * max_col
    # Обычные ячейки
    for cell in ws[row_idx]:
        col = cell.column - 1
        if 0 <= col < max_col:
            values[col] = cell.value
    # Развернуть merged-ranges на эту строку
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row_idx <= mr.max_row:
            top_left = ws.cell(row=mr.min_row, column=mr.min_col).value
            for col in range(mr.min_col - 1, mr.max_col):
                if 0 <= col < max_col:
                    values[col] = top_left
    return values


def _extract_headers(ws, header_row: int = 1) -> list[str]:
    """Composite-ключи из ДВУХ строк xlsx.

    Логика:
    - main и sub оба пустые → ""
    - только main → main
    - только sub → sub
    - main == sub (merged на 2 строки одно значение) → main
    - main != sub оба непустые → "main (sub)"  ← композит
    """
    max_col = ws.max_column
    main_values = _expand_merged_row(ws, header_row, max_col)
    try:
        sub_values = _expand_merged_row(ws, header_row + 1, max_col)
    except Exception:
        sub_values = [None] * max_col

    out: list[str] = []
    for i in range(max_col):
        main = _norm_header(main_values[i]) if main_values[i] not in (None, "") else ""
        sub = _norm_header(sub_values[i]) if sub_values[i] not in (None, "") else ""

        if not main and not sub:
            out.append("")
        elif not sub:
            out.append(main)
        elif not main:
            out.append(sub)
        elif main == sub:
            out.append(main)
        elif main in HEADER_MAP:
            # known single-row header — sub скорее всего contamination
            # (data row или повтор), берём только main
            out.append(main)
        else:
            out.append(f"{main} ({sub})")

    _parser_log.info(f"bank_statement_parser: extracted {len(out)} composite headers")
    known = sum(1 for h in out if h in HEADER_MAP)
    _parser_log.info(f"bank_statement_parser: {known}/{len(out)} headers found in HEADER_MAP")
    return out


# ---------------------------------------------------------------------------
# _build_row — заполнить ParsedRow из значений ячеек + headers
# ---------------------------------------------------------------------------

def _build_row(
    headers: list[str],
    values: list[Any],
) -> Optional[ParsedRow]:
    """Строит ParsedRow из одной строки данных.

    Возвращает None если строка полностью пустая.
    """
    if all(v is None or str(v).strip() == "" for v in values):
        return None

    row = ParsedRow()
    raw: dict[str, Any] = {}

    col_groups: dict[str, list[Any]] = {}
    for h, v in zip(headers, values):
        if not h:
            continue
        col_groups.setdefault(h, []).append(v)

    def _json_safe(v: Any) -> Any:
        # JSONB-колонка не сериализует datetime/date/Decimal — конвертим в строки
        if v is None:
            return None
        if isinstance(v, datetime):
            return v.isoformat()
        if isinstance(v, date):
            return v.isoformat()
        if isinstance(v, Decimal):
            return str(v)
        if isinstance(v, (str, int, float, bool)):
            return v
        return str(v)

    for h, vals in col_groups.items():
        non_null = [v for v in vals if v is not None and str(v).strip() != ""]
        if not non_null:
            raw[h] = None
        elif len(non_null) == 1:
            raw[h] = _json_safe(non_null[0])
        else:
            raw[h] = "\n".join(str(_json_safe(v)) for v in non_null)

    row.raw_json = raw
    row.source_row_hash = compute_row_hash(raw)

    # Legacy raw_json may have composite keys "MAIN (SUB)" where MAIN is the real header
    # but SUB was data row contamination. Build a normalized lookup view of raw.
    raw_norm: dict[str, Any] = {}
    for raw_key, raw_val in raw.items():
        normalized = _normalize_lookup_key(raw_key)
        if raw_val is not None and (normalized not in raw_norm or raw_norm[normalized] is None):
            raw_norm[normalized] = raw_val

    for h, field_name in HEADER_MAP.items():
        v = raw.get(h)
        if v is None:
            v = raw_norm.get(h)
        if v is None:
            continue

        if field_name == "expense_code_short_col":
            row.expense_code_short_col = str(v).strip() if v else None
        elif field_name == "expense_code_detail_col":
            row.expense_code_detail_col = str(v).strip() if v else None
        elif field_name == "payment_number":
            row.payment_number = _norm_inn(v)
        elif field_name == "payment_date":
            row.payment_date = _to_date(v)
        elif field_name == "status":
            row.status = str(v).strip().upper() if v else None
        elif field_name == "amount":
            row.amount = _to_decimal(v)
        elif field_name == "execution_datetime":
            row.execution_datetime = _to_datetime(v)
        elif field_name == "payer_inn":
            row.payer_inn = _norm_inn(v)
        elif field_name == "payer_kpp":
            row.payer_kpp = _norm_inn(v)
        elif field_name == "payer_name":
            row.payer_name = str(v).strip() if v else None
        elif field_name == "payer_account":
            row.payer_account = str(v).strip() if v else None
        elif field_name == "payer_block":
            block_text = str(v)
            if not row.payer_inn:
                inn_m = RX_INN.search(block_text)
                if inn_m:
                    row.payer_inn = inn_m.group(1)
        elif field_name == "payee_inn":
            row.payee_inn = _norm_inn(v)
        elif field_name == "payee_kpp":
            row.payee_kpp = _norm_inn(v)
        elif field_name == "payee_name":
            row.payee_name = str(v).strip() if v else None
        elif field_name == "payee_account":
            row.payee_account = str(v).strip() if v else None
        elif field_name == "payee_bik":
            row.payee_bik = str(v).strip() if v else None
        elif field_name == "payee_bank":
            row.payee_bank = str(v).strip() if v else None
        elif field_name == "payee_block":
            block_text = str(v)
            if not row.payee_inn:
                inn_m = RX_INN.search(block_text)
                if inn_m:
                    row.payee_inn = inn_m.group(1)
        elif field_name == "purpose_text":
            if not row.purpose_text:  # берём первое найденное
                row.purpose_text = str(v).strip() if v else None
        elif field_name == "basis_doc_text":
            row.basis_doc_text = str(v).strip() if v else None
        elif field_name == "subsidy_code":
            row.subsidy_code = str(v).strip() if v else None
        elif field_name == "external_doc_id":
            row.external_doc_id = _norm_inn(v)

    # Парсим purpose_text
    if row.purpose_text:
        pp = parse_purpose(row.purpose_text)
        row.parsed_contract_number = pp.get("contract_number")
        row.parsed_contract_date = pp.get("contract_date")
        row.parsed_kbk = pp.get("kbk")
        row.parsed_documents = extract_all_documents(row.purpose_text)
        # Заполняем parsed_contract_number из parsed_documents.contracts[0] если parse_purpose не нашёл
        if not row.parsed_contract_number and row.parsed_documents.get("contracts"):
            first = row.parsed_documents["contracts"][0]
            row.parsed_contract_number = first.get("number")
            if first.get("date"):
                try:
                    row.parsed_contract_date = datetime.strptime(first["date"], "%d.%m.%Y").date()
                except ValueError:
                    pass

    # Парсим basis_doc_text
    if row.basis_doc_text:
        row.basis_doc_number, row.basis_doc_date = parse_basis_doc(row.basis_doc_text)

    # Этап 3: код расходов (КРЦС) — колонки выписки приоритетнее regex из purpose_text
    from app.services.payment_basis import expense_code as _resolve_expense_code
    row.expense_code = _resolve_expense_code(row)

    # Статус
    status_norm = (row.status or "").upper().strip()
    if status_norm in EXECUTED_STATUSES:
        row.is_executed = True
    # Этап 1: отклонённые/аннулированные статусы больше НЕ пропускают импорт —
    # строка сохраняется как есть (status хранит фактический статус), иначе
    # аннулированный платёж негде увидеть. REJECTED_STATUSES остаётся
    # классификатором для UI/отчётов, но не поводом для skip_reason.
    # skip_reason остаётся общим механизмом для будущих причин пропуска строки.

    return row


# ---------------------------------------------------------------------------
# _is_numbering_row — строка нумерации колонок xlsx (1, 2, 3, ... 39)
# ---------------------------------------------------------------------------

def _is_numbering_row(values: list) -> bool:
    """Detect 'numbering row' — все непустые значения короткие числа 1-200.

    В некоторых xlsx после строки заголовков идёт строка вида:
    1, 2, 3, ..., 39 — нумерация позиций для ручного заполнения.
    Такую строку надо пропустить, иначе она попадает в реестр как платёж.
    """
    if not values:
        return False
    numeric_count = 0
    non_empty_count = 0
    for v in values:
        if v is None or str(v).strip() == "":
            continue
        non_empty_count += 1
        s = str(v).strip()
        try:
            n = int(float(s))
            if 1 <= n <= 200:
                numeric_count += 1
            else:
                return False  # число вне диапазона — не нумерация
        except (ValueError, OverflowError):
            return False  # хотя бы одно нечисло — это не строка нумерации
    # Минимум 5 числовых ячеек, и все непустые — числа 1-200
    return non_empty_count >= 5 and numeric_count == non_empty_count


# ---------------------------------------------------------------------------
# _find_first_data_row — пропустить sub-headers, numbering rows, empty
# ---------------------------------------------------------------------------

def _find_first_data_row(ws, header_row: int, headers: list[str]) -> int:
    """Найти первую строку с реальными данными после header_row.

    Пропускает sub-headers, numbering rows (24, 25, 26...), пустые строки.
    Признак data row: в колонке СУММА число >= 1, ИЛИ в НОМЕР ДОКУМЕНТА строка длиной >3.
    """
    sum_idx = next((i for i, h in enumerate(headers) if h.startswith('СУММА')), -1)
    docnum_idx = next((i for i, h in enumerate(headers) if h.startswith('НОМЕР ДОКУМЕНТА')), -1)

    for row_idx in range(header_row + 1, header_row + 12):
        try:
            row_vals = [c.value for c in ws[row_idx]]
        except Exception:
            continue

        if _is_numbering_row(row_vals):
            _parser_log.info(f"bank_statement_parser: skip numbering row at {row_idx}")
            continue

        # Пропустить пустую строку
        if all(v is None or str(v).strip() == "" for v in row_vals):
            continue

        # Hard signal: SUMMA >= 1
        if sum_idx >= 0 and sum_idx < len(row_vals):
            v = row_vals[sum_idx]
            if v is not None:
                try:
                    if float(v) >= 1:
                        _parser_log.info(f"bank_statement_parser: data starts at row {row_idx} (SUMMA={v})")
                        return row_idx
                except (ValueError, TypeError):
                    pass

        # Soft signal: docnum длиннее 3
        if docnum_idx >= 0 and docnum_idx < len(row_vals):
            v = row_vals[docnum_idx]
            if v is not None and len(str(v).strip()) > 3:
                _parser_log.info(f"bank_statement_parser: data starts at row {row_idx} (DOCNUM={v})")
                return row_idx

    fallback = header_row + 2
    _parser_log.warning(f"bank_statement_parser: no clear data row, fallback to {fallback}")
    return fallback


# ---------------------------------------------------------------------------
# parse_workbook — главная функция
# ---------------------------------------------------------------------------

def parse_workbook(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
) -> tuple[str, list[ParsedRow]]:
    """Главная функция парсера.

    1. Открывает xlsx
    2. Берёт sheet_name или первый активный
    3. Читает headers, нормализует, ищет в HEADER_MAP индексы
    4. Для каждой data row:
       - заполняет ParsedRow (известные поля)
       - сохраняет ВСЕ значения в raw_json по нормализованному ключу
       - парсит purpose_text → contract_number/date/kbk + все parsed_documents
       - парсит basis_doc_text → basis_doc_number/date
       - заполняет subsidy_code из «Аналитический код раздела...»
       - ставит is_executed=True если status IN EXECUTED_STATUSES
    5. Возвращает (sheet_name, list[ParsedRow])
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)  # без read_only=True для корректной работы с merged_cells

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        active_name = sheet_name
    else:
        ws = wb.active
        active_name = ws.title

    header_row = _find_header_row(ws)
    headers = _extract_headers(ws, header_row)
    data_start_row = _find_first_data_row(ws, header_row, headers)
    _parser_log.info(f"bank_statement_parser: parsing headers={len(headers)}, data starts at row {data_start_row}")

    parsed: list[ParsedRow] = []

    rows_iter = ws.iter_rows(min_row=data_start_row, values_only=True)
    for raw_values in rows_iter:
        values = list(raw_values)

        while len(values) < len(headers):
            values.append(None)

        if _is_numbering_row(values):  # ещё одна защита если случайно попалось
            _parser_log.info("bank_statement_parser: skip numbering row in data section")
            continue

        row = _build_row(headers, values)
        if row is None:
            continue
        parsed.append(row)

    wb.close()
    return active_name, parsed
