# -*- coding: utf-8 -*-
"""test_plan_graph_export_dashboard.py — дашборд на листе «Сводная» (владелец
08.10.2026, план binary-crunching-island.md раздел 3): выпадающий список +
4 графика (openpyxl.chart) над уже написанной таблицей «Сводная», без
второй формулы (ПРАВИЛО №6 — проверяем ТОЛЬКО форму, не бизнес-числа,
которые уже покрыты test_subsidy_summary_by_kind.py).

Строится через СИНТЕТИЧЕСКИЙ summary_by_kind, без БД (та же манера, что
test_plan_graph_export_flat_sheet.py)."""
from types import SimpleNamespace

import openpyxl

from app.services.item_type_split import ALL_KINDS
from app.services.plan_graph_export_summary_sheet import write_summary_sheet


def _sub():
    return SimpleNamespace(name="Тестовая субсидия", year=2026)


def _summary_by_kind(with_payroll=False):
    base = {
        "goods": {"feo": 1000.0, "plan": 800.0, "committed": 600.0, "monthly": 50.0,
                  "likely": 100.0, "nice": 50.0, "paid": 400.0, "delivered_unpaid": 200.0},
        "services": {"feo": 500.0, "plan": 300.0, "committed": 200.0, "monthly": 20.0,
                     "likely": 60.0, "nice": 20.0, "paid": 150.0, "delivered_unpaid": 50.0},
        "payroll": {k: 0.0 for k in (
            "feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid"
        )},
        "unspecified": {k: 0.0 for k in (
            "feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid"
        )},
    }
    if with_payroll:
        base["payroll"] = {"feo": 200.0, "plan": 150.0, "committed": 100.0, "monthly": 10.0,
                            "likely": 20.0, "nice": 20.0, "paid": 80.0, "delivered_unpaid": 20.0}
    return base


def _directions():
    return [
        {"name": "Направление А", "feo": 1000.0, "plan": 800.0, "ordered": 600.0, "paid": 400.0},
        {"name": "Направление Б", "feo": 500.0, "plan": 300.0, "ordered": 200.0, "paid": 150.0},
    ]


def _build_wb(with_payroll=False, directions=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_summary_sheet(wb, _sub(), _summary_by_kind(with_payroll), directions=directions)
    return wb["Сводная"]


def test_four_charts_present():
    sm = _build_wb(directions=_directions())
    assert len(sm._charts) == 4


def test_four_charts_present_without_directions():
    """Без направлений (None) — дашборд всё равно несёт 4 графика (график 4
    рисуется с заглушкой «Нет данных»)."""
    sm = _build_wb(directions=None)
    assert len(sm._charts) == 4


def test_data_validation_on_n3():
    sm = _build_wb(directions=_directions())
    dvs = list(sm.data_validations.dataValidation)
    assert len(dvs) == 1
    dv = dvs[0]
    assert dv.type == "list"
    assert any("N3" in str(rng) for rng in dv.sqref.ranges)


def test_data_validation_source_matches_total_row_without_payroll():
    """Без ФОТ (все нули — строка не выводится) — только «Товары»/«Услуги» +
    ИТОГО, total_row = 7 (строки 5,6 данные + 7 итого)."""
    sm = _build_wb(with_payroll=False, directions=_directions())
    dv = list(sm.data_validations.dataValidation)[0]
    formula = dv.formula1
    assert "$A$5:$A$7" in formula
    assert sm.cell(row=7, column=1).value == "ИТОГО"
    assert sm.cell(row=3, column=14).value == "ИТОГО"  # N3 default


def test_data_validation_source_matches_total_row_with_payroll():
    """С ФОТ (ненулевой) — три строки данных (goods/services/payroll, строки
    5-7; "unspecified" всегда нулевой на этой фикстуре и не выводится) +
    ИТОГО в строке 8."""
    sm = _build_wb(with_payroll=True, directions=_directions())
    dv = list(sm.data_validations.dataValidation)[0]
    formula = dv.formula1
    assert "$A$5:$A$8" in formula
    assert sm.cell(row=8, column=1).value == "ИТОГО"


def test_helper_formulas_reference_table_columns():
    sm = _build_wb(directions=_directions())
    # N5/O5 — «План по ФЭО» = INDEX на столбец B.
    assert sm.cell(row=5, column=14).value == "План по ФЭО, ₽"
    formula = sm.cell(row=5, column=15).value
    assert formula.startswith("=INDEX($B$5:$B$")
    assert "MATCH($N$3,$A$5:$A$" in formula

    # Заказано — столбец H.
    formula_ordered = sm.cell(row=7, column=15).value
    assert formula_ordered.startswith("=INDEX($H$5:$H$")

    # «Свободно» — MAX(0, ФЭО - План закупок), не формула INDEX.
    formula_free = sm.cell(row=13, column=15).value
    assert formula_free == "=MAX(0,O5-O6)"


def test_directions_table_matches_input_sums():
    directions = _directions()
    sm = _build_wb(directions=directions)
    # Таблица направлений начинается в строке helper_last_row+2 = 15, данные с 16.
    header_row = 15
    assert sm.cell(row=header_row, column=14).value == "Направление"
    assert sm.cell(row=16, column=14).value == "Направление А"
    assert sm.cell(row=16, column=15).value == 1000.0
    assert sm.cell(row=16, column=16).value == 800.0
    assert sm.cell(row=16, column=17).value == 600.0
    assert sm.cell(row=16, column=18).value == 400.0
    assert sm.cell(row=17, column=14).value == "Направление Б"

    # Сумма таблицы направлений == Σ входных directions (ПРАВИЛО №6 — число
    # из вызывающего кода, здесь просто сверяем, что оно переписано 1:1).
    total_feo = sum(d["feo"] for d in directions)
    written_feo = sm.cell(row=16, column=15).value + sm.cell(row=17, column=15).value
    assert written_feo == total_feo
