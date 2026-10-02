"""display_contractor_name (app/services/purchase_contractor_display.py) —
владелец, 30.09: «Если это авансовый отчёт, в столбце "Контрагент" должны
выводиться данные, кому его возмещать, а не у кого куплено» — и в реестре
закупок (GET /api/purchases/{id}, GET /api/purchases/), и в Excel-экспорте
(GET /api/purchases/export/excel). Один хелпер, оба читателя — ПРАВИЛО №6,
проверяем, что оба реально дают одно и то же ФИО на одних данных, а не только
что механизм где-то есть (см. VAULT-урок feedback_verify_user_scenario_not_mechanism).

Три случая:
  1. reimbursement_user_id задан → contractor_name = ФИО получателя возмещения,
     ВО ВСЕХ трёх читателях (helper напрямую, GET /{id}, экспорт).
  2. reimbursement_user_id пуст, service_note_by (автор/инициатор) задан →
     фолбэк на ФИО автора.
  3. Обычная (не авансовая) закупка — contractor_name не трогается хелпером
     (поведение как раньше, по шапке договора/contractor_id).
"""
import uuid
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.purchase import Purchase
from app.models.purchase_item import PurchaseItem
from app.services.purchase_contractor_display import display_contractor_name, seller_display_for_advance


async def _make_subsidy(db_session, test_org):
    from app.models.subsidy import Subsidy
    s = Subsidy(name=f"TestSubsidy-{uuid.uuid4().hex[:8]}", year=2026, budget=1_000_000,
                require_planned_dates=False, org_id=test_org.id)
    db_session.add(s)
    await db_session.commit()
    await db_session.refresh(s)
    return s


async def _make_user(db_session, test_org, full_name):
    from app.models.user import User
    from app.auth.jwt import hash_password
    username = f"test_user_{uuid.uuid4().hex[:8]}"
    u = User(
        username=username,
        password_hash=hash_password("testpass123"),
        role="employee",
        org_id=test_org.id,
        full_name=full_name,
    )
    db_session.add(u)
    await db_session.commit()
    await db_session.refresh(u)
    return u


@pytest.mark.asyncio
async def test_helper_prefers_reimbursement_then_author_then_none():
    advance = Purchase(purchase_method="advance", item_name="Аванс")
    single = Purchase(purchase_method="single", item_name="Обычная закупка")

    # 1) получатель возмещения задан — побеждает продавца и автора
    assert display_contractor_name(
        advance, contractor_name="ООО Продавец",
        reimbursement_user_name="Иванов И.И.", service_note_by_name="Петров П.П.",
    ) == "Иванов И.И."

    # 2) получателя нет — фолбэк на автора/инициатора
    assert display_contractor_name(
        advance, contractor_name="ООО Продавец",
        reimbursement_user_name=None, service_note_by_name="Петров П.П.",
    ) == "Петров П.П."

    # 3) ни того, ни другого нет — пусто (продавец из чеков сюда не подставляется)
    assert display_contractor_name(
        advance, contractor_name="ООО Продавец",
        reimbursement_user_name=None, service_note_by_name=None,
    ) is None

    # Обычная закупка — хелпер не трогает, отдаёт контрагента шапки как раньше
    assert display_contractor_name(
        single, contractor_name="ООО Продавец",
        reimbursement_user_name="Иванов И.И.", service_note_by_name="Петров П.П.",
    ) == "ООО Продавец"


@pytest.mark.asyncio
async def test_get_purchase_uses_reimbursement_user_name(client, db_session, test_org, auth_headers):
    recipient = await _make_user(db_session, test_org, "Сидоров С.С.")
    p = Purchase(
        item_name="Аванс на билеты",
        purchase_method="advance",
        status="planned",
        reimbursement_user_id=recipient.id,
    )
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    resp = await client.get(f"/api/purchases/{p.id}", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["contractor_name"] == "Сидоров С.С."
    assert data["reimbursement_user_name"] == "Сидоров С.С."


@pytest.mark.asyncio
async def test_get_purchase_falls_back_to_author_when_no_reimbursement(client, db_session, test_org, auth_headers):
    author = await _make_user(db_session, test_org, "Кузнецов К.К.")
    p = Purchase(
        item_name="Аванс без получателя",
        purchase_method="advance",
        status="planned",
        reimbursement_user_id=None,
        service_note_by=author.id,
    )
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    resp = await client.get(f"/api/purchases/{p.id}", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["contractor_name"] == "Кузнецов К.К."


@pytest.mark.asyncio
async def test_purchases_list_uses_reimbursement_user_name(client, db_session, test_org, admin_headers):
    # admin_headers (org_admin, org-lead) — список закупок фильтрует по видимости
    # (build_visibility_clause), закупка без assigned_user_id видна только
    # org-lead'ам (см. purchases.py::list_purchases, is_org_lead-фолбэк); тест
    # проверяет сериализацию contractor_name, а не саму видимость.
    recipient = await _make_user(db_session, test_org, "Морозова М.М.")
    subsidy = await _make_subsidy(db_session, test_org)
    p = Purchase(
        item_name="Аванс в списке",
        purchase_method="advance",
        status="planned",
        subsidy_id=subsidy.id,
        reimbursement_user_id=recipient.id,
    )
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    resp = await client.get("/api/purchases/", headers=admin_headers)
    assert resp.status_code == 200, resp.text
    rows = [r for r in resp.json() if r["id"] == p.id]
    assert len(rows) == 1
    assert rows[0]["contractor_name"] == "Морозова М.М."


@pytest.mark.asyncio
async def test_excel_export_uses_reimbursement_user_name(client, db_session, test_org, auth_headers):
    recipient = await _make_user(db_session, test_org, "Волкова В.В.")
    p = Purchase(
        item_name="Аванс для экспорта",
        purchase_method="advance",
        status="planned",
        reimbursement_user_id=recipient.id,
    )
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    resp = await client.get("/api/purchases/export/excel", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    col_idx = header.index("Контрагент")  # 0-based; openpyxl rows are 1-based cells
    values_by_item_name = {
        row[header.index("Наименование")]: row[col_idx]
        for row in ws.iter_rows(min_row=2, values_only=True)
    }
    assert values_by_item_name.get("Аванс для экспорта") == "Волкова В.В."


# ---------------------------------------------------------------------------
# seller_display_for_advance (owner, 02.10): реестр АВАНСОВЫХ
# (AdvanceReportsView.vue) обязан показывать в "Контрагент" продавца из
# чеков/позиций, а не "кому возмещать" — в отличие от общего реестра закупок
# и его Excel, которые продолжают звать display_contractor_name выше
# (не трогаем, см. тесты test_excel_export_uses_reimbursement_user_name и
# test_purchases_list_uses_reimbursement_user_name).
# ---------------------------------------------------------------------------

async def _make_contractor(db_session, name, inn=None):
    from app.models.contractor import Contractor
    c = Contractor(name=name, inn=inn)
    db_session.add(c)
    await db_session.commit()
    await db_session.refresh(c)
    return c


def test_seller_helper_single_seller_wins():
    advance = Purchase(purchase_method="advance", item_name="Аванс")
    assert seller_display_for_advance(
        advance, item_contractor_names=["ООО БЭСТ ПРАЙС"], header_contractor_name="Скурская Анна",
    ) == "ООО БЭСТ ПРАЙС"


def test_seller_helper_multiple_sellers_label():
    advance = Purchase(purchase_method="advance", item_name="Аванс")
    assert seller_display_for_advance(
        advance, item_contractor_names=["ООО А", "ООО Б", "ООО А"], header_contractor_name=None,
    ) == "Множественный контрагент"


def test_seller_helper_falls_back_to_header_contractor_when_no_item_sellers():
    advance = Purchase(purchase_method="advance", item_name="Аванс")
    assert seller_display_for_advance(
        advance, item_contractor_names=[None, None], header_contractor_name="Скурская Анна",
    ) == "Скурская Анна"


def test_seller_helper_not_advance_returns_none():
    single = Purchase(purchase_method="single", item_name="Обычная")
    assert seller_display_for_advance(
        single, item_contractor_names=["ООО А"], header_contractor_name="Иванов",
    ) is None


@pytest.mark.asyncio
async def test_advance_registry_single_seller_shown_as_contractor(client, db_session, test_org, auth_headers):
    """Воспроизводит РЕЕ-2026-00975: продавец из чеков (ООО «БЭСТ ПРАЙС»),
    а не получатель возмещения (Скурская Анна), должен попасть в
    multi_contractor_label — то, что AdvanceReportsView.vue показывает в
    колонке "Контрагент" реестра авансовых."""
    recipient = await _make_user(db_session, test_org, "Скурская Анна")
    seller = await _make_contractor(db_session, "ООО «БЭСТ ПРАЙС»", inn="7700000001")
    p = Purchase(
        item_name="Аванс РЕЕ-2026-00975",
        purchase_method="advance",
        status="planned",
        reimbursement_user_id=recipient.id,
        contractor_id=seller.id,
    )
    db_session.add(p)
    await db_session.commit()
    # ВАЖНО: НЕ делать db_session.refresh(p) здесь — session (эта же, см.
    # conftest.client) identity-map кэширует Purchase.items как «уже
    # загруженную» (пустую) коллекцию при первом касании релейшена, и
    # дальнейший selectinload через GET /api/purchases/{id} её не обновит
    # (p.id уже доступен сразу после commit() без refresh, asyncpg
    # RETURNING). items ниже создаются ДО первого обращения к p.items.
    item = PurchaseItem(purchase_id=p.id, item_name="Товар", contractor_id=seller.id)
    db_session.add(item)
    await db_session.commit()

    resp = await client.get(f"/api/purchases/{p.id}", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    # Общий реестр закупок — contractor_name всё ещё "кому возмещать" (не трогаем).
    assert data["contractor_name"] == "Скурская Анна"
    # Реестр авансовых берёт seller из multi_contractor_label.
    assert data["multi_contractor_label"] == "ООО «БЭСТ ПРАЙС»"


@pytest.mark.asyncio
async def test_advance_registry_multiple_sellers_label(client, db_session, test_org, auth_headers):
    seller_a = await _make_contractor(db_session, "ООО А", inn="7700000002")
    seller_b = await _make_contractor(db_session, "ООО Б", inn="7700000003")
    p = Purchase(
        item_name="Аванс с двумя продавцами",
        purchase_method="advance",
        status="planned",
    )
    db_session.add(p)
    await db_session.commit()
    # (см. комментарий выше — без refresh(p) до создания items)
    db_session.add_all([
        PurchaseItem(purchase_id=p.id, item_name="Товар 1", contractor_id=seller_a.id),
        PurchaseItem(purchase_id=p.id, item_name="Товар 2", contractor_id=seller_b.id),
    ])
    await db_session.commit()

    resp = await client.get(f"/api/purchases/{p.id}", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["multi_contractor_label"] == "Множественный контрагент"
