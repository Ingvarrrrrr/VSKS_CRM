"""Импорт дерева отделов и сотрудников из Excel (вынесено из departments.py,
Правило №5, сессия 2026-09-08 — см. докстринг departments.py).

Регистрируется на том же префиксе ``/api/departments`` рядом с ядром;
``/import/template`` и ``/import/excel`` на сегмент длиннее catch-all
``/{dept_id}`` — коллизий нет, порядок регистрации относительно ядра
значения не имеет.

Должность (Правило №6) при создании НОВОЙ строки membership пишется
напрямую (``UserOrganization(..., position=...)``) — установленный паттерн
(см. докстринг ``set_user_position`` в services/user_position.py: создание
membership не входит в его ответственность); правка существующей строки —
только через set_user_position, как и в departments_members.py.
"""
from io import BytesIO
from urllib.parse import quote as _url_quote

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.jwt import get_single_org_id
from app.auth.permissions import require_tab
from app.database import get_db
from app.models.department import Department
from app.models.user import User
from app.models.user_organization import UserOrganization
from app.services.org_assignment_dates import (
    first_row_assignment_dates,
    dept_transfer_date,
    position_change_date,
)
from app.services.user_position import set_user_position

router = APIRouter(prefix="/api/departments", tags=["departments"])


# ── Excel import ─────────────────────────────────────────────────────────────

@router.get("/import/template")
async def dept_import_template(_=Depends(require_tab('staff'))):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "openpyxl не установлен")

    wb = Workbook()
    ws = wb.active
    ws.title = "Отделы и сотрудники"

    headers = ["Отдел", "Родительский отдел", "ФИО сотрудника", "Должность", "Начальник (да/нет)", "Субсидия"]
    hdr_fill = PatternFill("solid", fgColor="1E40AF")
    hdr_font = Font(bold=True, color="FFFFFF")
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    examples = [
        ["Отдел закупок", "", "Иванов Иван", "Начальник отдела", "да", ""],
        ["Отдел закупок", "", "Петров Пётр", "Менеджер", "нет", ""],
        ["Склад", "", "Сидоров Сидор", "Кладовщик", "да", ""],
        ["ИТ-отдел", "", "Козлов Андрей", "Разработчик", "нет", ""],
        ["Сектор мониторинга", "Отдел закупок", "Фёдорова Мария", "Аналитик", "да", "ФАДМ_2026"],
    ]
    for ri, row in enumerate(examples, 2):
        for ci, val in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=val)

    for ci, w in enumerate([30, 25, 30, 25, 18, 20], 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_url_quote('Шаблон_импорта_отделов.xlsx', safe='-_.~')}"},
    )


@router.post("/import/excel")
async def import_departments_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_tab('staff')),
):
    """Import department tree from Excel."""
    if not (file.filename or '').lower().endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "Только .xlsx / .xls файлы")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(500, "openpyxl не установлен")

    content = await file.read()
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active

    org_id = get_single_org_id(current_user) or current_user.org_id

    # Load existing users by full_name for matching
    users_res = await db.execute(select(User).where(User.org_id == org_id))  # superadmin-bypass-ok: Excel import lookup by name, not a user-list endpoint
    users_by_name = {}
    for u in users_res.scalars().all():
        if u.full_name:
            users_by_name[u.full_name.strip().lower()] = u

    # Load existing subsidies for matching
    from app.models.subsidy import Subsidy
    subs_res = await db.execute(select(Subsidy))
    subs_by_name = {s.name.strip().lower(): s for s in subs_res.scalars().all() if s.name}

    # Parse header
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise HTTPException(400, "Файл пустой")

    def _norm(v):
        return str(v).strip().lower() if v else ''

    col = {}
    for i, h in enumerate(header_row):
        hs = _norm(h)
        if any(x in hs for x in ('отдел', 'department', 'подразделен')):
            col.setdefault('dept', i)
        if any(x in hs for x in ('родител', 'parent')):
            col.setdefault('parent', i)
        if any(x in hs for x in ('фио', 'сотрудник', 'имя', 'full_name')):
            col.setdefault('name', i)
        if any(x in hs for x in ('должност', 'position', 'позиц')):
            col.setdefault('position', i)
        if any(x in hs for x in ('начальник', 'head', 'руковод')):
            col.setdefault('head', i)
        if any(x in hs for x in ('субсид', 'subsidy')):
            col.setdefault('subsidy', i)

    if 'dept' not in col:
        raise HTTPException(400, "Не найдена колонка «Отдел»")

    def _cell(row, field):
        idx = col.get(field)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        return str(v).strip() if v else None

    # First pass: collect departments
    dept_cache = {}  # name -> Department
    created_depts = 0
    created_members = 0
    errors = []

    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # Create departments first
    for ri, row in enumerate(rows, 2):
        dept_name = _cell(row, 'dept')
        if not dept_name:
            continue
        key = dept_name.lower()
        if key not in dept_cache:
            subsidy_name = _cell(row, 'subsidy')
            subsidy_id = None
            if subsidy_name:
                sub = subs_by_name.get(subsidy_name.lower())
                if sub:
                    subsidy_id = sub.id

            # Check if already exists in DB
            existing = (await db.execute(
                select(Department).where(Department.org_id == org_id, func.lower(Department.name) == key)
            )).scalar_one_or_none()
            if existing:
                dept_cache[key] = existing
            else:
                dept = Department(name=dept_name.strip(), org_id=org_id, subsidy_id=subsidy_id)
                db.add(dept)
                await db.flush()
                dept_cache[key] = dept
                created_depts += 1

    # Set parent_id for departments
    for ri, row in enumerate(rows, 2):
        dept_name = _cell(row, 'dept')
        parent_name = _cell(row, 'parent')
        if dept_name and parent_name:
            dept = dept_cache.get(dept_name.lower())
            parent = dept_cache.get(parent_name.lower())
            if dept and parent and dept.id != parent.id:
                dept.parent_id = parent.id

    # Second pass: add members
    for ri, row in enumerate(rows, 2):
        dept_name = _cell(row, 'dept')
        user_name = _cell(row, 'name')
        if not dept_name or not user_name:
            continue

        dept = dept_cache.get(dept_name.lower())
        if not dept:
            errors.append({"row": ri, "error": f"Отдел «{dept_name}» не найден"})
            continue

        user = users_by_name.get(user_name.lower())
        if not user:
            errors.append({"row": ri, "error": f"Сотрудник «{user_name}» не найден в системе"})
            continue

        position = _cell(row, 'position')
        is_head = _cell(row, 'head')
        is_head_bool = is_head and is_head.lower() in ('да', 'yes', '1', 'true', 'head', 'начальник')

        # Add member via user_organizations (single source of truth)
        existing_m = (await db.execute(
            select(UserOrganization).where(
                UserOrganization.user_id == user.id,
                UserOrganization.org_id == dept.org_id,
                UserOrganization.dept_id == dept.id,
            )
        )).scalar_one_or_none()
        if not existing_m:
            # Заглушка «Без отдела» для этой же пары (user, org) — переносим её
            # hired_at/ставку в новую dept-строку и удаляем саму заглушку (та же
            # болезнь, что и в add_member: иначе остаётся висеть «Без отдела»).
            null_row = (await db.execute(
                select(UserOrganization).where(
                    UserOrganization.user_id == user.id,
                    UserOrganization.org_id == dept.org_id,
                    UserOrganization.dept_id.is_(None),
                )
            )).scalar_one_or_none()
            base_row = null_row or (await db.execute(
                select(UserOrganization)
                .where(
                    UserOrganization.user_id == user.id,
                    UserOrganization.org_id == dept.org_id,
                )
                .order_by(UserOrganization.id.asc())
                .limit(1)
            )).scalars().first()
            new_uo = UserOrganization(user_id=user.id, org_id=dept.org_id, dept_id=dept.id, position=position)
            if base_row is not None and base_row.hired_at is not None:
                new_uo.hired_at = base_row.hired_at
            if base_row is None:
                # Первая-ЕВЕР строка пары (user, org) — приём на работу, обе даты
                # назначения = дате приёма (владелец, 2026-09-01).
                dates = first_row_assignment_dates(
                    new_uo.hired_at, has_position=bool(position), has_dept=True,
                )
                new_uo.dept_assigned_at = dates.get("dept_assigned_at")
                new_uo.position_assigned_at = dates.get("position_assigned_at")
            else:
                # Уже трудоустроен — импорт добавляет его в новый отдел (перевод).
                new_uo.dept_assigned_at = dept_transfer_date(None)
                position_changed = bool(position) and position != base_row.position
                new_uo.position_assigned_at = position_change_date(
                    position_changed=position_changed,
                    current=base_row.position_assigned_at,
                    explicit=None,
                )
            if base_row is not None:
                if base_row.salary_amount is not None:
                    new_uo.salary_amount = base_row.salary_amount
                if base_row.employment_percent is not None:
                    new_uo.employment_percent = base_row.employment_percent
            if null_row is not None:
                await db.delete(null_row)
            db.add(new_uo)
            created_members += 1
        elif position:
            position_changed = position != existing_m.position
            await set_user_position(db, position, membership=existing_m)
            existing_m.position_assigned_at = position_change_date(
                position_changed=position_changed,
                current=existing_m.position_assigned_at,
                explicit=None,
            )

        if is_head_bool:
            dept.head_user_id = user.id

        # Update user.department
        user.department = dept.name

    await db.commit()
    return {
        "created_departments": created_depts,
        "created_members": created_members,
        "errors": errors,
    }
