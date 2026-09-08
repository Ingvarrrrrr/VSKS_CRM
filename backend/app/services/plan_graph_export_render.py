"""Рендер книги Excel плана-графика по дереву ФЭО (или плоскому снапшоту v1).

Вынесено из app/routers/subsidy_plan_graph_export.py (Правило №5, рефакторинг
2026-09-08). Используется эндпоинтом экспорта версии (снапшота) плана-графика:
GET .../plan-graph/versions/{version_id}/export.

Живой (текущий) экспорт .../plan-graph/export НЕ использует эту функцию — у
него собственный, более детальный рендерер с колонками каскада статусов,
факта по позициям и листом «Сводная»: app.services.plan_graph_export_xlsx.
"""
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None


def render_plan_graph_workbook(tree: list, items: list, meta: dict):
    """
    Common renderer for plan-graph Excel workbook.

    Args:
        tree: recursive list of FeoCategory nodes
              [{id, name, level, code, appendix, budget, planned_amount,
                planned_quantity, unit, children:[...]}]
              If empty — fallback to flat `items` (v1 backward-compat).
        items: flat list of snapshot items (v1 backward-compat)
               [{name, planned_amount, used_amount, residual}]
        meta: dict with keys:
              subsidy_name, subsidy_year (optional), effective_date (str|None),
              version_number (int|None), note (str|None), generated_at (str|None)

    Returns:
        openpyxl.Workbook
    """
    if openpyxl is None:
        raise RuntimeError("openpyxl не установлен")

    HEADER_FILL  = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT  = Font(color="FFFFFF", bold=True, size=9)
    L1_FILL      = PatternFill("solid", fgColor="DBEAFE")
    L1_FONT      = Font(bold=True, size=9)
    L2_FILL      = PatternFill("solid", fgColor="F0F9FF")
    L2_FONT      = Font(bold=True, size=9, color="0C4A6E")
    L3_FILL      = PatternFill("solid", fgColor="F0FDF4")
    L3_FONT      = Font(size=9, color="166534")
    ITEM_FONT    = Font(size=9)
    RED_FONT     = Font(size=9, color="EF4444", bold=True)
    META_FONT    = Font(size=9, italic=True, color="374151")
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_ALIGN   = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT_ALIGN  = Alignment(horizontal="right", vertical="center")
    THIN_BORDER  = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    HEADERS = [
        "№", "Направление расходов", "Тип расходов", "Наименование",
        "Ед.", "Кол-во план", "Плановая сумма, ₽",
        "Фактическая сумма, ₽", "Остаток, ₽",
        "% исполнения", "Исполнитель", "Статус",
    ]
    COL_WIDTHS = [5, 30, 25, 40, 8, 10, 18, 18, 18, 12, 30, 15]
    n_cols = len(HEADERS)
    last_col_letter = chr(64 + n_cols)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "План закупок"

    # ── Title row ─────────────────────────────────────────────────────────────
    subsidy_name = meta.get("subsidy_name", "")
    subsidy_year = meta.get("subsidy_year", "")
    title_text = f"ПЛАН-ГРАФИК — {subsidy_name}"
    if subsidy_year:
        title_text += f" ({subsidy_year})"
    ws.append([title_text] + [""] * (n_cols - 1))
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=12, color="1E3A5F")
    ws.merge_cells(f"A1:{last_col_letter}1")
    title_cell.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    # ── Meta info rows (version, effective_date, note, generated_at) ──────────
    meta_start_row = 2
    meta_rows = []
    if meta.get("version_number") is not None:
        meta_rows.append(f"Версия: {meta['version_number']}")
    if meta.get("effective_date"):
        meta_rows.append(f"Дата редакции: {meta['effective_date']}")
    if meta.get("note"):
        meta_rows.append(f"Примечание: {meta['note']}")
    if meta.get("generated_at"):
        meta_rows.append(f"Сформировано: {meta['generated_at']}")

    for i, mtext in enumerate(meta_rows):
        r = meta_start_row + i
        ws.append([mtext] + [""] * (n_cols - 1))
        cell = ws.cell(row=r, column=1)
        cell.font = META_FONT
        cell.alignment = LEFT_ALIGN
        ws.merge_cells(f"A{r}:{last_col_letter}{r}")
        ws.row_dimensions[r].height = 16

    # ── Column headers ────────────────────────────────────────────────────────
    header_row = meta_start_row + len(meta_rows)
    ws.append(HEADERS)
    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[header_row].height = 32
    ws.freeze_panes = f"A{header_row + 1}"

    row_num = header_row + 1
    seq = 0

    def _write_row(values, fill, font, height=20):
        nonlocal row_num
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.fill = fill
            cell.font = font
            cell.border = THIN_BORDER
            cell.alignment = RIGHT_ALIGN if col_idx >= 6 else LEFT_ALIGN
        ws.row_dimensions[row_num].height = height
        row_num += 1

    if tree:
        # ── Tree mode (schema_version=2) ──────────────────────────────────────
        def _traverse_node(node, direction_name="", type_name=""):
            nonlocal seq
            level = node.get("level", 1)
            name = node.get("name", "")
            code = node.get("code") or ""
            budget = node.get("budget")

            if level == 1:
                direction_name = name
                _write_row(
                    [code, name, "", "", "", "", budget or "", "", "", "", "", ""],
                    L1_FILL, L1_FONT, height=22,
                )
            elif level == 2:
                type_name = name
                _write_row(
                    ["", direction_name, name, "", "", "", budget or "", "", "", "", "", ""],
                    L2_FILL, L2_FONT,
                )
            elif level == 3:
                _write_row(
                    ["", direction_name, type_name, name, "", "", budget or "", "", "", "", "", ""],
                    L3_FILL, L3_FONT,
                )

            for child in node.get("children", []):
                _traverse_node(child, direction_name, type_name)

        for root in tree:
            _traverse_node(root)
    else:
        # ── Flat mode (v1 backward-compat) ─────────────────────────────────────
        for item in items:
            seq += 1
            planned = float(item.get("planned_amount") or 0)
            used = float(item.get("used_amount") or 0)
            residual = float(item.get("residual") or (planned - used))
            pct = round(used / planned * 100) if planned > 0 else 0
            status = "Выполнено" if pct >= 100 else ("В работе" if used > 0 else "Не начато")
            font = RED_FONT if used > planned else ITEM_FONT
            _write_row(
                [
                    seq, "", "", item.get("name", ""),
                    "", "",
                    round(planned, 2), round(used, 2), round(residual, 2),
                    f"{pct}%", "", status,
                ],
                PatternFill(), font,
            )

    return wb
