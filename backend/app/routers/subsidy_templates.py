"""Шаблоны документов субсидии — глобальные и per-subsidy overrides.

Вынесено из app/routers/subsidies.py (Правило №5, рефакторинг 2026-09-07):
list/upload/download/delete per-subsidy шаблонов, global-templates upload,
массовая нормализация (.docx normalize-all) + вспомогательные
_check_template_render_ok / _normalize_docx_template / _repair_docx_template.

Собственный APIRouter на том же префиксе /api/subsidies — регистрируется в
app/routes.py РЯДОМ с subsidies.router (см. комментарий про порядок там):
"/global-templates/*" и "/templates/normalize-all" — статические литеральные
пути, должны резолвиться раньше catch-all "/{subsidy_id}" в subsidies.router.

`_normalize_docx_template` используется лениво из documents.py (рендер
шаблона) — re-export оставлен в app/routers/subsidies.py для обратной
совместимости этого импорта.
"""
import os
import re
import shutil
import logging

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import FileResponse

from app.auth.permissions import require_tab, require_action

try:
    from docxtpl import DocxTemplate
except ImportError:
    DocxTemplate = None

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/subsidies", tags=["subsidies"])


# ── Per-subsidy document templates ──────────────────────────────────────────

TEMPLATES_BASE = "/app/templates"
SUBSIDY_TEMPLATES_BASE = "/app/uploads/templates"
SUPPORTED_DOC_TYPES = {
    "contract":                 "Договор",
    "contract_tz":              "ТЗ (общий шаблон)",
    # tech_spec slot removed from SubsidiesView UI 2026-04-21 — the endpoint
    # still resolves to contract_tz.docx (see documents.py DOC_TYPES fallback)
    # for any client that requests /documents/tech_spec directly.
    "service_note_delivery":    "СЗ на выдачу",
    "service_note_payment":     "СЗ на оплату",
    # Phase 19.05: split ТЗ and dedicated SZ на закупку
    "service_note_procurement": "СЗ на закупку",
    # Phase 19.07: СЗ на аванс
    "service_note_advance":     "СЗ на аванс",
    "tech_spec_request":        "ТЗ для запроса цен",
    "tech_spec_contract":       "ТЗ для договора",
    "approval_sheet":           "Лист согласования",
    "order_purchase":           "Приказ на закупку",
    # Phase 28: typed contract forms per-subsidy
    # Форма «услуги» объединена в один шаблон; большая/малая отчётность теперь
    # выбирается отдельной методичкой (methodology_large/small), приклеиваемой
    # к любому из семи договоров, а не отдельным шаблоном текста договора.
    "contract_services":             "Договор услуг",
    "contract_services_food":       "Договор услуг (питание)",
    "contract_goods_single":        "Договор поставки (разовый)",
    "contract_gph_individual":      "Договор ГПХ с физ.лицом (без РИД)",
    "contract_gph_individual_rid":  "Договор ГПХ с физ.лицом (+РИД)",
    "contract_repair_vehicle":      "Договор на ремонт ТС",
    "contract_repair_framework":    "Рамочный договор на ремонт ТС",
    # Методические рекомендации — приклеиваются к договору по Purchase.methodology
    "methodology_large":            "Методические рекомендации — большие",
    "methodology_small":            "Методические рекомендации — малые",
    # Fabrikant ЭТП package
    "fabrikant_instruction":        "Инструкция по заполнению заявки (Фабрикант)",
    "fabrikant_application_form":   "Форма заявки (Фабрикант)",
    "fabrikant_documentation":      "Документация к закупке (Фабрикант)",
    "fabrikant_contract_project":   "Проект договора (Фабрикант)",
}


def _check_template_render_ok(path: str) -> bool:
    """Try a trial render with empty context; return True if template is valid."""
    if not os.path.exists(path):
        return False
    try:
        if DocxTemplate is None:
            return True  # docxtpl not installed — assume ok
        tpl = DocxTemplate(path)
        tpl.render({})
        return True
    except Exception:
        return False


@router.get("/{subsidy_id}/templates")
async def list_subsidy_templates(
    subsidy_id: int,
    current_user=Depends(require_tab('subsidies')),
):
    """List which doc types have a subsidy-specific template override."""
    result = []
    subsidy_dir = os.path.join(SUBSIDY_TEMPLATES_BASE, "subsidies", str(subsidy_id))
    for doc_type, label in SUPPORTED_DOC_TYPES.items():
        path = os.path.join(subsidy_dir, f"{doc_type}.docx")
        global_path = os.path.join(TEMPLATES_BASE, f"{doc_type}.docx")
        has_custom = os.path.exists(path)
        render_ok: bool | None = None
        if has_custom:
            render_ok = _check_template_render_ok(path)
        result.append({
            "doc_type": doc_type,
            "label": label,
            "has_custom": has_custom,
            "has_global": os.path.exists(global_path),
            "render_ok": render_ok,
        })
    return result


def _normalize_docx_template(path: str) -> dict:
    """Strip Word-internal markers that split jinja-placeholders across runs.

    Word inserts <w:proofErr/>, bookmarks, comments, lastRenderedPageBreak
    between runs inside `{{ var }}` / `{% ... %}` blocks during user editing.
    For text placeholders docxtpl can recover, but for InlineImage the tag
    MUST be inside a single contiguous run — otherwise the drawing element
    is silently dropped.

    Touches ONLY xml files where jinja syntax may live (document/header/
    footer/footnotes/endnotes). Avoiding [Content_Types].xml and *.rels —
    rewriting those breaks the package and was the root cause of the
    Phase 26-VV revert.

    Writes the cleaned bytes through a tempfile + atomic replace so the
    on-disk file stays valid even if a write is interrupted.
    """
    import zipfile as _zipfile
    import tempfile as _tempfile

    JINJA_BEARING = ('word/document.xml',)
    JINJA_BEARING_PREFIXES = ('word/header', 'word/footer')
    JINJA_BEARING_SUFFIXES = ('word/footnotes.xml', 'word/endnotes.xml')

    def _is_jinja_bearing(name: str) -> bool:
        if name in JINJA_BEARING or name in JINJA_BEARING_SUFFIXES:
            return True
        return any(name.startswith(p) and name.endswith('.xml') for p in JINJA_BEARING_PREFIXES)

    stats = {"proofErr": 0, "bookmark": 0, "comment": 0, "pageBreak": 0}
    tmp_fd, tmp_path = _tempfile.mkstemp(suffix=".docx", dir=os.path.dirname(path) or None)
    os.close(tmp_fd)
    try:
        with _zipfile.ZipFile(path, 'r') as zin:
            with _zipfile.ZipFile(tmp_path, 'w', _zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if _is_jinja_bearing(item.filename):
                        xml = data.decode('utf-8', errors='replace')
                        before_proof = xml.count('<w:proofErr')
                        xml = re.sub(r'<w:proofErr\s[^/]*?/>', '', xml)
                        stats["proofErr"] += before_proof - xml.count('<w:proofErr')

                        before_bm = xml.count('<w:bookmark')
                        xml = re.sub(r'<w:bookmarkStart\s[^/]*?/>', '', xml)
                        xml = re.sub(r'<w:bookmarkEnd\s[^/]*?/>', '', xml)
                        stats["bookmark"] += before_bm - xml.count('<w:bookmark')

                        before_cm = xml.count('<w:comment')
                        xml = re.sub(r'<w:commentRangeStart\s[^/]*?/>', '', xml)
                        xml = re.sub(r'<w:commentRangeEnd\s[^/]*?/>', '', xml)
                        xml = re.sub(r'<w:commentReference\s[^/]*?/>', '', xml)
                        stats["comment"] += before_cm - xml.count('<w:comment')

                        before_pb = xml.count('<w:lastRenderedPageBreak')
                        xml = re.sub(r'<w:lastRenderedPageBreak\s*/>', '', xml)
                        stats["pageBreak"] += before_pb - xml.count('<w:lastRenderedPageBreak')

                        data = xml.encode('utf-8')
                    zout.writestr(item, data)
        os.replace(tmp_path, path)
        logger.info("docx normalize %s: %s", path, stats)
    except Exception as e:
        logger.warning("docx normalize failed for %s: %s", path, e)
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
    return stats


def _repair_docx_template(path: str) -> list[str]:
    """Fix common Jinja2 errors in uploaded .docx templates.

    Word often splits {{ variable }} across multiple XML runs, and users
    accidentally put Russian text like 'г.' inside Jinja tags.
    Returns list of fixes applied.
    """
    fixes = []
    try:
        from docx import Document
        doc = Document(path)

        def _fix_paragraph(para):
            if not para.runs:
                return False
            full = "".join(r.text for r in para.runs)
            if "{{" not in full and "{%" not in full:
                return False

            original = full
            # Fix: {{ var something_russian }} -> {{ var }} something_russian
            # Pattern: inside {{ }}, after a valid variable name, remove non-ASCII chars
            def _clean_tag(m):
                inner = m.group(1)
                # Split into variable name and trailing junk
                parts = inner.strip().split()
                if len(parts) <= 1:
                    return m.group(0)  # just {{ var }} — ok
                var_name = parts[0]
                # Check if first part looks like a variable (ascii, underscores, dots)
                if not re.match(r'^[a-zA-Z_][a-zA-Z0-9_.]*$', var_name):
                    return m.group(0)
                # Remaining parts — check if they contain non-ASCII (Russian text)
                rest = " ".join(parts[1:])
                if any(ord(c) > 127 for c in rest):
                    return "{{ " + var_name + " }} " + rest
                return m.group(0)

            fixed = re.sub(r'\{\{([^}]+)\}\}', _clean_tag, full)

            # Fix: {{полностью_русский текст}} — снимаем маркеры, оставляем текст
            def _strip_if_all_non_ascii(m):
                inner = m.group(1).strip()
                if not inner:
                    return m.group(0)
                # Если ни одного ASCII-идентификаторного символа — это ремарка, не переменная
                if not re.search(r'[a-zA-Z_]', inner):
                    return inner
                return m.group(0)
            fixed = re.sub(r'\{\{([^}]+)\}\}', _strip_if_all_non_ascii, fixed)

            # Fix: {% русский текст %} — снимаем маркеры если не Jinja keyword
            _JINJA_KEYWORDS = {
                'if', 'else', 'elif', 'endif', 'for', 'endfor', 'set',
                'with', 'endwith', 'tr', 'block', 'endblock', 'macro',
                'endmacro', 'include', 'extends', 'import', 'from',
                'do', 'raw', 'endraw', 'call', 'endcall', 'filter',
                'endfilter', 'autoescape', 'endautoescape',
            }
            def _strip_pct_if_no_keyword(m):
                inner = m.group(1).strip().lstrip('-').rstrip('-').strip()
                if not inner:
                    return m.group(0)
                first = inner.split()[0].lower() if inner.split() else ''
                # Если первое слово — не Jinja keyword И не похоже на ASCII identifier — ремарка
                if first not in _JINJA_KEYWORDS and not re.match(r'^[a-zA-Z_]', first):
                    return inner
                return m.group(0)
            fixed = re.sub(r'\{%([^%]+)%\}', _strip_pct_if_no_keyword, fixed)

            # Fix doubled text like "г. г." that can result from prior fixes
            fixed = re.sub(r'(\b\w+\.)\s+\1', r'\1', fixed)

            if fixed == original:
                return False
            # Write back: all text into first run, clear rest
            para.runs[0].text = fixed
            for r in para.runs[1:]:
                r.text = ""
            return True

        count = 0
        for para in doc.paragraphs:
            if _fix_paragraph(para):
                count += 1
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        if _fix_paragraph(para):
                            count += 1

        if count:
            doc.save(path)
            fixes.append(f"Исправлено {count} Jinja-тегов")

        # Validate the template renders
        from docxtpl import DocxTemplate
        tpl = DocxTemplate(path)
        tpl.get_undeclared_template_variables()

    except Exception as e:
        logger.warning("Template repair/validation warning for %s: %s", path, e)
        fixes.append(f"Предупреждение: {e}")

    return fixes


@router.put("/{subsidy_id}/templates/{doc_type}")
async def upload_subsidy_template(
    subsidy_id: int,
    doc_type: str,
    file: UploadFile = File(...),
    current_user=Depends(require_tab('subsidies')),
):
    """Upload a .docx template override for a specific subsidy and doc type."""
    if doc_type not in SUPPORTED_DOC_TYPES:
        raise HTTPException(400, f"Неизвестный тип документа: {doc_type}")
    if not file.filename.endswith(".docx"):
        raise HTTPException(400, "Допускаются только .docx файлы")

    subsidy_dir = os.path.join(SUBSIDY_TEMPLATES_BASE, "subsidies", str(subsidy_id))
    os.makedirs(subsidy_dir, exist_ok=True)
    dest = os.path.join(subsidy_dir, f"{doc_type}.docx")

    with open(dest, "wb") as f:
        shutil.copyfileobj(file.file, f)

    norm_stats = _normalize_docx_template(dest)
    repairs = _repair_docx_template(dest)

    # Trial render: validate docxtpl syntax before confirming upload
    if DocxTemplate is not None:
        try:
            tpl = DocxTemplate(dest)
            tpl.render({})
        except Exception as e:
            try:
                os.remove(dest)
            except OSError:
                pass
            raise HTTPException(
                status_code=400,
                detail=f"Шаблон некорректен: {type(e).__name__}: {str(e)[:200]}"
            )

    return {
        "ok": True,
        "doc_type": doc_type,
        "label": SUPPORTED_DOC_TYPES[doc_type],
        "repairs": repairs,
        "normalized": norm_stats,
    }


@router.get("/{subsidy_id}/templates/{doc_type}/download")
async def download_subsidy_template(
    subsidy_id: int,
    doc_type: str,
    current_user=Depends(require_tab('subsidies')),
):
    """Download the current subsidy-specific template (or global if no override)."""
    if doc_type not in SUPPORTED_DOC_TYPES:
        raise HTTPException(400, f"Неизвестный тип документа: {doc_type}")

    subsidy_path = os.path.join(SUBSIDY_TEMPLATES_BASE, "subsidies", str(subsidy_id), f"{doc_type}.docx")
    global_path = os.path.join(TEMPLATES_BASE, f"{doc_type}.docx")

    if os.path.exists(subsidy_path):
        path = subsidy_path
    elif os.path.exists(global_path):
        path = global_path
    else:
        raise HTTPException(404, "Шаблон не найден")

    response = FileResponse(
        path,
        filename=f"{doc_type}_{subsidy_id}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


@router.delete("/{subsidy_id}/templates/{doc_type}")
async def delete_subsidy_template(
    subsidy_id: int,
    doc_type: str,
    current_user=Depends(require_tab('subsidies')),
):
    """Delete subsidy-specific template override (falls back to global)."""
    if doc_type not in SUPPORTED_DOC_TYPES:
        raise HTTPException(400, f"Неизвестный тип документа: {doc_type}")

    path = os.path.join(SUBSIDY_TEMPLATES_BASE, "subsidies", str(subsidy_id), f"{doc_type}.docx")
    if not os.path.exists(path):
        raise HTTPException(404, "Индивидуальный шаблон не найден")

    os.remove(path)
    return {"ok": True}


# ── Global template upload (admin) ──────────────────────────────────────────

@router.put("/global-templates/{doc_type}")
async def upload_global_template(
    doc_type: str,
    file: UploadFile = File(...),
    current_user=Depends(require_action('subsidy.edit')),
):
    """Upload a global .docx template (superadmin/account_owner only)."""
    if doc_type not in SUPPORTED_DOC_TYPES:
        raise HTTPException(400, f"Неизвестный тип документа: {doc_type}")
    if not file.filename.endswith(".docx"):
        raise HTTPException(400, "Допускаются только .docx файлы")

    dest = os.path.join(TEMPLATES_BASE, f"{doc_type}.docx")
    with open(dest, "wb") as f:
        shutil.copyfileobj(file.file, f)

    norm_stats = _normalize_docx_template(dest)
    repairs = _repair_docx_template(dest)

    # Trial render: validate docxtpl syntax before confirming upload
    if DocxTemplate is not None:
        try:
            tpl = DocxTemplate(dest)
            tpl.render({})
        except Exception as e:
            try:
                os.remove(dest)
            except OSError:
                pass
            raise HTTPException(
                status_code=400,
                detail=f"Шаблон некорректен: {type(e).__name__}: {str(e)[:200]}"
            )

    return {"ok": True, "doc_type": doc_type, "repairs": repairs, "normalized": norm_stats}


@router.post("/templates/normalize-all")
async def normalize_all_existing_templates(
    current_user=Depends(require_action('subsidy.edit')),
):
    """One-shot migration: normalize every existing .docx template in place.

    Walks SUBSIDY_TEMPLATES_BASE and TEMPLATES_BASE, strips Word-internal
    markers from each file. Idempotent — safe to call multiple times.
    """
    processed = []
    errors = []

    def _walk_and_normalize(root: str):
        if not os.path.isdir(root):
            return
        for dirpath, _dirs, files in os.walk(root):
            for fn in files:
                if not fn.endswith(".docx"):
                    continue
                fpath = os.path.join(dirpath, fn)
                try:
                    stats = _normalize_docx_template(fpath)
                    processed.append({"path": fpath, "stripped": stats})
                    try:
                        repairs = _repair_docx_template(fpath)
                        if repairs:
                            processed[-1]["repairs"] = repairs
                    except Exception as re_err:
                        errors.append({"path": fpath, "error": f"repair: {re_err}"})
                except Exception as e:
                    errors.append({"path": fpath, "error": str(e)})

    _walk_and_normalize(SUBSIDY_TEMPLATES_BASE)
    _walk_and_normalize(TEMPLATES_BASE)

    return {"ok": True, "processed": len(processed), "errors": len(errors), "details": processed, "error_details": errors}

