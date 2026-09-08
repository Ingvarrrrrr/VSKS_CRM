"""Рендер живой (текущей) книги Excel плана-графика субсидии.

Вынесено из app/routers/subsidy_plan_graph_export.py::export_plan_graph_excel
(Правило №5, рефакторинг 2026-09-08). Принимает уже собранные данные из
app.services.plan_graph_export_data.gather_live_plan_graph_data — сам к БД не
обращается.

Отличается от app.services.plan_graph_export_render.render_plan_graph_workbook
(используется для экспорта снапшота версии): здесь каскад статусов
Заказано/Поставлено/Оплачено, факт по каждой закупленной позиции с гиперссылкой
на заказ, и лист «Сводная».
"""
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None


def build_live_plan_graph_xlsx(sub, base_url: str, data: dict):
    """Строит openpyxl.Workbook живого плана-графика.

    Args:
        sub: Subsidy — для заголовка (name, year).
        base_url: базовый URL запроса (для гиперссылок «Открыть» на заказ).
        data: dict из gather_live_plan_graph_data(...).

    Returns:
        openpyxl.Workbook

    Примечание: наличие openpyxl проверяет вызывающий роутер (дважды, включая
    унаследованный дефект-дубль — см. app/routers/subsidy_plan_graph_export.py);
    здесь отдельной проверки нет, чтобы не дублировать исключение сверх
    существовавших в дорефакторинговом коде.
    """
    cats = data["cats"]
    items_by_cat = data["items_by_cat"]
    used_map = data["used_map"]
    contractor_map = data["contractor_map"]
    status_sums_map = data["status_sums_map"]
    purchased_by_item = data["purchased_by_item"]
    cat_status_map = data["cat_status_map"]
    cat_monthly_map = data["cat_monthly_map"]
    cat_contractor_map = data["cat_contractor_map"]
    purchased_by_cat = data["purchased_by_cat"]
    unlinked_purchases = data["unlinked_purchases"]
    summary = data["summary"]

    def _st(item_id: int, *statuses: str) -> float:
        d = status_sums_map.get(item_id, {})
        return sum(d.get(s, 0.0) for s in statuses)

    HEADER_FILL  = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT  = Font(color="FFFFFF", bold=True, size=9)
    L1_FILL      = PatternFill("solid", fgColor="DBEAFE")
    L1_FONT      = Font(bold=True, size=9)
    L2_FILL      = PatternFill("solid", fgColor="F0F9FF")
    L2_FONT      = Font(bold=True, size=9, color="0C4A6E")
    L3_FILL      = PatternFill("solid", fgColor="F0FDF4")
    L3_FONT      = Font(size=9, color="166534")
    ITEM_FONT    = Font(size=9)
    RED_FONT     = Font(size=9, color="EF4444", bold=True)
    SUB_FONT     = Font(size=8, italic=True, color="6B7280")
    SUB_FILL     = PatternFill("solid", fgColor="FAFAFA")
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_ALIGN   = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT_ALIGN  = Alignment(horizontal="right", vertical="center")
    THIN_BORDER  = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    HEADERS = [
        "№", "Направление расходов", "Тип расходов", "Наименование",
        "Ед.", "Кол-во план", "Бюджет ФЭО, ₽", "Запланировано, ₽",
        "Договор, ₽", "Заказано, ₽", "Поставлено, ₽", "Оплачено, ₽",
        "Фактически (итого), ₽", "Остаток, ₽",
        "% исполнения", "Исполнитель", "Статус",
        "Ежемес. (регуляр.) платежи, ₽",
        "№ акта", "Документы",
    ]
    COL_WIDTHS = [5, 30, 25, 40, 8, 10, 18, 18, 18, 18, 18, 18, 20, 18, 12, 30, 15, 22, 16, 14]

    cats_by_parent: dict = {}
    for c in cats:
        cats_by_parent.setdefault(c.parent_id, []).append(c)

    # Роллап факта/бюджета вверх по дереву: направление (ур.1) и группа (ур.2)
    # показывают суммарные итоги по всем листовым категориям-потомкам.
    # Факт листа берём по feo_category_id (cat_status_map) — это включает все
    # закупки категории, поэтому суммирование по потомкам не задваивается.
    subtree_map: dict[int, dict] = {}

    def _compute_subtree(cat) -> dict:
        if cat.id in subtree_map:
            return subtree_map[cat.id]
        children = cats_by_parent.get(cat.id, [])
        status_agg: dict[str, float] = {}
        if not children:
            for k, v in cat_status_map.get(cat.id, {}).items():
                status_agg[k] = status_agg.get(k, 0.0) + v
            monthly = cat_monthly_map.get(cat.id, 0.0)
            if cat.budget is not None:
                budget = float(cat.budget)
            else:
                budget = sum(float(it.amount or 0) for it in items_by_cat.get(cat.id, []))
        else:
            budget = 0.0
            monthly = 0.0
            for ch in children:
                r = _compute_subtree(ch)
                for k, v in r["status"].items():
                    status_agg[k] = status_agg.get(k, 0.0) + v
                budget += r["budget"]
                monthly += r["monthly"]
        res = {"status": status_agg, "budget": budget, "monthly": monthly}
        subtree_map[cat.id] = res
        return res

    for c in cats:
        _compute_subtree(c)

    def _sub(cat_id: int, *statuses: str) -> float:
        d = subtree_map.get(cat_id, {}).get("status", {})
        return sum(d.get(s, 0.0) for s in statuses)

    LINK_FONT = Font(size=8, italic=True, color="2563EB", underline="single")

    def _cascade(raw_status: str, total: float):
        """Каскад суммы поставки по стадиям: сумма падает во все столбцы до текущей
        стадии включительно; недостигнутые = 0. Столбцы: Заказано/Поставлено/Оплачено."""
        t = round(total or 0, 2)
        ordered = t if raw_status in ("ordered", "delivered", "paid") else 0
        delivered = t if raw_status in ("delivered", "paid") else 0
        paid = t if raw_status == "paid" else 0
        return ordered, delivered, paid

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "План закупок"

    # Заголовок-титул в строке 1 (не через insert_rows — иначе гиперлинки не
    # сдвигаются вместе с ячейками и «съезжают» на строку выше).
    title_cell = ws.cell(row=1, column=1, value=f"ПЛАН-ГРАФИК — {sub.name} ({sub.year})")
    title_cell.font = Font(bold=True, size=12, color="1E3A5F")
    ws.merge_cells(f"A1:{chr(64 + len(HEADERS))}1")
    title_cell.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    for col_idx, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "A3"

    row_num = 3
    seq = 0

    def _write_row(values, fill, font, height=20):
        nonlocal row_num
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.fill = fill
            cell.font = font
            cell.border = THIN_BORDER
            cell.alignment = RIGHT_ALIGN if col_idx >= 6 else LEFT_ALIGN
        ws.row_dimensions[row_num].height = height
        row_num += 1

    def _set_doc_link(purchase_id):
        """Проставить гиперлинк «Открыть» в колонку 20 последней записанной строки —
        ведёт на страницу заказа, где под авторизацией доступны все документы."""
        if not purchase_id:
            return
        cell = ws.cell(row=row_num - 1, column=20)
        cell.value = "Открыть"
        cell.hyperlink = f"{base_url}/orders/{purchase_id}"
        cell.font = LINK_FONT

    E = ""  # пустая ячейка

    def _money_cols(cat_id: int, contractor: str) -> list:
        """Денежные колонки 7..18 из роллапа поддерева по статусам."""
        st = subtree_map.get(cat_id, {})
        budget = st.get("budget", 0.0)
        monthly = st.get("monthly", 0.0)
        plan = _sub(cat_id, "plan_schedule", "work_in_progress", "wishes")
        contract = _sub(cat_id, "contracted")
        ordered = _sub(cat_id, "ordered")
        delivered = _sub(cat_id, "delivered")
        paid = _sub(cat_id, "paid")
        used = plan + contract + ordered + delivered + paid
        resid = budget - used

        def _nz(v):
            return round(v, 2) if abs(v) > 0.005 else E

        return [
            round(budget, 2) if budget else E,
            _nz(plan), _nz(contract), _nz(ordered), _nz(delivered), _nz(paid),
            _nz(used),
            round(resid, 2) if (budget or abs(used) > 0.005) else E,
            (f"{round(used / budget * 100)}%" if budget > 0 else E),
            contractor, E,
            _nz(monthly),
        ]

    def _traverse(cat, direction_name="", type_name=""):
        nonlocal seq
        if cat.level == 1:
            direction_name = cat.name
            _write_row(
                [cat.code or "", cat.name, E, E, E, E] + _money_cols(cat.id, ""),
                L1_FILL, L1_FONT, height=22
            )
        elif cat.level == 2:
            type_name = cat.name
            _write_row(
                [E, direction_name, cat.name, E, E, E] + _money_cols(cat.id, ""),
                L2_FILL, L2_FONT
            )
        elif cat.level == 3:
            _write_row(
                [E, direction_name, type_name, cat.name, E, E]
                + _money_cols(cat.id, cat_contractor_map.get(cat.id, "")),
                L3_FILL, L3_FONT
            )
            # Под-строки: закупки, привязанные напрямую к категории (без плановой статьи).
            # По каждой фактической позиции — цена, кол-во, сумма, контрагент, № акта,
            # ссылка на документы + каскад Заказано/Поставлено/Оплачено.
            cat_purchases = purchased_by_cat.get(cat.id, [])
            cat_ord = cat_del = cat_paid = 0.0
            for pi in cat_purchases:
                price_note = f"    └ {pi['name']} — {round(pi['unit_price'], 2):,.2f} ₽/ед."
                status_txt = pi["status"]
                monthly_val = ""
                if pi.get("is_monthly"):
                    cnt = pi.get("monthly_count")
                    amt = pi.get("monthly_amount") or 0
                    status_txt = f"{status_txt} · ежемес." + (f" ×{cnt}" if cnt else "")
                    monthly_val = round(pi["total"], 2)
                    if amt:
                        price_note += f" ({round(amt, 2):,.2f} ₽/мес)"
                ord_a, del_a, paid_a = _cascade(pi["raw_status"], pi["total"])
                cat_ord += ord_a; cat_del += del_a; cat_paid += paid_a
                _write_row(
                    [
                        "", "", "", price_note,
                        pi["unit"], round(pi["qty"], 3),
                        "", "", "",
                        ord_a, del_a, paid_a,
                        round(pi["total"], 2), "",
                        "", pi["contractor"], status_txt,
                        monthly_val, pi["act_number"],
                    ],
                    SUB_FILL, SUB_FONT, height=16
                )
                _set_doc_link(pi["purchase_id"])
            if cat_purchases:
                _write_row(
                    [
                        "", "", "", "    Итого по факту:",
                        "", "",
                        "", "", "",
                        round(cat_ord, 2), round(cat_del, 2), round(cat_paid, 2),
                        "", "", "", "", "", "",
                    ],
                    SUB_FILL, Font(size=8, bold=True, color="374151"), height=16
                )
            for item in items_by_cat.get(cat.id, []):
                seq += 1
                feo_budget = float(item.amount or 0)
                plan_sum = _st(item.id, "plan_schedule", "work_in_progress", "wishes")
                contract_sum = _st(item.id, "contracted")
                ordered_sum = _st(item.id, "ordered")
                delivered_sum = _st(item.id, "delivered")
                paid_sum = _st(item.id, "paid")
                used = used_map.get(item.id, 0.0)
                residual = feo_budget - used
                pct = round(used / feo_budget * 100) if feo_budget > 0 else 0
                contractor = contractor_map.get(item.id, "")
                status = "Выполнено" if pct >= 100 else ("В работе" if used > 0 else "Не начато")
                font = RED_FONT if used > feo_budget else ITEM_FONT
                _write_row(
                    [
                        seq, direction_name, type_name, item.name,
                        item.unit or "", float(item.quantity or 0),
                        round(feo_budget, 2), round(plan_sum, 2),
                        round(contract_sum, 2), round(ordered_sum, 2),
                        round(delivered_sum, 2), round(paid_sum, 2),
                        round(used, 2), round(residual, 2),
                        f"{pct}%", contractor, status,
                    ],
                    PatternFill(), font
                )
                # Под-строки: реально закупленные позиции (что, по какой цене, № акта,
                # каскад статусов Заказано/Поставлено/Оплачено + ссылка на документы)
                for pi in purchased_by_item.get(item.id, []):
                    price_note = f"    └ {pi['name']} — {round(pi['unit_price'], 2):,.2f} ₽/ед."
                    ord_a, del_a, paid_a = _cascade(pi["raw_status"], pi["total"])
                    _write_row(
                        [
                            "", "", "", price_note,
                            pi["unit"], round(pi["qty"], 3),
                            "", "", "",
                            ord_a, del_a, paid_a,
                            round(pi["total"], 2), "",
                            "", pi["contractor"], pi["status"],
                            "", pi["act_number"],
                        ],
                        SUB_FILL, SUB_FONT, height=16
                    )
                    _set_doc_link(pi["purchase_id"])

        for child in cats_by_parent.get(cat.id, []):
            _traverse(child, direction_name, type_name)

    for root in cats_by_parent.get(None, []):
        _traverse(root)

    # ── Секция: закупки, привязанные к субсидии, но без категории ФЭО ────────
    if unlinked_purchases:
        _write_row(
            ["", "Закупки по субсидии без привязки к категории ФЭО", "", "", "", ""]
            + [""] * 12,
            L1_FILL, L1_FONT, height=22
        )
        u_ord = u_del = u_paid = 0.0
        for pi in unlinked_purchases:
            price_note = f"    └ {pi['name']} — {round(pi['unit_price'], 2):,.2f} ₽/ед."
            status_txt = pi["status"]
            monthly_val = ""
            if pi.get("is_monthly"):
                cnt = pi.get("monthly_count")
                amt = pi.get("monthly_amount") or 0
                status_txt = f"{status_txt} · ежемес." + (f" ×{cnt}" if cnt else "")
                monthly_val = round(pi["total"], 2)
                if amt:
                    price_note += f" ({round(amt, 2):,.2f} ₽/мес)"
            ord_a, del_a, paid_a = _cascade(pi["raw_status"], pi["total"])
            u_ord += ord_a; u_del += del_a; u_paid += paid_a
            _write_row(
                [
                    "", "", "", price_note,
                    pi["unit"], round(pi["qty"], 3),
                    "", "", "",
                    ord_a, del_a, paid_a,
                    round(pi["total"], 2), "",
                    "", pi["contractor"], status_txt,
                    monthly_val, pi["act_number"],
                ],
                SUB_FILL, SUB_FONT, height=16
            )
            _set_doc_link(pi["purchase_id"])
        _write_row(
            [
                "", "", "", "    Итого по факту:",
                "", "",
                "", "", "",
                round(u_ord, 2), round(u_del, 2), round(u_paid, 2),
                "", "", "", "", "", "",
            ],
            SUB_FILL, Font(size=8, bold=True, color="374151"), height=16
        )

    # ── Лист «Сводная»: деньги по Товары/Услуги × корзины обязательств ────────
    sm = wb.create_sheet("Сводная", 0)
    SUM_HEADERS = [
        "Категория",
        "Оплаченные+поставленные (треб. оплата), ₽",
        "Оплаченные, ₽",
        "Принятые, не оплаченные, ₽",
        "Планируемые, ₽",
        "Ежемесячные оплаты, ₽",
        "Скорее всего понадобятся, ₽",
        "Всего, ₽",
    ]
    SUM_WIDTHS = [16, 28, 18, 24, 20, 22, 26, 20]
    sm_title = sm.cell(row=1, column=1, value=f"СВОДНАЯ — {sub.name} ({sub.year})")
    sm_title.font = Font(bold=True, size=12, color="1E3A5F")
    sm.merge_cells(f"A1:{chr(64 + len(SUM_HEADERS))}1")
    sm_title.alignment = CENTER_ALIGN
    sm.row_dimensions[1].height = 28
    for ci, (h, w) in enumerate(zip(SUM_HEADERS, SUM_WIDTHS), 1):
        cell = sm.cell(row=2, column=ci, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        sm.column_dimensions[cell.column_letter].width = w
    sm.row_dimensions[2].height = 34

    _order = ["delivered", "paid", "accepted_unpaid", "planned", "monthly", "likely", "total"]

    def _sum_row(sm_row, label, b, fill, font):
        sm.cell(row=sm_row, column=1, value=label)
        for j, k in enumerate(_order, 2):
            c = sm.cell(row=sm_row, column=j, value=round(b[k], 2) if b[k] else "")
            c.number_format = "#,##0.00"
            c.alignment = RIGHT_ALIGN
            c.border = THIN_BORDER
        lc = sm.cell(row=sm_row, column=1)
        lc.fill = fill; lc.font = font; lc.border = THIN_BORDER
        for j in range(2, 9):
            sm.cell(row=sm_row, column=j).fill = fill

    _sum_row(3, "Услуги", summary["услуга"], L2_FILL, L2_FONT)
    _sum_row(4, "Товары", summary["товар"], L3_FILL, L3_FONT)
    empty_buckets_keys = ("paid", "delivered", "accepted_unpaid", "planned", "monthly", "likely", "total")
    total_b = {k: summary["услуга"][k] + summary["товар"][k] for k in empty_buckets_keys}
    _sum_row(5, "ИТОГО", total_b, L1_FILL, Font(bold=True, size=10))
    sm.freeze_panes = "A3"

    return wb
