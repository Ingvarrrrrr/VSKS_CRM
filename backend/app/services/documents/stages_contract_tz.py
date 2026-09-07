"""generate_document: append a ТЗ (items) table after a page break for
doc_type == "contract".

Split out of app/routers/documents.py::generate_document (wave 3f refactor).
Failure here is non-fatal by design (original just logs and returns buf
unchanged) — preserved exactly.
"""
from io import BytesIO

from app.services.documents.formatting import _fmt_date, _fmt_money_plain


def append_tz_table_for_contract(buf: BytesIO, doc_type: str, items_list: list, p) -> BytesIO:
    # For contract docs: append ТЗ table with items
    if doc_type == "contract" and items_list:
        try:
            from docx import Document as _DocxDoc
            from docx.shared import Pt, Cm, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement

            doc = _DocxDoc(buf)

            # Add page break before ТЗ
            doc.add_page_break()

            # Title
            title_para = doc.add_paragraph()
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title_para.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ")
            run.bold = True
            run.font.size = Pt(12)

            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.add_run(f"(Приложение к договору № {p.contract_number or '___'} от {_fmt_date(p.contract_date) or '__.__.____'})")

            doc.add_paragraph()  # spacer

            # Table with columns: №, Фото, Наименование и описание, Кол-во, Ед., Цена ед., Сумма
            table = doc.add_table(rows=1, cols=7)
            table.style = 'Table Grid'

            # Header row
            hdr_cells = table.rows[0].cells
            headers_tz = ["№", "Фото", "Наименование и описание", "Кол-во", "Ед.", "Цена ед., ₽", "Сумма, ₽"]
            col_widths_cm = [1.0, 2.5, 8.0, 1.8, 1.5, 2.8, 2.8]
            for i, (hdr_cell, hdr_text, w) in enumerate(zip(hdr_cells, headers_tz, col_widths_cm)):
                hdr_cell.width = Cm(w)
                para = hdr_cell.paragraphs[0]
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(hdr_text)
                run.bold = True
                run.font.size = Pt(9)
                # Blue background
                tc = hdr_cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), '1E40AF')
                shd.set(qn('w:color'), 'FFFFFF')
                shd.set(qn('w:val'), 'clear')
                tcPr.append(shd)
                # White font
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # Data rows
            total_sum = 0.0
            for item_data in items_list:
                row_cells = table.add_row().cells
                row_cells[0].text = str(item_data["num"])
                row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Photo cell: leave empty
                row_cells[1].text = ""
                # Name + description
                name_cell = row_cells[2]
                name_para = name_cell.paragraphs[0]
                name_run = name_para.add_run(str(item_data["name"]))
                name_run.bold = True
                name_run.font.size = Pt(9)
                if item_data.get("description"):
                    desc_para = name_cell.add_paragraph()
                    desc_run = desc_para.add_run(str(item_data["description"]))
                    desc_run.font.size = Pt(8)
                    desc_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

                qty_val = item_data.get("quantity", "")
                row_cells[3].text = str(qty_val) if qty_val != "" else "—"
                row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row_cells[4].text = str(item_data.get("unit", "") or "—")
                row_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row_cells[5].text = str(item_data.get("unit_price", "") or "—")
                row_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                row_cells[6].text = str(item_data.get("total_price", "") or "—")
                row_cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

                # Font size for all data cells
                for cell in row_cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if not run.font.size:
                                run.font.size = Pt(9)

                try:
                    price_str = str(item_data.get("total_price", "")).replace(" ", "").replace(",", ".").replace("₽", "").strip()
                    total_sum += float(price_str) if price_str else 0
                except Exception:
                    pass

            # Total row
            total_row_cells = table.add_row().cells
            # Merge cells 0-5
            total_row_cells[5].merge(total_row_cells[0])
            merged_para = total_row_cells[0].paragraphs[0]
            merged_run = merged_para.add_run("НМЦК итого:")
            merged_run.bold = True
            merged_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            total_row_cells[6].text = _fmt_money_plain(total_sum)
            if total_row_cells[6].paragraphs[0].runs:
                total_row_cells[6].paragraphs[0].runs[0].bold = True
            total_row_cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

            # Save back to buf
            buf2 = BytesIO()
            doc.save(buf2)
            buf2.seek(0)
            buf = buf2
        except Exception as tz_err:
            # Don't fail the whole request if ТЗ append fails — just log
            import traceback
            print(f"ТЗ append error: {tz_err}\n{traceback.format_exc()}")

    return buf
