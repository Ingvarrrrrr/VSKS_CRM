"""Импорт договоров из xlsx/xls/docx/pdf (включая OCR сканов) — preview + mapped.

Вынесено из app/routers/contracts.py (Правило №5, резка волны 2026-09-08).
Тот же префикс /api/contracts, оба пути POST "/import/preview" и
"/import/mapped" — литеральные двухсегментные, не пересекаются по форме ни с
одним путём ядра (PUT/DELETE "/{cid}") — порядок регистрации относительно
contracts.router не важен (см. комментарий в routes.py).
"""
from datetime import date as date_type
from decimal import Decimal
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth.permissions import require_tab
from app.database import get_db
from app.models.contract import Contract
from app.models.contractor import Contractor

try:
    from openpyxl import load_workbook as _load_workbook
except ImportError:
    _load_workbook = None

try:
    import pdfplumber as _pdfplumber
except ImportError:
    _pdfplumber = None

router = APIRouter(prefix="/api/contracts", tags=["contracts"])


# ── OCR helper ─────────────────────────────────────────────────────────────────

def _ocr_pdf_to_rows(content: bytes) -> tuple[list, str | None]:
    """Fallback: convert scanned PDF pages to images, run OCR, parse lines.
    Returns (rows, error_message). error_message is None on success."""
    try:
        from pdf2image import convert_from_bytes
        import pytesseract
    except ImportError:
        return [], "OCR-библиотеки не установлены (pdf2image, pytesseract). Обратитесь к администратору."

    try:
        images = convert_from_bytes(content, dpi=300)
    except Exception as e:
        return [], f"Не удалось преобразовать PDF в изображения для OCR: {e}"

    if not images:
        return [], "PDF не содержит страниц для распознавания."

    import re
    all_rows = []
    for img in images:
        try:
            text = pytesseract.image_to_string(img, lang='rus+eng')
        except Exception as e:
            return [], f"Ошибка OCR-распознавания: {e}"
        if not text:
            continue
        for line in text.strip().split('\n'):
            line = line.strip()
            if not line:
                continue
            # Split by tabs or multiple spaces
            cells = re.split(r'\t|  {2,}', line)
            cells = [c.strip() for c in cells if c.strip()]
            if cells:
                all_rows.append(cells)
    if not all_rows:
        return [], "OCR не нашёл текст на страницах. Возможно, качество скана слишком низкое."
    return all_rows, None


# ── File parsing helpers ───────────────────────────────────────────────────────

_CONTRACT_HINTS = (
    'номер', 'дата', 'контрагент', 'поставщик', 'исполнитель',
    'предмет', 'сумма', 'субсидия', 'статус', 'тип', 'метод',
    'начал', 'оконч', 'number', 'date', 'contractor', 'amount',
)


def _detect_contract_hdr(rows):
    best_score, best_idx = 0, 0
    for ri, row in enumerate(rows):
        norm = [str(h).strip().lower() if h is not None else "" for h in row]
        score = sum(1 for h in norm if h and any(x in h for x in _CONTRACT_HINTS))
        if score > best_score:
            best_score = score
            best_idx = ri
    return best_idx


def _parse_contract_file(fname: str, content: bytes):
    """Parse xlsx/xls/docx/pdf and return (all_rows, hdr_idx)."""
    fname = fname.lower()

    if fname.endswith('.xls') and not fname.endswith('.xlsx'):
        try:
            import xlrd
        except ImportError:
            raise HTTPException(500, "xlrd не установлен")
        try:
            wb = xlrd.open_workbook(file_contents=content)
        except Exception as e:
            raise HTTPException(400, f"Не удалось прочитать .xls: {e}")
        ws = wb.sheet_by_index(0)
        all_rows = [list(ws.row_values(i)) for i in range(ws.nrows)]

    elif fname.endswith('.xlsx'):
        if not _load_workbook:
            raise HTTPException(500, "openpyxl не установлен")
        try:
            wb = _load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise HTTPException(400, f"Не удалось прочитать .xlsx: {e}")
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

    elif fname.endswith(('.docx', '.doc')):
        try:
            from docx import Document
        except ImportError:
            raise HTTPException(500, "python-docx не установлен")
        try:
            doc = Document(BytesIO(content))
        except Exception as e:
            raise HTTPException(400, f"Не удалось прочитать .docx: {e}")
        if not doc.tables:
            raise HTTPException(400, "В документе не найдено таблиц")
        tbl = doc.tables[0]
        all_rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]

    elif fname.endswith('.pdf'):
        if not _pdfplumber:
            raise HTTPException(500, "pdfplumber не установлен")
        try:
            with _pdfplumber.open(BytesIO(content)) as pdf:
                all_rows = []
                has_text = False
                for page in pdf.pages:
                    tables = page.extract_tables()
                    if tables:
                        for tbl in tables:
                            all_rows.extend([r for r in tbl if r])
                if not all_rows:
                    for page in _pdfplumber.open(BytesIO(content)).pages:
                        text = page.extract_text()
                        if text:
                            has_text = True
                            for line in text.strip().split('\n'):
                                cells = [c.strip() for c in line.split('\t')]
                                if len(cells) < 2:
                                    cells = [c.strip() for c in line.split('  ') if c.strip()]
                                if cells:
                                    all_rows.append(cells)
        except Exception as e:
            raise HTTPException(400, f"Не удалось прочитать PDF-файл: {e}")
        if not all_rows:
            if not has_text:
                # Scanned PDF — try OCR
                all_rows, ocr_error = _ocr_pdf_to_rows(content)
                if not all_rows:
                    detail = ocr_error or "OCR не смог распознать таблицу."
                    raise HTTPException(
                        400,
                        f"Этот PDF — скан (изображение). {detail} "
                        "Попробуйте сохранить данные в Excel (.xlsx) или Word (.docx)."
                    )
            else:
                raise HTTPException(
                    400,
                    "В PDF найден текст, но не удалось распознать таблицу. "
                    "Попробуйте сохранить данные в Excel (.xlsx) и загрузить его."
                )
    else:
        raise HTTPException(400, "Неподдерживаемый формат. Используйте .xlsx, .xls, .docx, .doc или .pdf")

    if not all_rows:
        raise HTTPException(400, "Файл пустой или не содержит данных")

    hdr_idx = _detect_contract_hdr(all_rows)
    return all_rows, hdr_idx


def _parse_date(val) -> Optional[date_type]:
    if not val:
        return None
    if isinstance(val, date_type):
        return val
    s = str(val).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            from datetime import datetime
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_amount(val) -> Optional[Decimal]:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    s = str(val).strip().replace(' ', '').replace('\xa0', '').replace(',', '.')
    s = ''.join(c for c in s if c.isdigit() or c == '.')
    if not s:
        return None
    try:
        return Decimal(s)
    except Exception:
        return None


# ── Import endpoints ──────────────────────────────────────────────────────────

@router.post("/import/preview")
async def contracts_import_preview(
    file: UploadFile = File(...),
    current_user=Depends(require_tab('contracts')),
):
    """Parse uploaded file and return headers + sample rows for column mapping UI."""
    fname = (file.filename or '').lower()
    content = await file.read()
    all_rows, hdr_idx = _parse_contract_file(fname, content)
    hdr_rows = all_rows[hdr_idx:]
    if not hdr_rows:
        raise HTTPException(400, "Файл пустой")

    headers = [
        str(c).strip() if c else f"Столбец {j + 1}"
        for j, c in enumerate(hdr_rows[0])
    ]
    sample = [
        [str(c).strip() if c is not None else "" for c in row]
        for row in hdr_rows[1:min(6, len(hdr_rows))]
    ]
    return {
        "headers": headers,
        "sample": sample,
        "total_rows": max(len(all_rows) - hdr_idx - 1, 0),
        "header_row_offset": hdr_idx,
    }


@router.post("/import/mapped")
async def contracts_import_mapped(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_tab('contracts')),
    col_number: Optional[int] = Query(None),
    col_date: Optional[int] = Query(None),
    col_contractor: Optional[int] = Query(None),
    col_subject: Optional[int] = Query(None),
    col_max_amount: Optional[int] = Query(None),
    col_start_date: Optional[int] = Query(None),
    col_end_date: Optional[int] = Query(None),
    col_contract_type: Optional[int] = Query(None),
    col_purchase_method: Optional[int] = Query(None),
    col_item_type: Optional[int] = Query(None),
    col_status: Optional[int] = Query(None),
    col_notes: Optional[int] = Query(None),
    header_row_offset: int = Query(0),
    subsidy_id: Optional[int] = Query(None),
):
    """Import contracts using user-specified column mapping."""
    if col_number is None:
        raise HTTPException(400, "Не указан столбец «Номер договора»")

    fname = (file.filename or '').lower()
    content = await file.read()
    all_rows, _ = _parse_contract_file(fname, content)
    data_rows = all_rows[header_row_offset + 1:]

    col_map = {
        'number': col_number, 'date': col_date, 'contractor': col_contractor,
        'subject': col_subject, 'max_amount': col_max_amount,
        'start_date': col_start_date, 'end_date': col_end_date,
        'contract_type': col_contract_type, 'purchase_method': col_purchase_method,
        'item_type': col_item_type, 'status': col_status, 'notes': col_notes,
    }

    def _get(row, field):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        return str(v).strip() if v is not None else None

    # Pre-load contractors for matching by name/INN
    ctr_result = await db.execute(select(Contractor))
    all_contractors = ctr_result.scalars().all()
    ctr_by_name = {c.name.strip().lower(): c.id for c in all_contractors if c.name}
    ctr_by_inn = {c.inn.strip(): c.id for c in all_contractors if c.inn}

    created, skipped, errors = 0, 0, []
    for ri, row in enumerate(data_rows, start=1):
        num = _get(row, 'number')
        if not num:
            continue

        # Check duplicate by number
        dup = (await db.execute(
            select(Contract.id).where(Contract.number == num).limit(1)
        )).scalar_one_or_none()
        if dup:
            skipped += 1
            continue

        # Resolve contractor
        contractor_id = None
        ctr_val = _get(row, 'contractor')
        if ctr_val:
            ctr_lower = ctr_val.lower().strip()
            contractor_id = ctr_by_name.get(ctr_lower)
            if not contractor_id:
                contractor_id = ctr_by_inn.get(ctr_val.strip())

        # Resolve contract_type
        ct_raw = (_get(row, 'contract_type') or '').lower()
        if any(x in ct_raw for x in ('рамоч', 'накопит', 'framework')):
            contract_type = 'framework_cumulative'
        elif any(x in ct_raw for x in ('разов', 'единич', 'single')):
            contract_type = 'single'
        else:
            contract_type = 'single'

        c = Contract(
            number=num,
            date=_parse_date(_get(row, 'date')),
            contract_type=contract_type,
            contractor_id=contractor_id,
            subsidy_id=subsidy_id,
            subject=_get(row, 'subject'),
            max_amount=_parse_amount(_get(row, 'max_amount')),
            start_date=_parse_date(_get(row, 'start_date')),
            end_date=_parse_date(_get(row, 'end_date')),
            purchase_method=_get(row, 'purchase_method'),
            item_type=_get(row, 'item_type'),
            status=_get(row, 'status') or 'active',
            notes=_get(row, 'notes'),
        )
        db.add(c)
        created += 1

    if created:
        await db.commit()

    return {"created": created, "skipped": skipped, "errors": errors}
