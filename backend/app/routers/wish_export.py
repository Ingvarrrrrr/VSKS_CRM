"""GET /{wish_id}/export.xlsx — выгрузка позиций заявки в Excel (с фото или без).

Вынесено из app/routers/wishes.py (Правило №5, модульность; сессия 2026-09-06,
разрезание wishes.py по образцу purchases.py → purchase_ops.py/purchase_*.py)
БЕЗ ИЗМЕНЕНИЯ ПОВЕДЕНИЯ.

Хелперы ядра вызываются через `wishes_core.<имя>` — см. докстринг wish_transitions.py.
"""
import re as _re

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.auth.jwt import get_current_user, get_org_filter
from app.models.user import User
from app.models.wish_item import WishItem
from app.routers import wishes as wishes_core

router = APIRouter(prefix="/api/wishes", tags=["wishes"])


@router.get("/{wish_id}/export.xlsx")
async def export_wish_xlsx(
    wish_id: int,
    with_photos: bool = True,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    import httpx
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.comments import Comment
    from fastapi.responses import StreamingResponse
    import io
    from urllib.parse import quote
    from PIL import Image as PILImage
    try:
        from pillow_heif import register_heif_opener
        register_heif_opener()
    except Exception:
        pass

    wish = await wishes_core._load_wish(wish_id, db)
    org_ids = get_org_filter(current_user)
    if org_ids is not None and wish.org_id not in org_ids:
        raise HTTPException(status_code=403, detail="Нет доступа к этой заявке")
    # Позиции заявки + связанный товар (фото/категория/вид/описание/ссылки)
    res = await db.execute(
        select(WishItem)
        .where(WishItem.wish_id == wish_id)
        .options(selectinload(WishItem.product))
        .order_by(WishItem.id)
    )
    items = res.scalars().all()

    async def _image_bytes(p):
        if p is not None and getattr(p, "photo_data", None):
            return bytes(p.photo_data)
        url = (p.photo_url or p.photo_link) if p is not None else None
        if url and str(url).startswith("http"):
            try:
                async with httpx.AsyncClient(timeout=8.0, follow_redirects=True) as cl:
                    resp = await cl.get(url)
                    if resp.status_code == 200 and resp.content:
                        return resp.content
            except Exception:
                return None
        return None

    def _thumb_png(raw: bytes):
        try:
            im = PILImage.open(BytesIO(raw))
            if im.mode not in ("RGB", "RGBA"):
                im = im.convert("RGB")
            im.thumbnail((90, 90))
            out = BytesIO()
            im.save(out, format="PNG")
            out.seek(0)
            return out, im.width, im.height
        except Exception:
            return None

    from openpyxl.styles import Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Заявка"
    thin = Side(style="thin", color="C0C0C0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="FB923C")
    header_font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    link_font = Font(color="0563C1", underline="single")

    if with_photos:
        headers = [
            "№ п/п", "Наименование", "Фото", "Категория товара", "Вид",
            "Описание", "Ссылка пример", "Количество", "Ед. изм.",
            "Плановая Цена за ед", "Плановая Сумма",
        ]
        widths = [7, 34, 14, 20, 16, 50, 32, 12, 10, 16, 16]
    else:
        headers = [
            "№ п/п", "Наименование", "Категория товара", "Вид",
            "Описание", "Ссылка пример", "Количество", "Ед. изм.",
            "Плановая Цена за ед", "Плановая Сумма",
        ]
        widths = [7, 34, 20, 16, 50, 32, 12, 10, 16, 16]
    ncols = len(headers)
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = center; cell.border = border
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 32

    # Позиции колонок (без фото всё после "Наименование" сдвинуто влево)
    col_name = 2
    col_photo = 3 if with_photos else None
    base = 4 if with_photos else 3  # "Категория товара" и далее
    col_cat = base
    col_kind = base + 1
    col_descr = base + 2
    col_example = base + 3
    col_qty = base + 4
    col_unit = base + 5
    col_price = base + 6
    col_total = base + 7

    if with_photos:
        ws.cell(row=1, column=col_photo).comment = Comment(
            "Изображения встроены в файл и отображаются всегда — настройка безопасности Excel не требуется. "
            "Оригинальные ссылки на фото — на скрытом листе «Ссылки» (правый клик по ярлыку листа → Показать).",
            "GALA")

    def _photo_url(p):
        if not p:
            return ""
        return p.photo_url or p.photo_link or (f"/api/products/{p.id}/photo" if p.photo_data else "")

    def _example_url(p):
        if not p:
            return ""
        links = p.price_links if isinstance(p.price_links, list) else []
        for l in links:
            url = l.get("url") if isinstance(l, dict) else None
            if url and str(url).startswith("http"):
                return url
        cl = p.clarification_link or ""
        return cl if cl.startswith("http") else ""

    # Скрытый лист со ссылками — только в режиме с фото
    if with_photos:
        ws_links = wb.create_sheet("Ссылки")
        ws_links.append(["№ п/п", "Наименование", "Фото URL"])

    r = 2
    for i, it in enumerate(items, start=1):
        p = getattr(it, "product", None)
        name = it.item_name or (p.name if p else "") or ""
        category = (p.category if p else "") or ""
        kind = (p.product_type if p else "") or ""
        descr = (p.description if p else "") or ""
        photo = _photo_url(p)
        example = _example_url(p)
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=col_name, value=name)
        if with_photos:
            raw = await _image_bytes(p)
            if raw:
                t = _thumb_png(raw)
                if t:
                    bio, w_px, h_px = t
                    img = XLImage(bio)
                    img.width = w_px; img.height = h_px
                    ws.add_image(img, f"{get_column_letter(col_photo)}{r}")
                    ws.row_dimensions[r].height = max(
                        ws.row_dimensions[r].height or 0, h_px * 0.78 + 6
                    )
        ws.cell(row=r, column=col_cat, value=category)
        ws.cell(row=r, column=col_kind, value=kind)
        ws.cell(row=r, column=col_descr, value=descr)
        cex = ws.cell(row=r, column=col_example, value=example)
        if example.startswith("http"):
            cex.hyperlink = example; cex.font = link_font
        qty = float(it.quantity or 0)
        unit_price = float(it.unit_price or 0)
        sum_val = float(it.total_price or 0) or (qty * unit_price)
        ws.cell(row=r, column=col_qty, value=qty)
        ws.cell(row=r, column=col_unit, value=it.unit or "")
        price = ws.cell(row=r, column=col_price, value=unit_price)
        total = ws.cell(row=r, column=col_total, value=sum_val)
        price.number_format = '# ##0.00'
        total.number_format = '# ##0.00'
        left_cols = {col_name, col_descr}
        leftcenter_cols = {col_example}
        if with_photos:
            leftcenter_cols.add(col_photo)
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if c in left_cols:
                cell.alignment = left_top
            elif c in leftcenter_cols:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if with_photos:
            ws_links.append([i, name, photo or ""])
        r += 1

    if with_photos:
        ws_links.sheet_state = "hidden"
    last_row = max(r - 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last_row}"
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    _wish_title_raw = (wish.title or "").strip()
    _wish_title_clean = _re.sub(r'[\\/:*?"<>|\r\n]+', "", _wish_title_raw)
    _wish_title_clean = _re.sub(r'\s+', "_", _wish_title_clean)[:50]
    fname = f"Заявка_{_wish_title_clean}_{wish.id}.xlsx" if _wish_title_clean else f"Заявка_{wish.id}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname, safe='-_.~')}"},
    )
