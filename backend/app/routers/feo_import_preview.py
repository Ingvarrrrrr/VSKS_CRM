"""POST /api/feo-categories/import-preview — читает Excel/DOCX/XLS/PDF и
возвращает заголовки + примеры строк для ручного маппинга колонок (мастер
сопоставления перед /import-mapped).

Вынесено из app/routers/feo_import.py (Правило №5, разрезание 1088-строчного
роутера) без изменения поведения — самодостаточен (не вызывает _do_feo_import,
только читает файл и определяет заголовки). Путь статичный (без {cat_id}) —
регистрируется в app/routes.py рядом с feo_import.router и ДО
feo_categories.router (несёт catch-all GET/PUT/DELETE "/{cat_id}"), как и
остальные соседи feo_import_*.
"""
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from app.database import get_db
from app.auth.permissions import require_tab
from app.routers import feo_categories as fc

router = APIRouter(prefix="/api/feo-categories", tags=["feo_categories"])


@router.post("/import-preview")
async def feo_import_preview(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_tab('feo_categories')),
):
    """Read Excel/DOCX file and return headers + sample rows for column mapping."""
    # B2: файл ещё не связан с конкретной субсидией на этом шаге (маппинг колонок
    # выбирается ПОСЛЕ) — subsidy_id=None, право проверяется по любой доступной орге.
    await fc._require_feo_category_write(current_user, db, None)
    fname = (file.filename or "").lower()
    if not fname.endswith((".xlsx", ".xls", ".docx", ".doc", ".pdf")):
        raise HTTPException(400, "Поддерживаются файлы .xlsx, .xls, .docx, .pdf")

    content = await file.read()

    _FEO_HINTS = (
        "субсидия", "наименован", "направлен", "расходов", "уровень",
        "код", "финансирован", "количеств", "ед. изм", "ед.изм",
        "активн", "приложен", "бюджет", "плановый", "тип расх",
    )

    def _detect_hdr(rows):
        best_score, best_idx = 0, 0
        for ri, row in enumerate(rows[:20]):
            norm = [str(h).strip().lower() if h is not None else "" for h in row]
            score = sum(1 for h in norm if h and any(x in h for x in _FEO_HINTS))
            if score > best_score:
                best_score = score
                best_idx = ri
        return best_idx

    try:
        # ── PDF ──
        if fname.endswith(".pdf"):
            try:
                import pdfplumber
            except ImportError:
                raise HTTPException(500, "pdfplumber не установлен")
            pdf = pdfplumber.open(BytesIO(content))
            all_rows = []
            for page in pdf.pages:
                for t in (page.extract_tables() or []):
                    if t:
                        all_rows.extend([[str(c).strip() if c else "" for c in row] for row in t])
            pdf.close()
            if not all_rows:
                raise HTTPException(400, "Не удалось извлечь таблицы из PDF")
            hdr_idx = _detect_hdr(all_rows)
            headers = [str(h).strip() if h else f"Столбец {j+1}" for j, h in enumerate(all_rows[hdr_idx])]
            data = all_rows[hdr_idx + 1:]
            sample = [[str(c) if c else "" for c in r] for r in data[:5]]
            return {"sheets": [{"name": "PDF", "headers": headers, "sample": sample, "total_rows": len(data), "header_row_offset": hdr_idx}]}

        # ── DOCX ──
        if fname.endswith((".docx", ".doc")):
            try:
                from docx import Document as _DDoc
            except ImportError:
                raise HTTPException(500, "python-docx не установлен")
            doc = _DDoc(BytesIO(content))
            all_rows = []
            for table in doc.tables:
                for row in table.rows:
                    all_rows.append([cell.text.strip() for cell in row.cells])
            if not all_rows:
                for para in doc.paragraphs:
                    text = para.text.strip()
                    if text:
                        all_rows.append([text])
            if not all_rows:
                raise HTTPException(400, "Не удалось извлечь данные из документа")
            hdr_idx = _detect_hdr(all_rows)
            headers = [str(h).strip() if h else f"Столбец {j+1}" for j, h in enumerate(all_rows[hdr_idx])]
            data = all_rows[hdr_idx + 1:]
            sample = [[str(c) if c else "" for c in r] for r in data[:5]]
            return {"sheets": [{"name": "Document", "headers": headers, "sample": sample, "total_rows": len(data), "header_row_offset": hdr_idx}]}

        # ── XLS ──
        if fname.endswith(".xls"):
            try:
                import xlrd as _xlrd_mod
            except ImportError:
                raise HTTPException(500, "xlrd не установлен")
            wb_xls = _xlrd_mod.open_workbook(file_contents=content)
            sheets = []
            for sheet_name in wb_xls.sheet_names():
                ws_xls = wb_xls.sheet_by_name(sheet_name)
                all_rows = [list(ws_xls.row_values(i)) for i in range(ws_xls.nrows)]
                if not all_rows:
                    continue
                hdr_idx = _detect_hdr(all_rows)
                hdr_rows = all_rows[hdr_idx:]
                if not hdr_rows:
                    continue
                headers = [str(c).strip() if c else f"Столбец {j+1}" for j, c in enumerate(hdr_rows[0])]
                sample = [[str(c).strip() if c is not None else "" for c in row] for row in hdr_rows[1:min(6, len(hdr_rows))]]
                sheets.append({"name": sheet_name, "headers": headers, "sample": sample,
                               "total_rows": ws_xls.nrows - hdr_idx - 1, "header_row_offset": hdr_idx})

        # ── XLSX ──
        else:
            if load_workbook is None:
                raise HTTPException(500, "openpyxl не установлен")
            wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheets = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                all_rows = list(ws.iter_rows(values_only=True))
                if not all_rows:
                    continue
                hdr_idx = _detect_hdr(all_rows)
                hdr_rows = all_rows[hdr_idx:]
                if not hdr_rows:
                    continue
                headers = [str(c).strip() if c else f"Столбец {j+1}" for j, c in enumerate(hdr_rows[0])]
                sample = [[str(c).strip() if c is not None else "" for c in row] for row in hdr_rows[1:min(6, len(hdr_rows))]]
                sheets.append({"name": sheet_name, "headers": headers, "sample": sample,
                               "total_rows": len(all_rows) - hdr_idx - 1, "header_row_offset": hdr_idx})
            wb.close()

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Не удалось прочитать файл ({file.filename}): {e}")

    if not sheets:
        raise HTTPException(400, "Файл не содержит листов с данными")

    return {"sheets": sheets}
