"""Ручной проверочный скрипт (НЕ pytest, запускается напрямую python'ом внутри
контейнера) — прогоняет реальный файл «Абхазия ЦЭМАК» на ВРЕМЕННОЙ субсидии и
печатает числа для приёмки задачи 2026-09-16 (feo_amount/по ФЭО/предупреждения).

Раскладка колонок (0-based, дано владельцем):
  1 Ур.2, 2 Ур.3, 3 Ур.4, 4 Плановая позиция, 5 Товар/услуга,
  6-9 ФЭО кол-во/ед./цена/сумма, 10-13 план кол-во/ед./цена/сумма,
  14 Код, 15 Приложение, 16 Активна, 17 Финансирование(устар.).
  Строки данных 2..190 (Excel, 1-based) = openpyxl rows[1:190].

Временная субсидия создаётся и удаляется этим же скриптом.
"""
import asyncio
import sys

from openpyxl import load_workbook

sys.path.insert(0, "/app")

from sqlalchemy import select, text
from app.database import async_session as AsyncSessionLocal
from app.models.subsidy import Subsidy
from app.models.feo_category import FeoCategory
from app.models.feo_planned_item import FeoPlannedItem
from app.services.feo_import_core import _do_feo_import
from app.services.subsidy_budget import calculate_budget_from_categories

XLSX_PATH = "/tmp/abkhazia_cemak.xlsx"

COL = dict(
    lvl2=1, lvl3=2, lvl4=3, lvl5=4, item_type=5,
    feo_qty=6, feo_unit=7, feo_price=8, feo_sum=9,
    plan_qty=10, plan_unit=11, plan_price=12, plan_sum=13,
    code=14, appendix=15, active=16, budget=17,
)


async def main():
    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))
    data_rows = all_rows[1:190]  # строки 2..190 включительно (1-based Excel)
    print(f"Всего строк данных: {len(data_rows)}")

    async with AsyncSessionLocal() as db:
        sub = Subsidy(name="TestAbkhazia_2_verify", year=2026, budget=0, require_planned_dates=False)
        db.add(sub)
        await db.commit()
        await db.refresh(sub)
        sid = sub.id
        print(f"Временная субсидия id={sid}")

        try:
            result = await _do_feo_import(
                rows=data_rows,
                c_subsidy=None,
                c_lvl2=COL["lvl2"], c_lvl3=COL["lvl3"], c_lvl4=COL["lvl4"], c_lvl5=COL["lvl5"],
                c_qty=None, c_unit=None, c_item_amt=None,
                c_code=COL["code"], c_appendix=COL["appendix"], c_budget=COL["budget"], c_active=COL["active"],
                c_row_feo_qty=COL["feo_qty"], c_row_feo_unit=COL["feo_unit"],
                c_row_feo_price=COL["feo_price"], c_row_feo_sum=COL["feo_sum"],
                c_row_plan_qty=COL["plan_qty"], c_row_plan_unit=COL["plan_unit"],
                c_row_plan_price=COL["plan_price"], c_row_plan_sum=COL["plan_sum"],
                c_item_type=COL["item_type"],
                default_subsidy_id=sid,
                dry_run=False,
                db=db,
            )
            print("created/updated/skipped:", result["created"], result["updated"], result["skipped"])
            print("errors:", result["errors"][:5], "... total", len(result["errors"]))
            print("warnings by kind:")
            _counts = {}
            for w in result["warnings"]:
                _counts[w["kind"]] = _counts.get(w["kind"], 0) + 1
            for k, v in sorted(_counts.items()):
                print(f"  {k}: {v}")
            for w in result["warnings"]:
                if w["kind"] in ("level_column_empty_in_file", "code_column_holds_amounts", "item_name_equals_category"):
                    print(f"    -> {w['kind']}: {w['message']}")

            # Задача владельца 2026-09-16: категории, у которых есть И
            # собственная сумма, И собственные позиции с feo_amount —
            # по умолчанию побеждает собственная сумма (resolution='own'),
            # но группа должна быть видна человеку на предпросмотре.
            catsum_groups = result.get("category_sum_conflict_groups", [])
            print(f"\ncategory_sum_conflict_groups: {len(catsum_groups)}")
            for g in catsum_groups:
                print(f"  {g}")

            total_feo = await calculate_budget_from_categories(db, sid)
            print(f"\nПо ФЭО субсидии (calculate_budget_from_categories) = {total_feo:,.2f}".replace(",", " "))

            cats = (await db.execute(select(FeoCategory).where(FeoCategory.subsidy_id == sid))).scalars().all()
            by_name = {}
            for c in cats:
                by_name.setdefault(c.name, []).append(c)
            from app.services.subsidy_budget import _active_feo_items_with_amount, compute_budget_map
            items = await _active_feo_items_with_amount(db, [c.id for c in cats])
            bmap = compute_budget_map(cats, items)
            for name in ("Оборудование и снаряжение", "Катер"):
                for c in by_name.get(name, []):
                    print(f"по ФЭО «{name}» (id={c.id}, level={c.level}) = {bmap.get(c.id):,.2f}".replace(",", " "))
                    its = (await db.execute(select(FeoPlannedItem).where(FeoPlannedItem.feo_category_id == c.id))).scalars().all()
                    print(f"  items under «{name}»: {[(it.name, str(it.amount), str(it.feo_amount)) for it in its]}")

            for name in ("Спусковое устройство Венто",):
                found = []
                for c in cats:
                    its = (await db.execute(select(FeoPlannedItem).where(FeoPlannedItem.feo_category_id == c.id))).scalars().all()
                    for it in its:
                        if name in (it.name or ""):
                            found.append(it)
                for it in found:
                    print(
                        f"item «{it.name}»: qty={it.quantity} unit_price={it.unit_price} amount={it.amount} "
                        f"feo_qty={it.feo_quantity} feo_unit_price={it.feo_unit_price} feo_amount={it.feo_amount}"
                    )
        finally:
            await db.execute(text(
                "DELETE FROM feo_planned_items WHERE feo_category_id IN "
                "(SELECT id FROM feo_categories WHERE subsidy_id = :sid)"
            ), {"sid": sid})
            await db.execute(text("DELETE FROM feo_categories WHERE subsidy_id = :sid"), {"sid": sid})
            await db.execute(text("DELETE FROM subsidies WHERE id = :sid"), {"sid": sid})
            await db.commit()
            print(f"\nВременная субсидия id={sid} удалена")


if __name__ == "__main__":
    asyncio.run(main())
