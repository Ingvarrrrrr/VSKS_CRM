"""Рендер пакета документов для ЭТП «Росэлторг».

Пока один файл в пакете: xlsx «Позиции по лоту» по шаблону площадки
(backend/templates/roseltorg/lot_items_template.xlsx). По аналогии с
fabrikant_package.py — отдельный модуль (ПРАВИЛО №5: не дописывать в
fabrikant_package.py, у Росэлторга другая площадка и другой формат пакета).

ПРАВИЛО №6: сырые quantity/unit_price/description/name/unit позиции не
считаются здесь заново — _purchase_item_doc_values (contexts.py) — тот же
единственный расчёт, которым пользуется и _build_items_list_from_purchase_items
для docx-пакетов. Эта функция возвращает НЕФОРМАТИРОВАННЫЕ числа (Decimal/
float), которые нужны xlsx-ячейкам; build_tz_rows (app.services.tz_items) —
тот же единственный источник группировки/слияния дублей строк ТЗ.

ОКПД2 у позиции закупки/товара как отдельное поле в модели не хранится
(ни у Product, ни у PurchaseItem) — есть только okpd2_codes (справочник-
классификатор) и одноразовый payload["okpd2_code"] публикации на Фабрикант
(один код на ВСЮ закупку, не персистентный, см. app/routers/publications.py
и publications_fabrikant_soap.py::item.get("okpd2_code")). Для Росэлторга
колонка C «Код ОКПД2» заполняется пустой строкой на каждую позицию — то же
отсутствие источника, что и у Фабриканта, не второй придуманный расчёт.
"""

import os
import zipfile
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.purchase import Purchase
from app.models.purchase_item import PurchaseItem

from .contexts import _require_tz_duplicates_resolved_for_doc, _purchase_item_doc_values
from app.services.tz_items import build_tz_rows
from .doc_types import TEMPLATES_DIR

# TEMPLATES_DIR ("/app/templates" в контейнере backend, см. doc_types.py) —
# тот же единственный корень шаблонов, что у всех doc_type (ПРАВИЛО №6);
# backend/templates/roseltorg/... на хосте монтируется именно туда — не
# путь относительно __file__ (образ бэкенда не видит остальной репозиторий,
# см. project_backend_image_cannot_read_repo_root.md).
ROSELTORG_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "roseltorg", "lot_items_template.xlsx")

ROSELTORG_SHEET_NAME_INDEX = 0  # "Список позиций" — первый лист шаблона
ROSELTORG_FIRST_DATA_ROW = 4
ROSELTORG_MAX_ROWS = 500

ROSELTORG_ZIP_FILE_NAME = "Позиции_по_лоту.xlsx"


def build_roseltorg_lot_items_xlsx(rows: list[dict]) -> bytes:
    """Заполняет шаблон «Позиции по лоту» Росэлторга и возвращает xlsx-байты.

    rows: [{okpd2, name, quantity, unit, unit_price, description}, ...] —
    сырые значения (не отформатированные под docx-шаблоны).

    Чистая функция, без обращения к БД — тестируется без фикстур (см.
    backend/tests/test_roseltorg_package.py).
    """
    if len(rows) > ROSELTORG_MAX_ROWS:
        raise ValueError(
            f"Превышен лимит Росэлторга: максимум {ROSELTORG_MAX_ROWS} позиций "
            f"в лоте, передано {len(rows)}"
        )

    import openpyxl

    wb = openpyxl.load_workbook(ROSELTORG_TEMPLATE_PATH, data_only=False)
    ws = wb.worksheets[ROSELTORG_SHEET_NAME_INDEX]

    for i, row in enumerate(rows):
        r = ROSELTORG_FIRST_DATA_ROW + i

        okpd2 = str(row.get("okpd2") or "").strip()
        ws.cell(row=r, column=3, value=okpd2)  # C — Код ОКПД2

        name = str(row.get("name") or "").strip()[:250]
        ws.cell(row=r, column=4, value=name)  # D — Наименование*

        quantity = row.get("quantity")
        try:
            quantity = round(float(quantity), 2) if quantity is not None else 0.0
        except (TypeError, ValueError):
            quantity = 0.0
        if quantity <= 0:
            quantity = 0.0
        ws.cell(row=r, column=5, value=quantity)  # E — Количество*

        unit = str(row.get("unit") or "").strip()[:20] or "Условная единица"
        ws.cell(row=r, column=6, value=unit)  # F — Ед.измерения*

        unit_price = row.get("unit_price")
        try:
            unit_price = round(float(unit_price), 2) if unit_price is not None else 0.0
        except (TypeError, ValueError):
            unit_price = 0.0
        ws.cell(row=r, column=7, value=unit_price)  # G — Цена за единицу

        description = str(row.get("description") or "").strip()[:1000]
        ws.cell(row=r, column=8, value=description)  # H — Описание

        # Колонка A (ID) — оставляем пустой, колонку B (нумерация) не трогаем:
        # там формула шаблона (=IF(D{r}="", "", COUNTA($D4:D{r}))).

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def render_roseltorg_package_files(
    db: AsyncSession,
    purchase_id: int,
    current_user=None,
) -> tuple[list[tuple[str, bytes]], list[str]]:
    """Собирает «Позиции_по_лоту.xlsx» для покупки *purchase_id*.

    Returns:
        files:  list of (file_name, bytes) — пока ровно один файл при успехе.
        errors: list of human-readable error strings.
    """
    result = await db.execute(
        select(Purchase)
        .options(selectinload(Purchase.items).selectinload(PurchaseItem.product))
        .where(Purchase.id == purchase_id)
    )
    p = result.scalar_one_or_none()
    if not p:
        return [], [f"Закупка {purchase_id} не найдена"]

    # Тот же гейт, что у пакета Фабрикант (ПРАВИЛО №6: единственная проверка
    # нерешённых дублей строк ТЗ, app.services.documents.contexts) — пакет
    # Росэлторга тоже печатает purchase_items (план), до выбора поставщика.
    _require_tz_duplicates_resolved_for_doc(p)

    _tz = build_tz_rows(getattr(p, "items", None) or [], getattr(p, "tz_duplicate_decisions", None))

    rows: list[dict] = []
    errors: list[str] = []
    for row in _tz["rows"]:
        item = row["item"]
        try:
            raw = _purchase_item_doc_values(p, item, row)
        except Exception as e:
            errors.append(f"Позиция «{getattr(item, 'item_name', '') or item.id}»: {type(e).__name__}: {e}")
            continue
        rows.append({
            # ОКПД2 у позиции не хранится (см. докстринг модуля) — пустая
            # строка, не придуманное значение.
            "okpd2": "",
            "name": raw["name"],
            "quantity": raw["quantity"],
            "unit": raw["unit"],
            "unit_price": raw["unit_price"],
            "description": raw["description"],
        })

    if not rows:
        return [], errors or [f"У закупки {purchase_id} нет позиций для пакета Росэлторг"]

    xlsx_bytes = build_roseltorg_lot_items_xlsx(rows)
    return [(ROSELTORG_ZIP_FILE_NAME, xlsx_bytes)], errors


def build_roseltorg_zip_bytes(
    files: list[tuple[str, bytes]],
    errors: list[str],
) -> bytes:
    """Собирает ZIP пакета Росэлторг из результата render_roseltorg_package_files.

    По образцу build_fabrikant_zip_bytes (fabrikant_package.py) — errors.txt
    в архиве только если были ошибки.
    """
    zip_buf = BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for file_name, data in files:
            zf.writestr(file_name, data)
        if errors:
            zf.writestr("errors.txt", "\n\n".join(errors))
    return zip_buf.getvalue()
