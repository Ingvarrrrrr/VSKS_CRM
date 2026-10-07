"""app.utils.xlsx_row_height — высота строки по тексту (Задача C экспорта
плана-графика, владелец 07.10.2026). Чистые unit-тесты на openpyxl Workbook,
без БД."""
import openpyxl
import pytest

from app.utils.xlsx_row_height import apply_row_heights, row_height_for_row, wrapped_line_count


def test_wrapped_line_count_short_text_fits_one_line():
    assert wrapped_line_count("Короткий текст", col_width=40, font_size=9) == 1


def test_wrapped_line_count_empty_is_one_line():
    assert wrapped_line_count("", col_width=20) == 1
    assert wrapped_line_count(None, col_width=20) == 1


def test_wrapped_line_count_long_text_wraps_multiple_lines():
    long_text = "слово " * 40  # ~240 символов
    lines_wide = wrapped_line_count(long_text, col_width=40, font_size=9)
    lines_narrow = wrapped_line_count(long_text, col_width=10, font_size=9)
    assert lines_wide >= 2
    # уже столбец → больше строк переноса на тот же текст.
    assert lines_narrow > lines_wide


def test_wrapped_line_count_respects_explicit_newlines():
    text = "Первая строка\nВторая строка\nТретья строка"
    assert wrapped_line_count(text, col_width=100, font_size=9) == 3


def test_row_height_for_row_long_text_exceeds_minimum():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["A"].width = 15
    long_text = "Статья ФЭО с очень длинным названием, которое точно не влезет в одну строку при ширине колонки 15"
    ws.cell(row=1, column=1, value=long_text)

    height = row_height_for_row(ws, 1, min_height=14.0, font_size=9)
    assert height > 14.0


def test_row_height_for_row_short_text_is_minimum():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["A"].width = 40
    ws.cell(row=1, column=1, value="Короткая строка")

    # Одна строка переноса (короткий текст, широкий столбец) → высота =
    # line_height(font_size*1.35) + 3, т.к. это больше min_height=14.0 при
    # font_size=9 (12.15+3=15.15) — floor применяется только при явно
    # заниженном min_height/font_size.
    height = row_height_for_row(ws, 1, min_height=14.0, font_size=9)
    assert height == pytest.approx(9 * 1.35 + 3)


def test_apply_row_heights_sets_dimensions_for_range():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.column_dimensions["A"].width = 10
    ws.cell(row=1, column=1, value="A")
    ws.cell(row=2, column=1, value="слово " * 60)  # много слов → много строк переноса

    apply_row_heights(ws, range(1, 3))

    assert ws.row_dimensions[1].height == pytest.approx(9 * 1.35 + 3)
    assert ws.row_dimensions[2].height > ws.row_dimensions[1].height
