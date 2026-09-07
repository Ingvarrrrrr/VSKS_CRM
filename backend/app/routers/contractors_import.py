"""Импорт контрагентов из файлов: шаблон, Excel bulk-импорт, preview/mapped
multi-format импорт (xlsx/xls/docx/doc/pdf), разбор одиночной карточки.

ПЕРЕНЕСЕНО (не изменено) из app/routers/contractors.py при разрезании
монолитного роутера (Правило №5, сессия 2026-09-08). Тот же префикс
/api/contractors. Пути здесь двухсегментные (/import/template, /import/excel,
/import/preview, /import/mapped) либо односегментные POST (/parse-file) —
конфликта с GET /{cid} core-роутера нет (другой метод или число сегментов),
но регистрируется рядом с остальными contractors_* siblings для единообразия.

_parse_file_to_rows переиспользуется из app/services/contractor_file_parsing.py
(единственный парсер карточек контрагента — второго не заводим).
"""
from io import BytesIO
from typing import Optional
from urllib.parse import quote as _url_quote

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.contractor import Contractor
from app.auth.jwt import get_current_user, get_single_org_id
from app.auth.permissions import require_tab
from app.models.user import User
from app.services.contractor_file_parsing import _parse_file_to_rows

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

router = APIRouter(prefix="/api/contractors", tags=["contractors"])


@router.post("/parse-file")
async def parse_contractor_file(
    file: UploadFile = File(...),
    _=Depends(get_current_user),
):
    """Parse a contractor card file (xlsx, docx, pdf) and return extracted fields."""
    fname = (file.filename or '').lower()
    content = await file.read()
    try:
        all_rows, hdr_idx = _parse_file_to_rows(fname, content)
    except HTTPException as e:
        raise
    except Exception as e:
        raise HTTPException(400, f"Ошибка чтения файла: {e}")

    if not all_rows or len(all_rows) <= hdr_idx:
        raise HTTPException(400, "Файл не содержит данных")

    # If it's a kv-card (2 rows: header + values), extract directly
    headers = all_rows[hdr_idx] if len(all_rows) > hdr_idx else []
    values = all_rows[hdr_idx + 1] if len(all_rows) > hdr_idx + 1 else []

    result = {}
    field_hints = {
        'name': ('назван', 'наимен', 'name', 'органи', 'учредит'),
        'inn': ('инн', 'inn'),
        'kpp': ('кпп', 'kpp'),
        'ogrn': ('огрн', 'ogrn'),
        'address': ('юридич', 'адрес', 'address'),
        'postal_address': ('почтов', 'postal'),
        'phone': ('телефон', 'phone', 'тел'),
        'email': ('email', 'e-mail'),
        'signatory': ('подписант', 'signatory', 'директор', 'руководит', 'уполномоч'),
        'bik': ('бик', 'bik'),
        'settlement_account': ('расч', 'р/с', 'settlement'),
        'correspondent_account': ('корр', 'к/с', 'correspondent'),
        'bank_name': ('банк', 'bank'),
        'org_type': ('тип', 'форма', 'org_type'),
    }

    for j, h in enumerate(headers):
        h_str = str(h).strip().lower() if h else ''
        if not h_str or j >= len(values):
            continue
        val = str(values[j]).strip() if values[j] is not None else ''
        if not val or val.lower() in ('none', 'null', '-', '—', ''):
            continue
        for field, hints in field_hints.items():
            if field not in result and any(x in h_str for x in hints):
                # INN/KPP length limits
                if field == 'inn' and len(val.replace(' ', '')) > 12:
                    continue
                if field == 'kpp' and len(val.replace(' ', '')) > 9:
                    continue
                result[field] = val
                break

    return result


@router.get("/import/template")
async def contractors_import_template(_=Depends(require_tab('contractors'))):
    """Download xlsx template for bulk contractor import."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(500, "openpyxl не установлен")

    wb = Workbook()
    ws = wb.active
    ws.title = "Контрагенты"

    headers = [
        "Наименование", "ИНН", "КПП", "ОГРН",
        "Адрес местонахождения", "Почтовый адрес",
        "Подписант", "Основание",
        "Контактное лицо", "Телефон", "Email",
        "Расчётный счёт", "Банк", "БИК", "Корр. счёт",
        "Банковские реквизиты",
    ]
    required = {"Наименование", "ИНН"}

    header_fill = PatternFill("solid", fgColor="1E40AF")
    req_fill    = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(bold=True, color="FFFFFF")

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = req_fill if h in required else header_fill
        cell.alignment = Alignment(horizontal="center")

    # Example row
    example = [
        "ООО Пример", "1234567890", "123456789", "1234567890123",
        "г. Москва, ул. Примерная, д. 1", "г. Москва, ул. Почтовая, д. 2",
        "Иванов Иван Иванович, Генеральный директор", "Устава",
        "Петров Пётр Петрович", "+7 (999) 000-00-00", "example@mail.ru",
        "40702810000000000000", "ПАО Сбербанк", "044525225", "30101810400000000225",
        "",
    ]
    for ci, val in enumerate(example, 1):
        ws.cell(row=2, column=ci, value=val)

    col_widths = [35, 14, 12, 16, 40, 40, 45, 20, 30, 20, 25, 24, 30, 12, 24, 35]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{_url_quote('Шаблон_импорта_контрагентов.xlsx', safe='-_.~')}"},
    )


@router.post("/import/excel")
async def import_contractors_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_tab('contractors'))
):
    """Bulk import contractors from Excel. First row must be headers."""
    if not (file.filename or '').lower().endswith(('.xlsx', '.xls')):
        raise HTTPException(400, "Поддерживаются только файлы .xlsx / .xls")

    if not load_workbook:
        raise HTTPException(500, "openpyxl не установлен")

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")

    ws = wb.active

    # Detect header row
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise HTTPException(400, "Файл пустой")

    def _norm(v) -> str:
        return str(v).strip().lower() if v else ''

    col: dict[str, int] = {}
    for i, h in enumerate(header_row):
        h_str = _norm(h)
        if any(x in h_str for x in ('назван', 'наимен', 'name', 'органи')):
            col.setdefault('name', i)
        elif any(x in h_str for x in ('инн', 'inn', 'идентиф')):
            col.setdefault('inn', i)
        elif any(x in h_str for x in ('кпп', 'kpp')):
            col.setdefault('kpp', i)
        elif any(x in h_str for x in ('огрн', 'ogrn')):
            col.setdefault('ogrn', i)
        elif any(x in h_str for x in ('почтов', 'postal')):
            col.setdefault('postal_address', i)
        elif any(x in h_str for x in ('адрес', 'address')):
            col.setdefault('address', i)
        elif any(x in h_str for x in ('основан', 'basis', 'устав', 'действует')):
            col.setdefault('signatory_basis', i)
        elif any(x in h_str for x in ('подписант', 'signatory', 'директор', 'руководит')):
            col.setdefault('signatory', i)
        elif any(x in h_str for x in ('контакт', 'contact', 'лицо')):
            col.setdefault('contact_person', i)
        elif any(x in h_str for x in ('телефон', 'phone', 'тел.')):
            col.setdefault('phone', i)
        elif 'email' in h_str or 'e-mail' in h_str or 'mail' in h_str:
            col.setdefault('email', i)
        elif any(x in h_str for x in ('расч', 'р/с', 'р/с', 'settlement')):
            col.setdefault('settlement_account', i)
        elif any(x in h_str for x in ('корр', 'к/с', 'correspondent')):
            col.setdefault('correspondent_account', i)
        elif any(x in h_str for x in ('бик', 'bik')):
            col.setdefault('bik', i)
        elif any(x in h_str for x in ('банк', 'bank')):
            col.setdefault('bank_name', i)
        elif any(x in h_str for x in ('реквизит',)):
            col.setdefault('bank_details', i)

    if 'name' not in col:
        raise HTTPException(
            400,
            "Не найдена колонка с наименованием. "
            "Убедитесь что первая строка — заголовки с колонкой «Наименование» или «name»."
        )

    _limits = {
        'inn': 12, 'kpp': 9, 'ogrn': 20, 'bik': 20,
        'settlement_account': 100, 'correspondent_account': 100,
        'phone': 50, 'email': 255, 'contact_person': 255,
        'signatory': 255, 'signatory_basis': 500, 'bank_name': 500,
    }

    def _cell(row, field):
        idx = col.get(field)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        if v is None:
            return None
        s = str(v).strip()
        if not s or s.lower() in ('none', 'null', '-', '—'):
            return None
        # Clean ".0" suffix from numeric fields (Excel float issue)
        if field in ('inn', 'kpp', 'ogrn', 'bik') and s.endswith('.0'):
            s = s[:-2]
        limit = _limits.get(field)
        return s[:limit] if limit else s

    created = 0
    skipped = 0
    errors_list = []

    # Collect existing INNs for dedup
    inn_result = await db.execute(select(Contractor.inn).where(Contractor.inn.isnot(None)))
    existing_inns = {r[0] for r in inn_result}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = _cell(row, 'name')
        if not name:
            skipped += 1
            continue

        inn = _cell(row, 'inn')
        if inn and inn in existing_inns:
            skipped += 1
            continue

        c = Contractor(
            name=name,
            inn=inn,
            kpp=_cell(row, 'kpp'),
            ogrn=_cell(row, 'ogrn'),
            address=_cell(row, 'address'),
            postal_address=_cell(row, 'postal_address'),
            signatory=_cell(row, 'signatory'),
            signatory_basis=_cell(row, 'signatory_basis'),
            contact_person=_cell(row, 'contact_person'),
            phone=_cell(row, 'phone'),
            email=_cell(row, 'email'),
            settlement_account=_cell(row, 'settlement_account'),
            bank_name=_cell(row, 'bank_name'),
            bik=_cell(row, 'bik'),
            correspondent_account=_cell(row, 'correspondent_account'),
            bank_details=_cell(row, 'bank_details'),
        )
        db.add(c)
        if inn:
            existing_inns.add(inn)
        created += 1

    await db.commit()
    return {"created": created, "skipped": skipped}


# ---------------------------------------------------------------------------
# New multi-format import endpoints
# ---------------------------------------------------------------------------

@router.post("/import/preview")
async def contractors_import_preview(
    file: UploadFile = File(...),
    _=Depends(require_tab('contractors')),
):
    """Parse uploaded file and return headers + sample rows for column mapping UI."""
    fname = (file.filename or '').lower()
    content = await file.read()

    try:
        all_rows, hdr_idx = _parse_file_to_rows(fname, content)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")

    hdr_rows = all_rows[hdr_idx:]
    if not hdr_rows:
        raise HTTPException(400, "Файл пустой или не содержит данных после заголовка")

    headers = [
        str(c).strip() if c else f"Столбец {j + 1}"
        for j, c in enumerate(hdr_rows[0])
    ]
    sample = [
        [str(c).strip() if c is not None else "" for c in row]
        for row in hdr_rows[1:min(4, len(hdr_rows))]
    ]
    total_rows = len(all_rows) - hdr_idx - 1

    return {
        "headers": headers,
        "sample": sample,
        "total_rows": max(total_rows, 0),
        "header_row_offset": hdr_idx,
    }


@router.post("/import/mapped")
async def contractors_import_mapped(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_tab('contractors')),
    col_name: Optional[int] = Query(None),
    col_inn: Optional[int] = Query(None),
    col_kpp: Optional[int] = Query(None),
    col_ogrn: Optional[int] = Query(None),
    col_address: Optional[int] = Query(None),
    col_postal_address: Optional[int] = Query(None),
    col_signatory: Optional[int] = Query(None),
    col_signatory_basis: Optional[int] = Query(None),
    col_contact_person: Optional[int] = Query(None),
    col_phone: Optional[int] = Query(None),
    col_email: Optional[int] = Query(None),
    col_org_phone: Optional[int] = Query(None),
    col_org_email: Optional[int] = Query(None),
    col_settlement_account: Optional[int] = Query(None),
    col_bank_name: Optional[int] = Query(None),
    col_bik: Optional[int] = Query(None),
    col_correspondent_account: Optional[int] = Query(None),
    col_bank_details: Optional[int] = Query(None),
    col_org_type: Optional[int] = Query(None),
    col_manual_product_categories: Optional[int] = Query(None),
    header_row_offset: int = Query(0),
):
    """Import contractors using user-specified column mapping."""
    if col_name is None:
        raise HTTPException(400, "Не указан столбец «Наименование»")

    fname = (file.filename or '').lower()
    content = await file.read()

    try:
        all_rows, _ = _parse_file_to_rows(fname, content)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")

    # Skip header row
    skip = header_row_offset + 1
    data_rows = all_rows[skip:]

    _limits = {
        'inn': 12, 'kpp': 9, 'ogrn': 20, 'bik': 20,
        'settlement_account': 100, 'correspondent_account': 100,
        'phone': 50, 'email': 255, 'contact_person': 255,
        'signatory': 255, 'signatory_basis': 500, 'bank_name': 500,
    }

    col_map = {
        'name': col_name,
        'inn': col_inn,
        'kpp': col_kpp,
        'ogrn': col_ogrn,
        'address': col_address,
        'postal_address': col_postal_address,
        'signatory': col_signatory,
        'signatory_basis': col_signatory_basis,
        'contact_person': col_contact_person,
        'phone': col_phone,
        'email': col_email,
        'org_phone': col_org_phone,
        'org_email': col_org_email,
        'settlement_account': col_settlement_account,
        'bank_name': col_bank_name,
        'bik': col_bik,
        'correspondent_account': col_correspondent_account,
        'bank_details': col_bank_details,
        'org_type': col_org_type,
        'manual_product_categories': col_manual_product_categories,
    }

    def _get(row, field):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return None
        v = row[idx]
        if v is None:
            return None
        s = str(v).strip()
        if not s or s.lower() in ('none', 'null', '-', '—'):
            return None
        # Clean ".0" suffix from numeric fields (Excel float issue)
        if field in ('inn', 'kpp', 'ogrn', 'bik') and s.endswith('.0'):
            s = s[:-2]
        limit = _limits.get(field)
        return s[:limit] if limit else s

    # Collect existing contractors by INN for merge
    inn_result = await db.execute(
        select(Contractor).where(Contractor.inn.isnot(None))
    )
    existing_by_inn: dict[str, Contractor] = {c.inn: c for c in inn_result.scalars().all() if c.inn}

    _updatable_fields = [
        'kpp', 'ogrn', 'address', 'postal_address', 'signatory', 'signatory_basis',
        'contact_person', 'phone', 'email', 'org_phone', 'org_email',
        'settlement_account', 'bank_name', 'bik', 'correspondent_account',
        'bank_details', 'org_type',
    ]

    created = 0
    updated = 0
    skipped_empty = 0
    errors_list = []
    update_details = []

    for row_num, row in enumerate(data_rows, start=2):
        try:
            name = _get(row, 'name')
            if not name:
                skipped_empty += 1
                continue

            inn = _get(row, 'inn')

            # Parse categories: comma/semicolon separated string → JSON array
            cats_raw = _get(row, 'manual_product_categories')
            cats = None
            if cats_raw:
                cats = [c.strip() for c in cats_raw.replace(';', ',').split(',') if c.strip()]

            # If contractor with this INN exists — merge new data into it
            if inn and inn in existing_by_inn:
                existing = existing_by_inn[inn]
                changed_fields = []
                for field in _updatable_fields:
                    new_val = _get(row, field)
                    if new_val and not getattr(existing, field, None):
                        setattr(existing, field, new_val)
                        changed_fields.append(field)
                # Merge categories
                if cats:
                    old_cats = existing.manual_product_categories or []
                    merged = list(set(old_cats + cats))
                    if merged != old_cats:
                        existing.manual_product_categories = merged
                        changed_fields.append('categories')
                # Update name if existing is shorter/empty
                if name and (not existing.name or len(name) > len(existing.name)):
                    existing.name = name
                    changed_fields.append('name')
                if changed_fields:
                    updated += 1
                    update_details.append(f"ИНН {inn}: дополнены {', '.join(changed_fields)}")
                continue

            c = Contractor(
                name=name,
                inn=inn,
                kpp=_get(row, 'kpp'),
                ogrn=_get(row, 'ogrn'),
                address=_get(row, 'address'),
                postal_address=_get(row, 'postal_address'),
                signatory=_get(row, 'signatory'),
                signatory_basis=_get(row, 'signatory_basis'),
                contact_person=_get(row, 'contact_person'),
                phone=_get(row, 'phone'),
                email=_get(row, 'email'),
                org_phone=_get(row, 'org_phone'),
                org_email=_get(row, 'org_email'),
                settlement_account=_get(row, 'settlement_account'),
                bank_name=_get(row, 'bank_name'),
                bik=_get(row, 'bik'),
                correspondent_account=_get(row, 'correspondent_account'),
                bank_details=_get(row, 'bank_details'),
                org_type=_get(row, 'org_type'),
                org_id=get_single_org_id(current_user) or current_user.org_id,
                manual_product_categories=cats,
            )
            db.add(c)
            if inn:
                existing_by_inn[inn] = c
            created += 1
        except Exception as e:
            errors_list.append(f"Строка {row_num}: {str(e)}")

    await db.commit()
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped_empty,
        "skipped_empty": skipped_empty,
        "update_details": update_details[:50],
        "errors": errors_list,
    }
