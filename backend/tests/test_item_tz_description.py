"""Своё ТЗ позиции заявки/закупки (владелец, 14.10.2026, план
item-tz-description.md, Этап А):
  - app/services/item_tz.py::effective_item_tz — единственный источник формулы
    «tz_description позиции, если непустое, иначе product.description»;
  - миграция l2m3n4o5p6q7 добавляет wish_items.tz_description /
    purchase_items.tz_description;
  - Excel «умный» импорт распознаёт колонку «Характеристики» и отдаёт её как
    description (файл), раздельно от catalog_description (каталог);
  - импорт НИКОГДА не пишет описание в СУЩЕСТВУЮЩИЙ товар каталога — текст
    уходит в tz_description самой позиции;
  - лимит 200 с total_in_file/truncated — на предпросмотре.
"""
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from sqlalchemy import select, text

from app.models.product import Product
from app.models.purchase_item import PurchaseItem
from app.services.item_tz import effective_item_tz, effective_item_tz_44fz

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

_UNIQ = "ТестItemTz20261014"


# ---------------------------------------------------------------------------
# effective_item_tz — ОДНА функция-источник (unit-level, без БД)
# ---------------------------------------------------------------------------

def test_effective_item_tz_prefers_own_text():
    item = SimpleNamespace(tz_description="Своё ТЗ позиции", product=SimpleNamespace(description="Каталожное"))
    assert effective_item_tz(item) == "Своё ТЗ позиции"


def test_effective_item_tz_falls_back_to_product():
    item = SimpleNamespace(tz_description=None, product=SimpleNamespace(description="Каталожное"))
    assert effective_item_tz(item) == "Каталожное"


def test_effective_item_tz_blank_own_text_falls_back():
    """Пустая строка (не None) — тоже считается «нет своего ТЗ»."""
    item = SimpleNamespace(tz_description="   ", product=SimpleNamespace(description="Каталожное"))
    assert effective_item_tz(item) == "Каталожное"


def test_effective_item_tz_no_product_no_own_text():
    item = SimpleNamespace(tz_description=None, product=None)
    assert effective_item_tz(item) is None


def test_effective_item_tz_missing_tz_description_attr_falls_back_safely():
    """Объект без атрибута tz_description вовсе (напр. ContractItem) — не падает."""
    item = SimpleNamespace(product=SimpleNamespace(description="Каталожное"))
    assert effective_item_tz(item) == "Каталожное"


def test_effective_item_tz_44fz_uses_44fz_catalog_field():
    item = SimpleNamespace(tz_description=None, product=SimpleNamespace(description_44fz="44-ФЗ текст"))
    assert effective_item_tz_44fz(item) == "44-ФЗ текст"


def test_effective_item_tz_44fz_own_text_wins_over_44fz_catalog():
    item = SimpleNamespace(tz_description="Своё", product=SimpleNamespace(description_44fz="44-ФЗ текст"))
    assert effective_item_tz_44fz(item) == "Своё"


# ---------------------------------------------------------------------------
# Миграция l2m3n4o5p6q7 — колонки существуют
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_migration_added_tz_description_columns(db_session):
    rows = (await db_session.execute(text(
        "SELECT table_name, column_name FROM information_schema.columns "
        "WHERE table_name IN ('wish_items', 'purchase_items') AND column_name = 'tz_description'"
    ))).fetchall()
    tables_with_column = {r[0] for r in rows}
    assert tables_with_column == {"wish_items", "purchase_items"}


# ---------------------------------------------------------------------------
# Excel-парсер — колонка «Характеристики»
# ---------------------------------------------------------------------------

def _xlsx_with_description(rows: list[tuple]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Наименование", "Характеристики", "Количество", "Цена за ед."])
    for row in rows:
        ws.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_smart_xlsx_direct_parses_description_column():
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.services.items_import_parsing import _smart_import_xlsx_direct

    xlsx = _xlsx_with_description([
        (f"{_UNIQ} Перчатки", "Нитриловые, размер M, упаковка 100 шт", 10, 50),
    ])
    preview, columns, warnings = _smart_import_xlsx_direct(xlsx, fname="test.xlsx")
    assert len(preview) == 1
    assert preview[0]["item_name"] == f"{_UNIQ} Перчатки"
    assert preview[0]["description"] == "Нитриловые, размер M, упаковка 100 шт"
    assert "description" in columns
    # Колонка характеристик не должна путаться с item_type ('тип'/'вид').
    assert preview[0]["item_type"] == "товар"


def test_smart_xlsx_direct_skip_filter_unchanged():
    """Строка без qty/price/total всё ещё пропускается, даже если у неё
    заполнена колонка характеристик (владелец, 14.10.2026 — фильтр не менять)."""
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.services.items_import_parsing import _smart_import_xlsx_direct

    wb = Workbook()
    ws = wb.active
    ws.append(["Наименование", "Характеристики", "Количество", "Цена за ед."])
    ws.append([f"{_UNIQ} Раздел-заголовок", "Текст характеристик есть, но это подзаголовок", None, None])
    buf = BytesIO()
    wb.save(buf)
    preview, _columns, _warnings = _smart_import_xlsx_direct(buf.getvalue(), fname="test.xlsx")
    assert preview == []


# ---------------------------------------------------------------------------
# import-smart-nopid / import-mapped-nopid — description (файл) И
# catalog_description (каталог) раздельно
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_smart_nopid_returns_file_description_and_catalog_description_separately(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")

    catalog_name = f"{_UNIQ} Товар в каталоге"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ", price=Decimal("100"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла, другое", 2, 100)])
    resp = await client.post(
        "/api/purchases/items/import-smart-nopid",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    row = data["preview"][0]
    assert row["description"] == "ТЗ из файла, другое"
    assert row["catalog_description"] == "Каталожное ТЗ"
    assert "total_in_file" in data and "truncated" in data
    assert data["truncated"] is False
    assert data["total_in_file"] == 1


@pytest.mark.asyncio
async def test_mapped_nopid_returns_file_description_and_catalog_description_separately(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")

    catalog_name = f"{_UNIQ} Товар в каталоге Б"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ Б", price=Decimal("50"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла Б", 1, 50)])
    resp = await client.post(
        "/api/purchases/items/import-mapped-nopid",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={"col_item_name": 0, "col_description": 1, "col_quantity": 2, "col_unit_price": 3},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    row = data["items"][0]
    assert row["description"] == "ТЗ из файла Б"
    assert row["catalog_description"] == "Каталожное ТЗ Б"
    assert "total_in_file" in data and "truncated" in data


# ---------------------------------------------------------------------------
# Импорт В закупку (pid) — файл пишется в PurchaseItem.tz_description,
# СУЩЕСТВУЮЩИЙ товар каталога НЕ меняется
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_mapped_import_with_pid_does_not_touch_existing_product_description(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")

    from app.models.purchase import Purchase

    purchase = Purchase(item_name=f"{_UNIQ} Закупка", status="draft")
    db_session.add(purchase)
    await db_session.commit()
    await db_session.refresh(purchase)

    catalog_name = f"{_UNIQ} Товар для закупки"
    p = Product(name=catalog_name, category="Прочее", description="Старое каталожное описание", price=Decimal("10"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла для этой закупки", 3, 10)])
    resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-mapped",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={
            "col_item_name": 0, "col_description": 1, "col_quantity": 2, "col_unit_price": 3,
            "confirm": "true",
        },
    )
    assert resp.status_code == 200, resp.text

    await db_session.refresh(p)
    assert p.description == "Старое каталожное описание", "импорт не должен менять описание существующего товара"

    item = (await db_session.execute(
        select(PurchaseItem).where(
            PurchaseItem.purchase_id == purchase.id,
            PurchaseItem.item_name == catalog_name,
        )
    )).scalar_one()
    assert item.tz_description == "ТЗ из файла для этой закупки"
    assert item.product_id == p.id


# ---------------------------------------------------------------------------
# Лимит 200 + total_in_file/truncated на 201+ строках
# ---------------------------------------------------------------------------

def test_apply_import_preview_limit_truncates_at_200():
    from app.services.items_import_limits import apply_import_preview_limit

    rows = [{"n": i} for i in range(205)]
    limited, total, truncated = apply_import_preview_limit(rows)
    assert total == 205
    assert truncated is True
    assert len(limited) == 200


def test_apply_import_preview_limit_no_truncation_under_limit():
    from app.services.items_import_limits import apply_import_preview_limit

    rows = [{"n": i} for i in range(50)]
    limited, total, truncated = apply_import_preview_limit(rows)
    assert total == 50
    assert truncated is False
    assert len(limited) == 50


# ---------------------------------------------------------------------------
# Импорт В закупку (pid) — preview отдаёт catalog_description, confirm
# принимает tz_choices (окно сравнения, Этап Б бэкенд)
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_smart_import_pid_preview_returns_catalog_description(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.models.purchase import Purchase

    purchase = Purchase(item_name=f"{_UNIQ} Закупка smart pid preview", status="draft")
    db_session.add(purchase)
    await db_session.commit()
    await db_session.refresh(purchase)

    catalog_name = f"{_UNIQ} Товар smart pid"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ smart pid", price=Decimal("10"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла smart pid", 2, 10)])
    resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-smart",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={"confirm": "false"},
    )
    assert resp.status_code == 200, resp.text
    row = resp.json()["preview"][0]
    assert row["description"] == "ТЗ из файла smart pid"
    assert row["catalog_description"] == "Каталожное ТЗ smart pid"


@pytest.mark.asyncio
async def test_smart_import_pid_confirm_tz_choice_catalog_sets_null(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.models.purchase import Purchase
    import json as _json

    purchase = Purchase(item_name=f"{_UNIQ} Закупка smart pid confirm", status="draft")
    db_session.add(purchase)
    await db_session.commit()
    await db_session.refresh(purchase)

    catalog_name = f"{_UNIQ} Товар smart pid confirm"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ confirm", price=Decimal("10"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла confirm", 1, 10)])
    preview_resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-smart",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={"confirm": "false"},
    )
    row_num = preview_resp.json()["preview"][0]["row"]

    resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-smart",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={"confirm": "true", "tz_choices": _json.dumps({str(row_num): "catalog"})},
    )
    assert resp.status_code == 200, resp.text

    item = (await db_session.execute(
        select(PurchaseItem).where(
            PurchaseItem.purchase_id == purchase.id,
            PurchaseItem.item_name == catalog_name,
        )
    )).scalar_one()
    assert item.tz_description is None, "выбор catalog в tz_choices — tz_description NULL (действующее ТЗ подтянется из каталога)"


@pytest.mark.asyncio
async def test_smart_import_pid_confirm_default_uses_file_text(client, auth_headers, db_session):
    """Без tz_choices (или без ключа строки) — по умолчанию текст из файла
    (владелец, 14.10.2026: «по умолчанию — из файла»)."""
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.models.purchase import Purchase

    purchase = Purchase(item_name=f"{_UNIQ} Закупка smart pid default", status="draft")
    db_session.add(purchase)
    await db_session.commit()
    await db_session.refresh(purchase)

    catalog_name = f"{_UNIQ} Товар smart pid default"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ default", price=Decimal("10"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла default", 1, 10)])
    resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-smart",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={"confirm": "true"},
    )
    assert resp.status_code == 200, resp.text

    item = (await db_session.execute(
        select(PurchaseItem).where(
            PurchaseItem.purchase_id == purchase.id,
            PurchaseItem.item_name == catalog_name,
        )
    )).scalar_one()
    assert item.tz_description == "ТЗ из файла default"


@pytest.mark.asyncio
async def test_mapped_import_pid_preview_returns_description_and_catalog_description(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.models.purchase import Purchase

    purchase = Purchase(item_name=f"{_UNIQ} Закупка mapped pid preview", status="draft")
    db_session.add(purchase)
    await db_session.commit()
    await db_session.refresh(purchase)

    catalog_name = f"{_UNIQ} Товар mapped pid"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ mapped pid", price=Decimal("10"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла mapped pid", 2, 10)])
    resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-mapped",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={
            "col_item_name": 0, "col_description": 1, "col_quantity": 2, "col_unit_price": 3,
            "confirm": "false",
        },
    )
    assert resp.status_code == 200, resp.text
    row = resp.json()["preview"][0]
    assert row["description"] == "ТЗ из файла mapped pid"
    assert row["catalog_description"] == "Каталожное ТЗ mapped pid"


@pytest.mark.asyncio
async def test_mapped_import_pid_confirm_tz_choice_catalog_sets_null(client, auth_headers, db_session):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")
    from app.models.purchase import Purchase
    import json as _json

    purchase = Purchase(item_name=f"{_UNIQ} Закупка mapped pid confirm", status="draft")
    db_session.add(purchase)
    await db_session.commit()
    await db_session.refresh(purchase)

    catalog_name = f"{_UNIQ} Товар mapped pid confirm"
    p = Product(name=catalog_name, category="Прочее", description="Каталожное ТЗ mapped confirm", price=Decimal("10"))
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)

    xlsx = _xlsx_with_description([(catalog_name, "ТЗ из файла mapped confirm", 1, 10)])
    preview_resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-mapped",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={
            "col_item_name": 0, "col_description": 1, "col_quantity": 2, "col_unit_price": 3,
            "confirm": "false",
        },
    )
    row_num = preview_resp.json()["preview"][0]["row"]

    resp = await client.post(
        f"/api/purchases/{purchase.id}/items/import-mapped",
        headers=auth_headers,
        files={"file": ("tz.xlsx", xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        params={
            "col_item_name": 0, "col_description": 1, "col_quantity": 2, "col_unit_price": 3,
            "confirm": "true", "tz_choices": _json.dumps({str(row_num): "catalog"}),
        },
    )
    assert resp.status_code == 200, resp.text

    item = (await db_session.execute(
        select(PurchaseItem).where(
            PurchaseItem.purchase_id == purchase.id,
            PurchaseItem.item_name == catalog_name,
        )
    )).scalar_one()
    assert item.tz_description is None


@pytest.mark.asyncio
async def test_smart_nopid_reports_truncation_for_201_rows(client, auth_headers):
    if Workbook is None:
        pytest.skip("openpyxl не установлен")

    wb = Workbook()
    ws = wb.active
    ws.append(["Наименование", "Количество", "Цена за ед."])
    for i in range(201):
        ws.append([f"{_UNIQ} Позиция {i}", 1, 10])
    buf = BytesIO()
    wb.save(buf)

    resp = await client.post(
        "/api/purchases/items/import-smart-nopid",
        headers=auth_headers,
        files={"file": ("big.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["total_in_file"] == 201
    assert data["truncated"] is True
    assert len(data["preview"]) == 200
