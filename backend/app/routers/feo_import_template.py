"""GET /api/feo-categories/import/template — генерация Excel-шаблона импорта
категорий ФЭО.

Вынесено из app/routers/feo_import.py (Правило №5, разрезание 1088-строчного
роутера) без изменения поведения — самый крупный из пяти эндпоинтов (генерация
Excel: заголовки, примеры, подсказки, лист «Справочники», лист «Как
заполнять»), самодостаточен (не вызывает _do_feo_import). Путь статичный
(без {cat_id}) — регистрируется в app/routes.py рядом с feo_import.router и
ДО feo_categories.router (несёт catch-all GET/PUT/DELETE "/{cat_id}"), как и
остальные соседи feo_import_*.
"""
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    Workbook = None
    DataValidation = None

from app.database import get_db
from app.models.feo_category import FeoCategory
from app.auth.jwt import get_current_user
from app.auth.visibility import get_visible_subsidy_ids
from app.utils.http import content_disposition

router = APIRouter(prefix="/api/feo-categories", tags=["feo_categories"])


_content_disposition = content_disposition


@router.get("/import/template")
async def download_feo_template(
    subsidy_id: Optional[int] = Query(
        None,
        description=(
            "Если передан — шаблон делается ПОД ЭТУ субсидию: колонка «Субсидия» в "
            "примерах = её имя, строки-примеры используют её РЕАЛЬНЫЕ названия "
            "уровней, а колонки «Уровень 2/3/4» получают выпадающий список из "
            "реальных категорий этой субсидии (лист «Справочники»). Без параметра — "
            "общий шаблон с нейтральными примерами, без выпадающих списков по уровням."
        ),
    ),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Шаблон Excel для импорта категорий ФЭО (18 колонок, переверстан 2026-08-14;
    subsidy_id + выпадающие списки по категориям добавлены 2026-09-04).

    Причина переверстки (боевой файл владельца, «Субсидия ДНР 2.xlsx»): в
    37-колоночном шаблоне между заголовком уровня и его числовыми колонками
    стояло по семь числовых колонок — легко промахнуться и набрать название
    следующего уровня в числовой колонке предыдущего (см. level_name_in_number_column
    в _do_feo_import). Новый шаблон разводит уровни (только названия, 4 колонки)
    и числа (2 ОДНА-на-всю-строку пары колонок «по ФЭО»/«плана», а не по 12
    числовых колонок на каждый уровень) — числа сами прикрепляются к самому
    глубокому заполненному уровню строки, либо к плановой позиции, если она
    заполнена (колонка 5).

    Колонки: Субсидия | Уровень 2 | Уровень 3 | Уровень 4 | Плановая позиция
    (папка НЕ создаётся) | Товар/услуга | Количество/Ед.изм./Цена/Сумма по ФЭО
    (одна пара на строку) | Плановое количество/Ед.изм./Цена/Сумма плана (одна
    пара на строку) | Код | Приложение | Активна | Финансирование (устар.).
    Если уровень пропущен, содержимое нижнего поднимается на его место. Если в
    строке нет ни одного уровня, а плановая позиция заполнена — её название
    становится Уровнем 2. Сумма строки приоритетнее кол-во × цена; расхождение —
    предупреждение. Второй лист «Как заполнять» — текстовые правила.

    2026-09-04 (владелец): раньше примеры ВСЕГДА были зашиты строками с именем
    «ДНР_2026» и её реальными уровнями — при скачивании шаблона из ЛЮБОЙ другой
    субсидии (например, ЦентрПоиск_2026) пользователь путал чужие уровни/название
    со своими. Теперь: subsidy_id передан → примеры и списки строятся из дерева
    ЭТОЙ субсидии; не передан → нейтральный заполнитель «Название вашей субсидии»,
    никакого чужого имени/уровней.
    """
    if Workbook is None:
        raise HTTPException(500, "openpyxl не установлен")

    from app.models.subsidy import Subsidy

    subsidy: Optional[Subsidy] = None
    # db level 1/2/3 = UI «Уровень 2/3/4» (см. find_or_create/db_level в _do_feo_import).
    cats_by_level: dict[int, list[FeoCategory]] = {1: [], 2: [], 3: []}
    cats_by_id: dict[int, FeoCategory] = {}
    if subsidy_id is not None:
        subsidy = (await db.execute(select(Subsidy).where(Subsidy.id == subsidy_id))).scalar_one_or_none()
        if subsidy is None:
            raise HTTPException(404, "Субсидия не найдена")
        # Та же видимость, что и у списка/дерева категорий (list_categories/category_tree):
        # доступ к субсидии по вкладке ФЭО ИЛИ заявок ИЛИ закупок.
        vis = await get_visible_subsidy_ids(current_user, db, "feo_categories")
        if vis is not None:
            vis = vis | await get_visible_subsidy_ids(current_user, db, "purchases")
            vis = vis | await get_visible_subsidy_ids(current_user, db, "wishes")
            if subsidy_id not in vis:
                raise HTTPException(403, "Нет доступа к этой субсидии")
        cats_q = (
            select(FeoCategory)
            .where(FeoCategory.subsidy_id == subsidy_id, FeoCategory.is_active.is_(True))
            .order_by(FeoCategory.level, FeoCategory.sort_order.nulls_last(), FeoCategory.id)
        )
        all_cats = (await db.execute(cats_q)).scalars().all()
        cats_by_id = {c.id: c for c in all_cats}
        for c in all_cats:
            if c.level in cats_by_level:
                cats_by_level[c.level].append(c)

    wb = Workbook()
    ws = wb.active
    ws.title = "Категории ФЭО"
    headers = [
        "Субсидия",                                          # A   1
        "Уровень 2 (Направление расходов по ФЭО)",           # B   2
        "Уровень 3 (Тип расходов по ФЭО)",                    # C   3
        "Уровень 4 (Конкретизированный)",                     # D   4
        "Плановая позиция (папка НЕ создаётся)",              # E   5
        "Товар/услуга",                                       # F   6
        "Количество по ФЭО",                                  # G   7
        "Ед. изм. по ФЭО",                                    # H   8
        "Цена за единицу по ФЭО",                             # I   9
        "Сумма по ФЭО",                                       # J  10
        "Плановое количество",                                # K  11
        "Ед. изм. плана",                                     # L  12
        "Плановая цена за единицу",                           # M  13
        "Сумма плана",                                        # N  14
        "Код",                                                # O  15
        "Приложение",                                         # P  16
        "Активна",                                            # Q  17
        "Финансирование (устар., можно не заполнять)",        # R  18
    ]
    ws.append(headers)

    # Цветовое кодирование заголовков
    fill_cat  = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")   # синий — субсидия/уровни + ФЭО
    fill_plan = PatternFill(start_color="0891B2", end_color="0891B2", fill_type="solid")   # голубой — план
    fill_item = PatternFill(start_color="059669", end_color="059669", fill_type="solid")   # зелёный — плановая позиция
    fill_attr = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")   # фиолетовый — атрибуты
    font_w = Font(color="FFFFFF", bold=True, size=10)

    _blue_cols  = {1, 2, 3, 4, 7, 8, 9, 10}    # субсидия, уровни, «по ФЭО» (одна пара на строку)
    _cyan_cols  = {11, 12, 13, 14}             # план (одна пара на строку)
    _green_cols = {5, 6}                       # плановая позиция + тип

    for i, cell in enumerate(ws[1], start=1):
        if i in _blue_cols:
            cell.fill = fill_cat
        elif i in _cyan_cols:
            cell.fill = fill_plan
        elif i in _green_cols:
            cell.fill = fill_item
        else:
            cell.fill = fill_attr
        cell.font = font_w
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 40

    # Примеры (строки 2..N), 18 элементов каждая:
    # col: 1=Субс 2=Ур2 3=Ур3 4=Ур4 5=Плановая_позиция 6=Тип 7=feoQ 8=feoU 9=feoЦена 10=feoСумма 11=planQ 12=planU 13=planЦена 14=planСумма 15=Код 16=Прил 17=Акт 18=Финанс(устар.)
    #
    # 2026-09-04 (владелец): раньше здесь стояло жёстко «ДНР_2026» с её реальными
    # уровнями — при открытии шаблона из ЛЮБОЙ другой субсидии пользователь видел
    # чужие название/уровни и путался. Теперь: если subsidy_id передан и у субсидии
    # УЖЕ есть категории — примеры строятся из её РЕАЛЬНОГО дерева (по одному
    # примеру на глубину, где эта глубина существует). Иначе (subsidy_id не передан,
    # либо у субсидии дерево ещё пустое) — нейтральный нейтральный нейтральный заполнитель, никакой чужой субсидии.
    subsidy_label = subsidy.name if subsidy is not None else "Название вашей субсидии"

    def _path_names(cat: FeoCategory) -> tuple[str, str, str]:
        """(имя Ур.2, Ур.3, Ур.4) для цепочки cat → его родителей (пусто, если родителя нет на этом уровне)."""
        chain: list[FeoCategory] = []
        cur: Optional[FeoCategory] = cat
        seen: set[int] = set()
        while cur is not None and cur.id not in seen:
            seen.add(cur.id)
            chain.append(cur)
            cur = cats_by_id.get(cur.parent_id) if cur.parent_id else None
        chain.reverse()  # от корня (Ур.2) к листу
        names = [c.name for c in chain]
        while len(names) < 3:
            names.append("")
        return names[0], names[1], names[2]

    example_rows: list[list] = []
    has_real_tree = bool(cats_by_level[1] or cats_by_level[2] or cats_by_level[3])
    if subsidy is not None and has_real_tree:
        if cats_by_level[1]:
            leaf = cats_by_level[1][0]
            n1, n2, n3 = _path_names(leaf)
            example_rows.append([
                subsidy_label, n1, n2, n3, "", "", "", "", "", "2000000", "", "", "", "",
                leaf.code or "", leaf.appendix or "", "да" if leaf.is_active else "нет", "",
            ])
        if cats_by_level[2]:
            leaf = cats_by_level[2][0]
            n1, n2, n3 = _path_names(leaf)
            example_rows.append([
                subsidy_label, n1, n2, n3, "", "", "6", "шт", "150000", "", "6", "шт", "150000", "900000",
                leaf.code or "", leaf.appendix or "", "да" if leaf.is_active else "нет", "",
            ])
        if cats_by_level[3]:
            leaf = cats_by_level[3][0]
            n1, n2, n3 = _path_names(leaf)
            example_rows.append([
                subsidy_label, n1, n2, n3, "", "", "", "", "", "", "1", "усл", "178779.59", "178779.59",
                leaf.code or "", leaf.appendix or "", "да" if leaf.is_active else "нет", "",
            ])
    else:
        # Общий шаблон (нет subsidy_id) ИЛИ субсидия передана, но дерево у неё ещё
        # пустое — реальных уровней показать нечего, только условные подписи.
        example_rows.extend([
            [subsidy_label, "Пример: Техническое оснащение", "Пример: Оргтехника", "", "", "", "", "", "", "2000000", "", "", "", "", "01.01.01", "Прил. 1", "да", ""],
            [subsidy_label, "Пример: Техническое оснащение", "Пример: Оргтехника", "Пример: Закупка компьютеров", "", "", "6", "шт", "150000", "", "6", "шт", "150000", "900000", "01.01.02", "Прил. 1", "да", ""],
        ])
    # Эти две строки — условные во всех случаях (демонстрируют «Плановую позицию»
    # и план без разбивки на уровни, а не сами уровни ФЭО) — реальных категорий не касаются.
    example_rows.append([subsidy_label, "", "", "", "Пример: Услуга по заправке техники", "Услуга", "", "", "", "", "500", "л", "60", "30000", "", "", "да", ""])
    example_rows.append([subsidy_label, "Пример: Организация мероприятий", "Пример: Слёт студентов", "", "", "", "", "", "", "", "", "", "", "3500000", "02.02.01", "Прил. 2", "да", ""])

    for row in example_rows:
        ws.append(row)

    # Подсказки в строке 7
    hints = [
        "← Точное название как в системе",                                                              # A   1
        "← Направление расходов (создаётся если нет)",                                                  # B   2
        "← Тип расходов (если пусто — атрибуты к Ур.2); пропущенный уровень поднимается вверх",          # C   3
        "← Конкретизированный (если пусто — к Ур.3); пропущенный уровень поднимается вверх",             # D   4
        "← Плановая позиция — папка НЕ создаётся; если уровней в строке вовсе нет, станет Уровнем 2",    # E   5
        "← Товар / Услуга / Работа (выпадающий список)",                                                 # F   6
        "← Кол-во по ФЭО — одна колонка на строку, привязывается к самому глубокому уровню",             # G   7
        "← Ед. изм. по ФЭО",                                                                              # H   8
        "← Цена за единицу по ФЭО (руб.)",                                                                # I   9
        "← Сумма по ФЭО (руб.); если пусто — кол-во × цена; сумма приоритетнее",                          # J  10
        "← Плановое кол-во — одна колонка на строку; если задана «Плановая позиция», уходит в неё",       # K  11
        "← Ед. изм. плана",                                                                               # L  12
        "← Плановая цена за единицу (руб.)",                                                              # M  13
        "← Сумма плана (руб.); если пусто — кол-во × цена; сумма приоритетнее",                           # N  14
        "← Код категории",                                                                                # O  15
        "← Номер приложения",                                                                             # P  16
        "← да/нет",                                                                                       # Q  17
        "← устарело, можно не заполнять — используйте «Сумма по ФЭО»",                                    # R  18
    ]
    for col, hint in enumerate(hints, start=1):
        ws.cell(7, col).value = hint
        ws.cell(7, col).font = Font(italic=True, color="888888", size=8)

    # Ширины колонок (18 штук)
    col_widths = [18, 42, 42, 42, 42, 14, 16, 14, 16, 16, 16, 14, 16, 16, 10, 12, 10, 22]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"

    # Выпадающий список Товар/Услуга/Работа на колонку F (Товар/услуга), строки данных 2:1000.
    if DataValidation is not None:
        dv_item_type = DataValidation(
            type="list",
            formula1='"Товар,Услуга,Работа"',
            allow_blank=True,
            showErrorMessage=False,
            showInputMessage=True,
        )
        dv_item_type.error = "Значение не из списка — будет принято как есть"
        dv_item_type.errorTitle = "Нестандартное значение"
        dv_item_type.promptTitle = "Товар/услуга"
        dv_item_type.prompt = "Выберите Товар, Услуга или Работа"
        dv_item_type.sqref = "F2:F1000"
        ws.add_data_validation(dv_item_type)

    # Лист «Справочники» + выпадающие списки по колонкам Уровень 2/3/4 (только
    # когда subsidy_id передан и у субсидии УЖЕ есть хоть одна категория —
    # 2026-09-04). Список — ПОДСКАЗКА, не запрет: через этот шаблон заводят и
    # НОВЫЕ категории, поэтому showErrorMessage=False (как у dv_item_type выше).
    # Инлайн-список (formula1='"a,b,c"') не годится — Excel режет его на ~255
    # символах, категорий может быть больше, поэтому список живёт на отдельном
    # листе, а формула ссылается на диапазон.
    if subsidy is not None and has_real_tree and DataValidation is not None:
        ws3 = wb.create_sheet("Справочники")
        level_ref_cols = {1: "A", 2: "B", 3: "C"}   # колонки листа "Справочники"
        level_titles = {1: "Уровень 2", 2: "Уровень 3", 3: "Уровень 4"}
        level_target_cols = {1: "B", 2: "C", 3: "D"}  # колонки листа "Категории ФЭО"
        for lvl, col_letter in level_ref_cols.items():
            cell = ws3[f"{col_letter}1"]
            cell.value = level_titles[lvl]
            cell.font = Font(bold=True)
            ws3.column_dimensions[col_letter].width = 42
        for lvl, ref_col in level_ref_cols.items():
            names = sorted(dict.fromkeys(c.name for c in cats_by_level[lvl] if c.name))
            for i, name in enumerate(names, start=2):
                ws3[f"{ref_col}{i}"] = name
            if not names:
                continue  # у субсидии нет категорий этого уровня — список не вешаем
            last_row = len(names) + 1 + 50  # запас с рядом, чтобы список не обрезался
            dv_level = DataValidation(
                type="list",
                formula1=f"Справочники!${ref_col}$2:${ref_col}${last_row}",
                allow_blank=True,
                showErrorMessage=False,
                showInputMessage=True,
            )
            dv_level.error = "Значение не из списка — можно ввести и новую категорию"
            dv_level.errorTitle = "Нестандартное значение"
            dv_level.promptTitle = level_titles[lvl]
            dv_level.prompt = f"Выберите существующую категорию «{level_titles[lvl]}» или впишите новую"
            target_col = level_target_cols[lvl]
            dv_level.sqref = f"{target_col}2:{target_col}1000"
            ws.add_data_validation(dv_level)

    # Второй лист — текстовые правила заполнения.
    ws2 = wb.create_sheet("Как заполнять")
    ws2.column_dimensions["A"].width = 110
    rules = [
        "Как заполнять шаблон импорта направлений ФЭО",
        "",
        "1. Уровни (Уровень 2 → Уровень 3 → Уровень 4) идут СЛЕВА НАПРАВО. Уровень 2 — обязателен.",
        "2. Пропущенный уровень поджимается вверх: заполнены Уровень 2 и Уровень 4, а Уровень 3 пуст — Уровень 4 становится ребёнком Уровня 2 напрямую.",
        "3. «Плановая позиция» (колонка 5) — конкретный товар/услуга/работа ВНУТРИ категории. Заполненная «Плановая позиция» папку (категорию) НЕ создаёт — только позицию.",
        "4. Если в строке не заполнен НИ ОДИН из уровней (2/3/4), а «Плановая позиция» заполнена — её название становится Уровнем 2 (создаётся направлением), а позиция при этом не создаётся.",
        "5. «Количество/Ед.изм./Цена/Сумма по ФЭО» и «Плановое количество/Ед.изм./Цена/Сумма плана» — ОДНА пара колонок на всю строку, а не по уровням. Если «Плановая позиция» не задана — числа привязываются к самому глубокому заполненному уровню строки. Если позиция задана — числа плана описывают именно её.",
        "6. Сумма (по ФЭО / плана) ПРИОРИТЕТНЕЕ «количество × цена»: если сумма указана явно — используется она; расхождение с расчётом по кол-во × цена попадёт в предупреждения импорта.",
        "7. «Товар/услуга» (колонка 6) — тип плановой позиции: Товар, Услуга или Работа (выпадающий список в ячейке).",
        "8. Не меняйте заголовки колонок — импорт определяет их по названию, а не по порядку.",
        "9. Строка-подсказка (начинается с «←», строка 7 примера) в файл не попадает — служебная, игнорируется при импорте.",
    ]
    for line in rules:
        ws2.append([line])
    ws2["A1"].font = Font(bold=True, size=12)

    wb.active = 0
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition("Шаблон_импорта_направлений_ФЭО.xlsx")})
