"""Жалоба владельца 05.10.2026: шаг сопоставления колонок мастера «Импорт
ФЭО» (ручной маппинг, /import-mapped) предлагал устаревшие поля старого
37-колоночного формата (qty_lvl2/3/4, feo_qty_lvl*, feo_amount_lvl* и т.п.) и
НЕ предлагал колонки нового блока «Факт» (решение владельца 05.10.2026,
колонки U-AB шаблона) и «Нужность» (колонка T).

Три проверки:
  1. Заголовки актуального шаблона (`download_feo_template`, колонки A-AB) —
     ЗАФИКСИРОВАНЫ этим тестом. Если кто-то поменяет порядок/подписи колонок
     шаблона, не предупредив фронт (useFeoImport.ts::FEO_TARGET_FIELDS) — тест
     упадёт здесь, а не после жалобы пользователя на кривой список выбора.
  2. `resolve_fact_columns` (app/services/feo_import_fact_summary.py) — явное
     сопоставление колонки (от ручного маппинга) побеждает автоопределение по
     заголовку, и работает даже когда заголовки файла вообще не несут
     узнаваемых слов блока «Факт».
  3. `/import-mapped` (`import_feo_mapped`) с явно сопоставленными col_fact_*
     и col_need_level на файле с ПРОИЗВОЛЬНЫМИ заголовками (не совпадающими
     со словами автоопределения) — строит ТОТ ЖЕ план, что и обычный маппинг,
     плюс has_fact_columns=True/fact_rows>0 и need_level из явно
     сопоставленной колонки (а не 'likely' по умолчанию).

Гонять по одному узлу за вызов pytest в контейнере (флейк pytest-asyncio
«different loop» — project_pytest_asyncio_loop_flake, Lessons.md VSKS_CRM).
"""
import io
import uuid

import pytest
from fastapi import UploadFile
from openpyxl import Workbook
from starlette.requests import Request

from app.routers.feo_import_template import download_feo_template
from app.services.feo_import_fact_summary import FACT_FIELDS, resolve_fact_columns
from tests.test_feo_import_tree import _cleanup_subsidy, _get_categories, _get_items, _make_subsidy

# Ровно то, что строит download_feo_template (headers, feo_import_template.py) —
# единственный контракт, с которым фронт (FEO_TARGET_FIELDS) обязан совпадать.
EXPECTED_TEMPLATE_HEADERS = [
    "Субсидия",
    "Уровень 2 (Направление расходов по ФЭО)",
    "Уровень 3 (Тип расходов по ФЭО)",
    "Уровень 4 (Конкретизированный)",
    "Плановая позиция (папка НЕ создаётся)",
    "Товар/услуга/работа",
    "Количество по ФЭО",
    "Ед. изм. по ФЭО",
    "Цена за единицу по ФЭО",
    "Сумма по ФЭО",
    "Плановое количество",
    "Ед. изм. плана",
    "Плановая цена за единицу",
    "Сумма плана",
    "Код",
    "Приложение",
    "Активна",
    "Финансирование (устар., можно не заполнять)",
    "Комментарий",
    "Нужность",
    "Правильный статус",
    "Факт: Количество",
    "Факт: Цена",
    "Факт: Сумма",
    "Оплачено",
    "Законтрактовано",
    "Поставщик",
    "№ закупки",
]


@pytest.mark.asyncio
async def test_template_headers_match_fixed_contract(db_session, superadmin_user):
    """Шаблон (A-AB) совпадает с зафиксированным списком выше — именно этот
    список (без qty_lvl2/feo_qty_lvl*/feo_amount_lvl* и т.п.) должен быть
    источником для ручного выбора полей на фронте."""
    from openpyxl import load_workbook

    response = await download_feo_template(subsidy_id=None, db=db_session, current_user=superadmin_user)
    body = b"".join([chunk async for chunk in response.body_iterator])
    wb = load_workbook(io.BytesIO(body))
    ws = wb["Категории ФЭО"]
    headers = [c.value for c in ws[1]]
    assert headers == EXPECTED_TEMPLATE_HEADERS, headers
    assert len(headers) == 28


def test_resolve_fact_columns_explicit_overrides_header_detection():
    """Заголовки файла НЕ несут ни одного узнаваемого слова блока «Факт»
    (произвольные подписи «Col21».. «Col28») — автоопределение находит
    ничего, но явное сопоставление (`explicit`, как шлёт ручной маппинг)
    всё равно разрешает все 8 полей."""
    raw_headers = [f"col{i}" for i in range(1, 29)]
    explicit = {field: 20 + i for i, field in enumerate(FACT_FIELDS)}  # 20..27
    resolved = resolve_fact_columns(raw_headers, explicit)
    for i, field in enumerate(FACT_FIELDS):
        assert resolved[field] == 20 + i, (field, resolved)


def test_resolve_fact_columns_without_explicit_falls_back_to_header():
    """Без явного сопоставления (explicit=None/пустой) поведение — то же, что
    у старого detect_fact_columns_by_header (авто-путь /import не затронут)."""
    raw_headers = ["правильный статус", "факт: количество"]
    resolved = resolve_fact_columns(raw_headers, None)
    assert resolved["c_fact_status"] == 0
    assert resolved["c_fact_qty"] == 1
    assert resolved["c_fact_price"] is None


def _multipart_empty_request() -> Request:
    """Пустой multipart-запрос (без текстовых полей) — merge_form_over_query
    откатывается на значения, переданные напрямую как kwargs вызова
    (тот же приём, что test_feo_import_params_form_over_query.py)."""
    boundary = "testboundary123"
    body = f"--{boundary}--\r\n".encode("utf-8")

    async def receive():
        return {"type": "http.request", "body": body, "more_body": False}

    scope = {
        "type": "http",
        "method": "POST",
        "headers": [
            (b"content-type", f"multipart/form-data; boundary={boundary}".encode()),
            (b"content-length", str(len(body)).encode()),
        ],
    }
    return Request(scope, receive)


def _mk_xlsx_upload(headers, rows) -> UploadFile:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return UploadFile(file=bio, filename=f"import-{uuid.uuid4().hex[:6]}.xlsx")


@pytest.mark.asyncio
async def test_import_mapped_with_explicit_fact_and_need_level_mapping(db_session, superadmin_user):
    """Ручной маппинг /import-mapped: файл с ПРОИЗВОЛЬНЫМИ заголовками (не
    словами автоопределения) — явное сопоставление col_need_level/col_fact_*
    всё равно даёт нужность из файла (не дефолт 'likely') и has_fact_columns=
    True/fact_rows=1, план строится как обычно."""
    from app.routers.feo_import import import_feo_mapped
    from app.services.plan_need_level import NEED_LEVEL_LABELS

    subsidy = await _make_subsidy(db_session)
    # col idx:     0        1              2        3            4       5
    headers = ["Напр.", "Позиция", "Кол-во", "Цена", "Нужность!", "СтатусX"]
    row = [
        "Направление — тест факт-маппинга", "Позиция — тест факт-маппинга",
        "2", "500", NEED_LEVEL_LABELS["nice_to_have"], "Заключён договор",
    ]
    upload = _mk_xlsx_upload(headers, [row])
    request = _multipart_empty_request()
    try:
        result = await import_feo_mapped(
            request=request, file=upload, sheet_name="", header_row_offset=0,
            col_subsidy=-1, col_lvl2=0, col_lvl3=-1, col_lvl4=-1,
            col_lvl5=1, col_code=-1, col_appendix=-1, col_budget=-1,
            col_quantity=-1, col_unit=-1, col_item_amt=-1, col_active=-1,
            col_qty_lvl2=-1, col_qty_lvl3=-1, col_qty_lvl4=-1,
            col_unit_lvl2=-1, col_unit_lvl3=-1, col_unit_lvl4=-1,
            col_amt_lvl2=-1, col_amt_lvl3=-1, col_amt_lvl4=-1,
            col_feo_qty_lvl2=-1, col_feo_qty_lvl3=-1, col_feo_qty_lvl4=-1,
            col_feo_unit_lvl2=-1, col_feo_unit_lvl3=-1, col_feo_unit_lvl4=-1,
            col_feo_amount_lvl2=-1, col_feo_amount_lvl3=-1, col_feo_amount_lvl4=-1,
            col_feo_sum_lvl2=-1, col_feo_sum_lvl3=-1, col_feo_sum_lvl4=-1,
            col_plan_sum_lvl2=-1, col_plan_sum_lvl3=-1, col_plan_sum_lvl4=-1,
            col_item_price=-1,
            col_row_feo_qty=-1, col_row_feo_unit=-1, col_row_feo_price=-1, col_row_feo_sum=-1,
            col_row_plan_qty=2, col_row_plan_unit=-1, col_row_plan_price=3, col_row_plan_sum=-1,
            col_item_type=-1, col_comment=-1,
            col_need_level=4,
            col_fact_status=5, col_fact_qty=-1, col_fact_price=-1, col_fact_amount=-1,
            col_fact_paid=-1, col_fact_contracted=-1, col_fact_supplier=-1, col_fact_purchase_no=-1,
            default_subsidy_id=subsidy.id, dry_run=False,
            remap="", apply_remap=False, duplicate_resolutions="", item_type_decisions="",
            db=db_session, current_user=superadmin_user,
        )
        assert result["errors"] == [], result["errors"]
        assert result["has_fact_columns"] is True
        assert result["fact_rows"] == 1

        cats = await _get_categories(db_session, subsidy.id)
        lvl2_cat = next(c for c in cats if c.name == "Направление — тест факт-маппинга")
        items = await _get_items(db_session, lvl2_cat.id)
        assert len(items) == 1
        assert items[0].name == "Позиция — тест факт-маппинга"
        assert float(items[0].quantity) == 2
        assert float(items[0].unit_price) == 500
        assert getattr(items[0], "need_level", None) == "nice_to_have", (
            f"need_level={getattr(items[0], 'need_level', '<нет атрибута>')!r} — "
            "явно сопоставленная колонка «Нужность!» не прочиталась через col_need_level"
        )
    finally:
        await _cleanup_subsidy(db_session, subsidy.id)


@pytest.mark.asyncio
async def test_import_mapped_without_fact_mapping_reports_zero(db_session, superadmin_user):
    """Тот же маппинг, но без col_fact_* (все -1) и без колонки статуса в
    файле вовсе — has_fact_columns=False/fact_rows=0, план не меняется."""
    from app.routers.feo_import import import_feo_mapped

    subsidy = await _make_subsidy(db_session)
    headers = ["Напр.", "Позиция", "Кол-во", "Цена"]
    row = ["Направление — без факта", "Позиция — без факта", "1", "100"]
    upload = _mk_xlsx_upload(headers, [row])
    request = _multipart_empty_request()
    try:
        result = await import_feo_mapped(
            request=request, file=upload, sheet_name="", header_row_offset=0,
            col_subsidy=-1, col_lvl2=0, col_lvl3=-1, col_lvl4=-1,
            col_lvl5=1, col_code=-1, col_appendix=-1, col_budget=-1,
            col_quantity=-1, col_unit=-1, col_item_amt=-1, col_active=-1,
            col_qty_lvl2=-1, col_qty_lvl3=-1, col_qty_lvl4=-1,
            col_unit_lvl2=-1, col_unit_lvl3=-1, col_unit_lvl4=-1,
            col_amt_lvl2=-1, col_amt_lvl3=-1, col_amt_lvl4=-1,
            col_feo_qty_lvl2=-1, col_feo_qty_lvl3=-1, col_feo_qty_lvl4=-1,
            col_feo_unit_lvl2=-1, col_feo_unit_lvl3=-1, col_feo_unit_lvl4=-1,
            col_feo_amount_lvl2=-1, col_feo_amount_lvl3=-1, col_feo_amount_lvl4=-1,
            col_feo_sum_lvl2=-1, col_feo_sum_lvl3=-1, col_feo_sum_lvl4=-1,
            col_plan_sum_lvl2=-1, col_plan_sum_lvl3=-1, col_plan_sum_lvl4=-1,
            col_item_price=-1,
            col_row_feo_qty=-1, col_row_feo_unit=-1, col_row_feo_price=-1, col_row_feo_sum=-1,
            col_row_plan_qty=2, col_row_plan_unit=-1, col_row_plan_price=3, col_row_plan_sum=-1,
            col_item_type=-1, col_comment=-1,
            col_need_level=-1,
            col_fact_status=-1, col_fact_qty=-1, col_fact_price=-1, col_fact_amount=-1,
            col_fact_paid=-1, col_fact_contracted=-1, col_fact_supplier=-1, col_fact_purchase_no=-1,
            default_subsidy_id=subsidy.id, dry_run=False,
            remap="", apply_remap=False, duplicate_resolutions="", item_type_decisions="",
            db=db_session, current_user=superadmin_user,
        )
        assert result["errors"] == [], result["errors"]
        assert result["has_fact_columns"] is False
        assert result["fact_rows"] == 0
    finally:
        await _cleanup_subsidy(db_session, subsidy.id)
