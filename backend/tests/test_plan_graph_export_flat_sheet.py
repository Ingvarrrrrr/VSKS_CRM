"""Тесты владельца 08.10.2026 (план «план-график, замечания по экспорту»):
  п.1 — плановые позиции внутри статьи сортируются по наименьшему № закупки
        среди их закупок, фактические строки получают номера «N.k»;
  п.2 — лист «План закупок (по порядку)»: одна строка на закупленную позицию,
        без группировки, сортировка по № закупки глобально; хвост — позиции
        без закупок.
Строится через СИНТЕТИЧЕСКИЙ `data` (та же форма, что gather_live_plan_graph_data),
без БД — см. test_plan_graph_export_cascade.py."""
from types import SimpleNamespace

from app.services.plan_graph_export_columns import resolve_selected_keys
from app.services.plan_graph_export_rows import collect_plan_graph_rows
from app.services.plan_graph_export_flat_sheet import (
    DATA_FIRST_ROW, HEADER_ROW, write_flat_plan_graph_sheet,
)


def _cat(id_, level, parent_id=None, name="Статья"):
    return SimpleNamespace(id=id_, level=level, parent_id=parent_id, name=name, code=None, budget=None)


def _item(id_, feo_category_id, name, unit="шт", quantity=1, amount=1000):
    return SimpleNamespace(id=id_, feo_category_id=feo_category_id, name=name, unit=unit, quantity=quantity, amount=amount)


def _pi(purchase_id, name, unit_price=100.0, qty=1, raw_status="paid", contractor=""):
    return {
        "name": name, "unit": "шт", "qty": qty, "unit_price": unit_price,
        "total": unit_price * qty, "contractor": contractor, "raw_status": raw_status,
        "status": "Оплачено", "purchase_id": purchase_id, "act_number": "",
        "is_monthly": False, "monthly_count": None, "monthly_amount": 0.0,
    }


def _purchase(id_, purchase_number):
    return SimpleNamespace(
        id=id_, purchase_number=purchase_number, order_number="", registry_number=f"РЕЕ-{id_}",
        status="paid", subsidy_id=None, feo_category_id=None, item_name="", item_type="",
        unit="", planned_quantity=None, subject="", country_origin="", planned_unit_price=None,
        planned_total_price=None, total_nmck=None, contract_price=None, price_increase=None,
        purchase_method=None, purchase_basis=None, purchase_contract_type=None, framework_seq=None,
        contract_number="", contract_date=None, execution_term=None, execution_term_changed=None,
        delivery_date=None, contractor_id=None, reimbursement_user_id=None, service_note_by=None,
        responsible_person="", acceptance_doc_number="", acceptance_docs=None,
        payment_doc_number="", payment_doc_date=None, payment_amount=None, payment_federal=None,
        delivery_payment_amount=None, vat_applicable=False, vat_rate=None, vat_exemption_article=None,
        vat_mode=None, etp_url=None, region="", delivery_region="", delivery_location="",
        delivery_address="", final_unit_price=None, final_total_amount=None, contract_end_date=None,
        submission_deadline=None, commitment_quarter=None, planned_payment_month=None,
        stage_label=None, substatus=None,
    )


def _empty_ctx():
    return {
        "contractors": {}, "contractor_inns": {}, "subsidies": {}, "feo_categories": {},
        "payment_purposes": {}, "ru_map": {}, "sn_map": {}, "economy": {},
    }


def _scenario():
    """Статья (cat 3) с тремя плановыми позициями: item 20 (закупка №3),
    item 21 (закупка №1), item 22 — БЕЗ закупок."""
    cat1 = _cat(1, 1, None, name="Направление")
    cat2 = _cat(2, 2, 1, name="Тип")
    cat3 = _cat(3, 3, 2, name="Статья")
    item20 = _item(20, 3, "Позиция A")
    item21 = _item(21, 3, "Позиция B")
    item22 = _item(22, 3, "Позиция без закупок")

    purchase_rows_by_id = {101: _purchase(101, 3), 102: _purchase(102, 1)}
    data = {
        "cats": [cat1, cat2, cat3],
        "items_by_cat": {3: [item20, item21, item22]},
        "purchased_by_item": {
            20: [_pi(101, "Товар для A")],
            21: [_pi(102, "Товар для B")],
        },
        "purchased_by_cat": {},
        "unlinked_purchases": [],
        "purchase_rows_by_id": purchase_rows_by_id,
        "purchase_export_ctx": _empty_ctx(),
    }
    return data, item20, item21, item22


def test_items_sorted_by_min_purchase_number_and_subrows_numbered():
    data, item20, item21, item22 = _scenario()
    rows_info = collect_plan_graph_rows(data)

    ordered = rows_info["items_order_by_cat"][3]
    assert [it.id for it in ordered] == [21, 20, 22]  # закупка №1 раньше №3, без закупок — в конце

    assert rows_info["item_seq"][21] == 1
    assert rows_info["item_seq"][20] == 2
    assert rows_info["item_seq"][22] == 3

    row_b = rows_info["fact_rows_by_item"][21][0]
    assert row_b["num"] == "1.1"
    row_a = rows_info["fact_rows_by_item"][20][0]
    assert row_a["num"] == "2.1"

    assert rows_info["items_without_purchase"][0]["item"].id == 22
    assert rows_info["items_without_purchase"][0]["seq"] == 3


def test_new_columns_present_on_fact_row():
    data, *_ = _scenario()
    rows_info = collect_plan_graph_rows(data)
    row = rows_info["fact_rows_by_item"][20][0]
    assert row["composition"] == "Товар для A"
    assert row["unit_price"] == 100.0
    assert row["purchase_number"] == 3
    assert row["registry_number"] == "РЕЕ-101"


def test_flat_sheet_sorted_globally_by_purchase_number_with_tail():
    data, item20, item21, item22 = _scenario()
    selected = resolve_selected_keys(None)
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_flat_plan_graph_sheet(wb, data, selected, "http://example.test")

    ws = wb["План закупок (по порядку)"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[HEADER_ROW])}
    rows = list(ws.iter_rows(min_row=DATA_FIRST_ROW, values_only=False))

    def _val(row, label):
        return row[headers[label] - 1].value

    # Первая строка данных — закупка №1 (item B), затем закупка №3 (item A),
    # затем хвост "Позиция без закупок" без закупки.
    names = [_val(r, "Наименование") for r in rows]
    assert names[0] == "Позиция B"
    assert names[1] == "Позиция A"
    assert names[-1] == "Позиция без закупок"
    assert _val(rows[-1], "Статус") == "Не начато"
    assert _val(rows[0], "№") == "1.1"
    assert _val(rows[1], "№") == "2.1"

    # Автофильтр и закреплённая шапка — владелец п.2.
    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes == f"A{DATA_FIRST_ROW}"
