"""Тест реестра договоров на ФИКСТУРЕ БД с составом (PurchaseItem) — владелец/
QA 09.10.2026: предыдущие тесты листа «Реестр договоров»
(test_plan_graph_export_contracts_sheet.py) собирали `groups` вручную как
синтетические dict'и и НЕ проходили через gather_contracts_sheet_data с
реальными PurchaseItem — поэтому пропустили падение 500 на живых данных:
_items_with_stage() распаковывал cascade_by_stage() (5 значений) в 2
переменные (`d, pa = cascade_by_stage(...)`) → ValueError для ЛЮБОЙ закупки с
составом. Этот тест обязан проходить через настоящий PurchaseItem, чтобы
такое падение не повторилось.

Сценарии:
  1. Разовая закупка (paid) с 2 PurchaseItem — gather_contracts_sheet_data +
     write_contracts_sheet целиком; строки «Позиция» несут Сумма заявки/
     Поставлено/Оплачено по позиции, Σ = значению строки договора.
  2. Рамочная голова с одним заказом из 2 PurchaseItem — не падает, состав
     заказа выводится (без денег уровня заявки на позициях — это только для
     разовых, см. plan_graph_export_contracts_sheet.py).
  3. Путь роутера целиком (app.routers.subsidy_plan_graph_export.
     export_plan_graph_excel) на этой же фикстуре — gather_live_plan_graph_data
     → build_live_plan_graph_xlsx → gather_contracts_sheet_data →
     write_contracts_sheet, та же последовательность вызовов, что в роутере."""
import uuid

import pytest

from app.models.feo_category import FeoCategory
from app.models.purchase import Purchase
from app.models.purchase_item import PurchaseItem
from app.services.plan_graph_export_contracts_sheet import (
    DATA_FIRST_ROW, HEADER_ROW, gather_contracts_sheet_data, write_contracts_sheet,
)
from app.services.plan_graph_export_data import gather_live_plan_graph_data
from app.services.plan_graph_export_xlsx import build_live_plan_graph_xlsx


async def _make_subsidy(db_session, budget=5_000_000):
    from app.models.subsidy import Subsidy
    s = Subsidy(
        name=f"TestSubsidy-{uuid.uuid4().hex[:8]}", year=2026, budget=budget,
        require_planned_dates=False, status="approved",
    )
    db_session.add(s)
    await db_session.commit()
    await db_session.refresh(s)
    return s


async def _make_cat(db_session, subsidy_id, level, parent_id=None, name="Статья"):
    cat = FeoCategory(subsidy_id=subsidy_id, parent_id=parent_id, level=level, name=name)
    db_session.add(cat)
    await db_session.commit()
    await db_session.refresh(cat)
    return cat


async def _make_purchase(db_session, **kwargs) -> Purchase:
    defaults = dict(
        status="ordered", purchase_contract_type=None, parent_purchase_id=None,
        contract_id=None, contract_number=None, feo_category_id=None,
    )
    defaults.update(kwargs)
    p = Purchase(**defaults)
    db_session.add(p)
    await db_session.commit()
    await db_session.refresh(p)
    return p


async def _make_item(db_session, purchase_id, name, qty, unit_price, total_price) -> PurchaseItem:
    it = PurchaseItem(
        purchase_id=purchase_id, item_name=name, quantity=qty, unit="шт.",
        unit_price=unit_price, total_price=total_price,
    )
    db_session.add(it)
    await db_session.commit()
    await db_session.refresh(it)
    return it


@pytest.fixture
async def registry_fixture(db_session):
    sub = await _make_subsidy(db_session)
    cat = await _make_cat(db_session, sub.id, 3, name="Статья ФЭО")

    single = await _make_purchase(
        db_session, subsidy_id=sub.id, feo_category_id=cat.id,
        purchase_number=5001, status="paid", contract_number="ДОГ-5001",
        item_name="Разовая поставка с составом", planned_total_price=80000,
    )
    await _make_item(db_session, single.id, "Товар А", 3, 10000, 30000)
    await _make_item(db_session, single.id, "Товар Б", 1, 50000, 50000)

    head = await _make_purchase(
        db_session, subsidy_id=sub.id, feo_category_id=cat.id,
        purchase_number=5002, status="ordered",
        purchase_contract_type="framework_with_amount", contract_number="ДОГ-5002",
        item_name="Рамочный договор",
    )
    order = await _make_purchase(
        db_session, subsidy_id=sub.id, feo_category_id=cat.id,
        purchase_number=5003, order_number="1", status="paid",
        parent_purchase_id=head.id, item_name="Заказ 1", planned_total_price=20000,
    )
    await _make_item(db_session, order.id, "Деталь 1", 2, 5000, 10000)
    await _make_item(db_session, order.id, "Деталь 2", 1, 10000, 10000)

    return sub, single, head, order


@pytest.mark.asyncio
async def test_gather_does_not_raise_and_single_item_money_matches_head(db_session, registry_fixture):
    sub, single, head, order = registry_fixture

    # Прежде падало ValueError («too many values to unpack») в
    # _items_with_stage на ЛЮБОЙ закупке с PurchaseItem — главное в этом
    # тесте то, что вызов не бросает исключение.
    groups = await gather_contracts_sheet_data(db_session, sub.id)

    single_group = next(g for g in groups if g["head"]["purchase_number"] == 5001)
    items = single_group["head"]["items"]
    assert len(items) == 2
    assert {it["order_sum"] for it in items} == {30000.0, 50000.0}
    assert {it["order_delivered"] for it in items} == {30000.0, 50000.0}  # paid ⊇ delivered
    assert {it["order_paid"] for it in items} == {30000.0, 50000.0}
    assert sum(it["order_sum"] for it in items) == pytest.approx(single_group["head"]["contract_sum"])

    head_group = next(g for g in groups if g["head"]["purchase_number"] == 5002)
    assert len(head_group["orders"]) == 1
    assert len(head_group["orders"][0]["items"]) == 2


@pytest.mark.asyncio
async def test_write_contracts_sheet_from_real_fixture_item_rows_carry_order_money(db_session, registry_fixture):
    sub, single, head, order = registry_fixture
    groups = await gather_contracts_sheet_data(db_session, sub.id)

    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_contracts_sheet(wb, groups, title="test")
    ws = wb["Реестр договоров"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[HEADER_ROW])}

    def _rows_with(level, purchase_number=None):
        out = []
        for r in range(DATA_FIRST_ROW, ws.max_row + 1):
            if ws.cell(row=r, column=headers["Уровень"]).value != level:
                continue
            out.append(r)
        return out

    # Строка договора разовой закупки (5001) — найти по № закупки.
    single_row = next(
        r for r in range(DATA_FIRST_ROW, ws.max_row + 1)
        if ws.cell(row=r, column=headers["№ закупки"]).value == 5001
    )
    assert ws.cell(row=single_row, column=headers["Сумма заявки, ₽"]).value == ""

    item_rows = [r for r in range(single_row + 1, ws.max_row + 1)
                 if ws.cell(row=r, column=headers["Уровень"]).value == "Позиция"
                 and ws.cell(row=r, column=headers["№ договора"]).value == "ДОГ-5001"]
    assert len(item_rows) == 2
    item_sums = [ws.cell(row=r, column=headers["Сумма заявки, ₽"]).value for r in item_rows]
    assert sorted(item_sums) == [30000.0, 50000.0]
    assert sum(item_sums) == ws.cell(row=single_row, column=headers["Сумма договора, ₽"]).value

    # QA 09.10.2026: «Цена за ед., ₽» писалась с number_format General (189
    # ячеек на стенде) — теперь тот же "#,##0.00", что у денежных столбцов.
    price_col = headers["Цена за ед., ₽"]
    for r in item_rows:
        assert ws.cell(row=r, column=price_col).number_format == "#,##0.00"


@pytest.mark.asyncio
async def test_full_router_build_path_does_not_raise(db_session, registry_fixture):
    """Та же последовательность вызовов, что app.routers.
    subsidy_plan_graph_export.export_plan_graph_excel (без HTTP-слоя) —
    весь путь живого экспорта на фикстуре с составом не падает и даёт лист
    «Реестр договоров»."""
    sub, single, head, order = registry_fixture

    data = await gather_live_plan_graph_data(db_session, sub.id)
    wb = build_live_plan_graph_xlsx(sub, "http://example.test", data)
    contracts_data = await gather_contracts_sheet_data(db_session, sub.id)
    write_contracts_sheet(wb, contracts_data, title=f"{sub.name} ({sub.year}) — итого всего")

    assert "Реестр договоров" in wb.sheetnames

    import io
    buf = io.BytesIO()
    wb.save(buf)
    assert buf.tell() > 0
