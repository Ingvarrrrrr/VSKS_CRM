"""Сравнение редакций плана-графика (plan-graph versions compare).

Вынесено из app/routers/subsidies.py (Правило №5, рефакторинг 2026-09-07):
GET .../versions/compare, GET .../versions/compare.xlsx (2 редакции),
GET .../versions/export-multi.xlsx (N редакций + опционально текущая живая
ФЭО) + вспомогательные _build_compare_rows/_build_current_feo_tree/
_build_multi_compare_rows.

Node-matching (`_match_feo_nodes`, `_normalize_name`) — общий источник
истины с subsidy_plan_graph_versions.py (ПРАВИЛО №6), импортируются оттуда,
а не дублируются здесь.

Собственный APIRouter на префиксе /api/subsidies — регистрируется в
app/routes.py рядом с subsidies.router; ВАЖНО: `.../versions/compare*` и
`export-multi.xlsx` регистрируются ДО `.../versions/{version_id:int}` в
subsidy_plan_graph_versions.py (иначе catch-all int-конвертер перехватывает
литеральные сегменты "compare"/"export-multi.xlsx" — FastAPI/Starlette не
матчит "compare" на `{version_id:int}`, так что на практике порядок здесь не
критичен, но сохраняем для читаемости и на случай будущих правок пути).
"""
import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from app.utils.http import content_disposition as _content_disposition

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.auth.jwt import get_current_user, get_org_filter
from app.models.user import User
from app.models.subsidy import Subsidy
from app.models.feo_category import FeoCategory
from app.routers.subsidy_plan_graph_versions import _match_feo_nodes, _normalize_name

router = APIRouter(prefix="/api/subsidies", tags=["subsidies"])


def _build_compare_rows(tree_a: list, tree_b: list) -> list:
    """
    Build flat list of comparison rows for JSON/Excel output.
    Each row: {path, level, name_v1, name_v2, code, budget_v1, budget_v2,
               delta, delta_pct, status}
    status: unchanged | changed | new | removed | moved
    """
    match_result = _match_feo_nodes(tree_a, tree_b)
    match_by_a_id = {id(na): (nb, mtype) for na, nb, mtype in match_result["matches"]}
    matched_b_ids = {id(nb) for _, nb, _ in match_result["matches"]}

    rows = []

    def _walk_a(nodes, parent_path=()):
        for node in nodes:
            path = parent_path + (node.get("name", ""),)
            v1 = float(node.get("budget") or 0)
            mid = id(node)
            if mid in match_by_a_id:
                nb, mtype = match_by_a_id[mid]
                v2 = float(nb.get("budget") or 0)
                delta = v2 - v1
                delta_pct = round(delta / v1 * 100, 2) if v1 != 0 else None
                if mtype == "fallback":
                    status = "moved"
                elif abs(delta) < 0.01:
                    status = "unchanged"
                else:
                    status = "changed"
                rows.append({
                    "path": list(path),
                    "level": node.get("level", 0),
                    "name_v1": node.get("name"),
                    "name_v2": nb.get("name"),
                    "code": node.get("code") or nb.get("code"),
                    "budget_v1": round(v1, 2),
                    "budget_v2": round(v2, 2),
                    "delta": round(delta, 2),
                    "delta_pct": delta_pct,
                    "status": status,
                })
            else:
                rows.append({
                    "path": list(path),
                    "level": node.get("level", 0),
                    "name_v1": node.get("name"),
                    "name_v2": None,
                    "code": node.get("code"),
                    "budget_v1": round(v1, 2),
                    "budget_v2": 0.0,
                    "delta": round(-v1, 2),
                    "delta_pct": -100.0 if v1 != 0 else None,
                    "status": "removed",
                })
            _walk_a(node.get("children", []), path)

    def _walk_only_b(nodes):
        for node in nodes:
            if id(node) not in matched_b_ids:
                v2 = float(node.get("budget") or 0)
                rows.append({
                    "path": [node.get("name", "")],
                    "level": node.get("level", 0),
                    "name_v1": None,
                    "name_v2": node.get("name"),
                    "code": node.get("code"),
                    "budget_v1": 0.0,
                    "budget_v2": round(v2, 2),
                    "delta": round(v2, 2),
                    "delta_pct": None,
                    "status": "new",
                })

    _walk_a(tree_a)

    def _collect_b_nodes(nodes):
        result = []
        for n in nodes:
            result.append(n)
            result.extend(_collect_b_nodes(n.get("children", [])))
        return result

    all_b = _collect_b_nodes(tree_b)
    for nb in all_b:
        if id(nb) not in matched_b_ids:
            v2 = float(nb.get("budget") or 0)
            rows.append({
                "path": [nb.get("name", "")],
                "level": nb.get("level", 0),
                "name_v1": None,
                "name_v2": nb.get("name"),
                "code": nb.get("code"),
                "budget_v1": 0.0,
                "budget_v2": round(v2, 2),
                "delta": round(v2, 2),
                "delta_pct": None,
                "status": "new",
            })

    return rows


@router.get("/{subsidy_id}/plan-graph/versions/compare")
async def compare_plan_graph_versions(
    subsidy_id: int,
    v1: int = Query(..., description="ID первой версии"),
    v2: int = Query(..., description="ID второй версии"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Compare two plan-graph versions and return JSON diff."""
    from app.models.plan_graph_version import PlanGraphVersion as _PGV

    sub = (await db.execute(select(Subsidy).where(Subsidy.id == subsidy_id))).scalar_one_or_none()
    if not sub:
        raise HTTPException(404, "Субсидия не найдена")

    org_ids = get_org_filter(current_user)
    if org_ids is not None and sub.org_id not in org_ids:
        raise HTTPException(403, "Нет доступа к субсидии")

    ver1 = (await db.execute(
        select(_PGV).where(_PGV.id == v1, _PGV.subsidy_id == subsidy_id)
    )).scalar_one_or_none()
    ver2 = (await db.execute(
        select(_PGV).where(_PGV.id == v2, _PGV.subsidy_id == subsidy_id)
    )).scalar_one_or_none()

    if not ver1:
        raise HTTPException(404, f"Версия {v1} не найдена")
    if not ver2:
        raise HTTPException(404, f"Версия {v2} не найдена")

    snap1 = ver1.snapshot or {}
    snap2 = ver2.snapshot or {}
    tree1 = snap1.get("tree")
    tree2 = snap2.get("tree")

    if not tree1:
        raise HTTPException(
            422,
            f"Старая версия v{ver1.version_number} не содержит дерева ФЭО "
            f"(создана до Phase 12-05). Сравнение недоступно.",
        )
    if not tree2:
        raise HTTPException(
            422,
            f"Старая версия v{ver2.version_number} не содержит дерева ФЭО "
            f"(создана до Phase 12-05). Сравнение недоступно.",
        )

    rows = _build_compare_rows(tree1, tree2)

    def _ver_meta(ver, snap):
        return {
            "id": ver.id,
            "version_number": ver.version_number,
            "effective_date": ver.effective_date.isoformat() if ver.effective_date else snap.get("effective_date"),
            "note": ver.note,
            "created_at": ver.created_at.isoformat() if ver.created_at else None,
            "total_planned": snap.get("total_planned", 0),
        }

    return {
        "v1_meta": _ver_meta(ver1, snap1),
        "v2_meta": _ver_meta(ver2, snap2),
        "rows": rows,
    }


@router.get("/{subsidy_id}/plan-graph/versions/compare.xlsx")
async def compare_plan_graph_versions_excel(
    subsidy_id: int,
    v1: int = Query(..., description="ID первой версии"),
    v2: int = Query(..., description="ID второй версии"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Compare two plan-graph versions and return Excel diff."""
    if openpyxl is None:
        raise HTTPException(500, "openpyxl не установлен")

    from app.models.plan_graph_version import PlanGraphVersion as _PGV

    sub = (await db.execute(select(Subsidy).where(Subsidy.id == subsidy_id))).scalar_one_or_none()
    if not sub:
        raise HTTPException(404, "Субсидия не найдена")

    org_ids = get_org_filter(current_user)
    if org_ids is not None and sub.org_id not in org_ids:
        raise HTTPException(403, "Нет доступа к субсидии")

    ver1 = (await db.execute(
        select(_PGV).where(_PGV.id == v1, _PGV.subsidy_id == subsidy_id)
    )).scalar_one_or_none()
    ver2 = (await db.execute(
        select(_PGV).where(_PGV.id == v2, _PGV.subsidy_id == subsidy_id)
    )).scalar_one_or_none()

    if not ver1:
        raise HTTPException(404, f"Версия {v1} не найдена")
    if not ver2:
        raise HTTPException(404, f"Версия {v2} не найдена")

    snap1 = ver1.snapshot or {}
    snap2 = ver2.snapshot or {}
    tree1 = snap1.get("tree")
    tree2 = snap2.get("tree")

    if not tree1:
        raise HTTPException(
            422,
            f"Старая версия v{ver1.version_number} не содержит дерева ФЭО "
            f"(создана до Phase 12-05). Сравнение недоступно.",
        )
    if not tree2:
        raise HTTPException(
            422,
            f"Старая версия v{ver2.version_number} не содержит дерева ФЭО "
            f"(создана до Phase 12-05). Сравнение недоступно.",
        )

    rows = _build_compare_rows(tree1, tree2)

    # ── Build Excel ──────────────────────────────────────────────────────────
    HEADER_FILL   = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT   = Font(color="FFFFFF", bold=True, size=9)
    META_FONT     = Font(size=9, italic=True, color="374151")
    ITEM_FONT     = Font(size=9)
    CENTER_ALIGN  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_ALIGN    = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT_ALIGN   = Alignment(horizontal="right", vertical="center")
    THIN_BORDER   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    STATUS_FILLS = {
        "changed":   PatternFill("solid", fgColor="FFF3CD"),
        "new":       PatternFill("solid", fgColor="D4EDDA"),
        "removed":   PatternFill("solid", fgColor="F8D7DA"),
        "moved":     PatternFill("solid", fgColor="D1ECF1"),
        "unchanged": PatternFill(),
    }
    STATUS_LABELS = {
        "changed": "Изменено",
        "new": "Новое",
        "removed": "Удалено",
        "moved": "Перемещено",
        "unchanged": "Без изменений",
    }

    HEADERS = [
        "№", "Уровень", "Наименование", "Код",
        f"План v{ver1.version_number} (₽)", f"План v{ver2.version_number} (₽)",
        "Дельта (₽)", "Дельта (%)", "Статус",
    ]
    COL_WIDTHS = [5, 8, 50, 15, 18, 18, 18, 12, 15]
    n_cols = len(HEADERS)
    last_col = chr(64 + n_cols)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сравнение версий"

    # Title
    title_v1_date = ver1.effective_date.isoformat() if ver1.effective_date else (snap1.get("effective_date") or ver1.created_at.strftime("%Y-%m-%d") if ver1.created_at else "—")
    title_v2_date = ver2.effective_date.isoformat() if ver2.effective_date else (snap2.get("effective_date") or ver2.created_at.strftime("%Y-%m-%d") if ver2.created_at else "—")
    ws.append([f"СРАВНЕНИЕ ВЕРСИЙ — {sub.name}"] + [""] * (n_cols - 1))
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=12, color="1E3A5F")
    ws.merge_cells(f"A1:{last_col}1")
    title_cell.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    # Meta rows
    meta_texts = [
        f"v{ver1.version_number}: {title_v1_date}" + (f" — {ver1.note}" if ver1.note else ""),
        f"v{ver2.version_number}: {title_v2_date}" + (f" — {ver2.note}" if ver2.note else ""),
        f"Итого v{ver1.version_number}: {snap1.get('total_planned', 0):,.2f} ₽  |  "
        f"Итого v{ver2.version_number}: {snap2.get('total_planned', 0):,.2f} ₽  |  "
        f"Дельта: {(snap2.get('total_planned', 0) - snap1.get('total_planned', 0)):+,.2f} ₽",
    ]
    for i, mtext in enumerate(meta_texts):
        r = 2 + i
        ws.append([mtext] + [""] * (n_cols - 1))
        cell = ws.cell(row=r, column=1)
        cell.font = META_FONT
        cell.alignment = LEFT_ALIGN
        ws.merge_cells(f"A{r}:{last_col}{r}")
        ws.row_dimensions[r].height = 16

    # Column headers
    header_row = 2 + len(meta_texts)
    ws.append(HEADERS)
    for col_idx, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[header_row].height = 32
    ws.freeze_panes = f"A{header_row + 1}"

    # Data rows
    for seq_i, row in enumerate(rows, 1):
        level = row["level"]
        indent = "    " * (level - 1) if level > 0 else ""
        name = row["name_v2"] or row["name_v1"] or ""
        status = row["status"]
        fill = STATUS_FILLS.get(status, PatternFill())
        delta_pct_str = f"{row['delta_pct']:+.1f}%" if row["delta_pct"] is not None else "—"

        data_row_num = header_row + seq_i
        values = [
            seq_i, level, indent + name, row.get("code") or "",
            row["budget_v1"], row["budget_v2"],
            row["delta"], delta_pct_str,
            STATUS_LABELS.get(status, status),
        ]
        ws.append(values)
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=data_row_num, column=col_idx)
            cell.fill = fill
            cell.font = ITEM_FONT
            cell.border = THIN_BORDER
            if col_idx == 3:
                cell.alignment = LEFT_ALIGN
            elif col_idx >= 5:
                cell.alignment = RIGHT_ALIGN
            else:
                cell.alignment = CENTER_ALIGN
        ws.row_dimensions[data_row_num].height = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Субсидия_{subsidy_id}_сравнение_v{ver1.version_number}_и_v{ver2.version_number}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


# ── Multi-edition ФЭО compare (Phase 12-06) ──────────────────────────────────


async def _build_current_feo_tree(subsidy_id: int, db: AsyncSession) -> list:
    """Строит живое дерево FeoCategory для субсидии (аналог snapshot["tree"])."""
    cats = (await db.execute(
        select(FeoCategory).where(FeoCategory.subsidy_id == subsidy_id).order_by(FeoCategory.id)
    )).scalars().all()

    def _recurse(parent_id):
        nodes = []
        for c in cats:
            if c.parent_id == parent_id:
                nodes.append({
                    "id": c.id,
                    "name": c.name,
                    "level": c.level,
                    "code": c.code,
                    "budget": float(c.budget) if c.budget is not None else None,
                    "children": _recurse(c.id),
                })
        return nodes

    return _recurse(None)


def _build_multi_compare_rows(versions: list) -> list:
    """
    versions = [{"tree": [...]}, ...]  — уже в хронологическом порядке.
    Возвращает список строк: {"level", "name", "code", "budgets": [v0_val, v1_val, ...]}.
    Canonical key:
      - ("code", code)          если у узла есть code
      - ("path", level, path)   иначе (path = tuple нормализованных предков + имя узла)
    Ключи упорядочены по первому появлению (первая редакция задаёт порядок).
    """
    order: list = []            # список ключей в порядке появления
    meta: dict = {}             # key -> {"level", "name", "code"}
    budgets: dict = {}          # key -> {version_index: float}
    facts: dict = {}            # key -> {version_index: float} — фактически освоено (used_amount)

    def _dfs(node, path, ver_idx):
        norm_name = _normalize_name(node.get("name", ""))
        full_path = path + (norm_name,)

        code = (node.get("code") or "").strip()
        if code:
            key = ("code", code)
        else:
            key = ("path", node.get("level", 0), full_path)

        if key not in meta:
            order.append(key)
            meta[key] = {
                "level": node.get("level", 0),
                "name": node.get("name", ""),
                "code": code,
            }
        else:
            # Обновляем имя/код из последней редакции, если они непустые
            if node.get("name", ""):
                meta[key]["name"] = node["name"]
            if code:
                meta[key]["code"] = code

        if key not in budgets:
            budgets[key] = {}
        budgets[key][ver_idx] = float(node.get("budget") or 0)

        if key not in facts:
            facts[key] = {}
        facts[key][ver_idx] = float(node.get("used_amount") or 0)

        for child in node.get("children", []):
            _dfs(child, full_path, ver_idx)

    for ver_idx, ver in enumerate(versions):
        for root_node in ver.get("tree", []):
            _dfs(root_node, (), ver_idx)

    n_vers = len(versions)
    rows = []
    for key in order:
        m = meta[key]
        rows.append({
            "level": m["level"],
            "name": m["name"],
            "code": m["code"],
            "budgets": [budgets[key].get(i, 0.0) for i in range(n_vers)],
            "facts": [facts[key].get(i, 0.0) for i in range(n_vers)],
        })
    return rows


@router.get("/{subsidy_id}/plan-graph/versions/export-multi.xlsx")
async def export_plan_graph_versions_multi_excel(
    subsidy_id: int,
    ids: str = Query("", description="ID редакций через запятую"),
    include_current: bool = Query(False, description="Добавить колонку текущей живой ФЭО"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Экспорт N редакций ФЭО бок-о-бок в один Excel-файл с дельтами."""
    if openpyxl is None:
        raise HTTPException(500, "openpyxl не установлен")

    from app.models.plan_graph_version import PlanGraphVersion as _PGV
    from datetime import datetime as _dt

    # ── Авторизация ──────────────────────────────────────────────────────────
    sub = (await db.execute(select(Subsidy).where(Subsidy.id == subsidy_id))).scalar_one_or_none()
    if not sub:
        raise HTTPException(404, "Субсидия не найдена")

    org_ids = get_org_filter(current_user)
    if org_ids is not None and sub.org_id not in org_ids:
        raise HTTPException(403, "Нет доступа к субсидии")

    # ── Парсим id редакций ───────────────────────────────────────────────────
    try:
        id_list: list[int] = [int(x) for x in ids.split(",") if x.strip()]
    except ValueError:
        raise HTTPException(422, "Некорректный список id редакций")
    if not id_list and not include_current:
        raise HTTPException(422, "Выберите хотя бы одну редакцию")

    # ── Загружаем версии из БД ───────────────────────────────────────────────
    skipped_legacy: list[str] = []
    versions: list[dict] = []

    if id_list:
        ver_rows = (await db.execute(
            select(_PGV).where(_PGV.id.in_(id_list), _PGV.subsidy_id == subsidy_id)
        )).scalars().all()

        def _ver_date(v) -> date:
            if v.effective_date:
                return v.effective_date
            if v.created_at:
                return v.created_at.date()
            return date.today()

        # Хронологическая сортировка: сначала по дате, потом по номеру версии
        ver_rows_sorted = sorted(ver_rows, key=lambda v: (_ver_date(v), v.version_number))

        for ver in ver_rows_sorted:
            snap = ver.snapshot or {}
            tree = snap.get("tree")
            if not tree:
                skipped_legacy.append(f"v{ver.version_number}")
                continue
            d = _ver_date(ver)
            versions.append({
                "ver": ver,
                "label": f"v{ver.version_number}",
                "date_str": d.isoformat(),
                "date_display": d.strftime("%d.%m.%Y"),
                "note": ver.note,
                "tree": tree,
                "total_planned": snap.get("total_planned", 0),
            })

    # ── Текущая живая ФЭО ────────────────────────────────────────────────────
    if include_current:
        live_tree = await _build_current_feo_tree(subsidy_id, db)

        def _sum_tree(nodes):
            total = 0.0
            for n in nodes:
                total += float(n.get("budget") or 0)
                total += _sum_tree(n.get("children", []))
            return total

        today = date.today()
        versions.append({
            "ver": None,
            "label": "Текущая",
            "date_str": today.isoformat(),
            "date_display": today.strftime("%d.%m.%Y"),
            "note": None,
            "tree": live_tree,
            "total_planned": _sum_tree(live_tree),
        })

    if not versions:
        raise HTTPException(422, "Нет редакций с деревом ФЭО для выгрузки")

    rows = _build_multi_compare_rows(versions)

    # ── Стили (копируем из compare.xlsx) ────────────────────────────────────
    HEADER_FILL  = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT  = Font(color="FFFFFF", bold=True, size=9)
    META_FONT    = Font(size=9, italic=True, color="374151")
    ITEM_FONT    = Font(size=9)
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_ALIGN   = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT_ALIGN  = Alignment(horizontal="right", vertical="center")
    THIN_BORDER  = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    n_vers = len(versions)
    has_deltas = n_vers >= 2

    # ── Заголовки колонок ────────────────────────────────────────────────────
    fixed_headers = ["№", "Уровень", "Наименование", "Код"]
    plan_headers = [
        f"План ₽\n{v['date_display']}" for v in versions
    ]
    fact_headers = [
        f"Факт ₽\n{v['date_display']}" for v in versions
    ]
    delta_adj_headers = []
    if has_deltas:
        for i in range(1, n_vers):
            delta_adj_headers.append(
                f"Δ {versions[i]['date_display']}−{versions[i - 1]['date_display']}"
            )
    total_delta_headers = []
    if has_deltas:
        total_delta_headers = [
            f"Δ итог\n{versions[-1]['date_display']}−{versions[0]['date_display']}"
        ]

    all_headers = fixed_headers + plan_headers + fact_headers + delta_adj_headers + total_delta_headers
    n_cols = len(all_headers)
    last_col = chr(64 + n_cols) if n_cols <= 26 else (
        chr(64 + (n_cols - 1) // 26) + chr(64 + (n_cols - 1) % 26 + 1)
    )

    # Ширины колонок: фиксированные 5,8,50,15; план 18; факт 18; дельта 16
    col_widths = [5, 8, 50, 15] + [18] * n_vers + [18] * n_vers + [16] * (len(delta_adj_headers) + len(total_delta_headers))

    # ── Workbook ─────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Редакции ФЭО"

    # Строка-заголовок (merged)
    ws.append([f"РЕДАКЦИИ ФЭО — {sub.name}"] + [""] * (n_cols - 1))
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=12, color="1E3A5F")
    ws.merge_cells(f"A1:{last_col}1")
    title_cell.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 28

    # Мета-строки: по одной на каждую редакцию
    meta_lines = []
    for v in versions:
        line = f"{v['label']} — {v['date_display']}"
        if v["note"]:
            line += f" — {v['note']}"
        meta_lines.append(line)

    # Строка с итогами по редакциям
    totals_parts = [
        f"{v['label']}: {float(v['total_planned']):,.2f} ₽" for v in versions
    ]
    meta_lines.append("Итого: " + "  |  ".join(totals_parts))

    # Предупреждение о пропущенных legacy-версиях
    if skipped_legacy:
        meta_lines.append(
            f"⚠ Пропущены (нет дерева ФЭО): {', '.join(skipped_legacy)}"
        )

    for i, mtext in enumerate(meta_lines):
        r = 2 + i
        ws.append([mtext] + [""] * (n_cols - 1))
        cell = ws.cell(row=r, column=1)
        cell.font = META_FONT
        cell.alignment = LEFT_ALIGN
        ws.merge_cells(f"A{r}:{last_col}{r}")
        ws.row_dimensions[r].height = 16

    # Строка заголовков колонок
    header_row = 2 + len(meta_lines)
    ws.append(all_headers)
    for col_idx, (h, w) in enumerate(zip(all_headers, col_widths), 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[header_row].height = 32
    ws.freeze_panes = f"A{header_row + 1}"

    # ── Строки данных ────────────────────────────────────────────────────────
    GREEN_FONT = Font(size=9, color="1B7F3B")
    RED_FONT   = Font(size=9, color="B00020")

    for seq_i, row in enumerate(rows, 1):
        level = row["level"]
        indent = ("    " * (level - 1)) if level > 0 else ""
        name_str = indent + (row["name"] or "")
        code_str = row["code"] or ""
        bgets = row["budgets"]  # list[float] длиной n_vers
        fcts = row.get("facts") or [0.0] * n_vers  # фактически освоено по каждой редакции

        # Вычисляем дельта-значения
        adj_deltas = [bgets[i] - bgets[i - 1] for i in range(1, n_vers)] if has_deltas else []
        total_delta = [bgets[-1] - bgets[0]] if has_deltas else []

        row_values = [seq_i, level, name_str, code_str] + bgets + fcts + adj_deltas + total_delta
        data_row_num = header_row + seq_i
        ws.append(row_values)

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=data_row_num, column=col_idx)
            cell.border = THIN_BORDER

            if col_idx == 3:
                # Наименование — левое выравнивание
                cell.font = ITEM_FONT
                cell.alignment = LEFT_ALIGN
            elif col_idx <= 4:
                # №, Уровень, Код
                cell.font = ITEM_FONT
                cell.alignment = CENTER_ALIGN
            elif col_idx <= 4 + n_vers:
                # Плановые колонки
                cell.font = ITEM_FONT
                cell.alignment = RIGHT_ALIGN
            elif col_idx <= 4 + 2 * n_vers:
                # Фактические колонки (used_amount)
                cell.font = GREEN_FONT if isinstance(val, float) and val > 0.005 else ITEM_FONT
                cell.alignment = RIGHT_ALIGN
            else:
                # Дельта-колонки: цветной шрифт
                if isinstance(val, float) and val > 0.005:
                    cell.font = GREEN_FONT
                elif isinstance(val, float) and val < -0.005:
                    cell.font = RED_FONT
                else:
                    cell.font = ITEM_FONT
                cell.alignment = RIGHT_ALIGN

        ws.row_dimensions[data_row_num].height = 18

    # ── Сохраняем и отдаём ──────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"Субсидия_{subsidy_id}_ФЭО_редакции.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )
