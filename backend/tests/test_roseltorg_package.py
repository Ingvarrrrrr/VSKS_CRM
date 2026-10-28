# -*- coding: utf-8 -*-
"""Пакет документов для ЭТП «Росэлторг» — xlsx «Позиции по лоту».

Чистые функции build_roseltorg_lot_items_xlsx / build_roseltorg_zip_bytes,
offline, без БД/HTTP — по аналогии с tests/test_fabrikant_zip_names.py.
"""
import sys
import os
import zipfile
from io import BytesIO

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.services.documents.roseltorg_package import (
    build_roseltorg_lot_items_xlsx,
    build_roseltorg_zip_bytes,
    ROSELTORG_TEMPLATE_PATH,
    ROSELTORG_ZIP_FILE_NAME,
    ROSELTORG_MAX_ROWS,
)


def _orig_sheet2_hash():
    wb = openpyxl.load_workbook(ROSELTORG_TEMPLATE_PATH, data_only=False)
    ws2 = wb.worksheets[1]
    return ws2.cell(row=6, column=4).value


def test_two_rows_filled_correctly():
    rows = [
        {
            "okpd2": "13.25.13",
            "name": "Ткань хлопчатобумажная",
            "quantity": 10,
            "unit": "м",
            "unit_price": 123.456,
            "description": "Ткань для занавесей",
        },
        {
            "okpd2": "26.20.11",
            "name": "Ноутбук",
            "quantity": 2.5,
            "unit": "шт",
            "unit_price": 55000,
            "description": "",
        },
    ]
    out = build_roseltorg_lot_items_xlsx(rows)
    wb = openpyxl.load_workbook(BytesIO(out), data_only=False)
    ws = wb.worksheets[0]

    # Строка 4 — первая позиция
    assert ws.cell(row=4, column=3).value == "13.25.13"
    assert ws.cell(row=4, column=4).value == "Ткань хлопчатобумажная"
    assert ws.cell(row=4, column=5).value == 10
    assert ws.cell(row=4, column=6).value == "м"
    assert ws.cell(row=4, column=7).value == 123.46  # округление до 2 знаков
    assert ws.cell(row=4, column=8).value == "Ткань для занавесей"

    # Строка 5 — вторая позиция
    assert ws.cell(row=5, column=3).value == "26.20.11"
    assert ws.cell(row=5, column=4).value == "Ноутбук"
    assert ws.cell(row=5, column=5).value == 2.5
    assert ws.cell(row=5, column=6).value == "шт"
    assert ws.cell(row=5, column=7).value == 55000
    # Пустая строка: openpyxl не сохраняет "" в ячейке — после round-trip
    # читается как None (эквивалент пустого описания).
    assert ws.cell(row=5, column=8).value in ("", None)

    # Колонка A (ID) остаётся пустой
    assert ws.cell(row=4, column=1).value is None
    assert ws.cell(row=5, column=1).value is None


def test_column_b_formula_untouched():
    rows = [{"name": "Товар", "quantity": 1, "unit": "шт", "unit_price": 10, "description": "", "okpd2": ""}]
    out = build_roseltorg_lot_items_xlsx(rows)
    wb = openpyxl.load_workbook(BytesIO(out), data_only=False)
    ws = wb.worksheets[0]
    b4 = ws.cell(row=4, column=2).value
    assert isinstance(b4, str) and b4.startswith("=")


def test_sheet2_reference_untouched():
    rows = [{"name": "Товар", "quantity": 1, "unit": "шт", "unit_price": 10, "description": "", "okpd2": ""}]
    out = build_roseltorg_lot_items_xlsx(rows)
    wb = openpyxl.load_workbook(BytesIO(out), data_only=False)
    assert len(wb.worksheets) == 2
    ws2 = wb.worksheets[1]
    assert ws2.cell(row=6, column=4).value == _orig_sheet2_hash()


def test_default_unit_when_empty():
    rows = [{"name": "Товар", "quantity": 1, "unit": "", "unit_price": 10, "description": "", "okpd2": ""}]
    out = build_roseltorg_lot_items_xlsx(rows)
    wb = openpyxl.load_workbook(BytesIO(out), data_only=False)
    ws = wb.worksheets[0]
    assert ws.cell(row=4, column=6).value == "Условная единица"


def test_price_none_becomes_zero():
    rows = [{"name": "Товар", "quantity": 1, "unit": "шт", "unit_price": None, "description": "", "okpd2": ""}]
    out = build_roseltorg_lot_items_xlsx(rows)
    wb = openpyxl.load_workbook(BytesIO(out), data_only=False)
    ws = wb.worksheets[0]
    assert ws.cell(row=4, column=7).value == 0.0


def test_name_and_description_truncated():
    long_name = "Н" * 300
    long_desc = "О" * 1200
    rows = [{
        "name": long_name, "quantity": 1, "unit": "шт", "unit_price": 10,
        "description": long_desc, "okpd2": "",
    }]
    out = build_roseltorg_lot_items_xlsx(rows)
    wb = openpyxl.load_workbook(BytesIO(out), data_only=False)
    ws = wb.worksheets[0]
    assert ws.cell(row=4, column=4).value == long_name[:250]
    assert len(ws.cell(row=4, column=4).value) == 250
    assert ws.cell(row=4, column=8).value == long_desc[:1000]
    assert len(ws.cell(row=4, column=8).value) == 1000


def test_over_limit_raises_value_error():
    rows = [
        {"name": f"Товар {i}", "quantity": 1, "unit": "шт", "unit_price": 10, "description": "", "okpd2": ""}
        for i in range(ROSELTORG_MAX_ROWS + 1)
    ]
    with pytest.raises(ValueError):
        build_roseltorg_lot_items_xlsx(rows)


def test_zip_contains_lot_items_file_and_no_errors_file_when_no_errors():
    xlsx_bytes = build_roseltorg_lot_items_xlsx(
        [{"name": "Товар", "quantity": 1, "unit": "шт", "unit_price": 10, "description": "", "okpd2": ""}]
    )
    out = build_roseltorg_zip_bytes([(ROSELTORG_ZIP_FILE_NAME, xlsx_bytes)], [])
    with zipfile.ZipFile(BytesIO(out)) as zf:
        names = zf.namelist()
        assert ROSELTORG_ZIP_FILE_NAME in names
        assert "errors.txt" not in names


def test_zip_contains_errors_file_when_errors_present():
    xlsx_bytes = build_roseltorg_lot_items_xlsx(
        [{"name": "Товар", "quantity": 1, "unit": "шт", "unit_price": 10, "description": "", "okpd2": ""}]
    )
    out = build_roseltorg_zip_bytes([(ROSELTORG_ZIP_FILE_NAME, xlsx_bytes)], ["Позиция X: ошибка"])
    with zipfile.ZipFile(BytesIO(out)) as zf:
        names = zf.namelist()
        assert ROSELTORG_ZIP_FILE_NAME in names
        assert "errors.txt" in names
