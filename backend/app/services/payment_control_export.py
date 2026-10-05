"""Экспорт контрольного листа субсидии в Excel — план
.planning/quick/2026-10-06-statement-control/PLAN.md, п.3:
GET /api/subsidies/{id}/payment-control/export.xlsx.

ПРАВИЛО №6 — данные те же, что строит GET /payment-control
(app/services/subsidy_payment_control.py::build_payment_control), второй
расчёт здесь не заводится — выгрузка только форматирует уже готовый dict.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

_STATUS_LABELS = {
    "match": "Совпало",
    "amount_mismatch": "Расхождение суммы",
    "registry_only": "Не привязано",
    "duplicate": "Дубль",
    "not_reconciled": "Статья не сверяется",
    "declared_unconfirmed": "Отмечено, не подтверждено",
    "purchases_only": "Только в закупках",
}


def _header_row(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_payment_control_xlsx(data: dict, subsidy_name: str) -> BytesIO:
    wb = Workbook()

    # --- Лист 1: контрольный лист ------------------------------------
    ws = wb.active
    ws.title = "Контрольный лист"

    totals = data.get("totals", {})
    ws["A1"] = f"Сверка по выписке — {subsidy_name}"
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("A1:J1")
    ws["A2"] = (
        f"Привязано {totals.get('attached_count', 0)} из {totals.get('statement_count', 0)} платёжек; "
        f"не хватает {totals.get('unattached_count', 0)} на сумму "
        f"{totals.get('unattached_total', 0):,.2f} ₽".replace(",", " ")
    )
    ws.merge_cells("A2:J2")

    headers = ["№ п/п", "Дата", "Получатель", "ИНН", "Сумма, ₽", "Статья", "Статус сверки",
               "РЕЕ-закупка", "Номер закупки/заказа листа", "Разница, ₽"]
    _header_row(ws, 4, headers)

    row_idx = 5
    for r in data.get("rows", []):
        purchases = r.get("purchases") or []
        registry = ", ".join(p.get("registry_number") or str(p.get("id")) for p in purchases) or "—"
        sheet_ref = ", ".join(p.get("sheet_ref") for p in purchases if p.get("sheet_ref")) or "—"
        ws.cell(row=row_idx, column=1, value=r.get("payment_number") or "—")
        ws.cell(row=row_idx, column=2, value=r.get("payment_date") or "—")
        ws.cell(row=row_idx, column=3, value=r.get("payee_name") or "—")
        ws.cell(row=row_idx, column=4, value=r.get("payee_inn") or "—")
        ws.cell(row=row_idx, column=5, value=r.get("amount") or 0)
        ws.cell(row=row_idx, column=6, value=r.get("expense_name") or r.get("expense_code") or "—")
        ws.cell(row=row_idx, column=7, value=_STATUS_LABELS.get(r.get("status"), r.get("status")))
        ws.cell(row=row_idx, column=8, value=registry)
        ws.cell(row=row_idx, column=9, value=sheet_ref)
        ws.cell(row=row_idx, column=10, value=r.get("amount_diff") or 0)
        ws.cell(row=row_idx, column=5).number_format = "#,##0.00"
        ws.cell(row=row_idx, column=10).number_format = "#,##0.00"
        row_idx += 1

    for col, width in enumerate([12, 12, 28, 14, 14, 20, 20, 16, 26, 12], 1):
        ws.column_dimensions[chr(64 + col)].width = width

    # --- Лист 2: Не привязаны -----------------------------------------
    ws2 = wb.create_sheet("Не привязаны")
    _header_row(ws2, 1, ["№ п/п", "Дата", "Получатель", "ИНН", "Сумма, ₽", "Статья", "Подсказки"])
    row_idx = 2
    for r in data.get("rows", []):
        if r.get("status") != "registry_only":
            continue
        hints = r.get("near_miss") or []
        hint_text = "; ".join(
            f"{h.get('registry_number') or h.get('purchase_id')} (Δ {h.get('delta'):+.2f} ₽, {h.get('reason')})"
            for h in hints
        ) or "—"
        ws2.cell(row=row_idx, column=1, value=r.get("payment_number") or "—")
        ws2.cell(row=row_idx, column=2, value=r.get("payment_date") or "—")
        ws2.cell(row=row_idx, column=3, value=r.get("payee_name") or "—")
        ws2.cell(row=row_idx, column=4, value=r.get("payee_inn") or "—")
        ws2.cell(row=row_idx, column=5, value=r.get("amount") or 0)
        ws2.cell(row=row_idx, column=6, value=r.get("expense_name") or r.get("expense_code") or "—")
        ws2.cell(row=row_idx, column=7, value=hint_text)
        row_idx += 1
    for col, width in enumerate([12, 12, 28, 14, 14, 20, 50], 1):
        ws2.column_dimensions[chr(64 + col)].width = width

    # --- Лист 3: Отмечено, не подтверждено выпиской ---------------------
    ws3 = wb.create_sheet("Отмечено, не подтверждено")
    _header_row(ws3, 1, ["Документ", "Дата", "Сумма, ₽", "Закупка"])
    row_idx = 2
    for r in data.get("rows", []):
        if r.get("status") != "declared_unconfirmed":
            continue
        purchases = r.get("purchases") or []
        registry = ", ".join(p.get("registry_number") or str(p.get("id")) for p in purchases) or "—"
        ws3.cell(row=row_idx, column=1, value=r.get("payment_number") or "—")
        ws3.cell(row=row_idx, column=2, value=r.get("payment_date") or "—")
        ws3.cell(row=row_idx, column=3, value=r.get("amount") or 0)
        ws3.cell(row=row_idx, column=4, value=registry)
        row_idx += 1
    for col, width in enumerate([18, 12, 14, 28], 1):
        ws3.column_dimensions[chr(64 + col)].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
