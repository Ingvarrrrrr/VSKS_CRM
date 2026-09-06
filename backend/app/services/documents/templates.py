"""Резолв путей к .docx-шаблонам и pre-render проверки/подсказки.

Зависит только от doc_types.py (константы) — вниз по иерархии модулей."""


import os
import re
from typing import Optional

from fastapi import HTTPException

from .doc_types import (
    TEMPLATES_DIR,
    SUBSIDY_TEMPLATES_DIR,
    RECEIPTS_TABLE_MARKER,
    DOC_TYPES,
    DOC_TYPE_FALLBACK_FILES,
    _PURCHASE_METHOD_LABELS,
    _COMPETITIVE_FORM_LABELS,
    PURCHASE_METHOD_REQUIRED_DOC_TYPES,
    VAT_RATE_PRINTED_DOC_TYPES,
    VAT_EXEMPTION_ARTICLE_NPD,
    VAT_EXEMPTION_ARTICLE_GPH_INDIVIDUAL,
    _GPH_INDIVIDUAL_CONTRACT_FORMS,
)


def _sanitize_subject(text: str, max_len: int = 50) -> str:
    """Очищает произвольный текст для использования в имени файла.

    Убирает символы \\ / : * ? " < > | и переносы строк,
    схлопывает пробелы в «_», обрезает до max_len символов.
    """
    if not text:
        return ""
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', "", text)
    cleaned = re.sub(r'\s+', "_", cleaned.strip())
    return cleaned[:max_len]


def _insert_receipts_table_if_marker(
    doc,
    png_paths: list,
    *,
    cols: int = 2,
    col_width_cm: float = 8.0,
    img_width_cm: float = 7.5,
) -> bool:
    """Find paragraph containing RECEIPTS_TABLE_MARKER and replace it with
    a real 2-col docx table containing the receipt PNGs.

    Idempotent: if marker not present — no-op. If png_paths empty — marker
    is cleared but no table is inserted. Returns True if table was inserted.

    Why this exists: docxtpl InlineImage in a paragraph creates inline
    images that clip in narrow cells / overlap with surrounding text
    (Word renders very tall pictures behind subsequent paragraphs). A
    proper <w:tbl> with fixed col widths and an InlineShape inside each
    <w:tc> guarantees layout integrity.
    """
    from docx.shared import Cm as _Cm

    target_para = None
    for p in doc.paragraphs:
        if RECEIPTS_TABLE_MARKER in p.text:
            target_para = p
            break
    if target_para is None:
        return False

    if not png_paths:
        # marker present but no receipts — clear marker text and leave para
        for r in target_para.runs:
            if RECEIPTS_TABLE_MARKER in r.text:
                r.text = r.text.replace(RECEIPTS_TABLE_MARKER, "")
        return False

    rows = (len(png_paths) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.autofit = False
    for row in table.rows:
        for c_idx in range(cols):
            row.cells[c_idx].width = _Cm(col_width_cm)

    for idx, png_path in enumerate(png_paths):
        r, c = divmod(idx, cols)
        cell = table.rows[r].cells[c]
        # Cell creation always gives one empty paragraph — clear its runs
        # then add the picture inside a fresh run.
        para = cell.paragraphs[0]
        for run in list(para.runs):
            run.text = ""
        para.add_run().add_picture(png_path, width=_Cm(img_width_cm))

    # Move the freshly appended table from end-of-doc to the marker location,
    # then delete the marker paragraph.
    target_xml = target_para._p
    table_xml = table._tbl
    target_xml.addnext(table_xml)
    target_xml.getparent().remove(target_xml)
    return True


def _resolve_doc_template_path(doc_type: str, subsidy_id: Optional[int]) -> tuple[str, str, str]:
    """Резолвит путь к файлу шаблона doc_type с единым приоритетом:

      1) субсидийный override (uploads/templates/subsidies/<id>/<doc_type>.docx),
      2) глобальный файл backend/templates/<doc_type>.docx,
      3) DOC_TYPE_FALLBACK_FILES[doc_type], если глобальный файл отсутствует.

    Используется и для основного документа, и для приклеиваемой методички —
    единственное место с этой логикой, чтобы приоритет «сначала субсидия,
    иначе глобальный» не разъезжался между местами вызова.

    Возвращает (template_path, template_file, filename_base). template_path
    может не существовать — вызывающий код обязан проверить os.path.exists.
    """
    template_file, filename_base = DOC_TYPES[doc_type]
    template_path = os.path.join(TEMPLATES_DIR, template_file)

    if not os.path.exists(template_path):
        fallback = DOC_TYPE_FALLBACK_FILES.get(doc_type)
        if fallback:
            fallback_path = os.path.join(TEMPLATES_DIR, fallback)
            if os.path.exists(fallback_path):
                template_path = fallback_path
                template_file = fallback

    if subsidy_id:
        subsidy_template = os.path.join(
            SUBSIDY_TEMPLATES_DIR, "subsidies", str(subsidy_id), f"{doc_type}.docx"
        )
        if os.path.exists(subsidy_template):
            template_path = subsidy_template

    return template_path, template_file, filename_base


def _purchase_method_label(p) -> str:
    """Подпись способа закупки для п.2 приказа/листа согласования.

    Если способ — конкурентная процедура и заполнена уточняющая форма
    (competitive_form), подписью становится конкретная форма («Запрос цен» /
    «Аукцион (редукцион)» / «Конкурс»), а не общее «Конкурсная процедура».
    Форма не выбрана — старое поведение (общая подпись способа).
    """
    method = p.purchase_method or ""
    if method == "competitive":
        form = getattr(p, "competitive_form", None)
        form_label = _COMPETITIVE_FORM_LABELS.get(form or "")
        if form_label:
            return form_label
    return _PURCHASE_METHOD_LABELS.get(method, method)


def _resolve_vat_exemption_basis(p) -> str:
    """Основание «НДС не облагается» — введённое человеком ИЛИ определённое
    автоматически по данным закупки. Возвращает "" если основания нет вовсе.

    Приоритет: то, что ввёл человек в поле «Статья НК РФ» (p.vat_exemption_article),
    ВСЕГДА побеждает автоподстановку — автоматика не имеет права затирать
    ручной ввод.

    Автоопределение — только для двух случаев, где основание системе уже
    известно из самих данных закупки:
      - контрагент — самозанятый (Contractor.org_type == 'Самозанятый',
        это уже существующий в проекте, поддерживаемый признак: его
        выставляют вручную в карточке контрагента и автоматически по ответу
        реестра ФНС npd.nalog.ru при проверке ИНН, см. contractors.py::
        _check_npd_status). ИНН длиной 12 цифр сюда НЕ годится — так же
        выглядит и ИП, у которого своего освобождения от НДС может не быть;
      - договор ГПХ с физлицом (Purchase.contract_form in
        _GPH_INDIVIDUAL_CONTRACT_FORMS) — физлицо не является плательщиком НДС.

    Для всех остальных контрагентов (юрлицо/ИП без явно введённой статьи)
    возвращает "" — вызывающий код (см. _require_vat_rate_for_doc) обязан
    расценивать это как «основание не указано» и отказать в генерации.
    """
    manual = (getattr(p, "vat_exemption_article", None) or "").strip()
    if manual:
        return manual
    contractor = getattr(p, "contractor", None)
    if contractor is not None and (getattr(contractor, "org_type", None) or "") == "Самозанятый":
        return VAT_EXEMPTION_ARTICLE_NPD
    if (getattr(p, "contract_form", None) or "") in _GPH_INDIVIDUAL_CONTRACT_FORMS:
        return VAT_EXEMPTION_ARTICLE_GPH_INDIVIDUAL
    return ""


def _require_vat_rate_for_doc(p, doc_type: str) -> None:
    """422, если документ печатает данные о НДС, а они противоречивы или не
    заполнены — единая проверка на оба случая (владелец, 2026-09-04, «вторая
    половина правила»):

      1) НДС облагается (vat_applicable=True), а ставка не указана — раньше
         тихо подставлялась как 20% (`p.vat_rate or 20`, тот же `or` заодно
         подменял и явный 0% на 20%). Ноль — валидная ставка и НЕ считается
         «не указана»; «не указана» — это None.
      2) НДС НЕ облагается (vat_applicable=False, в модели это ещё и default,
         см. Purchase.vat_applicable), а основание освобождения не указано и
         не может быть определено автоматически (см. _resolve_vat_exemption_basis).
         На проде это оказалось массовым: vat_applicable=False почти нигде не
         выставляли осознанно, а документ печатал «НДС не облагается» без
         всякого основания — то же самое «система придумывает за пользователя»,
         только не про ставку, а про сам факт освобождения.

    Оба случая — один и тот же принцип («система ничего не придумывает про
    НДС сама»), поэтому проверка одна, не заводим второй механизм — только
    расширяем эту функцию.
    """
    if doc_type not in VAT_RATE_PRINTED_DOC_TYPES:
        return
    if p.vat_applicable:
        if p.vat_rate is not None:
            return
        raise HTTPException(
            422,
            detail={
                "code": "VAT_RATE_REQUIRED",
                "message": "Не указана ставка НДС",
                "hint": (
                    "Закупка отмечена как облагаемая НДС, но ставка НДС не заполнена. "
                    "Система не вправе подставлять её сама — ставка НДС разная у "
                    "разных поставщиков и договоров, а не всегда 20%.\n\n"
                    "Откройте карточку закупки и заполните поле «Ставка НДС» "
                    "(0, 10 или 20 — как указано у поставщика). Если НДС на самом "
                    "деле не облагается — снимите отметку «Облагается НДС» и "
                    "укажите статью НК РФ."
                ),
                "missing_fields": ["vat_rate"],
                "doc_type": doc_type,
            },
        )
    if _resolve_vat_exemption_basis(p):
        return
    raise HTTPException(
        422,
        detail={
            "code": "VAT_EXEMPTION_ARTICLE_REQUIRED",
            "message": "НДС отмечен как «не облагается», но не указано основание",
            "hint": _vat_exemption_missing_hint(p),
            "missing_fields": ["vat_exemption_article"],
            "doc_type": doc_type,
        },
    )


def _vat_exemption_missing_hint(p) -> str:
    """Текст подсказки для отказа VAT_EXEMPTION_ARTICLE_REQUIRED — единое
    место, используется и в _require_vat_rate_for_doc, и в per-doc проверке
    fabrikant_contract_project (см. render_fabrikant_package_files), чтобы
    формулировка не разъезжалась по двум местам.

    Владелец (2026-09-04): «человек должен из текста понять, что делать, а
    не искать» — если у контрагента закупки ИНН из 12 цифр, а тип не
    отмечен «Самозанятый», подсказка прямо называет вероятную причину и
    конкретное действие (карточка контрагента → «Проверить в реестре
    самозанятых»), а не просто «заполните поле».
    """
    parts = [
        "Закупка отмечена как НЕ облагаемая НДС, но основание освобождения "
        "не заполнено — система не вправе печатать «НДС не облагается» без "
        "основания. Для самозанятых исполнителей и договоров ГПХ с физлицом "
        "основание подставляется автоматически; для остальных — нужно указать "
        "его самим."
    ]
    contractor = getattr(p, "contractor", None)
    contractor_inn = (getattr(contractor, "inn", None) or "").strip() if contractor else ""
    contractor_org_type = (getattr(contractor, "org_type", None) or "").strip() if contractor else ""
    if len(contractor_inn) == 12 and contractor_org_type != "Самозанятый":
        parts.append(
            "У контрагента этой закупки ИНН из 12 цифр, а тип организации не "
            "отмечен как «Самозанятый» — это похоже на самозанятого или на ИП. "
            "Если это самозанятый: откройте карточку контрагента и нажмите "
            "«Проверить в реестре самозанятых» (или вручную выберите тип "
            "«Самозанятый» в поле «Форма организации») — основание для документа "
            "подставится само. Если это ИП на общей системе налогообложения — он "
            "вправе быть плательщиком НДС, тогда отметьте «Облагается НДС» и "
            "укажите ставку."
        )
    parts.append(
        "Выберите один из двух вариантов в карточке закупки:\n"
        "— укажите статью НК РФ, дающую освобождение от НДС, в поле «Статья НК РФ»;\n"
        "— либо отметьте «Облагается НДС» и заполните ставку (0, 10 или 20)."
    )
    return "\n\n".join(parts)


def _require_purchase_method_for_doc(p, doc_type: str) -> None:
    """422, если запрошен приказ о закупке/лист согласования, а способ закупки не выбран."""
    if doc_type not in PURCHASE_METHOD_REQUIRED_DOC_TYPES:
        return
    if p.purchase_method:
        return
    raise HTTPException(
        422,
        detail={
            "code": "PURCHASE_METHOD_REQUIRED",
            "message": "Не выбран способ закупки",
            "hint": (
                "Способ закупки — обязательный пункт приказа о закупке и листа "
                "согласования: без него документ считается недействительным. "
                "Откройте карточку закупки и заполните поле «Способ закупки». "
                "Если это конкурентная процедура, укажите ещё и её форму — "
                "запрос цен, аукцион или конкурс."
            ),
            "missing_fields": ["purchase_method"],
            "doc_type": doc_type,
        },
    )
