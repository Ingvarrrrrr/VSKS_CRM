"""GET /fact-import/template?layout=stages — ВТОРОЙ шаблон импорта факта, в
раскладке листа GoodsService владельца: стадия закупки в НЕСКОЛЬКИХ колонках
Да/Нет («Закупка проведена»/«Договор заключен»/«Товар заказан»/«Товар
поставлен»/«Оплачено») вместо одной «Правильный статус» (ту раскладку строит
template.py, layout='single').

Заголовки — РОВНО те строки, на которые матчит columns.py::_map_columns_
format (см. test_fact_import_stage_columns.py::_gs_header — тот же текст
буква-в-букву, без пустых колонок между ними: разбор ищет колонки по тексту
заголовка, не по фиксированной букве Excel). «Аванс (да/нет)» — columns.py
распознаёт любой заголовок, содержащий «аванс» (_map_columns_format), как и
в формате ХО (template.py) — включена сюда же, второй список допустимых
значений не завели (ПРАВИЛО №6).

Предзаполнение строк ТЕКУЩИМ планом субсидии — тот же plan_catalog.
load_plan_catalog и тот же хелпер entry_name_and_levels, что и template.py
(ПРАВИЛО №6 — один разбор каталога на строку шаблона, не второй обход
дерева ФЭО).
"""
from __future__ import annotations

from sqlalchemy.ext.asyncio import AsyncSession

from app.services.historical_fact_import import columns as columns_mod
from app.services.historical_fact_import.template_common import entry_name_and_levels
from app.services.plan_catalog import load_plan_catalog

# Заголовки листа GoodsService — порядок компактный (без пустых колонок,
# как у владельца между ними; columns.py ищет поля ПО ТЕКСТУ заголовка, не
# по фиксированному индексу/букве — порядок колонок на разбор не влияет).
_HEADER = [
    "Номер закупки",                                                              # 0  purchase_no
    "Наименование товара",                                                        # 1  plan_item_name
    "Контрагент",                                                                 # 2  supplier
    "Плановое количество",                                                        # 3  plan_qty
    "Ед. изм. (ОКЕИ)",                                                            # 4  unit
    "Плановая цена",                                                              # 5  plan_price
    "Плановая стоимость за объём",                                                # 6  plan_amount
    "Итоговая цена  за единицу товара",                                           # 7  fact_price
    "Итоговая цена за объём",                                                     # 8  fact_amount
    "Товары или Услуги",                                                          # 9  item_type
    "Направление расходов по категориям в соответсвии с ФЭО",                     # 10 path_l2
    "Тип расходов в соответсвии с ФЭО Конкретизированный",                        # 11 path_l3
    "Направление расходов из приложений к ФЭО",                                   # 12 path_l4
    "Тип расходов конкретизированный (сам придумываю общий тип для cебя)",        # 13 path_l5
    "Закупка проведена",                                                          # 14 stage_purchase_done
    "Договор заключен (устно или письменно)",                                     # 15 stage_contracted
    "Товар заказан",                                                              # 16 stage_ordered
    "Товар поставлен",                                                            # 17 stage_delivered
    "Оплачено",                                                                   # 18 stage_paid (флаг — рядом "Оплачено в деньгах")
    "Оплачено в деньгах",                                                         # 19 paid (сумма платежа)
    "Аванс (да/нет)",                                                             # 20 advance
]

_COL_PURCHASE_NO = 0
_COL_PLAN_ITEM_NAME = 1
_COL_PLAN_QTY = 3
_COL_PLAN_PRICE = 5
_COL_PLAN_AMOUNT = 6
_COL_PATH_L2 = 10
_COL_PATH_L5 = 13  # l2..l5 — 4 подряд идущие колонки

# Колонки со списком Да/Нет (не блокирующий выбор — showErrorMessage=False,
# как и остальные dropdown'ы шаблона, lesson feedback_excel_template_dropdowns).
_YES_NO_COLS = (14, 15, 16, 17, 18, 20)


async def build_stages_template_workbook(db: AsyncSession, subsidy_id: int) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    catalog = await load_plan_catalog(db, subsidy_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Импорт факта"
    ws.append(_HEADER)

    for entry in catalog:
        plan_item_name, levels = entry_name_and_levels(entry, depth=4)  # l2/l3/l4/l5

        row = [None] * len(_HEADER)
        row[_COL_PLAN_ITEM_NAME] = plan_item_name
        row[_COL_PLAN_QTY] = entry.get("quantity")
        row[_COL_PLAN_PRICE] = entry.get("unit_price")
        row[_COL_PLAN_AMOUNT] = entry.get("amount")
        row[_COL_PATH_L2:_COL_PATH_L5 + 1] = levels
        ws.append(row)

    # «Справочники» — один список Да/Нет, на который ссылаются ВСЕ колонки
    # стадий + «Аванс» (ПРАВИЛО №6 — одна строка-источник, не копия текста
    # «Да»/«Нет» в каждом DataValidation).
    ref = wb.create_sheet("Справочники")
    ref.append(["Да/Нет"])
    ref.append(["Да"])
    ref.append(["Нет"])

    last_row = max(ws.max_row, 2)
    dv = DataValidation(
        type="list",
        formula1="Справочники!$A$2:$A$3",
        allow_blank=True,
        showErrorMessage=False,  # не блокирует ручной ввод — lesson feedback_excel_template_dropdowns
    )
    for col in _YES_NO_COLS:
        letter = columns_mod._col_letter(col)
        dv.add(f"{letter}2:{letter}{last_row + 200}")  # запас строк под дозапись вручную
    ws.add_data_validation(dv)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
