"""GET /api/feo-categories/export — экспорт дерева категорий ФЭО одной
субсидии в Excel.

Вынесено из app/routers/feo_import.py (Правило №5, разрезание 1088-строчного
роутера) без изменения поведения — самодостаточен (не связан с импортом,
только читает дерево категорий и суммы закупок). Путь статичный (без
{cat_id}) — регистрируется в app/routes.py рядом с feo_import.router и ДО
feo_categories.router (несёт catch-all GET/PUT/DELETE "/{cat_id}"), как и
остальные соседи feo_import_*.
"""
from io import BytesIO

from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    Workbook = None

from app.database import get_db
from app.models.feo_category import FeoCategory
from app.auth.permissions import require_tab
from app.utils.http import content_disposition

router = APIRouter(prefix="/api/feo-categories", tags=["feo_categories"])


_content_disposition = content_disposition


@router.get("/export")
async def export_feo_to_excel(
    subsidy_id: int = Query(...),
    db: AsyncSession = Depends(get_db),
    _=Depends(require_tab('feo_categories')),
):
    """Экспорт дерева категорий ФЭО в Excel."""
    if Workbook is None:
        raise HTTPException(500, "openpyxl не установлен")
    from app.models.subsidy import Subsidy
    from app.models.purchase import Purchase

    sub = (await db.execute(select(Subsidy).where(Subsidy.id == subsidy_id))).scalar_one_or_none()
    if not sub:
        raise HTTPException(404, "Субсидия не найдена")

    cats = (await db.execute(
        select(FeoCategory).where(FeoCategory.subsidy_id == subsidy_id).order_by(FeoCategory.sort_order.nulls_last(), FeoCategory.id)
    )).scalars().all()

    # purchase totals
    pt_rows = (await db.execute(
        select(Purchase.feo_category_id, func.coalesce(func.sum(Purchase.planned_total_price), 0).label("total"))
        .where(Purchase.subsidy_id == subsidy_id)
        .where(Purchase.feo_category_id.isnot(None))
        .group_by(Purchase.feo_category_id)
    )).all()
    purchase_totals = {r.feo_category_id: float(r.total) for r in pt_rows}

    # Build tree
    by_id = {c.id: {"cat": c, "children": []} for c in cats}
    roots = []
    for c in cats:
        if c.parent_id and c.parent_id in by_id:
            by_id[c.parent_id]["children"].append(by_id[c.id])
        else:
            roots.append(by_id[c.id])

    def calc_budget(node):
        c = node["cat"]
        if not node["children"]:
            return float(c.budget) if c.budget is not None else None
        child_sum = sum(v for ch in node["children"] if (v := calc_budget(ch)) is not None)
        return child_sum if any(calc_budget(ch) is not None for ch in node["children"]) else (float(c.budget) if c.budget is not None else None)

    def calc_purchased(node):
        c = node["cat"]
        if not node["children"]:
            return purchase_totals.get(c.id, 0.0)
        return sum(calc_purchased(ch) for ch in node["children"])

    wb = Workbook()
    ws = wb.active
    ws.title = sub.name[:31]

    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    l1_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    l2_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    bold_font = Font(bold=True)
    semi_font = Font(bold=False)

    ws.append(["Наименование", "Код", "Прил.", "Финансирование по ФЭО (₽)", "Фактически запланировано (₽)", "Активна"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    def write_node(node, depth=0):
        c = node["cat"]
        indent = "  " * depth
        budget = calc_budget(node)
        purchased = calc_purchased(node)
        row = [
            indent + c.name,
            c.code or "",
            c.appendix or "",
            budget if budget is not None else "",
            purchased if purchased > 0 else "",
            "Да" if c.is_active else "Нет",
        ]
        ws.append(row)
        r = ws.max_row
        fill = l1_fill if c.level == 1 else (l2_fill if c.level == 2 else None)
        font = bold_font if c.level == 1 else (semi_font)
        for col in range(1, 7):
            cell = ws.cell(r, col)
            if fill:
                cell.fill = fill
            cell.font = font
            if col in (4, 5) and isinstance(cell.value, float):
                cell.number_format = '#,##0.00'
        for ch in node["children"]:
            write_node(ch, depth + 1)

    for root in roots:
        write_node(root)

    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 10
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_name = sub.name.replace(" ", "_").replace("/", "-")[:40]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(f"ФЭО_{safe_name}.xlsx")},
    )
