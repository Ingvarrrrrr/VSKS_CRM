"""app.services.plan_graph_export_xlsx — Задачи B/C/D экспорта плана-графика
(владелец 07.10.2026): каскад сумм по стадиям накопительный, отменённые
закупки не считаются, выбор столбцов, лист «Сводная» — формулы в D/G/I и
ИТОГО. Строится через СИНТЕТИЧЕСКИЙ `data`-словарь (та же форма, что отдаёт
gather_live_plan_graph_data) — без БД, проверяется только рендер книги."""
from types import SimpleNamespace

import pytest

from app.services.plan_graph_export_columns import resolve_selected_keys
from app.services.plan_graph_export_xlsx import (
    DATA_FIRST_ROW, HEADER_ROW, build_live_plan_graph_xlsx, cascade_by_stage,
)


def _make_sub():
    return SimpleNamespace(name="Тестовая субсидия", year=2026)


def _make_cat(id_, level, parent_id=None, name="Статья", code=None, budget=None):
    return SimpleNamespace(id=id_, level=level, parent_id=parent_id, name=name, code=code, budget=budget)


def _make_item(id_, feo_category_id, name="Позиция", unit="шт", quantity=2, amount=1000):
    return SimpleNamespace(id=id_, feo_category_id=feo_category_id, name=name, unit=unit, quantity=quantity, amount=amount)


def _base_data():
    cat1 = _make_cat(1, 1, None, name="Направление", code="01")
    cat2 = _make_cat(2, 2, 1, name="Тип расходов")
    cat3 = _make_cat(3, 3, 2, name="Статья ФЭО")
    item = _make_item(10, 3)

    return {
        "cats": [cat1, cat2, cat3],
        "item_ids": [10],
        "cat_ids": [1, 2, 3],
        "items_by_cat": {3: [item]},
        "used_map": {10: 1000.0},
        "contractor_map": {10: "ООО Ромашка"},
        "status_sums_map": {10: {"paid": 1000.0}},
        "purchased_by_item": {10: [{
            "name": "Купленное", "unit": "шт", "qty": 2, "unit_price": 500.0,
            "total": 1000.0, "contractor": "ООО Ромашка", "raw_status": "paid",
            "status": "Оплачено", "purchase_id": 55, "act_number": "A-1",
        }]},
        "cat_status_map": {},
        "cat_monthly_map": {},
        "cat_contractor_map": {},
        "purchased_by_cat": {},
        "unlinked_purchases": [{
            "name": "Без категории (отменена)", "unit": "шт", "qty": 1, "unit_price": 500.0,
            "total": 500.0, "contractor": "X", "raw_status": "cancelled",
            "status": "Отменено", "purchase_id": 77, "act_number": "",
            "is_monthly": False, "monthly_count": None, "monthly_amount": 0.0,
        }],
    }


def test_cascade_by_stage_is_cumulative_for_paid():
    planned, contract, ordered, delivered, paid = cascade_by_stage("paid", 1000.0)
    assert (planned, contract, ordered, delivered, paid) == (1000.0, 1000.0, 1000.0, 1000.0, 1000.0)


def test_cascade_by_stage_partial_stage_ordered():
    planned, contract, ordered, delivered, paid = cascade_by_stage("ordered", 500.0)
    assert (planned, contract, ordered) == (500.0, 500.0, 500.0)
    assert (delivered, paid) == (0.0, 0.0)


def test_cascade_by_stage_cancelled_is_all_zero():
    assert cascade_by_stage("cancelled", 999.0) == (0.0, 0.0, 0.0, 0.0, 0.0)


def test_build_live_plan_graph_xlsx_item_row_cumulative_and_cancelled_excluded():
    wb = build_live_plan_graph_xlsx(_make_sub(), "http://example.test", _base_data())
    ws = wb["План закупок (по направлениям)"]

    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[HEADER_ROW])}
    rows = list(ws.iter_rows(min_row=DATA_FIRST_ROW, values_only=False))
    item_row = next(r for r in rows if r[headers["Наименование"] - 1].value == "Позиция")

    def _val(row, label):
        return row[headers[label] - 1].value

    # Владелец 08.10.2026, часть B: плановая позиция ИМЕЕТ закупки — её
    # денежные столбцы-стадии теперь формула =SUBTOTAL(109,...) по строке(ам)
    # её собственных закупок, не литеральное число.
    for label in ("Запланировано, ₽", "Договор, ₽", "Заказано, ₽", "Поставлено, ₽", "Оплачено, ₽", "Фактически (итого), ₽"):
        assert str(_val(item_row, label)).startswith("=SUBTOTAL(109,")

    # Проверка «значения совпадают с прежними числами» (владелец) — сама
    # закупка под позицией оплачена на 1000 ₽, накопительно во всех стадиях;
    # SUBTOTAL(109,...) её строки корректно её суммирует (единственная
    # строка в диапазоне).
    sub_row = next(
        r for r in rows if "Купленное" in str(r[headers["Состав закупки"] - 1].value or "")
    )
    assert _val(sub_row, "Запланировано, ₽") == 1000.0
    assert _val(sub_row, "Фактически (итого), ₽") == 1000.0
    assert _val(sub_row, "Уровень") == "Закупка"
    assert _val(item_row, "Уровень") == "Позиция"

    # Отменённая закупка (секция «без категории») — все денежные столбцы 0,
    # не "" (видна, но не считается). Название закупленной позиции теперь
    # живёт в "Состав закупки" (владелец 08.10.2026 — своё значение в своём
    # столбце, а не текстом внутри "Наименование").
    cancelled_row = next(
        r for r in rows if "Без категории (отменена)" in str(r[headers["Состав закупки"] - 1].value or "")
    )
    assert _val(cancelled_row, "Запланировано, ₽") == 0.0
    assert _val(cancelled_row, "Фактически (итого), ₽") == 0.0


def test_hierarchical_group_rows_use_subtotal_and_level_column():
    """Владелец 08.10.2026, часть B: направление/тип/статья/плановая позиция
    («группы») больше не несут литеральное число денежных столбцов-стадий —
    это формула =SUBTOTAL(109,...) по их собственному диапазону строк;
    служебный столбец «Уровень» заполнен на каждом уровне; «Бюджет ФЭО»/
    «Остаток»/«% исполнения» остаются значениями (не формулой)."""
    wb = build_live_plan_graph_xlsx(_make_sub(), "http://example.test", _base_data())
    ws = wb["План закупок (по направлениям)"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[HEADER_ROW])}
    rows = list(ws.iter_rows(min_row=DATA_FIRST_ROW, values_only=False))

    def _val(row, label):
        return row[headers[label] - 1].value

    level_values = {_val(r, "Уровень") for r in rows}
    assert {"Направление", "Тип", "Статья", "Позиция", "Закупка"} <= level_values

    l1_row = next(r for r in rows if _val(r, "Уровень") == "Направление")
    l2_row = next(r for r in rows if _val(r, "Уровень") == "Тип")
    l3_row = next(r for r in rows if _val(r, "Уровень") == "Статья")
    item_row = next(r for r in rows if _val(r, "Уровень") == "Позиция")

    for group_row in (l1_row, l2_row, l3_row, item_row):
        assert str(_val(group_row, "Запланировано, ₽")).startswith("=SUBTOTAL(109,")
        assert str(_val(group_row, "Оплачено, ₽")).startswith("=SUBTOTAL(109,")

    # «Бюджет ФЭО»/«Остаток»/«% исполнения» — значения, не формулы (ФЭО
    # статьи не обязан совпадать с суммой позиций).
    assert not str(_val(l1_row, "Бюджет ФЭО, ₽")).startswith("=")
    assert not str(_val(l3_row, "Остаток, ₽")).startswith("=")

    # Проверка «значения групп совпадают с прежними числами на тех же
    # данных»: единственная закупка статьи оплачена на 1000 ₽ — это ровно
    # то число, которое раньше было литералом в L1/L2/L3/Позиция; теперь
    # SUBTOTAL(109,...) её диапазона физически суммирует ЭТУ строку (нижний,
    # не-SUBTOTAL уровень «Закупка», игнорируя вложенные SUBTOTAL — так
    # работает Excel/Google Sheets).
    leaf_row = next(r for r in rows if _val(r, "Уровень") == "Закупка" and _val(r, "Запланировано, ₽") == 1000.0)
    assert _val(leaf_row, "Оплачено, ₽") == 1000.0

    # Часть B — верхние строки 1/2: «Итого всего» (SUMIFS по Уровень=
    # "Закупка" для денежных столбцов-стадий, SUMIFS по "Направление" для
    # «Бюджет ФЭО»), «Итого по фильтру» (SUBTOTAL); «Остаток»/«% исполнения» —
    # пусто в обеих строках (сумма по ним бессмысленна).
    planned_ci = headers["Запланировано, ₽"]
    assert str(ws.cell(row=1, column=planned_ci).value).startswith("=SUMIFS(")
    assert '"Закупка"' in ws.cell(row=1, column=planned_ci).value
    assert str(ws.cell(row=2, column=planned_ci).value).startswith("=SUBTOTAL(109,")

    feo_ci = headers["Бюджет ФЭО, ₽"]
    assert str(ws.cell(row=1, column=feo_ci).value).startswith("=SUMIFS(")
    assert '"Направление"' in ws.cell(row=1, column=feo_ci).value
    assert ws.cell(row=2, column=feo_ci).value in (None, "")

    residual_ci = headers["Остаток, ₽"]
    assert ws.cell(row=1, column=residual_ci).value in (None, "")
    assert ws.cell(row=2, column=residual_ci).value in (None, "")

    assert ws.freeze_panes == f"A{DATA_FIRST_ROW}"
    assert ws.auto_filter.ref is not None


def test_column_selection_limits_headers_and_keeps_name():
    selected = resolve_selected_keys("paid,status")
    assert "name" in selected  # всегда присутствует
    assert selected.index("paid") < selected.index("status")  # фиксированный порядок PLAN_GRAPH_COLUMNS

    wb = build_live_plan_graph_xlsx(
        _make_sub(), "http://example.test", _base_data(), selected_columns=selected,
    )
    ws = wb["План закупок (по направлениям)"]
    header_values = [c.value for c in ws[HEADER_ROW] if c.value]
    assert header_values == ["Наименование", "Оплачено, ₽", "Статус"]


def test_resolve_selected_keys_empty_returns_default_only():
    """Владелец 07.10.2026: группа «Договор и оплата» — не в наборе по
    умолчанию; пустой ?columns= больше не означает «все столбцы»."""
    from app.services.plan_graph_export_columns import ALL_KEYS, DEFAULT_KEYS
    assert resolve_selected_keys(None) == list(DEFAULT_KEYS)
    assert resolve_selected_keys("") == list(DEFAULT_KEYS)
    assert set(DEFAULT_KEYS) < set(ALL_KEYS)  # новая группа существует, но не default


def test_columns_catalog_marks_default_flag():
    from app.services.plan_graph_export_columns import columns_catalog
    catalog = {c["key"]: c for c in columns_catalog()}
    assert catalog["name"]["default"] is True
    assert catalog["contract_number"]["default"] is False
    assert catalog["payment_doc_number"]["default"] is False
    assert catalog["contract_number"]["group"] == "Договор и оплата"


def test_new_contract_payment_columns_appear_only_when_selected():
    """Новый столбец («№ договора») отсутствует в выгрузке по умолчанию и
    появляется только если его явно выбрали через ?columns=."""
    data = _base_data()
    purchase = SimpleNamespace(
        id=55, purchase_number="55", order_number="", registry_number="РЕЕ-55",
        status="paid", subsidy_id=None, feo_category_id=None,
        item_name="", item_type="", unit="", planned_quantity=None,
        subject="", country_origin="", planned_unit_price=None, planned_total_price=None,
        total_nmck=None, contract_price=None, price_increase=None,
        purchase_method=None, purchase_basis=None, purchase_contract_type=None,
        framework_seq=None, contract_number="ДОГ-55/2026", contract_date=None,
        execution_term=None, execution_term_changed=None, delivery_date=None,
        contractor_id=None, reimbursement_user_id=None, service_note_by=None,
        responsible_person="", acceptance_doc_number="A-1", acceptance_docs=None,
        payment_doc_number="ПП-100", payment_doc_date=None, payment_amount=1000.0,
        payment_federal=None, delivery_payment_amount=None,
        vat_applicable=False, vat_rate=None, vat_exemption_article=None, vat_mode=None,
        etp_url=None, region="", delivery_region="", delivery_location="", delivery_address="",
        final_unit_price=None, final_total_amount=None, contract_end_date=None,
        submission_deadline=None, commitment_quarter=None, planned_payment_month=None,
        stage_label=None, substatus=None,
    )
    data["purchase_rows_by_id"] = {55: purchase}
    data["purchase_export_ctx"] = {
        "contractors": {}, "contractor_inns": {}, "subsidies": {}, "feo_categories": {},
        "payment_purposes": {55: "За услуги"}, "ru_map": {}, "sn_map": {}, "economy": {},
    }

    # По умолчанию (без явного выбора столбцов) — новой колонки нет вовсе.
    wb_default = build_live_plan_graph_xlsx(_make_sub(), "http://example.test", data)
    ws_default = wb_default["План закупок (по направлениям)"]
    headers_default = [c.value for c in ws_default[HEADER_ROW] if c.value]
    assert "№ договора" not in headers_default
    assert "ПП: №" not in headers_default

    # Явно выбрали — колонка появляется, в под-строке оплаченной закупки
    # заполнена значением с самой закупки, в строке плановой позиции — пусто.
    selected = resolve_selected_keys("name,composition,paid,contract_number,payment_doc_number,payment_purpose")
    wb = build_live_plan_graph_xlsx(
        _make_sub(), "http://example.test", data, selected_columns=selected,
    )
    ws = wb["План закупок (по направлениям)"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(ws[HEADER_ROW])}
    assert "№ договора" in headers
    assert "ПП: №" in headers

    rows = list(ws.iter_rows(min_row=DATA_FIRST_ROW, values_only=False))

    def _val(row, label):
        return row[headers[label] - 1].value

    item_row = next(r for r in rows if _val(r, "Наименование") == "Позиция")
    assert _val(item_row, "№ договора") == ""
    assert _val(item_row, "ПП: №") == ""

    sub_row = next(
        r for r in rows if "Купленное" in str(_val(r, "Состав закупки") or "")
    )
    assert _val(sub_row, "№ договора") == "ДОГ-55/2026"
    assert _val(sub_row, "ПП: №") == "ПП-100"
    assert _val(sub_row, "Назначение платежа") == "За услуги"


def test_summary_sheet_has_formulas_for_computed_columns_and_total():
    summary_by_kind = {
        "goods": {
            "feo": 100.0, "plan": 90.0, "committed": 40.0, "monthly": 10.0,
            "likely": 20.0, "nice": 20.0, "paid": 15.0, "delivered_unpaid": 5.0,
        },
        "services": {
            "feo": 0.0, "plan": 0.0, "committed": 0.0, "monthly": 0.0,
            "likely": 0.0, "nice": 0.0, "paid": 0.0, "delivered_unpaid": 0.0,
        },
        "payroll": {k: 0.0 for k in ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")},
        "unspecified": {k: 0.0 for k in ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")},
    }
    wb = build_live_plan_graph_xlsx(
        _make_sub(), "http://example.test", _base_data(), summary_by_kind=summary_by_kind,
    )
    assert "Сводная" in wb.sheetnames
    sm = wb["Сводная"]

    # Строка 5 — «Товары» (goods всегда показан), строка 6 — «Услуги».
    goods_row = 5
    assert sm.cell(row=goods_row, column=1).value == "Товары"
    assert sm.cell(row=goods_row, column=4).value == f"=C{goods_row}-H{goods_row}"
    assert sm.cell(row=goods_row, column=7).value == f"=D{goods_row}-E{goods_row}-F{goods_row}"
    assert sm.cell(row=goods_row, column=9).value == f"=J{goods_row}+K{goods_row}"

    services_row = 6
    assert sm.cell(row=services_row, column=1).value == "Услуги"

    # «ФОТ»/«Без типа» все нули — строки не рисуются.
    row_labels = [sm.cell(row=r, column=1).value for r in range(5, sm.max_row + 1)]
    assert "ФОТ" not in row_labels
    assert "Без типа" not in row_labels

    total_row = services_row + 1
    assert sm.cell(row=total_row, column=1).value == "ИТОГО"
    assert sm.cell(row=total_row, column=2).value == f"=SUM(B5:B{services_row})"


def test_no_summary_sheet_when_summary_by_kind_is_none():
    wb = build_live_plan_graph_xlsx(_make_sub(), "http://example.test", _base_data(), summary_by_kind=None)
    assert "Сводная" not in wb.sheetnames


def test_summary_sheet_contract_remaining_column_from_balance_groups():
    """Владелец 08.10.2026, часть C: «Остаток на договорах, ₽» — Σ remaining
    групп data["contract_balance_groups"] по виду (kind), не второй расчёт."""
    summary_by_kind = {
        "goods": {k: 0.0 for k in ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")},
        "services": {k: 0.0 for k in ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")},
        "payroll": {k: 0.0 for k in ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")},
        "unspecified": {k: 0.0 for k in ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")},
    }
    data = _base_data()
    data["contract_balance_groups"] = [
        {"head_id": 1, "contract_sum": 1000.0, "ordered": 500.0, "remaining": 500.0, "kind": "goods", "purchase_ids": [1, 2]},
        {"head_id": 3, "contract_sum": 200.0, "ordered": 200.0, "remaining": 0.0, "kind": "services", "purchase_ids": [3]},
        {"head_id": 4, "contract_sum": 300.0, "ordered": 100.0, "remaining": 200.0, "kind": "goods", "purchase_ids": [4]},
    ]
    wb = build_live_plan_graph_xlsx(
        _make_sub(), "http://example.test", data, summary_by_kind=summary_by_kind,
    )
    sm = wb["Сводная"]
    headers = {cell.value: idx + 1 for idx, cell in enumerate(sm[3])}
    col = headers["Остаток на договорах, ₽"]
    assert sm.cell(row=5, column=1).value == "Товары"
    assert sm.cell(row=5, column=col).value == 700.0  # 500 + 200, обе группы "goods"
    assert sm.cell(row=6, column=1).value == "Услуги"
    assert sm.cell(row=6, column=col).value == 0.0
    from openpyxl.utils import get_column_letter
    total_row = 7
    assert sm.cell(row=total_row, column=1).value == "ИТОГО"
    letter = get_column_letter(col)
    assert sm.cell(row=total_row, column=col).value == f"=SUM({letter}5:{letter}6)"
