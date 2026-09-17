"""Точечный фикс раскладки таблицы позиций в contract_tz.docx — без регенерации.

Владелец (2026-09-17): «переменная в таблице должна переноситься по строкам,
иначе в шаблоне получается слишком широкая» — {{item.num}} и {{item.quantity}}
без пробелов внутри длиннее, чем нужная ширина колонки (2 и 5 символов у
реального значения), а Word по умолчанию (autofit по содержимому) растягивает
столбец под весь плейсхолдер целиком, игнорируя уже выставленные per-cell
tcW.

ПОЧЕМУ ОТДЕЛЬНЫЙ СКРИПТ, А НЕ ПЕРЕГЕНЕРАЦИЯ generate_templates.py::make_contract_tz
(общее правило проекта — не патчить бинарник руками, перегенерировать тем же
генератором): при попытке перегенерировать этим скриптом contract_tz.docx
ТЕРЯЕТ {{item.form_summary}}/{{item.menu_lines}} — этот генератор устарел и
не знает о фиче, которая в реальном (текущем, закоммиченном) бинарнике уже
есть и покрыта test_contract_templates.py. Ни один из трёх найденных
генераторов (generate_templates.py, make_real_templates.py, make_tech_spec.py
делает только tech_spec_*) не воспроизводит текущий contract_tz.docx один в
один — до тех пор, пока такой генератор не появится, единственный безопасный
способ — точечная правка уже существующей структуры без переписывания
контента: читаем уже верные per-cell ширины (row 0, выставлены когда-то
_set_col_width) и просто заставляем Word их уважать (tblLayout=fixed +
синхронный tblGrid), не трогая ни один плейсхолдер/фото/форму.

Запуск:
    py backend/templates/fix_contract_tz_table_layout.py
"""
import os
from docx import Document
from docx.oxml.ns import qn

TEMPLATES_DIR = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.join(TEMPLATES_DIR, "contract_tz.docx")


def _cell_text(cell) -> str:
    return "".join(p.text for p in cell.paragraphs)


def main():
    doc = Document(PATH)

    items_table = None
    for tbl in doc.tables:
        # Строка данных содержит {{item.num}} — это и есть таблица позиций
        # (единственная в документе с таким плейсхолдером).
        if any("{{item.num}}" in _cell_text(c) for row in tbl.rows for c in row.cells):
            items_table = tbl
            break
    if items_table is None:
        raise SystemExit("Таблица позиций ({{item.num}}) не найдена — контент документа изменился, проверьте руками")

    # Ширины уже верные — они выставлены на ячейках заголовка (row 0) генератором
    # когда-то давно (_set_col_width), но append()-ом ПОВЕРХ дефолтного tcW,
    # который add_table() кладёт при создании таблицы — в каждой ячейке заголовка
    # остаётся ДВА <w:tcW> подряд (первый — дефолтный 1/N ширины, второй —
    # настоящий, от _set_col_width). Нужен ПОСЛЕДНИЙ. Достаём его оттуда, а не
    # дублируем числами в этом скрипте (Правило №6 — один источник, документ
    # сам себе источник).
    header_cells = items_table.rows[0].cells
    widths_twips = []
    for c in header_cells:
        tcPr = c._tc.tcPr
        tcWs = tcPr.findall(qn('w:tcW')) if tcPr is not None else []
        if not tcWs or tcWs[-1].get(qn('w:w')) is None:
            raise SystemExit(f"У ячейки заголовка нет tcW — ширины не заданы, править нечего: {_cell_text(c)!r}")
        widths_twips.append(tcWs[-1].get(qn('w:w')))

    if len(set(widths_twips)) == 1:
        raise SystemExit("Все ширины колонок одинаковые — похоже, читаем не ту таблицу или её ширины не выставлены")

    items_table.autofit = False  # <w:tblLayout w:type="fixed"/>
    grid = items_table._tbl.find(qn('w:tblGrid'))
    if grid is None:
        raise SystemExit("У таблицы нет tblGrid — неожиданная структура")
    grid_cols = grid.findall(qn('w:gridCol'))
    if len(grid_cols) != len(widths_twips):
        raise SystemExit(f"Число колонок в tblGrid ({len(grid_cols)}) не совпадает с числом ячеек заголовка ({len(widths_twips)})")
    for gc, w in zip(grid_cols, widths_twips):
        gc.set(qn('w:w'), w)

    doc.save(PATH)
    print(f"OK: {PATH} — tblLayout=fixed, tblGrid синхронизирован с существующими ширинами {widths_twips}")


if __name__ == "__main__":
    main()
