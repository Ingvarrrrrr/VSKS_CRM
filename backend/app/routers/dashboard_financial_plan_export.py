"""GET /api/dashboard/financial-plan/export.xlsx (/details/export.xlsx) — сосед
dashboard.py (Правило №5, резка 1641→core, 2026-09-08).

Excel-выгрузка того же финплана, что и dashboard_financial_plan.py: полная
таблица по всем периодам либо одна группа (период+категория) — содержимое
drill-down диалога.
"""
from typing import Optional
from fastapi import APIRouter, Depends, Query
from sqlalchemy import select
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
from app.database import get_db
from app.models.purchase import Purchase
from app.models.user import User
from app.auth.jwt import get_current_user, get_org_filter
from app.auth.visibility import get_visible_subsidy_ids
from app.routers.dashboard import _UNSET, _apply_purchase_org_filter, _effective_float, _expected_payment_date

router = APIRouter(prefix="/api/dashboard", tags=["dashboard"])


@router.get("/financial-plan/export.xlsx")
async def export_financial_plan_xlsx(
    granularity: str = Query("month", regex="^(month|quarter)$"),
    subsidy_id: Optional[int] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    scope: Optional[str] = Query(None),
):
    """Полная выгрузка всех закупок с группировкой по периоду и категории plan/committed."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from datetime import datetime as dt

    org_ids = get_org_filter(current_user)
    dash_sids = _UNSET
    if scope in ("dashboard", "plan"):
        dash_sids = await get_visible_subsidy_ids(current_user, db, scope)

    PLAN_STATUSES = {"planned", "wishes", "plan_schedule"}
    COMMITTED_STATUSES = {"contracted", "ordered", "delivered", "paid", "work_in_progress"}

    q = select(Purchase).where(Purchase.status.in_(PLAN_STATUSES | COMMITTED_STATUSES))
    if subsidy_id:
        q = q.where(Purchase.subsidy_id == subsidy_id)
    q = _apply_purchase_org_filter(q, current_user, org_ids, subsidy_ids=dash_sids)
    q = q.options(selectinload(Purchase.contractor))

    rows = (await db.execute(q)).scalars().all()

    STATUS_LABELS = {
        "planned": "Запланирован", "wishes": "Заявка",
        "plan_schedule": "План закупок",
        "contracted": "Заключён договор", "ordered": "Заказано", "delivered": "Поставлено",
        "paid": "Оплачено", "work_in_progress": "Ведётся работа",
    }

    grouped: dict = {}
    for p in rows:
        d = _expected_payment_date(p)
        if not d:
            continue
        amount = _effective_float(p)
        if amount == 0:
            continue
        if granularity == "month":
            period_key = f"{d.year}-{d.month:02d}"
        else:
            period_key = f"{d.year}-Q{(d.month - 1) // 3 + 1}"
        category = "plan" if p.status in PLAN_STATUSES else "committed"
        grouped.setdefault((period_key, category), []).append({
            "purchase_number": p.purchase_number or "",
            "subject": p.subject or "",
            "contractor_name": (p.contractor.name if p.contractor else "—"),
            "amount": amount,
            "status": STATUS_LABELS.get(p.status, p.status),
            "expected_date": d.isoformat(),
            "contract_number": p.contract_number or "",
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Финплан"

    ws['A1'] = f"Финансовый план — {'по месяцам' if granularity == 'month' else 'по кварталам'}"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:H1')

    headers = ["Период", "Категория", "№ закупки", "Предмет", "Контрагент", "Дата", "Статус", "Сумма, ₽"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 4
    period_totals: dict = {}

    for (period_key, category) in sorted(grouped.keys()):
        items = grouped[(period_key, category)]
        cat_label = "Плановые" if category == "plan" else "Принятые обязательства"
        cat_color = "FEF3C7" if category == "plan" else "D1FAE5"

        for item in items:
            ws.cell(row=row_idx, column=1, value=period_key)
            ws.cell(row=row_idx, column=2, value=cat_label)
            ws.cell(row=row_idx, column=3, value=item["purchase_number"])
            ws.cell(row=row_idx, column=4, value=item["subject"])
            ws.cell(row=row_idx, column=5, value=item["contractor_name"])
            ws.cell(row=row_idx, column=6, value=item["expected_date"])
            ws.cell(row=row_idx, column=7, value=item["status"])
            ws.cell(row=row_idx, column=8, value=item["amount"])
            for col in range(1, 9):
                ws.cell(row=row_idx, column=col).fill = PatternFill(
                    start_color=cat_color, end_color=cat_color, fill_type="solid"
                )
            ws.cell(row=row_idx, column=8).number_format = '#,##0.00 ₽'
            row_idx += 1

            period_totals.setdefault(period_key, {"plan": 0.0, "committed": 0.0})
            period_totals[period_key][category] += item["amount"]

    row_idx += 1
    totals_title = ws.cell(row=row_idx, column=1, value="ИТОГО ПО ПЕРИОДАМ")
    totals_title.font = Font(bold=True, size=12)
    row_idx += 1
    ws.cell(row=row_idx, column=1, value="Период").font = Font(bold=True)
    ws.cell(row=row_idx, column=2, value="Плановые").font = Font(bold=True)
    ws.cell(row=row_idx, column=3, value="Принятые").font = Font(bold=True)
    ws.cell(row=row_idx, column=4, value="Всего").font = Font(bold=True)
    row_idx += 1

    for pk in sorted(period_totals.keys()):
        t = period_totals[pk]
        ws.cell(row=row_idx, column=1, value=pk)
        plan_cell = ws.cell(row=row_idx, column=2, value=t["plan"])
        plan_cell.number_format = '#,##0.00 ₽'
        committed_cell = ws.cell(row=row_idx, column=3, value=t["committed"])
        committed_cell.number_format = '#,##0.00 ₽'
        total_cell = ws.cell(row=row_idx, column=4, value=t["plan"] + t["committed"])
        total_cell.number_format = '#,##0.00 ₽'
        row_idx += 1

    for col_letter, width in [('A', 12), ('B', 22), ('C', 12), ('D', 50), ('E', 30), ('F', 12), ('G', 18), ('H', 16)]:
        ws.column_dimensions[col_letter].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    from urllib.parse import quote
    filename = f"Финплан_{granularity}_{dt.now().strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='-_.~')}"},
    )


@router.get("/financial-plan/details/export.xlsx")
async def export_financial_plan_details_xlsx(
    period: str = Query(...),
    category: str = Query(..., regex="^(plan|committed)$"),
    granularity: str = Query("month", regex="^(month|quarter)$"),
    subsidy_id: Optional[int] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    scope: Optional[str] = Query(None),
):
    """Excel выгрузка одной группы (период+категория) — содержимое drill-down диалога."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    org_ids = get_org_filter(current_user)
    dash_sids = _UNSET
    if scope in ("dashboard", "plan"):
        dash_sids = await get_visible_subsidy_ids(current_user, db, scope)

    PLAN_STATUSES = {"planned", "wishes", "plan_schedule"}
    COMMITTED_STATUSES = {"contracted", "ordered", "delivered", "paid", "work_in_progress"}
    target_statuses = PLAN_STATUSES if category == "plan" else COMMITTED_STATUSES

    q = select(Purchase).where(Purchase.status.in_(target_statuses))
    if subsidy_id:
        q = q.where(Purchase.subsidy_id == subsidy_id)
    q = _apply_purchase_org_filter(q, current_user, org_ids, subsidy_ids=dash_sids)
    q = q.options(selectinload(Purchase.contractor))

    rows = (await db.execute(q)).scalars().all()

    STATUS_LABELS = {
        "planned": "Запланирован", "wishes": "Заявка",
        "plan_schedule": "План закупок",
        "contracted": "Заключён договор", "ordered": "Заказано", "delivered": "Поставлено",
        "paid": "Оплачено", "work_in_progress": "Ведётся работа",
    }

    items = []
    for p in rows:
        d = _expected_payment_date(p)
        if not d:
            continue
        if granularity == "month":
            row_period = f"{d.year}-{d.month:02d}"
        else:
            row_period = f"{d.year}-Q{(d.month - 1) // 3 + 1}"
        if row_period != period:
            continue
        amount = _effective_float(p)
        if amount == 0:
            continue
        items.append({
            "purchase_number": p.purchase_number or "",
            "subject": p.subject or "",
            "contractor_name": (p.contractor.name if p.contractor else "—"),
            "amount": amount,
            "status": STATUS_LABELS.get(p.status, p.status),
            "expected_date": d.isoformat(),
            "contract_number": p.contract_number or "",
        })

    items.sort(key=lambda r: (r["expected_date"], -r["amount"]))

    wb = Workbook()
    ws = wb.active
    cat_label = "Плановые" if category == "plan" else "Принятые обязательства"
    ws.title = f"{cat_label[:20]} {period}"

    ws['A1'] = f"Финплан — {cat_label} — {period}"
    ws['A1'].font = Font(size=14, bold=True)
    ws.merge_cells('A1:G1')

    headers = ["№", "Предмет", "Контрагент", "Дата", "Статус", "№ договора", "Сумма, ₽"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total = 0.0
    for i, item in enumerate(items, start=4):
        ws.cell(row=i, column=1, value=item["purchase_number"])
        ws.cell(row=i, column=2, value=item["subject"])
        ws.cell(row=i, column=3, value=item["contractor_name"])
        ws.cell(row=i, column=4, value=item["expected_date"])
        ws.cell(row=i, column=5, value=item["status"])
        ws.cell(row=i, column=6, value=item["contract_number"])
        amt_cell = ws.cell(row=i, column=7, value=item["amount"])
        amt_cell.number_format = '#,##0.00 ₽'
        total += item["amount"]

    total_row = 4 + len(items) + 1
    ws.cell(row=total_row, column=1, value="ИТОГО").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=7, value=total)
    total_cell.font = Font(bold=True)
    total_cell.number_format = '#,##0.00 ₽'
    total_color = "FEF3C7" if category == "plan" else "D1FAE5"
    total_cell.fill = PatternFill(start_color=total_color, end_color=total_color, fill_type="solid")

    for col_letter, width in [('A', 12), ('B', 50), ('C', 30), ('D', 12), ('E', 18), ('F', 16), ('G', 16)]:
        ws.column_dimensions[col_letter].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    from urllib.parse import quote
    filename = f"Финплан_{period}_{category}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='-_.~')}"},
    )
