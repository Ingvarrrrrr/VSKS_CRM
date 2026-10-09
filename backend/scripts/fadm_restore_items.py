#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Восстановление разбивки позиций закупок субсидии «ФАДМ 2026_2» (id=89 на
проде) по образцу старой субсидии «ФАДМ_2026» (id=7) — задача 09.10.2026.

ПОВОД. Субсидия 89 собрана скриптом scripts/fadm_sheet_load/sheet_v2_build.py
только из строк листа GoodsService — у многих закупок там ОДНА агрегированная
позиция (например «ПРИОБРЕТЕНИЕ ХОЗЯЙСТВЕННЫХ ТОВАРОВ», 63 688,81), хотя у той
же закупки в субсидии 7 — настоящая разбивка (16 позиций на ту же сумму
63 688,81, РЕЕ-2026-00964). Владелец собирается удалить субсидию 7 — разбивку
нужно перенести в 89 ДО удаления. Сопоставление закупок 7<->89 с уверенностью
и предложенным действием сделано read-only анализом (см. задачу «Таблица
сопоставления» этой же сессии) и лежит в xlsx, лист «Пары» (+ «Состав» — уже
выгруженные позиции старых закупок для справки/ревью, этот скрипт их
заново НЕ читает, читает purchase_items из БД напрямую).

ЧТО ДЕЛАЕТ (для каждой пары листа «Пары» с действием «разбить на N позиций» И
колонкой «Подтверждено» = «да»):
  1. Проверяет, что у закупки-получателя (89) позиций МЕНЬШЕ, чем у закупки-
     источника (7), и Σ total_price позиций 7 == Σ total_price позиций 89 "в
     копейку" (Decimal, допуск 0.01 ₽) — иначе пара ПРОПУСКАЕТСЯ с причиной.
  2. Резервная копия ТЕКУЩИХ позиций закупки 89 — в таблицу
     bk_20261009_fadm_restore_items (CREATE TABLE IF NOT EXISTS ... AS SELECT
     *, пара/момент копирования — отдельными ALTER TABLE ADD COLUMN IF NOT
     EXISTS, чтобы повторный прогон на разные пары не плодил вторую такую
     таблицу).
  3. Определяет feo_planned_item_id/feo_category_id, которые получат НОВЫЕ
     (клонированные) позиции: берёт их у удаляемых агрегированных позиций 89.
     Если среди них встречается НЕСКОЛЬКО разных непустых значений одной из
     этих колонок — пара ПРОПУСКАЕТСЯ с причиной (задание: «если несколько
     разных — пропуск»); если ровно одно — все новые позиции получают его.
  4. Удаляет агрегированные позиции закупки 89, клонирует позиции закупки 7
     (app.services.subsidy_copy.clone_items.clone_purchase_items — ПРАВИЛО
     №6, тот же клонировщик, что copy_purchases.py использует при копировании
     целой закупки) с overrides=(feo_planned_item_id, feo_category_id) из п.3.
  5. НИЧЕГО не меняет в шапке закупки 89 (договор/суммы/статус/платежи/акты)
     — recalc_purchase_money НЕ вызывается специально.
  6. Одна транзакция на пару (отдельная db-сессия) — commit только с --apply,
     иначе rollback. Печатает «до/после» app.services.purchase_amounts.
     load_purchase_amounts(89) и app.services.type_totals.subsidy_type_totals
     (план субсидии-получателя) — должны совпасть до/после (разбивка суммы не
     меняет), это и проверяется в выводе.

ЗАПУСК (внутри контейнера backend; backend/scripts не смонтирован volume'ом —
сначала docker cp):
    MSYS_NO_PATHCONV=1 docker cp backend/scripts/fadm_restore_items.py \\
        vsks-crm-backend_a-1:/app/scripts/fadm_restore_items.py
    docker exec -w /app vsks-crm-backend_a-1 python scripts/fadm_restore_items.py \\
        --pairs-xlsx /tmp/fadm_restore_pairs.xlsx --dry-run

--apply требует явного одобрения владельца ПОСЛЕ показа --dry-run отчёта.
"""
from __future__ import annotations

import argparse
import asyncio
import sys
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Optional

from sqlalchemy import select, text

from app.database import async_session
from app.models.purchase import Purchase
from app.models.purchase_item import PurchaseItem
from app.services.purchase_amounts import load_purchase_amounts
from app.services.subsidy_copy.clone_items import clone_purchase_items
from app.services.type_totals import subsidy_type_totals

BACKUP_TABLE = "bk_20261009_fadm_restore_items"
AMOUNT_TOLERANCE = Decimal("0.01")


@dataclass
class PairPlan:
    reg7: str
    reg89: str


@dataclass
class PairOutcome:
    reg7: str
    reg89: str
    status: str  # 'ok' | 'skipped'
    reason: str = ""
    n_old_items: int = 0
    n_new_items_before: int = 0
    subsidy_id: Optional[int] = None
    amount_before: Optional[Decimal] = None
    amount_after: Optional[Decimal] = None
    plan_before: Optional[dict] = None
    plan_after: Optional[dict] = None


def read_confirmed_pairs(xlsx_path: str) -> list[PairPlan]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Пары"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    required = ["Рег.№ ФАДМ_2026", "Рег.№ ФАДМ 2026_2", "Предлагаемое действие", "Подтверждено"]
    missing = [c for c in required if c not in idx]
    if missing:
        raise ValueError(f"В листе «Пары» не найдены колонки: {missing}")

    out: list[PairPlan] = []
    for row in rows[1:]:
        if row is None or all(v is None for v in row):
            continue
        action = str(row[idx["Предлагаемое действие"]] or "").strip()
        confirmed = str(row[idx["Подтверждено"]] or "").strip().lower()
        if action != "разбить на N позиций" or confirmed != "да":
            continue
        reg7 = str(row[idx["Рег.№ ФАДМ_2026"]] or "").strip()
        reg89 = str(row[idx["Рег.№ ФАДМ 2026_2"]] or "").strip()
        if not reg7 or not reg89:
            continue
        out.append(PairPlan(reg7=reg7, reg89=reg89))
    return out


def validate_split(
    *, n_old: int, n_new: int, sum_old: Decimal, sum_new: Decimal,
    fpi_ids: set, cat_ids: set,
) -> Optional[str]:
    """Чистая (без БД) проверка условий «можно разбить» — вынесена отдельно,
    чтобы её можно было юнит-тестировать без контейнера/БД (см.
    tests/test_fadm_restore_items.py). Возвращает None, если разбивка
    разрешена, иначе — причину пропуска (как в отчёте скрипта)."""
    if n_old == 0:
        return "у закупки-источника нет позиций — нечего переносить"
    if n_new >= n_old:
        return (
            f"у закупки-получателя позиций ({n_new}) не меньше, чем у источника "
            f"({n_old}) — пропуск (условие задания)"
        )
    if abs(sum_old - sum_new) > AMOUNT_TOLERANCE:
        return f"суммы не совпадают: Σ_источник={sum_old} != Σ_получатель={sum_new}"
    if len(fpi_ids) > 1 or len(cat_ids) > 1:
        return (
            f"у агрегированных позиций получателя несколько РАЗНЫХ "
            f"feo_planned_item_id/feo_category_id ({fpi_ids} / {cat_ids}) — пропуск (условие задания)"
        )
    return None


async def _load_purchase(db, registry_number: str, subsidy_id: int) -> Optional[Purchase]:
    return (await db.execute(
        select(Purchase).where(Purchase.registry_number == registry_number, Purchase.subsidy_id == subsidy_id)
    )).scalars().first()


async def _ensure_backup_table(db) -> None:
    await db.execute(text(
        f"CREATE TABLE IF NOT EXISTS {BACKUP_TABLE} AS SELECT * FROM purchase_items WHERE false"
    ))
    await db.execute(text(
        f"ALTER TABLE {BACKUP_TABLE} ADD COLUMN IF NOT EXISTS backed_up_at timestamptz DEFAULT now()"
    ))
    await db.execute(text(
        f"ALTER TABLE {BACKUP_TABLE} ADD COLUMN IF NOT EXISTS restore_reg7 text"
    ))
    await db.execute(text(
        f"ALTER TABLE {BACKUP_TABLE} ADD COLUMN IF NOT EXISTS restore_reg89 text"
    ))


async def _backup_items(db, purchase_id: int, reg7: str, reg89: str) -> None:
    await db.execute(text(
        f"INSERT INTO {BACKUP_TABLE} "
        "SELECT pi.*, now(), :reg7, :reg89 FROM purchase_items pi WHERE pi.purchase_id = :pid"
    ), {"pid": purchase_id, "reg7": reg7, "reg89": reg89})


async def process_pair(pair: PairPlan, *, source_subsidy_id: int, target_subsidy_id: int,
                        apply: bool) -> PairOutcome:
    async with async_session() as db:
        try:
            out = PairOutcome(reg7=pair.reg7, reg89=pair.reg89, status="skipped")

            p7 = await _load_purchase(db, pair.reg7, source_subsidy_id)
            p89 = await _load_purchase(db, pair.reg89, target_subsidy_id)
            if not p7:
                out.reason = f"закупка {pair.reg7} не найдена в субсидии {source_subsidy_id}"
                return out
            if not p89:
                out.reason = f"закупка {pair.reg89} не найдена в субсидии {target_subsidy_id}"
                return out
            out.subsidy_id = target_subsidy_id

            old_items = (await db.execute(
                select(PurchaseItem).where(PurchaseItem.purchase_id == p7.id)
            )).scalars().all()
            new_items = (await db.execute(
                select(PurchaseItem).where(PurchaseItem.purchase_id == p89.id)
            )).scalars().all()
            out.n_old_items = len(old_items)
            out.n_new_items_before = len(new_items)

            sum_old = sum((Decimal(str(it.total_price)) for it in old_items if it.total_price is not None), Decimal("0"))
            sum_new = sum((Decimal(str(it.total_price)) for it in new_items if it.total_price is not None), Decimal("0"))
            fpi_ids = {it.feo_planned_item_id for it in new_items if it.feo_planned_item_id is not None}
            cat_ids = {it.feo_category_id for it in new_items if it.feo_category_id is not None}

            skip_reason = validate_split(
                n_old=len(old_items), n_new=len(new_items), sum_old=sum_old, sum_new=sum_new,
                fpi_ids=fpi_ids, cat_ids=cat_ids,
            )
            if skip_reason:
                out.reason = skip_reason
                return out
            target_fpi = next(iter(fpi_ids), None)
            target_cat = next(iter(cat_ids), None)

            amounts_before = await load_purchase_amounts(db, [p89.id])
            out.amount_before = amounts_before[p89.id].effective if p89.id in amounts_before else None
            plan_before_all = await subsidy_type_totals(db, [target_subsidy_id])
            out.plan_before = plan_before_all.get(target_subsidy_id)

            await _ensure_backup_table(db)
            await _backup_items(db, p89.id, pair.reg7, pair.reg89)

            old_item_ids = [it.id for it in new_items]
            await db.execute(
                PurchaseItem.__table__.delete().where(PurchaseItem.id.in_(old_item_ids))
            )
            await db.flush()

            overrides = {}
            if target_fpi is not None:
                overrides["feo_planned_item_id"] = target_fpi
            if target_cat is not None:
                overrides["feo_category_id"] = target_cat
            await clone_purchase_items(db, p7.id, p89.id, overrides=overrides)
            await db.flush()

            amounts_after = await load_purchase_amounts(db, [p89.id])
            out.amount_after = amounts_after[p89.id].effective if p89.id in amounts_after else None
            plan_after_all = await subsidy_type_totals(db, [target_subsidy_id])
            out.plan_after = plan_after_all.get(target_subsidy_id)

            out.status = "ok"

            if apply:
                await db.commit()
            else:
                await db.rollback()
            return out
        except Exception:
            await db.rollback()
            raise


def render_report(outcomes: list[PairOutcome], *, apply: bool) -> str:
    lines = []
    mode = "APPLY (сохранено)" if apply else "DRY-RUN (откат в конце каждой пары)"
    lines.append(f"=== Восстановление позиций ФАДМ 2026_2 — {mode} ===")
    ok = [o for o in outcomes if o.status == "ok"]
    skipped = [o for o in outcomes if o.status == "skipped"]
    lines.append(f"Пар обработано: {len(outcomes)}; успешно: {len(ok)}; пропущено: {len(skipped)}")
    lines.append("")
    for o in ok:
        amt_match = "OK" if (o.amount_before == o.amount_after) else (
            f"РАЗОШЛОСЬ: {o.amount_before} -> {o.amount_after}"
        )
        plan_match = "OK" if (o.plan_before == o.plan_after) else (
            f"РАЗОШЁЛСЯ: {o.plan_before} -> {o.plan_after}"
        )
        lines.append(
            f"  {o.reg7} -> {o.reg89}: {o.n_old_items} позиций вместо {o.n_new_items_before}; "
            f"effective до/после: {amt_match}; план субсидии до/после: {plan_match}"
        )
    if skipped:
        lines.append("")
        lines.append("Пропущено:")
        for o in skipped:
            lines.append(f"  {o.reg7} -> {o.reg89}: {o.reason}")
    return "\n".join(lines)


async def main_async(args: argparse.Namespace) -> int:
    pairs = read_confirmed_pairs(args.pairs_xlsx)
    if not pairs:
        print("В листе «Пары» нет строк с действием «разбить на N позиций» и «Подтверждено»=«да».",
              file=sys.stderr)
        return 1
    outcomes = []
    for pair in pairs:
        outcome = await process_pair(
            pair, source_subsidy_id=args.source_subsidy_id, target_subsidy_id=args.target_subsidy_id,
            apply=args.apply,
        )
        outcomes.append(outcome)
    print(render_report(outcomes, apply=args.apply))
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--pairs-xlsx", required=True, help="Путь к fadm_restore_pairs.xlsx (внутри контейнера)")
    parser.add_argument("--source-subsidy-id", type=int, default=7, help="Субсидия-источник (ФАДМ_2026), по умолчанию 7")
    parser.add_argument("--target-subsidy-id", type=int, default=89, help="Субсидия-получатель (ФАДМ 2026_2), по умолчанию 89")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--dry-run", action="store_true", help="Всё в транзакции на пару, откат (по умолчанию)")
    group.add_argument("--apply", action="store_true", help="Сохранить изменения (commit на пару)")
    args = parser.parse_args()
    args.apply = bool(args.apply)  # --dry-run — поведение по умолчанию при отсутствии --apply

    exit_code = asyncio.run(main_async(args))
    sys.exit(exit_code)


if __name__ == "__main__":
    main()
