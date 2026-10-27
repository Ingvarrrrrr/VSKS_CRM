"""Шаблон импорта факта в раскладке GoodsService (layout='stages', владелец,
27.10.2026) — второй шаблон рядом со старым ХО (layout='single',
template.py). Проверяет:
  1) build_stages_template_workbook собирает лист с РОВНО тем заголовком,
     который columns.py распознаёт (та же проверка, что
     test_fact_import_stage_columns.py::_gs_header, но из сгенерированного
     файла, не ручной копии) — стадии/«Оплачено в деньгах»/4 уровня ФЭО/план
     замаплены;
  2) строка плана (FeoCategory → FeoPlannedItem, тот же plan_catalog, что
     читает matching.py, ПРАВИЛО №6) предзаполнена именем/количеством/
     ценой/суммой/путём ФЭО;
  3) заполнение стадий Да/Нет в строке файла даёт ожидаемый статус через
     тот же разбор, что test_goodsservice_status_from_stage_flags_all_au_
     values (statuses.resolve_status после rows.parse_rows);
  4) багфикс template.py (layout='single') — колонка «Субсидия» несёт имя
     РЕАЛЬНОЙ субсидии, не захардкоженную «ХО_2026».
"""
import uuid

import pytest
from openpyxl import load_workbook

from app.services.historical_fact_import import columns as columns_mod
from app.services.historical_fact_import import rows as rows_mod
from app.services.historical_fact_import import statuses as statuses_mod
from app.services.historical_fact_import.template import build_template_workbook
from app.services.historical_fact_import.template_stages import build_stages_template_workbook


async def _make_subsidy(db_session, org_id, name=None):
    from app.models.subsidy import Subsidy
    s = Subsidy(name=name or f"FactImport-Stages-{uuid.uuid4().hex[:8]}", year=2026, org_id=org_id)
    db_session.add(s)
    await db_session.commit()
    await db_session.refresh(s)
    return s


async def _make_leaf_category(db_session, subsidy_id, name, parent_id=None, level=3):
    from app.models.feo_category import FeoCategory
    c = FeoCategory(subsidy_id=subsidy_id, parent_id=parent_id, level=level, name=name)
    db_session.add(c)
    await db_session.commit()
    await db_session.refresh(c)
    return c


async def _make_planned_item(db_session, category_id, name, amount, unit_price, quantity):
    from app.models.feo_planned_item import FeoPlannedItem
    it = FeoPlannedItem(feo_category_id=category_id, name=name, amount=amount,
                         unit_price=unit_price, quantity=quantity, is_active=True)
    db_session.add(it)
    await db_session.commit()
    await db_session.refresh(it)
    return it


def _read_sheet_rows(content: bytes) -> list:
    from io import BytesIO
    wb = load_workbook(BytesIO(content))
    ws = wb["Импорт факта"]
    return [list(row) for row in ws.iter_rows(values_only=True)]


@pytest.mark.asyncio
async def test_stages_template_header_matches_parser_and_prefills_plan(db_session, test_org):
    subsidy = await _make_subsidy(db_session, test_org.id)
    cat = await _make_leaf_category(db_session, subsidy.id, "Командировочные расходы")
    await _make_planned_item(db_session, cat.id, "Билеты", amount=50000, unit_price=5000, quantity=10)

    content = await build_stages_template_workbook(db_session, subsidy.id)
    all_rows = _read_sheet_rows(content)
    header, data_rows = all_rows[0], all_rows[1:]

    detected = {"format": "columns", "header_row": 1, "rows": all_rows, "header": header}
    cols = columns_mod.build_columns(detected)
    by_field = {c["field"]: c["letter"] for c in cols if c["field"]}

    # Все поля, которые разбирает лист GoodsService, замаплены.
    for field in (
        "purchase_no", "plan_item_name", "supplier", "plan_qty", "unit", "plan_price",
        "plan_amount", "fact_price", "fact_amount", "item_type",
        "path_l2", "path_l3", "path_l4", "path_l5",
        "stage_purchase_done", "stage_contracted", "stage_ordered", "stage_delivered",
        "stage_paid", "paid", "advance",
    ):
        assert field in by_field, f"поле {field} не замаплено из сгенерированного шаблона"

    # Строка плана предзаполнена: имя, количество, цена, сумма, путь ФЭО.
    plan_item_col = [c for c in cols if c["field"] == "plan_item_name"][0]["index"]
    qty_col = [c for c in cols if c["field"] == "plan_qty"][0]["index"]
    price_col = [c for c in cols if c["field"] == "plan_price"][0]["index"]
    amount_col = [c for c in cols if c["field"] == "plan_amount"][0]["index"]
    # Единственная категория без предков — путь из 1 части ложится в САМЫЙ
    # БЛИЖНИЙ к листу уровень (path_l5), а не path_l2 (entry_name_and_levels
    # заполняет уровни с конца, как и в template.py — см. template_common.py).
    l5_col = [c for c in cols if c["field"] == "path_l5"][0]["index"]

    assert len(data_rows) == 1
    row = data_rows[0]
    assert row[plan_item_col] == "Билеты"
    assert row[qty_col] == 10
    assert row[price_col] == 5000
    assert row[amount_col] == 50000
    assert row[l5_col] == "Командировочные расходы"


@pytest.mark.asyncio
async def test_stages_template_stage_flags_resolve_to_expected_status(db_session, test_org):
    subsidy = await _make_subsidy(db_session, test_org.id)
    cat = await _make_leaf_category(db_session, subsidy.id, "Категория")
    await _make_planned_item(db_session, cat.id, "Товар", amount=1000, unit_price=1000, quantity=1)

    content = await build_stages_template_workbook(db_session, subsidy.id)
    all_rows = _read_sheet_rows(content)
    header = all_rows[0]
    detected = {"format": "columns", "header_row": 1, "rows": all_rows, "header": header}
    cols = columns_mod.build_columns(detected)
    by_field = {c["field"]: c["index"] for c in cols if c["field"]}

    # Строим строку «всё оплачено» вручную по индексам стадий из сгенерированного заголовка.
    row = list(all_rows[1])
    row[by_field["stage_purchase_done"]] = "Да"
    row[by_field["stage_contracted"]] = "Да"
    row[by_field["stage_ordered"]] = "Да"
    row[by_field["stage_delivered"]] = "Да"
    row[by_field["stage_paid"]] = "Да"
    row[by_field["paid"]] = 1000

    detected_filled = {"format": "columns", "header_row": 1, "rows": [header, row], "header": header}
    parsed = rows_mod.parse_rows(detected_filled, cols, header_row=1)
    assert len(parsed) == 1
    resolved = statuses_mod.resolve_status(parsed[0]["status_raw"])
    assert resolved["target_status"] == "paid"
    assert parsed[0]["paid"] == 1000


@pytest.mark.asyncio
async def test_single_template_subsidy_column_uses_real_name_not_hardcoded(db_session, test_org):
    subsidy = await _make_subsidy(db_session, test_org.id, name="Субсидия-Минтруд-2026")
    cat = await _make_leaf_category(db_session, subsidy.id, "Категория")
    await _make_planned_item(db_session, cat.id, "Товар", amount=1000, unit_price=1000, quantity=1)

    content = await build_template_workbook(db_session, subsidy.id)
    all_rows = _read_sheet_rows(content)
    data_row = all_rows[1]
    assert data_row[0] == "Субсидия-Минтруд-2026"
    assert data_row[0] != "ХО_2026"
