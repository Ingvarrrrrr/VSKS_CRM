"""plan_graph_export_summary_sheet.py — рендер листа «Сводная» живого
экспорта плана-графика (Задача А, владелец 07.10.2026, план .planning/quick/
2026-10-07-plan-graph-export/PLAN.md). Вынесено из plan_graph_export_xlsx.py
(Правило №5). Лист НИЧЕГО не считает сам — только рисует уже готовые числа
из app.services.subsidy_summary_by_kind.subsidy_summary_by_kind.

Раскладка по строке владельца: заголовок-титул, баннер «Запланированные
траты» над D..G, заголовки, строка «Как считается», затем по строке на вид
(«Товары»/«Услуги» всегда, «ФОТ»/«Без типа» — только если хоть одно число
ненулевое), «ИТОГО» с формулами СУММ(...). D/G/I — формулы Excel (не
значения), чтобы при ручной правке других столбцов пересчитывались сами."""
from __future__ import annotations

try:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover — openpyxl отсутствует, роутер уже отказал раньше
    pass

from app.services.item_type_split import ALL_KINDS

_KIND_LABELS = {"goods": "Товары", "services": "Услуги", "payroll": "ФОТ", "unspecified": "Без типа"}
_METRIC_KEYS = ("feo", "plan", "committed", "monthly", "likely", "nice", "paid", "delivered_unpaid")
# Строки для «Товары»/«Услуги» показываются всегда; «ФОТ»/«Без типа» —
# только если у вида есть хоть одно ненулевое число (ALL_KINDS уже идёт в
# порядке goods, services, payroll, unspecified — ровно порядок владельца).
_ALWAYS_SHOWN_KINDS = {"goods", "services"}

HEADERS = [
    "Товары или Услуги",
    "План по ФЭО, ₽",
    "План закупок (плановые позиции), ₽",
    "Планируется заказать, ₽",
    "Ежемесячные платежи (договор есть, заказ не оформлен), ₽",
    "Реальные будущие траты — скорее всего понадобятся, ₽",
    "Выдуманные будущие траты — хотелось бы, двигать как угодно, ₽",
    "Заказано, ₽",
    "Поставлено, ₽",
    "Оплачено, ₽",
    "Поставлено, не оплачено (треб. оплата), ₽",
    # Владелец 08.10.2026, часть C — Σ остатков договоров (app.services.
    # contract_balances.contract_balances) по виду ГОЛОВЫ договора (не
    # второй расчёт — читает уже посчитанные группы).
    "Остаток на договорах, ₽",
]
_HOW_CALCULATED = [
    "Как считается",
    "Бюджет ФЭО по статьям, разнесён по видам позиций",
    "Сумма плановых позиций субсидии",
    "= План закупок − Заказано (формула)",
    "Заказы рамочных договоров в статусе «Договор»: договор есть, заказ ещё не оформлен",
    "Остаток плановых позиций с отметкой «Скорее всего понадобится», ещё не заказанный (без ежемесячных)",
    "= Планируется заказать − Ежемесячные − Реальные (формула): позиции «Хотелось бы»",
    "Законтрактовано: разовые — с этапа «Договор», рамочные — суммы заказов",
    "= Оплачено + Поставлено, не оплачено (формула)",
    "Сумма оплат по отметке в закупках",
    "Остаток к оплате по поставленным закупкам",
    "Сумма договоров минус заявки по ним (рамочные без суммы договора — 0)",
]
COL_WIDTHS = [16, 20, 20, 20, 20, 20, 20, 20, 20, 20, 20, 20]

NUM_FMT = "#,##0.00"


def write_summary_sheet(wb, sub, summary_by_kind: dict, contract_remaining_by_kind: dict = None) -> None:
    """Пишет лист «Сводная» (вставляется первым, index 0). `summary_by_kind` —
    результат subsidy_summary_by_kind(db, subsidy_id). `contract_remaining_by_kind`
    (владелец 08.10.2026, часть C) — {kind: Σ остатков договоров}, из
    app.services.contract_balances.contract_balances (группы, агрегированные
    по kind вызывающим кодом, plan_graph_export_data.gather_live_plan_graph_data)
    — столбец «Остаток на договорах, ₽»; None/пусто — столбец весь 0."""
    HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=9)
    BANNER_FILL = PatternFill("solid", fgColor="FFF2CC")
    D_FILL = PatternFill("solid", fgColor="FFFF00")
    GOODS_FILL = PatternFill("solid", fgColor="F0F9FF")
    SERVICES_FILL = PatternFill("solid", fgColor="F0FDF4")
    OTHER_FILL = PatternFill("solid", fgColor="FAFAFA")
    TOTAL_FILL = PatternFill("solid", fgColor="DBEAFE")
    SUB_FONT = Font(size=8, italic=True, color="6B7280")
    CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT_ALIGN = Alignment(horizontal="right", vertical="center", wrap_text=True)
    THIN_BORDER = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    sm = wb.create_sheet("Сводная", 0)
    n_cols = len(HEADERS)
    last_col_letter = get_column_letter(n_cols)

    # Строка 1 — титул.
    title = sm.cell(row=1, column=1, value=f"СВОДНАЯ — {sub.name} ({sub.year})")
    title.font = Font(bold=True, size=12, color="1E3A5F")
    sm.merge_cells(f"A1:{last_col_letter}1")
    title.alignment = CENTER_ALIGN
    sm.row_dimensions[1].height = 28

    # Строка 2 — баннер «Запланированные траты» над D..G, остальное пусто.
    sm.merge_cells("D2:G2")
    banner = sm.cell(row=2, column=4, value="Запланированные траты")
    banner.alignment = CENTER_ALIGN
    banner.font = Font(bold=True, size=9)
    for ci in range(1, n_cols + 1):
        sm.cell(row=2, column=ci).fill = BANNER_FILL
        sm.cell(row=2, column=ci).border = THIN_BORDER

    # Строка 3 — заголовки столбцов (D — отдельная жёлтая заливка владельца).
    for ci, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = sm.cell(row=3, column=ci, value=h)
        cell.fill = D_FILL if ci == 4 else HEADER_FILL
        cell.font = Font(bold=True, size=9) if ci == 4 else HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
        sm.column_dimensions[cell.column_letter].width = w

    # Строка 4 — «Как считается».
    for ci, text in enumerate(_HOW_CALCULATED, 1):
        cell = sm.cell(row=4, column=ci, value=text)
        cell.font = SUB_FONT
        cell.alignment = LEFT_ALIGN
        cell.border = THIN_BORDER

    row = 5
    first_data_row = row
    _KIND_FILL = {"goods": GOODS_FILL, "services": SERVICES_FILL, "payroll": OTHER_FILL, "unspecified": OTHER_FILL}
    for kind in ALL_KINDS:
        d = summary_by_kind.get(kind) or {}
        if kind not in _ALWAYS_SHOWN_KINDS and not any(abs(d.get(k, 0.0)) > 0.005 for k in _METRIC_KEYS):
            continue

        label_cell = sm.cell(row=row, column=1, value=_KIND_LABELS[kind])
        sm.cell(row=row, column=2, value=round(d.get("feo", 0.0), 2))
        sm.cell(row=row, column=3, value=round(d.get("plan", 0.0), 2))
        sm.cell(row=row, column=4, value=f"=C{row}-H{row}")
        sm.cell(row=row, column=5, value=round(d.get("monthly", 0.0), 2))
        sm.cell(row=row, column=6, value=round(d.get("likely", 0.0), 2))
        sm.cell(row=row, column=7, value=f"=D{row}-E{row}-F{row}")
        sm.cell(row=row, column=8, value=round(d.get("committed", 0.0), 2))
        sm.cell(row=row, column=9, value=f"=J{row}+K{row}")
        sm.cell(row=row, column=10, value=round(d.get("paid", 0.0), 2))
        sm.cell(row=row, column=11, value=round(d.get("delivered_unpaid", 0.0), 2))
        sm.cell(row=row, column=12, value=round((contract_remaining_by_kind or {}).get(kind, 0.0), 2))

        fill = _KIND_FILL[kind]
        for ci in range(1, n_cols + 1):
            c = sm.cell(row=row, column=ci)
            c.fill = fill
            c.border = THIN_BORDER
            if ci >= 2:
                c.number_format = NUM_FMT
                c.alignment = RIGHT_ALIGN
            else:
                c.alignment = LEFT_ALIGN
        label_cell.font = Font(bold=True, size=9)
        row += 1

    last_data_row = row - 1
    total_row = row
    sm.cell(row=total_row, column=1, value="ИТОГО")
    for ci in range(2, n_cols + 1):
        col_letter = get_column_letter(ci)
        if last_data_row >= first_data_row:
            formula = f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})"
        else:
            formula = 0
        c = sm.cell(row=total_row, column=ci, value=formula)
        c.number_format = NUM_FMT
        c.alignment = RIGHT_ALIGN
        c.border = THIN_BORDER
        c.fill = TOTAL_FILL
    total_label = sm.cell(row=total_row, column=1)
    total_label.font = Font(bold=True, size=10)
    total_label.border = THIN_BORDER
    total_label.fill = TOTAL_FILL

    sm.freeze_panes = "A5"
