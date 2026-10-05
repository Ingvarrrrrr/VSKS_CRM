"""Необязательный блок «Факт (если закупка уже прошла)» шаблона ФЭО —
решение владельца 05.10.2026: субсидию, у которой часть закупок уже прошла,
грузят ОДНИМ файлом.

Проверяет:
  1. Шаблон (`download_feo_template`) содержит 8 новых колонок блока «Факт»
     и выпадающий список статусов на листе «Справочники».
  2. Импорт ФЭО (`import_feo_from_excel`) файла с заполненными колонками
     факта строит ТОТ ЖЕ план, что и без них, и отдаёт
     has_fact_columns=True/fact_rows>0 в ответе.
  3. Импорт ФЭО файла БЕЗ блока «Факт» (старый формат) — has_fact_columns=
     False/fact_rows=0, план не меняется.
  4. `/fact-import/preview` (build_preview) на ТОМ ЖЕ файле распознаёт
     колонки блока «Факт» автоматически (без ручного mapping) и находит
     строку со статусом.

Гонять по одному узлу за вызов pytest в контейнере (флейк pytest-asyncio
«different loop» — project_pytest_asyncio_loop_flake, Lessons.md VSKS_CRM).
"""
import io
import uuid

import pytest
from fastapi import UploadFile
from openpyxl import Workbook

from app.routers.feo_import import import_feo_from_excel
from app.routers.feo_import_template import download_feo_template
from app.services.historical_fact_import.statuses import STATUS_CHOICES
from tests.test_feo_import_tree import _cleanup_subsidy, _get_categories, _get_items, _make_subsidy

_FACT_HEADERS = [
    "Правильный статус", "Факт: Количество", "Факт: Цена", "Факт: Сумма",
    "Оплачено", "Законтрактовано", "Поставщик", "№ закупки",
]


def _mk_xlsx_upload(headers, rows) -> UploadFile:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return UploadFile(file=bio, filename=f"import-{uuid.uuid4().hex[:6]}.xlsx")


@pytest.mark.asyncio
async def test_template_has_fact_block_and_status_dropdown(db_session, superadmin_user):
    """Шаблон несёт 8 колонок блока «Факт» и лист «Справочники» со всеми 6
    статусами GALA (dropdown-источник для «Правильный статус»)."""
    from openpyxl import load_workbook

    response = await download_feo_template(subsidy_id=None, db=db_session, current_user=superadmin_user)
    body = b"".join([chunk async for chunk in response.body_iterator])
    wb = load_workbook(io.BytesIO(body))
    ws = wb["Категории ФЭО"]
    headers = [c.value for c in ws[1]]
    assert headers[-8:] == [
        "Правильный статус", "Факт: Количество", "Факт: Цена", "Факт: Сумма",
        "Оплачено", "Законтрактовано", "Поставщик", "№ закупки",
    ], headers

    ws3 = wb["Справочники"]
    assert ws3["G1"].value == "Правильный статус"
    status_labels_on_sheet = [ws3.cell(r, 7).value for r in range(2, 2 + len(STATUS_CHOICES))]
    assert status_labels_on_sheet == [c["label"] for c in STATUS_CHOICES]


async def _import_rows(db_session, superadmin_user, headers, rows):
    subsidy = await _make_subsidy(db_session)
    full_headers = ["Субсидия"] + headers
    full_rows = [[subsidy.name] + row for row in rows]
    upload = _mk_xlsx_upload(full_headers, full_rows)
    result = await import_feo_from_excel(
        file=upload, dry_run=False, remap="", apply_remap=False,
        duplicate_resolutions="", item_type_decisions="",
        db=db_session, current_user=superadmin_user,
    )
    return subsidy, result


_PLAN_HEADERS = ["Уровень 2", "Плановая позиция", "Плановое количество", "Плановая цена за единицу"]
_PLAN_ROW = ["Направление — тест блока факт", "Позиция — тест блока факт", "2", "500"]


@pytest.mark.asyncio
async def test_import_with_fact_columns_builds_same_plan_and_reports_counts(db_session, superadmin_user):
    """Файл С заполненным блоком «Факт» строит тот же план, что и без него,
    и отдаёт has_fact_columns=True/fact_rows=1 — новые колонки игнорируются
    при построении дерева (Правило №6: единственный источник факта — Импорт
    факта, этот импорт их только считает)."""
    headers = _PLAN_HEADERS + _FACT_HEADERS
    row = _PLAN_ROW + ["Заключён договор", "2", "500", "1000", "500", "1000", "ООО Тест", "РЕЕ-1"]
    subsidy, result = await _import_rows(db_session, superadmin_user, headers, [row])
    try:
        assert result["errors"] == [], result["errors"]
        assert result["has_fact_columns"] is True
        assert result["fact_rows"] == 1

        cats = await _get_categories(db_session, subsidy.id)
        lvl2_cat = next(c for c in cats if c.name == "Направление — тест блока факт")
        items = await _get_items(db_session, lvl2_cat.id)
        assert len(items) == 1
        assert items[0].name == "Позиция — тест блока факт"
        assert float(items[0].quantity) == 2
        assert float(items[0].unit_price) == 500
    finally:
        await _cleanup_subsidy(db_session, subsidy.id)


@pytest.mark.asyncio
async def test_import_without_fact_columns_reports_zero(db_session, superadmin_user):
    """Старый формат файла (без блока «Факт») — has_fact_columns=False,
    fact_rows=0, план не меняется."""
    subsidy, result = await _import_rows(db_session, superadmin_user, _PLAN_HEADERS, [_PLAN_ROW])
    try:
        assert result["errors"] == [], result["errors"]
        assert result["has_fact_columns"] is False
        assert result["fact_rows"] == 0
    finally:
        await _cleanup_subsidy(db_session, subsidy.id)


@pytest.mark.asyncio
async def test_import_fact_row_with_plan_status_not_counted(db_session, superadmin_user):
    """Блок «Факт» присутствует (колонки есть), но статус строки — «План
    закупок» (ещё не факт) -> has_fact_columns=True, fact_rows=0."""
    headers = _PLAN_HEADERS + _FACT_HEADERS
    row = _PLAN_ROW + ["План закупок", "", "", "", "", "", "", ""]
    subsidy, result = await _import_rows(db_session, superadmin_user, headers, [row])
    try:
        assert result["errors"] == [], result["errors"]
        assert result["has_fact_columns"] is True
        assert result["fact_rows"] == 0
    finally:
        await _cleanup_subsidy(db_session, subsidy.id)


@pytest.mark.asyncio
async def test_fact_import_preview_autodetects_feo_template_columns(db_session, superadmin_user):
    """Тот же файл (план ФЭО шаблона + заполненный блок «Факт») — «Импорт
    факта» (build_preview) распознаёт колонки АВТОМАТИЧЕСКИ (mapping=None) и
    находит строку со статусом, не требуя ручного сопоставления."""
    from app.services.historical_fact_import.preview import build_preview

    subsidy = await _make_subsidy(db_session)
    headers = ["Субсидия"] + _PLAN_HEADERS + _FACT_HEADERS
    row = [subsidy.name] + _PLAN_ROW + ["Заключён договор", "2", "500", "1000", "500", "1000", "ООО Тест", "РЕЕ-1"]
    upload = _mk_xlsx_upload(headers, [row])
    content = await upload.read()
    try:
        preview = await build_preview(
            db_session, subsidy.id, content, upload.filename or "", None, None, None,
        )
        fields_by_header = {c["header"]: c["field"] for c in preview["columns"] if c["field"]}
        assert fields_by_header.get("Уровень 2") == "path_l2"
        assert fields_by_header.get("Плановая позиция") == "plan_item_name"
        assert fields_by_header.get("Плановое количество") == "plan_qty"
        assert fields_by_header.get("Плановая цена за единицу") == "plan_price"
        assert fields_by_header.get("Правильный статус") == "status_raw"
        assert fields_by_header.get("Факт: Количество") == "fact_qty"
        assert fields_by_header.get("Факт: Цена") == "fact_price"
        assert fields_by_header.get("Факт: Сумма") == "fact_amount"
        assert fields_by_header.get("Законтрактовано") == "contracted"
        assert fields_by_header.get("Поставщик") == "supplier"

        rows_with_status = [r for r in preview["rows"] if r["status_raw"]]
        assert len(rows_with_status) == 1, preview["rows"]
        assert rows_with_status[0]["status"] == "contracted"
        assert rows_with_status[0]["name"] == "Позиция — тест блока факт"
    finally:
        await _cleanup_subsidy(db_session, subsidy.id)
