"""test_amount_chain_remnants.py — ПРАВИЛО №6, волна 4b-2c (2026-09-07):
остатки старых truthy-цепочек «сумма закупки» вне app.services.purchase_amounts,
найденные QA в contracts.py / reports.py / subsidy_plan_graph_export.py.

Общий сценарий регрессии для всех трёх мест: contract_price=0 (легитимный
НОЛЬ, не «пусто») + planned_total_price=1000. Старый код (`x or y or 0`,
Python truthy) молча проваливал 0 дальше по цепочке и показывал 1000. Новый
код (contract_amount()/purchase_amounts()/load_purchase_amounts() —
единственный источник, `is not None`) обязан вернуть 0.
"""
import io
import uuid
from decimal import Decimal

import pytest


async def _make_subsidy(db_session, org_id=None, **kwargs):
    from app.models.subsidy import Subsidy
    s = Subsidy(
        name=f"TestSubsidy-{uuid.uuid4().hex[:8]}", year=2026, budget=1_000_000,
        require_planned_dates=False, org_id=org_id, **kwargs,
    )
    db_session.add(s)
    await db_session.commit()
    await db_session.refresh(s)
    return s


# ---------------------------------------------------------------------------
# 1) contracts.py::_enrich_contract_from_purchases — было `total_nmck or
#    contract_price or planned_total_price`
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_contracts_enrich_default_not_truthy_on_zero_contract_price(
    db_session, make_purchase,
):
    from app.models.contract import Contract
    from app.routers.contracts import _enrich_contract_from_purchases

    contract = Contract(number=f"C-{uuid.uuid4().hex[:8]}", contract_type="single", max_amount=None)
    db_session.add(contract)
    await db_session.commit()
    await db_session.refresh(contract)

    p = await make_purchase(
        contract_id=contract.id,
        contract_price=Decimal("0"),
        planned_total_price=Decimal("1000"),
        total_nmck=None,
    )

    filled = await _enrich_contract_from_purchases(contract, db_session)
    assert filled > 0
    # НЕ 1000 — старый `total_nmck or contract_price or planned_total_price`
    # (truthy) провалил бы легитимный 0 и подставил бы план.
    assert contract.max_amount == Decimal("0")


@pytest.mark.asyncio
async def test_contracts_enrich_falls_back_to_plan_when_no_contract_price(
    db_session, make_purchase,
):
    """Контроль: когда contract_price реально ПУСТ (None, не 0) — фолбэк на
    purchase_amounts(p).plan по-прежнему работает (не регресс в другую сторону)."""
    from app.models.contract import Contract
    from app.routers.contracts import _enrich_contract_from_purchases

    contract = Contract(number=f"C-{uuid.uuid4().hex[:8]}", contract_type="single", max_amount=None)
    db_session.add(contract)
    await db_session.commit()
    await db_session.refresh(contract)

    p = await make_purchase(
        contract_id=contract.id,
        contract_price=None,
        planned_total_price=Decimal("1000"),
        total_nmck=None,
    )

    filled = await _enrich_contract_from_purchases(contract, db_session)
    assert filled > 0
    assert contract.max_amount == Decimal("1000")


# ---------------------------------------------------------------------------
# 2) reports.py::export_subsidy_report_xlsx, колонка L «Сумма договора, руб.»
#    — было `contract_price or planned_total_price`
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_reports_xlsx_contract_column_not_truthy_on_zero(
    client, db_session, make_purchase, auth_headers,
):
    from openpyxl import load_workbook
    from app.services.purchase_amounts import contract_amount

    subsidy = await _make_subsidy(db_session)
    p = await make_purchase(
        status="contracted",
        subsidy_id=subsidy.id,
        contract_price=Decimal("0"),
        planned_total_price=Decimal("1000"),
    )
    expected = contract_amount(p)
    assert expected == Decimal("0")  # sanity: сценарий действительно бьёт по truthy-багу

    resp = await client.get(f"/api/reports/subsidy/{subsidy.id}/xlsx", headers=auth_headers)
    assert resp.status_code == 200, resp.text

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    # reports.py: row = 6 — первая (и здесь единственная) строка данных, колонка L = 12.
    cell_val = ws.cell(row=6, column=12).value
    # НЕ "1 000.00" — contract_amount(), не подмена план-суммой.
    assert cell_val == "0.00", f"колонка L = {cell_val!r}, ожидалось '0.00' (не 1000 от плана)"


# ---------------------------------------------------------------------------
# 3) subsidy_plan_graph_export.py::export_plan_graph_excel, секция «закупки
#    без категории ФЭО» — было `final_total_amount or planned_total_price or 0`
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_subsidy_plan_graph_export_matches_purchase_amounts(
    client, db_session, make_purchase, test_org, auth_headers,
):
    from openpyxl import load_workbook
    from app.services.purchase_amounts import purchase_amounts

    subsidy = await _make_subsidy(db_session, org_id=test_org.id)
    # make_purchase() хардкодит item_name="Test purchase" (не переопределяется через
    # kwargs — Purchase() падает на "multiple values for keyword argument"); эта
    # закупка — единственная итемлесс-закупка данной субсидии в изолированной
    # транзакции теста, так что фиксированное имя уникально идентифицирует строку.
    p = await make_purchase(
        status="contracted",
        subsidy_id=subsidy.id,
        contract_price=Decimal("0"),
        planned_total_price=Decimal("1000"),
        final_total_amount=None,
    )
    expected = purchase_amounts(p).effective
    assert expected == Decimal("0")  # sanity: тот же сценарий 0-vs-план

    resp = await client.get(f"/api/subsidies/{subsidy.id}/plan-graph/export", headers=auth_headers)
    assert resp.status_code == 200, resp.text

    wb = load_workbook(io.BytesIO(resp.content))
    # wb.active — сводная сводка («Сводная»); строки закупок — на листе «План закупок».
    ws = wb["План закупок"]
    # Найти строку с нашей закупкой по имени (секция «без категории ФЭО» пишет
    # "    └ {name} — ...руб/ед." в колонку D) и прочитать колонку M
    # («Заказано/факт» — pi["total"], та самая сумма) на той же строке.
    found_row = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "Test purchase" in cell.value:
                found_row = cell.row
                break
        if found_row:
            break
    assert found_row is not None, "строка с тестовой закупкой не найдена в выгрузке"
    total_val = ws.cell(row=found_row, column=13).value  # column M
    # НЕ 1000 (planned_total_price) — final_total_amount убран из формулы,
    # эффективная сумма по стадии = contract_price = 0.
    assert float(total_val or 0) == 0.0, f"колонка M = {total_val!r}, ожидалось 0 (не 1000 от плана)"
