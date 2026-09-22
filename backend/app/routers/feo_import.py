"""Импорт категорий ФЭО из Excel (по заголовку и по пользовательскому
маппингу колонок) — ядро app/routers/feo_import.py.

Разрезано из app/routers/feo_categories.py (Правило №5, модульность) без
изменения поведения. Дальнейшее разрезание (сессия 2026-09-08, Правило №5,
файл разросся до 1088 строк): шаблон (/import/template), предпросмотр
(/import-preview) и экспорт (/export) — самодостаточные эндпоинты, не
вызывающие _do_feo_import, — уехали соседями app/routers/feo_import_template.py,
feo_import_preview.py, feo_import_export.py. Ссылочные хелперы
_relink_feo_category/_feo_category_load (чистая логика, без FastAPI-декораторов)
уехали в app/services/feo_import_links.py и ре-экспортируются здесь ниже —
и feo_import_remap.py/feo_import_report.py (`from app.routers.feo_import import
_feo_category_load, _relink_feo_category`, ЛОКАЛЬНО внутри функции), и
feo_categories.py (лениво, PEP 562 `__getattr__`) продолжают резолвить эти
имена через ЭТОТ модуль независимо от того, где физически лежит их тело.

Здесь остаются только /import и /import-mapped — они РАЗДЕЛЯЮТ движок
_do_feo_import (app/services/feo_import_engine.py) и объёмный разбор колонок
(find_col/per-level фолбэки), поэтому держатся вместе, а не режутся дальше.

Тяжёлое ядро парсинга (_do_feo_import) вынесено в
app/services/feo_import_engine.py — этот модуль только определяет колонки
файла и передаёт готовые индексы туда. Гейты записи зовутся через
`from app.routers import feo_categories as fc` — так monkeypatch
`fc._require_feo_category_write` в тестах продолжает работать независимо от
того, в каком файле реально живёт вызывающий обработчик.
"""
from io import BytesIO

from fastapi import APIRouter, Depends, Query, HTTPException, Request, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from app.database import get_db
from app.auth.permissions import require_tab
from app.routers import feo_categories as fc
from app.services.feo_import_common import get_cell, parse_numbering_cell
from app.services.feo_import_engine import _do_feo_import
from app.services.feo_import_links import _relink_feo_category, _feo_category_load  # noqa: F401 (re-export)
from app.services.feo_import_params import resolve_feo_import_mapped_params

# Нумерация строк A–D (22.09, боевой случай ДНР_2026) — главный источник
# иерархии, когда явно НЕ задана маппингом (/import-mapped может передать
# col_num1..col_num4 напрямую): для /import (импорт по заголовку) колонки
# нумерации не имеют заголовка вовсе — они авто-детектятся по СОДЕРЖИМОМУ:
# ведущие (с колонки 0) столбцы файла, где КАЖДАЯ непустая ячейка данных —
# целое число ИЛИ склеенный точками путь вида "2.2.1" (см.
# feo_import_common.parse_numbering_cell — тот же боевой файл ДНР_2026
# кладёт в один столбец то целое, то "2.2"/"2.3"/"2.2.1", см. докстринг
# feo_import_numbering.py::parse_row_path), до первой не подходящей под это
# колонки. Единственное место такого детекта (Правило №6) — используется и
# /import ниже, и parse_row_path (тот же критерий «валидная ячейка
# нумерации»).
_MAX_NUMBERING_COLS = 4


def detect_numbering_columns(data_rows: list, boundary_col: int | None) -> list[int]:
    """Ведущие (с колонки 0) столбцы file[0:boundary_col], которые целиком
    состоят из пустых/целых/склеенных-точками значений — колонки нумерации
    A–D. `boundary_col` обычно индекс колонки «Субсидия» (нумерация всегда
    стоит ДО неё в реальном файле) — без неё (None) используется c_lvl2.
    Меньше 2 таких ведущих колонок — нумерации в файле нет, возвращается
    пустой список (обычный построчный импорт без изменений)."""
    if not boundary_col:
        return []
    max_col = min(boundary_col, _MAX_NUMBERING_COLS)
    cols: list[int] = []
    for col in range(max_col):
        has_value = False
        ok = True
        for row in data_rows:
            v = get_cell(row, col)
            if v is None:
                continue
            has_value = True
            if parse_numbering_cell(v) is None:
                ok = False
                break
        if ok and has_value:
            cols.append(col)
        else:
            break
    return cols if len(cols) >= 2 else []

router = APIRouter(prefix="/api/feo-categories", tags=["feo_categories"])


@router.post("/import")
async def import_feo_from_excel(
    file: UploadFile = File(...),
    dry_run: bool = Query(False),
    remap: str = Query(""),
    apply_remap: bool = Query(False),
    duplicate_resolutions: str = Query(""),
    # Решения человека по конфликтам «тип позиции из файла vs тип товара
    # каталога» (владелец, 22.09) — JSON {номер_строки: "file"|"catalog"},
    # отдельный канал от duplicate_resolutions (см. feo_import_item_types.py).
    item_type_decisions: str = Query(""),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_tab('feo_categories')),
):
    """Импорт категорий ФЭО из Excel.
    Формат: Субсидия | Уровень 2 | Уровень 3 | Уровень 4 | Код | Приложение | Финансирование | Активна
    Каждая строка задаёт путь в иерархии. Промежуточные узлы создаются автоматически.
    Код/Приложение/Финансирование/Активна применяются к самому глубокому указанному уровню.
    dry_run=true: возвращает {created, updated, skipped, errors, warnings} без записи в БД.
    Переезд (remap) несопоставленных узлов и удаление опустевших старых узлов выполняются
    только при apply_remap=true; иначе выполняется только анализ (unmatched/new_paths).
    Возвращает {created, updated, skipped, errors, warnings}."""
    # B2: субсидия — построчная колонка внутри файла (может касаться нескольких
    # субсидий за один импорт), единого subsidy_id на уровне запроса нет.
    # Это ТОЛЬКО дешёвый предварительный фильтр («есть ли право хоть где-то» —
    # отсекает пользователей без feo_category.edit вообще, до чтения файла).
    # Авторитетная проверка — ПО КАЖДОЙ реально резолвящейся субсидии файла —
    # происходит внутри _do_feo_import (_require_feo_import_write, ДО первой
    # записи в БД); дыра импорта была именно в отсутствии этой второй проверки.
    await fc._require_feo_category_write(current_user, db, None)
    if load_workbook is None:
        raise HTTPException(500, "openpyxl не установлен")
    if not (file.filename or "").lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Поддерживаются только .xlsx и .xls")

    content = await file.read()
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(400, "Файл пустой")

    # Определяем индексы столбцов по заголовку строки 1
    raw_headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

    # Каждая колонка достаётся ровно одному полю
    _used_cols: set[int] = set()

    def find_col(keywords: list[str]) -> int | None:
        for kw in keywords:
            for i, h in enumerate(raw_headers):
                if i in _used_cols:
                    continue
                if kw in h:
                    _used_cols.add(i)
                    return i
        return None

    c_subsidy  = find_col(["субсидия"])
    c_lvl2     = find_col(["уровень 2", "направление расходов", "level 2"])
    c_lvl3     = find_col(["уровень 3", "тип расходов", "level 3"])
    c_lvl4     = find_col(["уровень 4", "конкретизир", "level 4"])
    c_lvl5     = find_col(["плановая позиция", "уровень 5", "плановый товар", "level 5"])
    c_qty      = find_col(["количество (ур.5)", "количество ур.5", "кол-во (ур.5)", "кол-во ур.5"])
    # Новый шаблон: «кол-во по фэо» — feo_quantity
    c_feo_qty_lvl2  = find_col(["кол-во по фэо (ур.2)", "кол-во по фэо ур.2"])
    c_feo_qty_lvl3  = find_col(["кол-во по фэо (ур.3)", "кол-во по фэо ур.3"])
    c_feo_qty_lvl4  = find_col(["кол-во по фэо (ур.4)", "кол-во по фэо ур.4"])
    c_feo_unit_lvl2 = find_col(["ед. изм. по фэо (ур.2)", "ед. изм. по фэо ур.2"])
    c_feo_unit_lvl3 = find_col(["ед. изм. по фэо (ур.3)", "ед. изм. по фэо ур.3"])
    c_feo_unit_lvl4 = find_col(["ед. изм. по фэо (ур.4)", "ед. изм. по фэо ур.4"])
    c_feo_amt_lvl2  = find_col(["стоимость по фэо (ур.2)", "стоимость по фэо ур.2"])
    c_feo_amt_lvl3  = find_col(["стоимость по фэо (ур.3)", "стоимость по фэо ур.3"])
    c_feo_amt_lvl4  = find_col(["стоимость по фэо (ур.4)", "стоимость по фэо ур.4"])
    # Сумма по ФЭО (итог строки) — НОВЫЕ колонки; НЕ путать со стоимостью за ед.
    c_feo_sum_lvl2  = find_col(["сумма по фэо (ур.2)", "сумма по фэо ур.2"])
    c_feo_sum_lvl3  = find_col(["сумма по фэо (ур.3)", "сумма по фэо ур.3"])
    c_feo_sum_lvl4  = find_col(["сумма по фэо (ур.4)", "сумма по фэо ур.4"])
    # Плановое кол-во (CRM-план)
    c_qty_lvl2 = find_col(["плановое кол-во (ур.2)", "плановое кол-во ур.2", "кол-во (ур.2)", "кол-во ур.2", "количество (ур.2)"])
    c_qty_lvl3 = find_col(["плановое кол-во (ур.3)", "плановое кол-во ур.3", "кол-во (ур.3)", "кол-во ур.3", "количество (ур.3)"])
    c_qty_lvl4 = find_col(["плановое кол-во (ур.4)", "плановое кол-во ур.4", "кол-во (ур.4)", "кол-во ур.4", "количество (ур.4)"])
    c_unit_lvl2 = find_col(["ед. изм. плана (ур.2)", "ед. изм. плана ур.2", "ед. изм. (ур.2)", "ед.изм. ур.2", "единица ур.2"])
    c_unit_lvl3 = find_col(["ед. изм. плана (ур.3)", "ед. изм. плана ур.3", "ед. изм. (ур.3)", "ед.изм. ур.3", "единица ур.3"])
    c_unit_lvl4 = find_col(["ед. изм. плана (ур.4)", "ед. изм. плана ур.4", "ед. изм. (ур.4)", "ед.изм. ур.4", "единица ур.4"])
    # Плановая стоимость за ед. — НЕ путать с «сумма плана»
    c_amt_lvl2 = find_col(["плановая стоимость за ед. (ур.2)", "плановая стоимость (ур.2)", "стоимость за ед. (ур.2)", "стоимость ур.2"])
    c_amt_lvl3 = find_col(["плановая стоимость за ед. (ур.3)", "плановая стоимость (ур.3)", "стоимость за ед. (ур.3)", "стоимость ур.3"])
    c_amt_lvl4 = find_col(["плановая стоимость за ед. (ур.4)", "плановая стоимость (ур.4)", "стоимость за ед. (ур.4)", "стоимость ур.4"])
    # Сумма плана — отдельные колонки; старые алиасы «плановая сумма / сумма ур.N» сюда, а НЕ в c_amt_lvl*
    c_plan_sum_lvl2 = find_col(["сумма плана (ур.2)", "плановая сумма (ур.2)", "сумма ур.2"])
    c_plan_sum_lvl3 = find_col(["сумма плана (ур.3)", "плановая сумма (ур.3)", "сумма ур.3"])
    c_plan_sum_lvl4 = find_col(["сумма плана (ур.4)", "плановая сумма (ур.4)", "сумма ур.4"])
    # Новый плоский 18-колоночный шаблон (2026-08-14): числа по ФЭО/плану — ОДНА
    # пара колонок на всю строку (не по уровням), см. _do_feo_import, блок «плоские
    # числа». Объявлены ПОСЛЕ всех per-level find_col выше и ДО generic-фолбэков
    # ниже (c_qty/c_unit) — иначе более общие ключи перехватили бы эти колонки
    # раньше специфичных. На старых 37-колоночных файлах колонки этих названий нет
    # (там «Кол-во по ФЭО (Ур.N)» уже разобран per-level выше и помечен used) — эти
    # find_col останутся None, старое поведение не меняется.
    c_row_feo_qty    = find_col(["количество по фэо"])
    c_row_feo_unit   = find_col(["ед. изм. по фэо"])
    c_row_feo_price  = find_col(["цена за единицу по фэо", "цена за ед. по фэо"])
    c_row_feo_sum    = find_col(["сумма по фэо"])
    c_row_plan_qty   = find_col(["плановое количество"])
    c_row_plan_unit  = find_col(["ед. изм. плана"])
    c_row_plan_price = find_col(["плановая цена за единицу", "плановая цена за ед."])
    c_row_plan_sum   = find_col(["сумма плана"])
    # Заголовок колонки переименован в «Тип (товар/услуга/работа)» (владелец,
    # 21.09, раздел W2 плана corrections-21-09.md) — старые файлы с «Товар/
    # услуга»/«Тип позиции» продолжают импортироваться (обратная совместимость).
    # Точные/специфичные варианты — ПЕРВЫМИ, generic «тип» — ПОСЛЕДНИМ: к этому
    # моменту c_lvl3 (find_col выше, «уровень 3»/«тип расходов») уже забрал
    # свою колонку в _used_cols, единственный другой заголовок в шаблоне,
    # содержащий «тип» — так что generic-фолбэк не перехватывает чужую колонку.
    c_item_type      = find_col(["тип (товар/услуга/работа)", "товар/услуга", "тип позиции", "тип"])
    # Fallback: generic qty column if no specific level columns present
    if c_qty is None and c_qty_lvl2 is None and c_qty_lvl3 is None and c_qty_lvl4 is None and c_feo_qty_lvl2 is None and c_feo_qty_lvl3 is None and c_feo_qty_lvl4 is None:
        c_qty = find_col(["количество", "кол-во", "qty"])
    c_unit      = find_col(["ед. измерения (ур.5)", "ед. изм. (ур.5)", "единица ур.5", "ед. изм", "единица изм", "ед.изм"])
    c_item_price = find_col(["цена за ед. (ур.5)", "цена за ед. ур.5"])
    c_item_amt  = find_col(["сумма по позиции (ур.5)", "сумма (ур.5)", "сумма ур", "плановая стоимость за ед. (ур.5)", "плановая стоимость (ур.5)", "стоимость за ед. (ур.5)", "стоимость ур.5", "сумма плановая"])
    c_code      = find_col(["код"])
    c_appendix  = find_col(["приложение"])
    c_budget    = find_col(["финансирование", "бюджет", "budget"])
    c_active    = find_col(["активна", "активен", "active"])

    _num_cols = detect_numbering_columns(rows[1:], c_subsidy if c_subsidy is not None else c_lvl2)
    c_num1 = _num_cols[0] if len(_num_cols) > 0 else None
    c_num2 = _num_cols[1] if len(_num_cols) > 1 else None
    c_num3 = _num_cols[2] if len(_num_cols) > 2 else None
    c_num4 = _num_cols[3] if len(_num_cols) > 3 else None

    return await _do_feo_import(
        rows=rows[1:],
        c_subsidy=c_subsidy, c_lvl2=c_lvl2, c_lvl3=c_lvl3, c_lvl4=c_lvl4,
        c_num1=c_num1, c_num2=c_num2, c_num3=c_num3, c_num4=c_num4,
        c_lvl5=c_lvl5, c_qty=c_qty, c_unit=c_unit, c_item_amt=c_item_amt,
        c_code=c_code, c_appendix=c_appendix, c_budget=c_budget, c_active=c_active,
        c_qty_lvl2=c_qty_lvl2, c_qty_lvl3=c_qty_lvl3, c_qty_lvl4=c_qty_lvl4,
        c_unit_lvl2=c_unit_lvl2, c_unit_lvl3=c_unit_lvl3, c_unit_lvl4=c_unit_lvl4,
        c_amt_lvl2=c_amt_lvl2, c_amt_lvl3=c_amt_lvl3, c_amt_lvl4=c_amt_lvl4,
        c_feo_qty_lvl2=c_feo_qty_lvl2, c_feo_qty_lvl3=c_feo_qty_lvl3, c_feo_qty_lvl4=c_feo_qty_lvl4,
        c_feo_unit_lvl2=c_feo_unit_lvl2, c_feo_unit_lvl3=c_feo_unit_lvl3, c_feo_unit_lvl4=c_feo_unit_lvl4,
        c_feo_amt_lvl2=c_feo_amt_lvl2, c_feo_amt_lvl3=c_feo_amt_lvl3, c_feo_amt_lvl4=c_feo_amt_lvl4,
        c_feo_sum_lvl2=c_feo_sum_lvl2, c_feo_sum_lvl3=c_feo_sum_lvl3, c_feo_sum_lvl4=c_feo_sum_lvl4,
        c_plan_sum_lvl2=c_plan_sum_lvl2, c_plan_sum_lvl3=c_plan_sum_lvl3, c_plan_sum_lvl4=c_plan_sum_lvl4,
        c_item_price=c_item_price,
        c_row_feo_qty=c_row_feo_qty, c_row_feo_unit=c_row_feo_unit,
        c_row_feo_price=c_row_feo_price, c_row_feo_sum=c_row_feo_sum,
        c_row_plan_qty=c_row_plan_qty, c_row_plan_unit=c_row_plan_unit,
        c_row_plan_price=c_row_plan_price, c_row_plan_sum=c_row_plan_sum,
        c_item_type=c_item_type,
        db=db, dry_run=dry_run,
        user=current_user, remap=remap, apply_remap=apply_remap,
        duplicate_resolutions=duplicate_resolutions,
        item_type_decisions=item_type_decisions,
    )


@router.post("/import-mapped")
async def import_feo_mapped(
    request: Request,
    file: UploadFile = File(...),
    sheet_name: str = Query(""),
    header_row_offset: int = Query(0),
    col_subsidy: int = Query(-1),
    col_lvl2: int = Query(-1),
    col_lvl3: int = Query(-1),
    col_lvl4: int = Query(-1),
    col_lvl5: int = Query(-1),
    col_code: int = Query(-1),
    col_appendix: int = Query(-1),
    col_budget: int = Query(-1),
    col_quantity: int = Query(-1),
    col_unit: int = Query(-1),
    col_item_amt: int = Query(-1),
    col_active: int = Query(-1),
    col_qty_lvl2: int = Query(-1),
    col_qty_lvl3: int = Query(-1),
    col_qty_lvl4: int = Query(-1),
    col_unit_lvl2: int = Query(-1),
    col_unit_lvl3: int = Query(-1),
    col_unit_lvl4: int = Query(-1),
    col_amt_lvl2: int = Query(-1),
    col_amt_lvl3: int = Query(-1),
    col_amt_lvl4: int = Query(-1),
    col_feo_qty_lvl2: int = Query(-1),
    col_feo_qty_lvl3: int = Query(-1),
    col_feo_qty_lvl4: int = Query(-1),
    col_feo_unit_lvl2: int = Query(-1),
    col_feo_unit_lvl3: int = Query(-1),
    col_feo_unit_lvl4: int = Query(-1),
    col_feo_amount_lvl2: int = Query(-1),
    col_feo_amount_lvl3: int = Query(-1),
    col_feo_amount_lvl4: int = Query(-1),
    # Новые параметры — сумма по ФЭО, сумма плана, цена за ед. Ур.5
    col_feo_sum_lvl2: int = Query(-1),
    col_feo_sum_lvl3: int = Query(-1),
    col_feo_sum_lvl4: int = Query(-1),
    col_plan_sum_lvl2: int = Query(-1),
    col_plan_sum_lvl3: int = Query(-1),
    col_plan_sum_lvl4: int = Query(-1),
    col_item_price: int = Query(-1),
    # Новый плоский 18-колоночный шаблон (2026-08-14) — одна пара колонок «по
    # ФЭО»/«плана» на всю строку, плюс тип плановой позиции.
    col_row_feo_qty: int = Query(-1),
    col_row_feo_unit: int = Query(-1),
    col_row_feo_price: int = Query(-1),
    col_row_feo_sum: int = Query(-1),
    col_row_plan_qty: int = Query(-1),
    col_row_plan_unit: int = Query(-1),
    col_row_plan_price: int = Query(-1),
    col_row_plan_sum: int = Query(-1),
    col_item_type: int = Query(-1),
    # Нумерация строк A–D (22.09) — явный маппинг из мастера, минуя авто-детект
    # (тот применяется только в /import, где колонки не размечены пользователем).
    col_num1: int = Query(-1),
    col_num2: int = Query(-1),
    col_num3: int = Query(-1),
    col_num4: int = Query(-1),
    default_subsidy_id: int = Query(-1),
    dry_run: bool = Query(False),
    remap: str = Query(""),
    apply_remap: bool = Query(False),
    duplicate_resolutions: str = Query(""),
    # Решения человека по конфликтам «тип позиции из файла vs тип товара
    # каталога» (владелец, 22.09) — см. пояснение у /import выше.
    item_type_decisions: str = Query(""),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_tab('feo_categories')),
):
    """Импорт категорий ФЭО с пользовательским маппингом столбцов.
    dry_run=true: вся обработка выполняется, транзакция откатывается; возвращает предупреждения.
    Переезд (remap) несопоставленных узлов и удаление опустевших старых узлов выполняются
    только при apply_remap=true; иначе выполняется только анализ (unmatched/new_paths).

    Баг владельца 2026-09-17 (HTTP 414 на субсидии "Центрпоиск_3"): фронт
    (useFeoImport.ts) теперь шлёт все параметры ниже ЧЕРЕЗ ТЕЛО multipart-
    формы вместе с файлом — все `col_*`, remap, duplicate_resolutions на
    больших субсидиях уходили на десятки КБ в query-строке, что выше лимита
    nginx `large_client_header_buffers`. Query(...) в сигнатуре ниже оставлен
    ТОЛЬКО ради обратной совместимости (PWA-кеш на проде мог сохранить
    старый фронт, шлющий эти же поля в query) — реальные значения теперь
    приходят через resolve_feo_import_mapped_params (Правило №6, один
    парсер и для формы, и для query, см. app/services/feo_import_params.py).
    """
    _p = await resolve_feo_import_mapped_params(
        request,
        sheet_name=sheet_name, header_row_offset=header_row_offset,
        col_subsidy=col_subsidy, col_lvl2=col_lvl2, col_lvl3=col_lvl3, col_lvl4=col_lvl4, col_lvl5=col_lvl5,
        col_code=col_code, col_appendix=col_appendix, col_budget=col_budget,
        col_quantity=col_quantity, col_unit=col_unit, col_item_amt=col_item_amt, col_active=col_active,
        col_qty_lvl2=col_qty_lvl2, col_qty_lvl3=col_qty_lvl3, col_qty_lvl4=col_qty_lvl4,
        col_unit_lvl2=col_unit_lvl2, col_unit_lvl3=col_unit_lvl3, col_unit_lvl4=col_unit_lvl4,
        col_amt_lvl2=col_amt_lvl2, col_amt_lvl3=col_amt_lvl3, col_amt_lvl4=col_amt_lvl4,
        col_feo_qty_lvl2=col_feo_qty_lvl2, col_feo_qty_lvl3=col_feo_qty_lvl3, col_feo_qty_lvl4=col_feo_qty_lvl4,
        col_feo_unit_lvl2=col_feo_unit_lvl2, col_feo_unit_lvl3=col_feo_unit_lvl3, col_feo_unit_lvl4=col_feo_unit_lvl4,
        col_feo_amount_lvl2=col_feo_amount_lvl2, col_feo_amount_lvl3=col_feo_amount_lvl3, col_feo_amount_lvl4=col_feo_amount_lvl4,
        col_feo_sum_lvl2=col_feo_sum_lvl2, col_feo_sum_lvl3=col_feo_sum_lvl3, col_feo_sum_lvl4=col_feo_sum_lvl4,
        col_plan_sum_lvl2=col_plan_sum_lvl2, col_plan_sum_lvl3=col_plan_sum_lvl3, col_plan_sum_lvl4=col_plan_sum_lvl4,
        col_item_price=col_item_price,
        col_row_feo_qty=col_row_feo_qty, col_row_feo_unit=col_row_feo_unit,
        col_row_feo_price=col_row_feo_price, col_row_feo_sum=col_row_feo_sum,
        col_row_plan_qty=col_row_plan_qty, col_row_plan_unit=col_row_plan_unit,
        col_row_plan_price=col_row_plan_price, col_row_plan_sum=col_row_plan_sum,
        col_item_type=col_item_type,
        col_num1=col_num1, col_num2=col_num2, col_num3=col_num3, col_num4=col_num4,
        default_subsidy_id=default_subsidy_id,
        dry_run=dry_run, remap=remap, apply_remap=apply_remap,
        duplicate_resolutions=duplicate_resolutions,
        item_type_decisions=item_type_decisions,
    )
    # Переменные, используемые ДО итогового вызова _do_feo_import ниже
    # (гейт прав, выбор листа/строки заголовка) — реальные значения теперь
    # только из `_p`; остальные ~40 col_* читаются из `_p` прямо в месте
    # вызова _do_feo_import, без промежуточного переприсваивания.
    sheet_name = _p["sheet_name"]
    header_row_offset = _p["header_row_offset"]
    col_subsidy = _p["col_subsidy"]
    col_lvl2 = _p["col_lvl2"]
    default_subsidy_id = _p["default_subsidy_id"]
    dry_run = _p["dry_run"]
    remap = _p["remap"]
    apply_remap = _p["apply_remap"]
    duplicate_resolutions = _p["duplicate_resolutions"]
    item_type_decisions = _p["item_type_decisions"]

    if col_lvl2 < 0:
        raise HTTPException(400, "Не указан обязательный столбец: Уровень 2")
    if col_subsidy < 0 and default_subsidy_id <= 0:
        raise HTTPException(400, "Укажите столбец Субсидия или выберите субсидию назначения")

    # B2: если субсидия назначения ОДНА на весь импорт (нет построчной колонки
    # «Субсидия», задан только default_subsidy_id) — проверяем право именно по ней
    # уже здесь (строже и точнее, отказ до чтения файла). Если субсидия построчная
    # (col_subsidy задан) — единого subsidy_id нет, здесь только дешёвый
    # предварительный фильтр «право хоть где-то» (как в /import). В ОБОИХ случаях
    # авторитетная проверка ПО КАЖДОЙ реально резолвящейся субсидии файла (включая
    # default_subsidy_id как построчный фолбэк при заданном col_subsidy — именно
    # тут была дыра) происходит внутри _do_feo_import (_require_feo_import_write),
    # ДО первой записи в БД.
    _gate_subsidy_id = default_subsidy_id if (col_subsidy < 0 and default_subsidy_id > 0) else None
    await fc._require_feo_category_write(current_user, db, _gate_subsidy_id)

    fname = (file.filename or "").lower()
    content = await file.read()

    try:
        if fname.endswith(".xls"):
            try:
                import xlrd as _xlrd_mod
            except ImportError:
                raise HTTPException(500, "xlrd не установлен")
            wb_xls = _xlrd_mod.open_workbook(file_contents=content)
            ws_names = wb_xls.sheet_names()
            target_sheet = sheet_name if sheet_name in ws_names else ws_names[0]
            ws_xls = wb_xls.sheet_by_name(target_sheet)
            all_rows = [list(ws_xls.row_values(i)) for i in range(ws_xls.nrows)]
        elif fname.endswith(".pdf"):
            try:
                import pdfplumber
            except ImportError:
                raise HTTPException(500, "pdfplumber не установлен")
            pdf = pdfplumber.open(BytesIO(content))
            all_rows = []
            for page in pdf.pages:
                for t in (page.extract_tables() or []):
                    if t:
                        all_rows.extend([[str(c).strip() if c else "" for c in row] for row in t])
            pdf.close()
        elif fname.endswith((".docx", ".doc")):
            try:
                from docx import Document as _DDoc
            except ImportError:
                raise HTTPException(500, "python-docx не установлен")
            doc = _DDoc(BytesIO(content))
            all_rows = []
            for table in doc.tables:
                for row in table.rows:
                    all_rows.append([cell.text.strip() for cell in row.cells])
        else:
            if load_workbook is None:
                raise HTTPException(500, "openpyxl не установлен")
            wb = load_workbook(BytesIO(content), data_only=True)
            ws_names = wb.sheetnames
            target_sheet = sheet_name if sheet_name in ws_names else ws_names[0]
            ws = wb[target_sheet]
            all_rows = list(ws.iter_rows(values_only=True))
            wb.close()
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Не удалось прочитать файл: {e}")

    if len(all_rows) <= header_row_offset + 1:
        raise HTTPException(400, "Файл пустой или не содержит данных после строки заголовка")

    data_rows = all_rows[header_row_offset + 1:]

    # Оставшиеся ~40 col_* читаются напрямую из `_p` (уже разрешённых форма/query,
    # см. resolve_feo_import_mapped_params выше) — им не нужно промежуточное
    # переприсваивание локальной переменной, они используются только здесь.
    def _c(name: str):
        v = _p[name]
        return v if v >= 0 else None

    return await _do_feo_import(
        rows=data_rows,
        c_subsidy=_c("col_subsidy"),
        c_lvl2=col_lvl2,
        c_lvl3=_c("col_lvl3"),
        c_lvl4=_c("col_lvl4"),
        c_lvl5=_c("col_lvl5"),
        c_qty=_c("col_quantity"),
        c_unit=_c("col_unit"),
        c_item_amt=_c("col_item_amt"),
        c_code=_c("col_code"),
        c_appendix=_c("col_appendix"),
        c_budget=_c("col_budget"),
        c_active=_c("col_active"),
        c_qty_lvl2=_c("col_qty_lvl2"),
        c_qty_lvl3=_c("col_qty_lvl3"),
        c_qty_lvl4=_c("col_qty_lvl4"),
        c_unit_lvl2=_c("col_unit_lvl2"),
        c_unit_lvl3=_c("col_unit_lvl3"),
        c_unit_lvl4=_c("col_unit_lvl4"),
        c_amt_lvl2=_c("col_amt_lvl2"),
        c_amt_lvl3=_c("col_amt_lvl3"),
        c_amt_lvl4=_c("col_amt_lvl4"),
        c_feo_qty_lvl2=_c("col_feo_qty_lvl2"),
        c_feo_qty_lvl3=_c("col_feo_qty_lvl3"),
        c_feo_qty_lvl4=_c("col_feo_qty_lvl4"),
        c_feo_unit_lvl2=_c("col_feo_unit_lvl2"),
        c_feo_unit_lvl3=_c("col_feo_unit_lvl3"),
        c_feo_unit_lvl4=_c("col_feo_unit_lvl4"),
        c_feo_amt_lvl2=_c("col_feo_amount_lvl2"),
        c_feo_amt_lvl3=_c("col_feo_amount_lvl3"),
        c_feo_amt_lvl4=_c("col_feo_amount_lvl4"),
        c_feo_sum_lvl2=_c("col_feo_sum_lvl2"),
        c_feo_sum_lvl3=_c("col_feo_sum_lvl3"),
        c_feo_sum_lvl4=_c("col_feo_sum_lvl4"),
        c_plan_sum_lvl2=_c("col_plan_sum_lvl2"),
        c_plan_sum_lvl3=_c("col_plan_sum_lvl3"),
        c_plan_sum_lvl4=_c("col_plan_sum_lvl4"),
        c_item_price=_c("col_item_price"),
        c_row_feo_qty=_c("col_row_feo_qty"),
        c_row_feo_unit=_c("col_row_feo_unit"),
        c_row_feo_price=_c("col_row_feo_price"),
        c_row_feo_sum=_c("col_row_feo_sum"),
        c_row_plan_qty=_c("col_row_plan_qty"),
        c_row_plan_unit=_c("col_row_plan_unit"),
        c_row_plan_price=_c("col_row_plan_price"),
        c_row_plan_sum=_c("col_row_plan_sum"),
        c_item_type=_c("col_item_type"),
        c_num1=_c("col_num1"), c_num2=_c("col_num2"), c_num3=_c("col_num3"), c_num4=_c("col_num4"),
        default_subsidy_id=default_subsidy_id if default_subsidy_id > 0 else None,
        db=db, dry_run=dry_run,
        user=current_user, remap=remap, apply_remap=apply_remap,
        duplicate_resolutions=duplicate_resolutions,
        item_type_decisions=item_type_decisions,
    )
