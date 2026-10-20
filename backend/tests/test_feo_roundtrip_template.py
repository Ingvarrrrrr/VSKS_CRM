"""Тест GET /api/feo-categories/roundtrip/template — шаблон формата
«Выгрузить ФЭО» для загрузки обратно (владелец, 20.10.2026, часть «шаблон»).

Сценарии:
  1) файл скачивается, основной лист не содержит ни одной непустой
     «№ в системе», есть лист «Как заполнять».
  2) содержимое шаблона разбирается parse_roundtrip_workbook + build_preview
     на ПУСТОЙ субсидии -> errors=0, все строки примера — created (категории
     и позиции), уровни верные.
  3) apply строит дерево в БД, совпадающее с примером (количество категорий/
     позиций и вложенность).

Запуск: docker compose -p vsks_crm exec backend_a pytest
tests/test_feo_roundtrip_template.py -q
"""
import io

import pytest
from sqlalchemy import select

from app.models.subsidy import Subsidy
from app.models.feo_category import FeoCategory
from app.models.feo_planned_item import FeoPlannedItem
from app.auth.jwt import create_access_token


def _headers_for(user):
    token = create_access_token({"sub": user.username, "org_id": user.org_id})
    return {"Authorization": f"Bearer {token}"}


async def _mk_subsidy(db_session, org_id, name):
    sub = Subsidy(name=name, org_id=org_id, year=2026)
    db_session.add(sub)
    await db_session.commit()
    await db_session.refresh(sub)
    return sub


@pytest.mark.asyncio
async def test_roundtrip_template_downloads_with_howto_sheet_and_empty_row_ids(
    db_session, test_org, make_user, client,
):
    owner = await make_user(role="superadmin", org_id=test_org.id)
    headers = _headers_for(owner)

    resp = await client.get("/api/feo-categories/roundtrip/template", headers=headers)
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Как заполнять" in wb.sheetnames
    assert wb.active.title != "Как заполнять"

    ws = wb.active
    hmap = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=3, column=c).value
        if v:
            hmap[str(v).strip()] = c
    assert "№ в системе" in hmap
    assert "Тип строки" in hmap

    row_id_col = hmap["№ в системе"]
    for r in range(4, ws.max_row + 1):
        name_val = ws.cell(row=r, column=hmap["Наименование"]).value
        if not name_val or not str(name_val).strip():
            continue
        assert ws.cell(row=r, column=row_id_col).value in (None, ""), \
            f"строка {r} несёт непустую «№ в системе»"


@pytest.mark.asyncio
async def test_roundtrip_template_parses_and_previews_as_all_created(
    db_session, test_org, make_user, client,
):
    """Шаблон, разобранный на ПУСТОЙ субсидии, должен дать errors=0 и
    created на все содержательные строки примера (категории+позиции)."""
    from app.services.feo_roundtrip_import.parse import parse_roundtrip_workbook
    from app.services.feo_roundtrip_import.diff import build_preview

    owner = await make_user(role="superadmin", org_id=test_org.id)
    headers = _headers_for(owner)
    sub = await _mk_subsidy(db_session, test_org.id, "RT-template-empty")

    resp = await client.get("/api/feo-categories/roundtrip/template", headers=headers)
    assert resp.status_code == 200, resp.text
    content = resp.content

    parsed = parse_roundtrip_workbook(content)
    assert parsed.errors == []
    assert parsed.meta is not None
    assert parsed.meta.subsidy_id is None  # шаблон не привязан к субсидии

    preview = await build_preview(db_session, sub.id, parsed.rows, parsed.meta)
    assert preview.errors == []
    assert len(preview.created) == len(parsed.rows)
    assert preview.missing == []

    created_categories = [c for c in preview.created if c["kind"] == "category"]
    created_positions = [c for c in preview.created if c["kind"] == "position"]
    assert len(created_categories) >= 5  # 2 направления + подкатегории + вложенная
    assert len(created_positions) >= 5


@pytest.mark.asyncio
async def test_roundtrip_template_apply_builds_matching_tree(
    db_session, test_org, make_user, client,
):
    from app.services.feo_roundtrip_import.parse import parse_roundtrip_workbook
    from app.services.feo_roundtrip_import.diff import build_preview

    owner = await make_user(role="superadmin", org_id=test_org.id)
    headers = _headers_for(owner)
    sub = await _mk_subsidy(db_session, test_org.id, "RT-template-apply")
    sub_id = sub.id

    resp = await client.get("/api/feo-categories/roundtrip/template", headers=headers)
    content = resp.content

    files = {"file": ("template.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    data = {"subsidy_id": str(sub_id)}
    resp2 = await client.post("/api/feo-categories/roundtrip/apply", headers=headers, files=files, data=data)
    assert resp2.status_code == 200, resp2.text
    result = resp2.json()
    assert result["created"] >= 10
    assert result["updated"] == 0
    assert result["moved"] == 0

    db_session.expire_all()
    cats = (await db_session.execute(select(FeoCategory).where(FeoCategory.subsidy_id == sub_id))).scalars().all()
    items = (await db_session.execute(
        select(FeoPlannedItem).join(FeoCategory, FeoPlannedItem.feo_category_id == FeoCategory.id)
        .where(FeoCategory.subsidy_id == sub_id)
    )).scalars().all()

    # Независимая проверка через свежий разбор файла — Σ категорий/позиций
    # в файле == Σ в БД (ПРАВИЛО №6: одно и то же дерево, не два расчёта).
    parsed = parse_roundtrip_workbook(content)
    preview = await build_preview(db_session, sub_id, parsed.rows, parsed.meta)
    expected_categories = sum(1 for rr in preview.resolved if rr.kind == "category")
    expected_positions = sum(1 for rr in preview.resolved if rr.kind == "position")
    assert len(cats) == expected_categories
    assert len(items) == expected_positions

    by_name = {c.name: c for c in cats}
    assert "Оснащение полевого лагеря" in by_name
    assert "Обучение добровольцев" in by_name
    lager = by_name["Оснащение полевого лагеря"]
    assert lager.level == 1
    tents = by_name["Палатки и спальное"]
    assert tents.level == 2
    assert tents.parent_id == lager.id
    kitchen = by_name["Кухонный инвентарь"]
    assert kitchen.level == 3
    assert kitchen.parent_id == by_name["Освещение и кухня"].id

    item_names = {it.name for it in items}
    assert "Палатка 4-местная" in item_names
    inactive_items = [it for it in items if not it.is_active]
    assert len(inactive_items) >= 1  # пример «Активна = Нет» применился
